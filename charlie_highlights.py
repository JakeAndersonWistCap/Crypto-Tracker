"""charlie_highlights.py - THE "HIGHLIGHTS" FIRST TAB OF THE CHARLIE EDITION (P2g-3, Jake 2026-10-29).

docs/charlie_edition_spec.md section 2. `add_highlights(wb, ctx)` adds one sheet to the reader workbook being built:
a title, the as-of line and two lines copied from About, up to eight cards (C1..C8), a coverage block and one closing
line. Additive only: nothing on any existing tab is read for anything but its values, and nothing there is written.

  * Every number on the tab is a FORMULA pointing at the reader cell it quotes. The cell is found at build time by tab +
    project row + header text (Lookup.ref), never by a typed address; it carries the source cell's number format, its
    confidence fill (the one fill meaning, copied, never added) and a comment naming the source.
  * A card has a build-time condition, evaluated on the SAME cell values its chips quote (xlcalc). A card whose
    condition fails is omitted and the others reflow; fewer than `min_cards` rendered raises HighlightsError listing
    every failed condition with its values. The condition functions (c1_holds ... c8_split) are plain functions of
    numbers, unit-tested on their own.
  * Parked projects (reader.parked_reason) are in no card and no mini-table; they are listed in the coverage block only.
  * Text is fixed. Where a sentence adds a clause, the clause's condition is evaluated here from the quoted cells.
    No number is typed into any sentence, and no word refers to an earlier build (FORBIDDEN, scanned by the tests).
  * Fill is never used for decoration: font colour, borders and bars only.
"""
from __future__ import annotations

import math
import re
from dataclasses import dataclass, field

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.properties import PageSetupProperties

import reader
import xlcalc

SHEET = "Highlights"
TITLE = "Token metrics — highlights"
FOOTER = "Confidence is shown by cell fill throughout; hover any number for its source."
FONT = "Calibri"
NAVY, TEAL, GREY, RULE = "14213D", "1F8A8A", "595959", "BFBFBF"
MIN_CARDS = 5
CHIPS_PER_ROW = 4
CARD_W = 8                                   # columns in a card (four chips of two columns)
LEFT0, GUTTER, RIGHT0 = 1, 9, 10             # first column of the left card, the gutter column, first of the right card
LAST_COL = RIGHT0 + CARD_W - 1               # 17
# the words that would date the page or refer to an earlier build; the tests scan every string of the tab for them
FORBIDDEN = re.compile(r"\b(still|now|newly)\b|improved since|\b(earlier|previous|last|prior) (build|run|version)\b"
                       r"|\bupdated since\b|yesterday", re.I)


class HighlightsError(RuntimeError):
    """The tab cannot be built: too few cards, or a chip that evaluates to an error or blank."""


class Missing(Exception):
    """A tab, project row or header the card quotes is not on the reader."""


class Omit(Exception):
    """A card's condition fails (or a cell it quotes is missing): the reason is kept for the build's report."""


# ======================================================================================================== conditions
def num(v):
    """A number, or None (a code such as 'n/a', blank, an Excel error and a bool are none)."""
    if isinstance(v, bool) or isinstance(v, xlcalc.XLError) or not isinstance(v, (int, float)):
        return None
    return None if (isinstance(v, float) and math.isnan(v)) else float(v)


def c1_holds(key, q1, q2, q3) -> bool:
    """C1: the key is at least 1.0 and strictly above each earlier quarter, which fall in order (key > -1q > -2q > -3q)."""
    k, a, b, c = num(key), num(q1), num(q2), num(q3)
    return None not in (k, a, b, c) and k >= 1.0 and k > a > b > c


def c2_holds(key, q3) -> bool:
    """C2: the key is at least 10 and at least 1.3 times its value three quarters back."""
    k, c = num(key), num(q3)
    return None not in (k, c) and k >= 10.0 and k >= 1.3 * c


def c3_pick(nhy: dict):
    """C3: (project, how many are positive) - the project with the highest Net holder yield among those given (already
    non-parked), only when that value is above 0; (None, n) when none is."""
    vals = {n: num(v) for n, v in nhy.items()}
    vals = {n: v for n, v in vals.items() if v is not None}
    pos = sum(1 for v in vals.values() if v > 0)
    if not vals:
        return None, 0
    best = max(vals, key=lambda n: vals[n])            # first of equals, in row order
    return (best if vals[best] > 0 else None), pos


def c4_holds(rank, in_a3: bool, key) -> bool:
    """C4: Summary's rank in archetype is 1 (in A3) and the A3 key is at least 0.9."""
    r, k = num(rank), num(key)
    return bool(in_a3) and r == 1.0 and k is not None and k >= 0.9


def is_largest(name: str, values: dict) -> bool:
    """True when `name` holds the (single) largest number of `values`."""
    vals = {n: num(v) for n, v in values.items()}
    vals = {n: v for n, v in vals.items() if v is not None}
    return name in vals and all(vals[name] > v for n, v in vals.items() if n != name)


def c5_holds(growth, net_supply) -> bool:
    """C5: revenue growth above 0.5 (50%) and a net supply change below 0."""
    g, s = num(growth), num(net_supply)
    return None not in (g, s) and g > 0.5 and s < 0


def c6_holds(latest, earlier) -> bool:
    """C6: the latest quarter of the key is 0 and the largest of the three earlier quarters is above 0."""
    lat = num(latest)
    e = [num(x) for x in earlier]
    e = [x for x in e if x is not None]
    return lat == 0.0 and bool(e) and max(e) > 0


def is_second_best(name: str, values: dict) -> bool:
    """True when exactly one project has a larger number than `name`."""
    vals = {n: num(v) for n, v in values.items()}
    vals = {n: v for n, v in vals.items() if v is not None}
    return name in vals and sum(1 for v in vals.values() if v > vals[name]) == 1


def c7_holds(name: str, growth: dict) -> bool:
    """C7: `name` has the highest revenue growth (strictly) among the projects given."""
    return is_largest(name, growth)


def c8_split(values: dict):
    """C8: ([three lowest], [three highest]) as (project, value), the highest first; None when fewer than six have a
    value (the two lists would overlap)."""
    vals = [(n, num(v)) for n, v in values.items()]
    vals = [(n, v) for n, v in vals if v is not None]
    if len(vals) < 6:
        return None
    asc = sorted(vals, key=lambda t: t[1])
    return asc[:3], sorted(vals, key=lambda t: -t[1])[:3]


# ================================================================================================= number display
def _split_sections(fmt: str) -> list:
    out, cur, q, br = [], "", False, False
    for ch in fmt:
        if ch == '"':
            q = not q
        elif ch == "[" and not q:
            br = True
        elif ch == "]" and not q:
            br = False
        if ch == ";" and not q and not br:
            out.append(cur)
            cur = ""
        else:
            cur += ch
    return out + [cur]


def _cond(sec: str):
    m = re.match(r"\s*\[(>=|<=|<>|>|<|=)\s*(-?[\d.]+)\]", sec)
    if not m:
        return None
    op, n = m.group(1), float(m.group(2))
    return {">=": lambda x: x >= n, "<=": lambda x: x <= n, ">": lambda x: x > n, "<": lambda x: x < n,
            "=": lambda x: x == n, "<>": lambda x: x != n}[op]


def _render_section(sec: str, x: float, sign: str) -> str:
    s = re.sub(r"\[[^\]]*\]", "", sec)
    items, i, run, pct = [], 0, "", False
    while i < len(s):
        ch = s[i]
        if ch in "#0?,.":
            run += ch
            i += 1
            continue
        if run:
            items.append(("num", run))
            run = ""
        if ch == '"':
            j = s.find('"', i + 1)
            j = len(s) if j < 0 else j
            items.append(("lit", s[i + 1:j]))
            i = j + 1
            continue
        if ch in "\\_*":
            if ch == "\\" and i + 1 < len(s):
                items.append(("lit", s[i + 1]))
            i += 2
            continue
        if ch == "%":
            pct = True
        items.append(("lit", ch))
        i += 1
    if run:
        items.append(("num", run))
    out = []
    for kind, t in items:
        if kind == "lit":
            out.append(t)
            continue
        scale = 1
        while t.endswith(","):
            t = t[:-1]
            scale *= 1000
        dec = len(t.split(".")[1]) if "." in t else 0
        y = x * (100 if pct else 1) / scale
        out.append(f"{y:,.{dec}f}" if "," in t else f"{y:.{dec}f}")
    return sign + "".join(out)


def format_value(v, fmt) -> str:
    """A cell value as Excel displays it under number format `fmt` (the formats the reader writes: 0.0%, 0.0"x", the
    dollar scale formats, #,##0, status zeros). A text value is shown as it is."""
    if isinstance(v, xlcalc.XLError):
        return v.code
    n = num(v)
    if n is None:
        return "" if v is None else str(v)
    fmt = fmt or "General"
    if fmt == "General":
        return f"{n:.6g}"
    secs = _split_sections(fmt)
    conds = [_cond(s) for s in secs]
    sign = "-" if n < 0 else ""
    chosen, x = secs[0], abs(n)
    if any(conds):
        chosen = next((s for s, c in zip(secs, conds) if c and c(n)), secs[-1])
    elif len(secs) >= 2:
        if n < 0:
            chosen, sign = secs[1], ""
        elif n == 0 and len(secs) >= 3:
            chosen = secs[2]
    return _render_section(chosen, x, sign)


# ================================================================================================== finding cells
@dataclass
class Ref:
    tab: str
    row: int
    col: int
    project: str
    header: str
    value: object
    fmt: str
    cell: object

    @property
    def coord(self) -> str:
        return f"{get_column_letter(self.col)}{self.row}"

    @property
    def formula(self) -> str:
        return f"={quote_sheet(self.tab)}!{self.coord}"


def quote_sheet(name: str) -> str:
    return f"'{name}'" if re.search(r"[^A-Za-z0-9_]", name) else name


class Lookup:
    """Cells of the finished reader by tab + project row + header text, valued by xlcalc. Header row 4, project rows
    from row 5 (a row with its name in column A and anything beside it; a section row holds the name alone)."""

    def __init__(self, wb, parked=None):
        self.wb = wb
        self.X = xlcalc.Workbook(wb)
        self.parked = parked or reader.parked_reason
        self._hdr: dict = {}
        self._rows: dict = {}

    def has(self, tab: str) -> bool:
        return tab in self.wb.sheetnames

    def headers(self, tab: str) -> dict:
        if tab not in self._hdr:
            if not self.has(tab):
                raise Missing(f"tab {tab!r}")
            ws = self.wb[tab]
            h: dict = {}
            for c in range(1, ws.max_column + 1):
                v = ws.cell(4, c).value
                if isinstance(v, str):
                    h.setdefault(v, []).append(c)
            self._hdr[tab] = h
        return self._hdr[tab]

    def rows(self, tab: str) -> dict:
        if tab not in self._rows:
            if not self.has(tab):
                raise Missing(f"tab {tab!r}")
            ws = self.wb[tab]
            out: dict = {}
            for r in range(5, ws.max_row + 1):
                a = ws.cell(r, 1).value
                if a and any(ws.cell(r, c).value is not None for c in range(2, ws.max_column + 1)):
                    out.setdefault(str(a), r)
            self._rows[tab] = out
        return self._rows[tab]

    def ref(self, tab: str, project: str, header: str) -> Ref:
        cols = self.headers(tab).get(header)
        if not cols:
            raise Missing(f"{tab}: no header {header!r}")
        if len(cols) > 1:
            raise Missing(f"{tab}: header {header!r} is not unique")
        row = self.rows(tab).get(project)
        if row is None:
            raise Missing(f"{tab}: no row for {project}")
        c = self.wb[tab].cell(row, cols[0])
        if c.value is None:
            raise Missing(f"{tab}: {project} / {header} is empty")
        return Ref(tab, row, cols[0], project, header, self.X.value(tab, row, cols[0]), c.number_format, c)

    def column(self, tab: str, header: str, include_parked: bool = False) -> dict:
        """{project: value} down a column, parked projects left out."""
        cols = self.headers(tab).get(header)
        if not cols or len(cols) > 1:
            raise Missing(f"{tab}: header {header!r} is {'missing' if not cols else 'not unique'}")
        return {n: self.X.value(tab, r, cols[0]) for n, r in self.rows(tab).items()
                if include_parked or not self.parked(n)}


# ======================================================================================================= the cards
@dataclass
class Chip:
    label: str
    ref: Ref


@dataclass
class Card:
    cid: str
    project: str
    headline: str
    chips: list
    reading: str
    caveat: str
    link_tab: str
    link_row: int
    facts: dict = field(default_factory=dict)       # the condition's values, for the report
    table: dict | None = None                       # C8: {"low": [(name, ref)], "high": [...], "header": str, "ref_header"}


_STEMS: dict = {}


def key_headers(tab: str):
    """(the tab's key header, [its key's three earlier-quarter headers]) from reader's own TabSpec and _traj_header."""
    if tab not in _STEMS:
        import build_workbook as bw                        # noqa: PLC0415
        spec = next(s for s in reader.tab_specs(bw) if s.tab == tab)
        stem = spec.key_traj[1]
        _STEMS[tab] = (spec.key_header, [reader._traj_header(stem, lab) for lab in reader.TRAJ_LABELS[1:]])
    return _STEMS[tab]


def _key_chips(L: Lookup, tab: str, project: str) -> list:
    head, earlier = key_headers(tab)
    return [L.ref(tab, project, h) for h in [head] + earlier]


def _show(v) -> str:
    n = num(v)
    return f"{n:.4g}" if n is not None else repr(v)


def _fail(cid: str, project: str, why: str, **vals):
    shown = ", ".join(f"{k}={_show(v)}" for k, v in vals.items())
    raise Omit(f"{cid} {project}: condition not met - {why}" + (f" [{shown}]" if shown else ""))


def _chips(refs) -> list:
    return [Chip(r.header, r) for r in refs]


def card_c1(L: Lookup) -> Card:
    p, tab = "GEODNET", reader.A2_TAB
    k = _key_chips(L, tab, p)
    hvy = L.ref(reader.COMPARISON, p, "Holder value yield")
    v = [r.value for r in k]
    if not c1_holds(*v):
        _fail("C1", p, "needs key >= 1.0 and key > -1q > -2q > -3q", key=v[0], q1=v[1], q2=v[2], q3=v[3])
    return Card("C1", p, f"{p} — Customer revenue exceeds the emissions that pay for it", _chips(k + [hvy]),
                "The key has risen every quarter shown and sits above the crossover.",
                "Net holder yield is negative: dilution exceeds what reaches holders.", tab, k[0].row,
                {"key": v[0], "-1q": v[1], "-2q": v[2], "-3q": v[3]})


def card_c2(L: Lookup) -> Card:
    p, tab = "Morpho", reader.A2_TAB
    k = _key_chips(L, tab, p)
    rev = L.ref(reader.SUMMARY, p, "Protocol revenue, annualised")
    tvl = L.ref(reader.BUSINESS, p, "TVL change, selected period")
    if not c2_holds(k[0].value, k[3].value):
        _fail("C2", p, "needs key >= 10 and key >= 1.3 x its value three quarters back", key=k[0].value,
              q3=k[3].value)
    return Card("C2", p, f"{p} — Customer revenue is many times the emissions that pay for it",
                _chips([k[0], k[3], rev, tvl]),
                "The business earns far more than it pays out in tokens; none of that revenue reaches the token yet "
                "(protocol revenue is a measured value; see the cell comment).",
                "Holder yield is n/a until revenue reaches the token; this is value created, not yet captured.",
                tab, k[0].row, {"key": k[0].value, "-3q": k[3].value})


def card_c3(L: Lookup) -> Card:
    nhy = L.column(reader.COMPARISON, "Net holder yield")
    p, npos = c3_pick(nhy)
    if p is None:
        _fail("C3", "(highest Net holder yield)", "no non-parked project has a Net holder yield above 0",
              positive_projects=npos)
    C = reader.COMPARISON
    refs = [L.ref(C, p, "Net holder yield"), L.ref(C, p, "Holder value yield"),
            L.ref(C, p, "Total dilution, % of float"), L.ref(C, p, "Lock rate, % of circulating")]
    lock = num(refs[3].value)
    clause = "the only project in the set where that holds" if npos == 1 else "the highest in the set"
    reading = f"Measured returns to holders exceed measured dilution; {clause}"
    if lock is not None and lock > 0.5:
        reading += "; most of the circulating supply is locked"
    reading += "."
    if p == "Sky":
        caveat = ("The buyback retires well below the amount emitted, and revenue growth is flat over the selected "
                  "period.")
    else:
        caveat = ("Net holder yield is the measured holder value yield less the measured dilution; hover each number "
                  "for its source and confidence.")
    return Card("C3", p, f"{p} — Holders are paid more than they are diluted", _chips(refs), reading, caveat, C,
                refs[0].row, {"net holder yield": refs[0].value, "positive projects": npos, "lock rate": lock})


def card_c4(L: Lookup) -> Card:
    p, tab = "Hyperliquid", reader.A3_TAB
    k = _key_chips(L, tab, p)
    S = reader.SUMMARY
    rank = L.ref(S, p, "Rank in archetype")
    rev = L.ref(S, p, "Protocol revenue, annualised")
    gro = L.ref(S, p, "Revenue growth, selected period")
    sec = [lab for lab, lo, hi in reader._summary_sections(L.wb[S]) if lo <= rank.row <= hi]
    in_a3 = bool(sec) and sec[0].startswith("A3")
    if not c4_holds(rank.value, in_a3, k[0].value):
        _fail("C4", p, "needs Summary rank in archetype 1 (in A3) and the A3 key >= 0.9", rank=rank.value,
              in_A3=in_a3, key=k[0].value)
    reading = "Ranked first in A3, at the break-even line between retired and emitted tokens"
    if is_largest(p, L.column(S, "Protocol revenue, annualised")):
        reading += ", on the largest revenue base in the set"
    return Card("C4", p, f"{p} — Buyback retires as many tokens as the protocol emits",
                _chips(k + [rev, gro]), reading + ".",
                "The key has narrowed over three quarters, net holder yield is negative, and the A4 crossover is "
                "blocked for lack of same-span coverage.", tab, k[0].row, {"rank": rank.value, "key": k[0].value})


def card_c5(L: Lookup) -> Card:
    p, S = "Uniswap", reader.SUMMARY
    rev = L.ref(S, p, "Protocol revenue, annualised")
    gro = L.ref(S, p, "Revenue growth, selected period")
    burn = L.ref(reader.A4_TAB, p, "Permanent burn yield")
    net = L.ref(S, p, "Net supply change, % a year")
    if not c5_holds(gro.value, net.value):
        _fail("C5", p, "needs revenue growth > 0.5 and net supply change < 0", growth=gro.value, net_supply=net.value)
    return Card("C5", p, f"{p} — Protocol revenue has scaled and supply is shrinking",
                _chips([rev, gro, burn, net]),
                "Revenue is wired to the token and the circulating supply is falling.",
                "The A4 key is n/a (nothing issued to divide by) and net holder yield is negative.", S, rev.row,
                {"growth": gro.value, "net supply change": net.value})


def card_c6(L: Lookup) -> Card:
    p, tab = "Fluid", reader.A3_TAB
    k = _key_chips(L, tab, p)
    C = reader.COMPARISON
    cover = L.ref(C, p, "Subsidy cover (x)")
    nhy = L.ref(C, p, "Net holder yield")
    v = [r.value for r in k]
    if not c6_holds(v[0], v[1:]):
        _fail("C6", p, "needs the latest key = 0 and the largest earlier quarter > 0", latest=v[0], q1=v[1], q2=v[2],
              q3=v[3])
    reading = "Retired ÷ emissions fell to zero over four quarters while revenue grew and subsidy cover is high"
    if is_second_best(p, L.column(C, "Net holder yield")):
        reading += ", and net holder yield is second best in the set"
    return Card("C6", p, f"{p} — The tracker shows the buyback pause", _chips(k + [cover, nhy]), reading + ".",
                "A pause is a policy choice; the key does not judge intent.", tab, k[0].row,
                {"latest": v[0], "-1q": v[1], "-2q": v[2], "-3q": v[3]})


def card_c7(L: Lookup) -> Card:
    p, S = "Near", reader.SUMMARY
    gro = L.ref(S, p, "Revenue growth, selected period")
    tvl = L.ref(reader.BUSINESS, p, "TVL change, selected period")
    price = L.ref(reader.BUSINESS, p, "Price change, selected period")
    nhy = L.ref(reader.COMPARISON, p, "Net holder yield")
    if not c7_holds(p, L.column(S, "Revenue growth, selected period")):
        _fail("C7", p, "needs the highest revenue growth among non-parked projects", growth=gro.value)
    return Card("C7", p, f"{p} — Fastest-growing business in the set", _chips([gro, tvl, price, nhy]),
                "Revenue and TVL are growing faster than any other project over the selected period.",
                "The price moved with it over the same window, so part of the growth is priced in.", S, gro.row,
                {"growth": gro.value})


def card_c8(L: Lookup) -> Card:
    C, head = reader.COMPARISON, "Total dilution, % of float"
    split = c8_split(L.column(C, head))
    if split is None:
        raise Omit(f"C8 Dilution league: condition not met - fewer than six non-parked projects have a numeric {head}")
    low, high = split
    return Card("C8", "Dilution league", "Dilution league", [],
                "The same measures that show improvement show who pays for growth with new tokens.", "", C,
                L.rows(C)[low[0][0]], {"lowest": low, "highest": high},
                {"low": [L.ref(C, n, head) for n, _ in low], "high": [L.ref(C, n, head) for n, _ in high],
                 "header": head})


CARD_BUILDERS = (card_c1, card_c2, card_c3, card_c4, card_c5, card_c6, card_c7, card_c8)


def build_cards(L: Lookup):
    """([Card], [omitted reasons]) in order C1..C8. A tab, row or header a card quotes being absent omits that card."""
    cards, omitted = [], []
    for fn in CARD_BUILDERS:
        try:
            cards.append(fn(L))
        except Omit as e:
            omitted.append(str(e))
        except Missing as e:
            omitted.append(f"{fn.__name__[5:].upper()}: a cell it quotes is not on the reader - {e}")
    return cards, omitted


# ========================================================================================================== styling
def _font(size=10, bold=False, color="000000", italic=False, underline=None) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)


_THIN = Side(style="thin", color=RULE)
_TEAL = Side(style="thin", color=TEAL)
_NAVY = Side(style="medium", color=NAVY)


def _border(cell, **sides) -> None:
    b = cell.border
    cell.border = Border(left=sides.get("left", b.left), right=sides.get("right", b.right),
                         top=sides.get("top", b.top), bottom=sides.get("bottom", b.bottom))


def _outline(ws, r1, c1, r2, c2, side=_THIN) -> None:
    for c in range(c1, c2 + 1):
        _border(ws.cell(r1, c), top=side)
        _border(ws.cell(r2, c), bottom=side)
    for r in range(r1, r2 + 1):
        _border(ws.cell(r, c1), left=side)
        _border(ws.cell(r, c2), right=side)


def _put(ws, r, c1, c2, value, font, align=None, height=None):
    if c2 > c1:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    cell = ws.cell(r, c1, value)
    cell.font = font
    cell.alignment = align or Alignment(wrap_text=True, vertical="top")
    if height:
        ws.row_dimensions[r].height = max(height, ws.row_dimensions[r].height or 0)
    return cell


def _lines(text: str, chars: int) -> int:
    return max(1, math.ceil(len(text) / chars))


def _chunks(items: list, n: int) -> list:
    return [items[i:i + n] for i in range(0, len(items), n)] or [[]]


def _chip(ws, r_label, r_val, c0, chip: Chip):
    ref = chip.ref
    _put(ws, r_label, c0, c0 + 1, chip.label, _font(8, color=GREY), Alignment(wrap_text=True, vertical="bottom",
                                                                           horizontal="center"), 24)
    code = num(ref.value) is None
    cell = _put(ws, r_val, c0, c0 + 1, ref.formula,
                _font(11, italic=True, color=GREY) if code else _font(16, bold=True, color=NAVY),
                Alignment(horizontal="center", vertical="center"), 28)
    cell.number_format = ref.fmt
    src = ref.cell
    if src.fill is not None and src.fill.fill_type == "solid":          # the confidence fill, copied; never added
        fg = src.fill.fgColor.rgb
        for c in (c0, c0 + 1):
            ws.cell(r_val, c).fill = PatternFill("solid", fgColor=fg)
    cell.comment = Comment(f"Source: {ref.tab}, {ref.project} — {ref.header}. Hover the same cell on that tab "
                           "for its working.", "Highlights", width=300, height=80)
    for c in (c0, c0 + 1):
        _border(ws.cell(r_val, c), bottom=_TEAL)
    return cell


def _link(ws, r, c0, tab, row, c1=None) -> None:
    cell = _put(ws, r, c0, c1 or c0 + CARD_W - 1, f"Open {tab}", _font(10, color=TEAL, underline="single"),
                Alignment(horizontal="left", vertical="center"), 16)
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"{quote_sheet(tab)}!A{row}", display=f"Open {tab}")


# ========================================================================================================== layout
def _write_pair(ws, row: int, pair: list, sink: list) -> int:
    """Two cards (or one) side by side from `row`; returns the next free row. Both share row heights."""
    bands = max(len(_chunks(c.chips, CHIPS_PER_ROW)) for c in pair)
    r_head = row
    r_bands = [(row + 1 + 2 * b, row + 2 + 2 * b) for b in range(bands)]
    r_read = row + 1 + 2 * bands
    r_cav, r_link, r_gap = r_read + 1, r_read + 2, r_read + 3
    for card, c0 in zip(pair, (LEFT0, RIGHT0)):
        _put(ws, r_head, c0, c0 + CARD_W - 1, card.headline, _font(12, bold=True, color=NAVY),
             Alignment(wrap_text=True, vertical="center"), 20 if len(card.headline) <= 62 else 34)
        _border(ws.cell(r_head, c0), bottom=_TEAL)
        for c in range(c0, c0 + CARD_W):
            _border(ws.cell(r_head, c), bottom=_TEAL)
        for b, chunk in enumerate(_chunks(card.chips, CHIPS_PER_ROW)):
            for k, chip in enumerate(chunk):
                sink.append((_chip(ws, r_bands[b][0], r_bands[b][1], c0 + 2 * k, chip), f"{card.cid} {card.project}"))
        for r, label, text in ((r_read, "Reading", card.reading), (r_cav, "Caveat", card.caveat)):
            _put(ws, r, c0, c0 + CARD_W - 1, f"{label}: {text}", _font(10), None, 14 * _lines(f"{label}: {text}", 100) + 3)
        _link(ws, r_link, c0, card.link_tab, card.link_row)
        _outline(ws, r_head, c0, r_link, c0 + CARD_W - 1)
    ws.row_dimensions[r_gap].height = 10
    return r_gap + 1


def _write_table_card(ws, row: int, card: Card, sink: list) -> int:
    """C8, full width: lowest three on the left, highest three on the right, as name | value rows."""
    t = card.table
    last = LAST_COL
    _put(ws, row, LEFT0, last, card.headline, _font(12, bold=True, color=NAVY), Alignment(vertical="center"), 20)
    for c in range(LEFT0, last + 1):
        _border(ws.cell(row, c), bottom=_TEAL)
    halves = (("Lowest: " + t["header"], t["low"], LEFT0), ("Highest: " + t["header"], t["high"], RIGHT0))
    for title, refs, c0 in halves:
        _put(ws, row + 1, c0, c0 + CARD_W - 1, title, _font(8, color=GREY), Alignment(vertical="bottom"), 20)
        for i, ref in enumerate(refs):
            r = row + 2 + i
            n = _put(ws, r, c0, c0 + 4, f"={quote_sheet(ref.tab)}!A{ref.row}", _font(11, bold=True, color=NAVY),
                     Alignment(vertical="center"), 20)
            n.number_format = "General"
            sink.append((n, f"{card.cid} {ref.project} (name)"))
            v = _put(ws, r, c0 + 5, c0 + 6, ref.formula, _font(14, bold=True, color=NAVY),
                     Alignment(horizontal="center", vertical="center"), 22)
            v.number_format = ref.fmt
            sink.append((v, f"{card.cid} {ref.project}"))
            if ref.cell.fill is not None and ref.cell.fill.fill_type == "solid":
                for c in (c0 + 5, c0 + 6):
                    ws.cell(r, c).fill = PatternFill("solid", fgColor=ref.cell.fill.fgColor.rgb)
            v.comment = Comment(f"Source: {ref.tab}, {ref.project} — {ref.header}. Hover the same cell on that "
                                "tab for its working.", "Highlights", width=300, height=80)
            for c in range(c0, c0 + 7):
                _border(ws.cell(r, c), bottom=_THIN)
    r = row + 5
    _put(ws, r, LEFT0, last, f"Reading: {card.reading}", _font(10), None, 17)
    _link(ws, r + 1, LEFT0, card.link_tab, card.link_row, last)
    _outline(ws, row, LEFT0, r + 1, last)
    ws.row_dimensions[r + 2].height = 10
    return r + 3


def _find_about(wb, pred) -> int | None:
    ab = wb["About"]
    for r in range(1, ab.max_row + 1):
        v = ab.cell(r, 1).value
        if isinstance(v, str) and pred(v):
            return r
    return None


def _first_sentence_formula(about_row: int, text: str) -> str:
    m = re.search(r"[.!?](\s|$)", text)
    return f"=LEFT(About!A{about_row},{m.start() + 1})" if m and m.start() + 1 < len(text) else f"=About!A{about_row}"


def _top(ws, wb) -> int:
    """Title, as-of and the two lines copied from About, by formula."""
    _put(ws, 1, LEFT0, LAST_COL, TITLE, _font(18, bold=True, color=NAVY), Alignment(vertical="center"), 30)
    ab = wb["About"]
    r = _find_about(wb, lambda v: v.startswith("Data as of "))
    if r is None:
        raise HighlightsError("About has no 'Data as of' line to quote")
    _put(ws, 2, LEFT0, LAST_COL, _first_sentence_formula(r, ab.cell(r, 1).value), _font(10, color=GREY),
         Alignment(vertical="center"), 16)
    r = _find_about(wb, lambda v: v == "What this file measures")
    if r is None:
        raise HighlightsError("About has no 'What this file measures' section to quote")
    body = ab.cell(r + 1, 1).value or ""
    _put(ws, 3, LEFT0, LAST_COL, _first_sentence_formula(r + 1, body), _font(10), None,
         14 * _lines(body.split(". ")[0], 150) + 3)
    r = _find_about(wb, lambda v: v.startswith("Fill is confidence"))
    if r is None:
        raise HighlightsError("About has no 'Fill is confidence' line to quote")
    _put(ws, 4, LEFT0, LAST_COL, f"=About!A{r}", _font(10), None, 14 * _lines(ab.cell(r, 1).value, 150) + 3)
    ws.row_dimensions[5].height = 8
    return 6


def _first_sentence(text: str, cap: int = 220) -> str:
    m = re.search(r"[.!?](\s|$)", text)
    s = text[:m.start() + 1] if m else text
    return s if len(s) <= cap else s[:cap].rsplit(" ", 1)[0] + "…"


def reason_of(comment_text: str) -> str:
    """The lead of a Summary key-cell comment: its first paragraph (the paragraph after 'No measured key.' when that is
    all the first says), cut at the first sentence end."""
    paras = [p.strip() for p in (comment_text or "").split("\n\n") if p.strip() and not p.strip().startswith("Working:")]
    if not paras:
        return ""
    text = paras[0]
    if text.rstrip(".").lower() == "no measured key" and len(paras) > 1:
        text = paras[1]
    return _first_sentence(text.replace("\n", " "))


def unranked_rows(L: Lookup) -> list:
    """[(project, code Ref, reason)] for every Summary row whose rank is not a number: unranked or parked, in Summary
    order. The code is the key cell where it reads a code, else the rank cell; the reason is that cell's comment lead
    (a parked project's comment opens with its park reason)."""
    S = reader.SUMMARY
    ws = L.wb[S]
    hdr = L.headers(S)
    kc, rc = hdr["Archetype key"][0], hdr["Rank in archetype"][0]
    out = []
    for name, r in L.rows(S).items():
        rank = L.X.value(S, r, rc)
        if num(rank) is not None:
            continue
        key = L.X.value(S, r, kc)
        col = kc if isinstance(key, str) else rc
        cell = ws.cell(r, col)
        ref = Ref(S, r, col, name, ws.cell(4, col).value, L.X.value(S, r, col), cell.number_format, cell)
        why = reason_of(cell.comment.text if cell.comment else "")
        if FORBIDDEN.search(why or ""):
            why = "No measured key; the Summary key cell's comment gives the reason."
        out.append((name, ref, why or "No measured key."))
    return out


def _write_coverage(ws, wb, L: Lookup, row: int) -> int:
    ab = wb["About"]
    r0 = _find_about(wb, lambda v: v.startswith("Honest coverage"))
    if r0 is None:
        raise HighlightsError("About has no 'Honest coverage' block to copy")
    _put(ws, row, LEFT0, LAST_COL, "Coverage", _font(12, bold=True, color=NAVY), Alignment(vertical="center"), 22)
    for c in range(LEFT0, LAST_COL + 1):
        _border(ws.cell(row, c), bottom=_TEAL)
    row += 1
    _put(ws, row, LEFT0, LAST_COL, f"=About!A{r0}", _font(9, color=GREY), Alignment(vertical="center"), 16)
    row += 1
    ncols = sum(1 for c in range(1, ab.max_column + 1) if ab.cell(r0 + 1, c).value is not None)
    phys = [c for c in range(LEFT0 + 2, LAST_COL + 1) if c != GUTTER]          # the numeric columns skip the gutter
    r = r0 + 1
    while r <= ab.max_row and ab.cell(r, 1).value is not None:
        head = r == r0 + 1
        last = str(ab.cell(r, 1).value).startswith("All ")
        bold = head or last
        for j in range(1, ncols + 1):
            src = ab.cell(r, j)
            if src.value is None:
                continue
            if j == 1:
                cell = _put(ws, row, LEFT0, LEFT0 + 1, f"=About!A{r}", _font(9 if head else 10, bold, NAVY if head else "000000"),
                            Alignment(wrap_text=True, vertical="center" if head else "top"))
            else:
                c = phys[j - 2]
                cell = ws.cell(row, c, f"=About!{get_column_letter(j)}{r}")
                cell.font = _font(9 if head else 10, bold, NAVY if head else "000000")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.number_format = src.number_format
        for c in [LEFT0, LEFT0 + 1] + phys[:ncols - 1]:
            _border(ws.cell(row, c), bottom=_NAVY if head else _THIN)
        ws.row_dimensions[row].height = 34 if head else 16
        row += 1
        r += 1
    row += 1
    _put(ws, row, LEFT0, LAST_COL, "Projects without a rank, and parked projects, with the code in the key cell and the "
         "reason its comment gives", _font(10, bold=True, color=NAVY), Alignment(vertical="center"), 18)
    row += 1
    for c0, c1, text in ((LEFT0, LEFT0 + 2, "Project"), (LEFT0 + 3, LEFT0 + 4, "Code"), (LEFT0 + 5, LAST_COL, "Reason")):
        _put(ws, row, c0, c1, text, _font(9, bold=True, color=NAVY), Alignment(vertical="center"), 16)
    for c in range(LEFT0, LAST_COL + 1):
        _border(ws.cell(row, c), bottom=_NAVY)
    row += 1
    for name, ref, why in unranked_rows(L):
        n = _put(ws, row, LEFT0, LEFT0 + 2, f"=Summary!A{ref.row}", _font(10), Alignment(vertical="top"))
        n.number_format = "General"
        code = _put(ws, row, LEFT0 + 3, LEFT0 + 4, ref.formula, _font(9, italic=True, color=GREY),
                    Alignment(horizontal="center", vertical="top"))
        code.number_format = ref.fmt
        _put(ws, row, LEFT0 + 5, LAST_COL, why, _font(9), None, 13 * _lines(why, 125) + 3)
        for c in range(LEFT0, LAST_COL + 1):
            _border(ws.cell(row, c), bottom=_THIN)
        row += 1
    row += 1
    _put(ws, row, LEFT0, LAST_COL, FOOTER, _font(10, italic=True, color=GREY), Alignment(vertical="center"), 18)
    return row + 1


def _page(ws) -> None:
    ws.sheet_properties.tabColor = TEAL
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    for c in range(1, LAST_COL + 1):
        ws.column_dimensions[get_column_letter(c)].width = 3 if c == GUTTER else 9.5


# ===================================================================================================== verification
def number_cells(ws) -> list:
    """Every cell of the tab that quotes a number: a formula whose target is another tab's cell (not About's text)."""
    out = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=") and "About!" not in c.value:
                out.append(c)
    return out


def verify(wb, chips: list) -> list:
    """Evaluate every chip / table value formula; returns [(cell, value, shown)]. Raises HighlightsError on an Excel
    error, a blank or None."""
    X = xlcalc.Workbook(wb)
    ws = wb[SHEET]
    bad, out = [], []
    for cell, label in chips:
        v = X.value(SHEET, cell.row, cell.column)
        if isinstance(v, xlcalc.XLError) or v is None or (isinstance(v, str) and not v.strip()):
            bad.append(f"{label} ({SHEET}!{cell.coordinate} = {cell.value}) evaluates to {xlcalc.display(v)!r}")
        out.append((cell, v, format_value(v, cell.number_format)))
    for c in number_cells(ws):                                           # every other formula: not an error
        v = X.value(SHEET, c.row, c.column)
        if isinstance(v, xlcalc.XLError):
            bad.append(f"{SHEET}!{c.coordinate} = {c.value} evaluates to {v.code}")
    if bad:
        raise HighlightsError("Highlights: " + str(len(bad)) + " formula(s) are not usable:\n  " + "\n  ".join(bad[:25]))
    return out


# ==================================================================================================== the entry point
def add_highlights(wb, ctx=None, min_cards: int = MIN_CARDS, parked=None) -> dict:
    """Add the Highlights sheet, first, to the reader workbook `wb`. `ctx` is reader.Ctx (not needed: everything is read
    from the finished tabs). Returns {"cards": [...], "omitted": [...], "lines": [printable lines]}; raises
    HighlightsError when fewer than `min_cards` render or a chip is an error or blank."""
    L = Lookup(wb, parked)
    cards, omitted = build_cards(L)
    if len(cards) < min_cards:
        raise HighlightsError(f"Highlights: only {len(cards)} of {len(CARD_BUILDERS)} cards render, {min_cards} are "
                              "needed. Cards that did not render:\n  " + "\n  ".join(omitted)
                              + "\nCards that did: " + (", ".join(c.cid for c in cards) or "none"))
    if SHEET in wb.sheetnames:
        raise HighlightsError("the workbook already has a Highlights sheet")
    ws = wb.create_sheet(SHEET, 0)
    _page(ws)
    for w in wb.worksheets:
        w.sheet_view.tabSelected = False
    ws.sheet_view.tabSelected = True
    wb.active = 0
    row = _top(ws, wb)
    chip_cells: list = []
    plain = [c for c in cards if c.table is None]
    for i in range(0, len(plain), 2):
        row = _write_pair(ws, row, plain[i:i + 2], chip_cells)
    for card in (c for c in cards if c.table is not None):
        row = _write_table_card(ws, row, card, chip_cells)
    _write_coverage(ws, wb, L, row)
    shown = {c.coordinate: (v, s) for c, v, s in verify(wb, chip_cells)}
    lines = report_lines(cards, omitted, ws, shown)
    return {"cards": [c.cid for c in cards], "omitted": omitted, "lines": lines, "shown": shown, "sheet": ws}


def report_lines(cards: list, omitted: list, ws, shown: dict) -> list:
    """The rendered cards as text: headline, each chip 'label: value as displayed', Reading, Caveat; then the omissions."""
    out = ["", "HIGHLIGHTS (the Charlie edition's first tab): %d card(s) rendered" % len(cards)]

    def disp(ref: Ref) -> str:
        return format_value(ref.value, ref.fmt)

    for c in cards:
        out.append("")
        out.append(f"  [{c.cid}] {c.headline}")
        for chip in c.chips:
            out.append(f"      {chip.label}: {disp(chip.ref)}")
        if c.table:
            for title, refs in (("lowest", c.table["low"]), ("highest", c.table["high"])):
                out.append(f"      {title} {c.table['header']}: " + "; ".join(
                    f"{ws.parent[r.tab].cell(r.row, 1).value} {disp(r)}" for r in refs))
        out.append(f"      Reading: {c.reading}")
        if c.caveat:
            out.append(f"      Caveat: {c.caveat}")
        out.append(f"      Open {c.link_tab}")
    for o in omitted:
        out.append(f"  omitted - {o}")
    return out


def all_strings(ws) -> list:
    """Every string on the tab: cell text, formulas' literal text, comments - for the forbidden-word scan."""
    out = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str):
                out.append(c.value)
            if c.comment is not None:
                out.append(c.comment.text)
            if c.hyperlink is not None:
                out.append(c.hyperlink.display or "")
    return out


__all__ = ["add_highlights", "HighlightsError", "format_value", "FORBIDDEN", "FOOTER", "SHEET", "openpyxl"]
