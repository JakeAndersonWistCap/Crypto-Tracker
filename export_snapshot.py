#!/usr/bin/env python3
"""
export_snapshot.py — one row per project, every headline metric, for the scoring theory (Jake, 2026-10-11, D2).

    python export_snapshot.py                  build today's workbook from metrics.db (a temp file), evaluate it and
                                               write snapshots/snapshot_<date>.csv
    python export_snapshot.py --workbook X     snapshot an existing workbook (its Credibility tab must be evaluable)
    python export_snapshot.py --out DIR        write into DIR instead of snapshots/

THE COLUMNS, per project row: project, asof; then for every headline figure, four columns —
    <tab> | <header>                      the value the tab shows (evaluated; text such as "BLOCKED — …" kept as is)
    <tab> | <header> | verdict            its Credibility verdict (PASS / CHECK / FRESH / … as the tab says)
    <tab> | <header> | near_limit         1 when the verdict says NEAR LIMIT (a pass within 80-100% of tolerance)
    <tab> | <header> | coverage_days      the Q0 days its series covered (Data: q0_covered_days), blank for a stock
The headlines are Master's columns (each mapped to its A-tab cell by cross_tab.MASTER_SOURCES, whose verdict it
carries) and every BOLD headline cell of A1-A4 (the cells the Credibility tab judges, build_workbook._HEADLINE_CELLS).
Then, per Trajectory-tab row (theory layer, Phase 1): traj | <tab> | <headline id> | Q0..Q-4, T12M, QoQ, 4Q growth,
quarters improved, trough, years to benchmark and the row's credibility.
Nothing is computed here that the workbook does not show: the snapshot is the workbook, flattened. windows/
daily_run.ps1 runs it as the LAST step. METRICS.md (generate_metrics_doc.py) defines every column.
"""
from __future__ import annotations

import argparse
import csv
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent


def _grid(ws):
    """(header row, {col: header}, {project: row}) for the first table of a Master / A-tab sheet."""
    hr = next((r for r in range(1, 12) if ws.cell(r, 1).value == "Project"), None)
    if hr is None:
        return None, {}, {}
    heads = {c: str(ws.cell(hr, c).value) for c in range(1, ws.max_column + 1) if ws.cell(hr, c).value}
    rows = {}
    for r in range(hr + 1, ws.max_row + 1):
        p = ws.cell(r, 1).value
        if not p:
            break
        rows[str(p)] = r
    return hr, heads, rows


def _coverage(wb, value) -> dict:
    """{(project, metric): q0 covered days} from the Data tab."""
    ws = wb["Data"]
    cols = {str(ws.cell(1, c).value or ""): c for c in range(1, ws.max_column + 1)}
    cov = next((c for h, c in cols.items() if h.startswith("Q0: days of the 90")), None)
    out = {}
    if cov is None:
        return out
    for r in range(2, ws.max_row + 1):
        k = str(ws.cell(r, 1).value or "")
        if "|" in k:
            v = value("Data", r, cov)
            out[tuple(k.split("|", 1))] = v if isinstance(v, (int, float)) else None
    return out


def snapshot_rows(wb, value, cred_rows: list, cred_tab: list, headline_cells: list, asof) -> list[dict]:
    """The rows (dicts, ordered columns) for every project on Master."""
    import cross_tab
    import config
    verdict = {}                                   # (project, sheet title, cell) -> verdict
    for r, t in zip(cred_rows, cred_tab):
        if r.get("tab") and r.get("cell"):
            verdict[(r["project"], r["tab"], r["cell"])] = t[10] if len(t) > 10 else None
    cov = _coverage(wb, value)
    _hr, mheads, mrows = _grid(wb["Master"])
    cells = {}
    for ws in wb.worksheets:
        pre = ws.title[:2]
        if pre in cross_tab.A_TABS:
            _h, heads, rows = _grid(ws)
            for p, r in rows.items():
                for c, h in heads.items():
                    cells[(pre, p, h)] = (ws.title, ws.cell(r, c).coordinate)
    by_proj: dict = {}
    for hc in headline_cells:
        by_proj.setdefault(hc["project"], []).append(hc)
    out = []
    for p, mr in mrows.items():
        row = {"project": p, "asof": str(asof)[:10]}
        arch = (config.PROJECT_BY_NAME.get(p) or {}).get("archetypes") or []

        def put(label, val, sheet, cell, metric):
            v = verdict.get((p, sheet, cell)) if sheet else None
            row[label] = val
            row[f"{label} | verdict"] = v if v is not None else ""
            row[f"{label} | near_limit"] = 1 if v and "NEAR LIMIT" in str(v).upper() else 0
            c = cov.get((p, metric)) if metric else None
            row[f"{label} | coverage_days"] = c if c is not None else ""
        for c, h in mheads.items():
            if h in ("Project", "Symbol"):
                continue
            src = cross_tab.resolve(cross_tab.sources_for(h, p), arch, cells, p)
            hc = next((x for x in by_proj.get(p, []) if src and x["sheet"] == src[0] and x["cell"] == src[1]), None)
            put(f"Master | {h}", value("Master", mr, c), src[0] if src else None, src[1] if src else None,
                (hc or {}).get("metric"))
        for hc in by_proj.get(p, []):
            ws = wb[hc["sheet"]]
            cell = ws[hc["cell"]]
            put(f"{hc['sheet'][:2]} | {hc['header']}", value(hc["sheet"], cell.row, cell.column), hc["sheet"],
                hc["cell"], hc.get("metric"))
        # SUPPORTING FIGURES THE THEORY NEEDS that are not Credibility headlines yet (unlock dilution, Jake 2026-10-11
        # A2): the first A-tab holding each, by header
        import build_workbook as bw
        for h in bw.EXTRA_SNAPSHOT_HEADERS:
            for tab in cross_tab.A_TABS:
                hit = cells.get((tab, p, h))
                if hit:
                    ws = wb[hit[0]]
                    cell = ws[hit[1]]
                    put(f"{tab} | {h}", value(hit[0], cell.row, cell.column), None, None,
                        "unlock_release_tokens")
                    break
        # TRAJECTORIES (theory layer, Phase 1): every Trajectory-tab row of the project — Q0..Q-4, T12M, QoQ, 4Q growth,
        # quarters improved, trough, years to benchmark — keyed by tab and headline id
        for col, val in _trajectory(wb, value).get(p, {}).items():
            row[col] = val
        out.append(row)
    return out


TRAJ_FIELDS = ("Q0", "Q-1", "Q-2", "Q-3", "Q-4", "T12M", "QoQ", "4Q growth (annual, log-linear)",
               "Quarters improved (direction-aware)", "Trough of the last 4", "Years to benchmark",
               "Credibility (Q0 = headline cell)")


def _trajectory(wb, value) -> dict:
    """{project: {"traj | <tab> | <headline id> | <field>": value}} from the Trajectory tab (evaluated)."""
    if "Trajectory" not in wb.sheetnames:
        return {}
    import trajectory                                     # noqa: PLC0415
    ws = wb["Trajectory"]
    hr = next((r for r in range(1, 12) if ws.cell(r, 1).value == "Project"), None)
    out: dict = {}
    for r in range(hr + 1, ws.max_row + 1) if hr else ():
        p = ws.cell(r, 1).value
        if not p:
            continue
        tab, hid = ws.cell(r, trajectory.COL["Tab"]).value, ws.cell(r, trajectory.COL["Headline id"]).value
        for f in TRAJ_FIELDS:
            v = value("Trajectory", r, trajectory.COL[f])
            out.setdefault(str(p), {})[f"traj | {tab} | {hid} | {f}"] = "" if v is None else v
    return out


def write_csv(rows: list[dict], path: Path) -> Path:
    cols: list = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in cols})
    return path


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workbook")
    ap.add_argument("--out", default=str(ROOT / "snapshots"))
    ap.add_argument("--asof")
    a = ap.parse_args(argv)
    import openpyxl
    import pandas as pd
    import xlcalc
    import build_workbook as bw
    from credibility_report import read_tab
    asof = pd.Timestamp(a.asof) if a.asof else pd.Timestamp.now("UTC").tz_localize(None).normalize()
    with tempfile.TemporaryDirectory() as tmp:
        path = a.workbook
        if not path:
            import store as store_mod
            if not Path(store_mod.DB_PATH).exists():
                print(f"no {store_mod.DB_PATH} — run token_metrics.py first")
                return 1
            path = str(Path(tmp) / "snapshot.xlsx")
            st = store_mod.Store(store_mod.DB_PATH)
            try:
                bw.build_workbook(st, path, asof=asof)
            finally:
                st.close()
        _summary, tab = read_tab(path, evaluate=True)
        wb = openpyxl.load_workbook(path)
        X = xlcalc.Workbook(wb)
        rows = snapshot_rows(wb, lambda s, r, c: xlcalc.display(X.value(s, r, c)), list(bw.CREDIBILITY_ROWS), tab,
                             list(bw._HEADLINE_CELLS), asof)
    out = write_csv(rows, Path(a.out) / f"snapshot_{asof.date()}.csv")
    print(f"snapshot: {len(rows)} project(s), {max((len(r) for r in rows), default=0)} column(s) -> {out}")
    return 0


if __name__ == "__main__":
    import progress                                    # stdout line-buffered when piped (2026-10-11)
    progress.line_buffered()
    sys.exit(main())
