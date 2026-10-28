"""ROUND P2e-2 (Jake 2026-10-28): THE TRAJECTORY BLOCK TRACKS THE KEY.

Block (c) of each reader archetype tab followed a series other than the rank key (A2 customer revenue $ per token
emitted while the key is customer revenue $ ÷ emissions $; A3 retired ÷ (emissions + unlocks) while the key is retired
÷ emissions; A4 net supply change while the key is burn ÷ issuance; A1 fees ÷ issuance while the key is settlement per
token, 1-year change). Now its first columns are THE KEY ITSELF for the latest quarter and 1, 2, 3 quarters back, plus
the change vs the previous quarter — the key's own formula (A2/A4: the key cell's existing Trajectory row; A3: its new
row a3_retire_vs_emissions; A1: reader.settlement_key with its end moved back). A quarter the key cannot form reads
the key's own code and reason. The other series follow, each header saying what it is. Built once on the committed
real dumps (fixtures/real_*.json). No network.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import reader  # noqa: E402

TABS = ("A1 Infrastructure", "A2 Coordination", "A3 Revenue Buyback", "A4 Permanent Burn")
KEY_STEM = {"A1 Infrastructure": "Settlement/token 1-yr chg", "A2 Coordination": "Revenue ÷ emissions (x)",
            "A3 Revenue Buyback": "Retired ÷ emissions (x)", "A4 Permanent Burn": "Burn ÷ issuance (x)"}
OTHER_STEM = {"A1 Infrastructure": "Fees ÷ issuance", "A2 Coordination": "$ per token emitted",
              "A3 Revenue Buyback": "Retired ÷ (emis + unlocks)", "A4 Permanent Burn": "Net supply change (tokens)"}
KEY_ID = {"A2 Coordination": "a2_rev_vs_emissions_value", "A3 Revenue Buyback": "a3_retire_vs_emissions",
          "A4 Permanent Burn": "a4_crossover"}


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    import build_workbook as bw
    import measuring_guard as mg
    import store as sm
    from real_dump import to_store
    d = tmp_path_factory.mktemp("p2e2")
    db = d / "metrics.db"
    asofs = []
    for pl in mg._payloads():
        to_store(pl, db)
        asofs.append(pl["meta"]["asof"])
    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(sm, "DB_PATH", str(db))
        mp.setitem(sys.modules, "recalc", None)
        mp.setenv("TOKEN_METRICS_STORE_VALUES", "1")
        st = sm.Store(str(db))
        try:
            bw.build_workbook(st, d / "audit.xlsx", asof=pd.Timestamp(max(asofs)), reader_path=d / "reader.xlsx")
        finally:
            st.close()
    assert "error" not in bw.READER_STATS, bw.READER_STATS
    return {"reader": openpyxl.load_workbook(d / "reader.xlsx"), "audit": d / "audit.xlsx",
            "traj": list(bw.TRAJECTORY_ROWS)}


def _value(wb, cell):
    """A reader cell's number (its `_data` reference resolved), or its code."""
    v = cell.value
    m = re.fullmatch(r"=_data!\$?([A-Z]+)\$?(\d+)", str(v)) if isinstance(v, str) else None
    return wb["_data"][f"{m.group(1)}{m.group(2)}"].value if m else v


def _heads(ws):
    return {ws.cell(4, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(4, c).value}


def _rows(ws):
    return {ws.cell(r, 1).value: r for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value is not None}


def _block_c(ws):
    """Block (c)'s headers in order (the block label is on row 3 of its first column only)."""
    out, cur = [], None
    for c in range(1, ws.max_column + 1):
        cur = ws.cell(3, c).value or cur
        if cur == "(c) Trajectory":
            out.append(ws.cell(4, c).value)
    return out


@pytest.mark.parametrize("tab", TABS)
def test_block_c_is_the_key_first_then_the_other_series_each_named(built, tab):
    ws = built["reader"][tab]
    labels = [f", {x}" for x in reader.TRAJ_LABELS]
    want = ([KEY_STEM[tab] + x for x in labels] + [reader.KEY_CHANGE_HEADER]
            + [OTHER_STEM[tab] + x for x in labels])
    got = _block_c(ws)
    assert got[:len(want)] == want, got
    assert len(got) == len(want) + 1 and "change" in got[-1], got
    assert all(len(h) <= reader.MAX_HEADER for h in got), [h for h in got if len(h) > reader.MAX_HEADER]
    assert len(set(_heads(ws))) == len(_heads(ws))
    defs = {(row[0].value, row[2].value) for row in built["reader"]["Definitions"].iter_rows(min_row=4) if row[0].value}
    assert all((tab, h) in defs for h in got), [h for h in got if (tab, h) not in defs]


@pytest.mark.parametrize("tab", TABS)
def test_the_latest_quarter_is_the_key_cell(built, tab):
    """THE TEST (P2e-2): every project with a numeric key — its latest-quarter key trajectory value equals the key cell;
    a key that is a code reads the same code in the latest quarter."""
    wb = built["reader"]
    ws = wb[tab]
    h = _heads(ws)
    kc, qc = 3, h[f"{KEY_STEM[tab]}, {reader.TRAJ_LABELS[0]}"]
    numeric = 0
    for name, r in _rows(ws).items():
        k, q = _value(wb, ws.cell(r, kc)), _value(wb, ws.cell(r, qc))
        if isinstance(k, (int, float)):
            numeric += 1
            assert isinstance(q, (int, float)), (tab, name, k, q)
            assert q == pytest.approx(k, rel=1e-9, abs=1e-12), (tab, name, k, q)
        else:
            assert q == k and q in reader.CODES, (tab, name, k, q)
    assert numeric or tab == "A1 Infrastructure", tab       # A1's key needs two years of settlement history


def test_each_key_has_its_own_trajectory_row_and_q0_is_the_key_cell(built):
    """A2 and A4 read the key cell's existing Trajectory row (its Q0 PASSes against the cell); A3's key has its own row
    now (a3_retire_vs_emissions), with a trend cell beside the key on the audit A3 tab."""
    rows = {(t["tab"], t["project"], t["hid"]): t for t in built["traj"]}
    for tab, hid in KEY_ID.items():
        mine = [t for (tb, _p, hd), t in rows.items() if tb == tab[:2] and hd == hid]
        assert mine, (tab, hid)
        assert all(t["verdict"].startswith("PASS") for t in mine), [(t["key"], t["verdict"]) for t in mine
                                                                   if not t["verdict"].startswith("PASS")]
        assert all(str(t["header"]).startswith(("Customer revenue ÷ emissions value", "RETIRED ÷ EMISSIONS",
                                                "BURN ÷ ISSUANCE")) for t in mine)
    a = openpyxl.load_workbook(built["audit"], read_only=True)
    ws = a["A3 Revenue Buyback"]
    hdr = next(r for r in ws.iter_rows(min_row=1, max_row=12, values_only=True) if r and r[0] == "Project")
    i = next(j for j, v in enumerate(hdr) if str(v or "").startswith("RETIRED ÷ EMISSIONS"))
    assert str(hdr[i + 1]).startswith("Trend — a3_retire_vs_emissions ")


def test_quarters_the_key_cannot_form_read_its_code_and_reason_never_zero(built):
    wb = built["reader"]
    traj = {t["key"]: t for t in built["traj"]}
    for tab in TABS:
        ws = wb[tab]
        h = _heads(ws)
        for name, r in _rows(ws).items():
            for i, lab in enumerate(reader.TRAJ_LABELS):
                cell = ws.cell(r, h[f"{KEY_STEM[tab]}, {lab}"])
                v = _value(wb, cell)
                if isinstance(v, (int, float)):
                    t = traj.get(f"{tab[:2]}|{name}|{KEY_ID[tab]}") if tab in KEY_ID else None
                    if t is not None:      # a number only where the key's own formula gives one, never 0 for unknown
                        assert t[("Q0", "Q-1", "Q-2", "Q-3")[i]] == pytest.approx(v), (tab, name, lab)
                    continue
                assert v in reader.CODES, (tab, name, lab, v)
                assert cell.comment is not None and len(cell.comment.text) > 40, (tab, name, lab)
    # the fixture's own cases: Near's A3 key two quarters back is BLOCKED by unequal coverage (rule 9), said so
    ws = wb["A3 Revenue Buyback"]
    c = ws.cell(_rows(ws)["Near"], _heads(ws)[f"{KEY_STEM['A3 Revenue Buyback']}, {reader.TRAJ_LABELS[2]}"])
    assert c.value == "blocked" and "unequal coverage" in c.comment.text
    # Plume's A1 key (settlement held from 2026-05-01): the key's own reason in every quarter — two years of history
    # wanting while the quarter's end is held, the quarter's end itself not held before that
    ws = wb["A1 Infrastructure"]
    for i, lab in enumerate(reader.TRAJ_LABELS):
        c = ws.cell(_rows(ws)["Plume"], _heads(ws)[f"{KEY_STEM['A1 Infrastructure']}, {lab}"])
        want = "two years of history" if i < 2 else "the quarter's end"
        assert c.value == "n/a" and want in c.comment.text and "held from 2026-05-01" in c.comment.text, \
            (lab, c.comment.text)


# ---- A1: the key's own function, its end moved back (no build) ----------------------------------------------------
def _groups(days: int, end: str = "2026-10-12", s0: float = 1e9, growth: float = 0.001, circ: float = 1e6):
    idx = pd.date_range(end=pd.Timestamp(end), periods=days)
    s = pd.DataFrame({"date": idx, "value": [s0 * (1 + growth) ** i for i in range(days)], "source": "artemis:x"})
    c = pd.DataFrame({"date": idx, "value": [circ] * days, "source": "chain:x"})
    return {("X", "settlement_volume_365d_usd"): s, ("X", "circulating_supply_used"): c}


def test_a1_key_quarters_are_the_key_at_each_quarter_end():
    g = _groups(365 + 3 * 90 + 10)
    t = reader.settlement_key_trajectory(g, "X", pd.Timestamp("2026-10-13"), 90)
    assert t["vals"][0] == reader.settlement_key(g, "X")[0]            # Q0 IS the key cell
    end = pd.Timestamp("2026-10-12")
    for k in range(4):
        e, b = end - pd.Timedelta(days=90 * k), end - pd.Timedelta(days=90 * k + 365)
        s = g[("X", "settlement_volume_365d_usd")].set_index("date")["value"]
        assert t["vals"][k] == pytest.approx(s[e] / s[b] - 1), k
        assert t["why"][k] == f"{b.date()} to {e.date()}"
    kind, v, _why = reader._quarter_change(t, pd.Timestamp("2026-10-13"))
    assert kind == "ok" and v == pytest.approx((t["vals"][0] - t["vals"][1]) / abs(t["vals"][1]))


def test_a1_a_quarter_history_does_not_reach_reads_n_a_with_the_reason():
    g = _groups(365 + 100)                                              # Q0 and Q-1 form; Q-2, Q-3 do not
    t = reader.settlement_key_trajectory(g, "X", pd.Timestamp("2026-10-13"), 90)
    assert t["vals"][0] is not None and t["vals"][1] is not None
    assert t["vals"][2] is None and t["vals"][3] is None
    assert "two years of history" in t["why"][2]
    # never a 0 for unknown, and no change where a quarter is missing
    t2 = dict(t, vals=[t["vals"][0], None, None, None])
    assert reader._quarter_change(t2, pd.Timestamp("2026-10-13"))[0] == "n/a"


def test_a1_a_change_of_measuring_point_inside_the_span_blocks_the_change():
    g = _groups(365 + 3 * 90 + 10)
    s = g[("X", "settlement_volume_365d_usd")]
    s.loc[s["date"] >= pd.Timestamp("2026-06-01"), "source"] = "artemis:y"   # a new measuring point, undeclared
    t = reader.settlement_key_trajectory(g, "X", pd.Timestamp("2026-10-13"), 90)
    assert "2026-06-01" in t["break"]
    assert reader._quarter_change(t, pd.Timestamp("2026-10-13"))[0] == "blocked"
