"""
tests/test_trajectory.py — the theory layer's trajectory engine (Phase 1, Jake 2026-10-12): every headline on Master and
A1-A4 over five consecutive trailing 90-day blocks by its OWN formula (trajectory.py, build_workbook.shifted), the
derived statistics, the years to benchmark, the trend cells, the snapshot columns and logs/history_gaps.csv. Real dumps.
"""
from __future__ import annotations

import math
import sys
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

DUMPS = ("Aerodrome", "Chainlink", "Ethereum", "Maple", "Pendle", "Plume")


def _built(tmp_path, monkeypatch, project):
    import openpyxl
    import store as sm
    import build_workbook as bw
    from real_dump import load, to_store
    payload = load(project)
    monkeypatch.setitem(sys.modules, "recalc", None)
    monkeypatch.setattr(bw, "_SCOPE", bw._SCOPE)
    db = tmp_path / "metrics.db"
    to_store(payload, db)
    monkeypatch.setattr(sm, "DB_PATH", str(db))
    st = sm.Store(str(db))
    try:
        bw.build_workbook(st, tmp_path / "wb.xlsx", asof=pd.Timestamp(payload["meta"]["asof"]), only=[project])
    finally:
        st.close()
    return openpyxl.load_workbook(tmp_path / "wb.xlsx"), list(bw.TRAJECTORY_ROWS)


@pytest.mark.parametrize("project", DUMPS)
def test_q0_is_the_headline_cell_and_every_headline_has_a_trend_beside_it(tmp_path, monkeypatch, project):
    """Q0 is the headline's own formula unshifted: every row PASSes against its cell; every bold A-tab headline and
    every listed Master column has a trend cell immediately to its right."""
    import build_workbook as bw
    wb, rows = _built(tmp_path, monkeypatch, project)
    assert rows, "no trajectory rows"
    bad = [(r["key"], r["verdict"]) for r in rows if not r["verdict"].startswith("PASS")]
    assert not bad, bad
    for hc in bw._HEADLINE_CELLS:
        ws = wb[hc["sheet"]]
        hr = next(r for r in range(1, 12) if ws.cell(r, 1).value == "Project")
        col = ws[hc["cell"]].column
        assert str(ws.cell(hr, col + 1).value).startswith("Trend — "), (hc["sheet"], hc["header"])
    assert "Trajectory" in wb.sheetnames


def test_shifted_quarters_are_the_data_tabs_own_earlier_windows(tmp_path, monkeypatch):
    """Master's Revenue Q0 is a plain pull of Data q0: its Q-1..Q-4 must be Data's q1..q4 for revenue_usd, exactly."""
    import xlcalc
    import build_workbook as bw
    wb, rows = _built(tmp_path, monkeypatch, "Chainlink")
    X = xlcalc.Workbook(wb)
    ws = wb["Data"]
    r = next(i for i in range(2, ws.max_row + 1) if ws.cell(i, 1).value == "Chainlink|revenue_usd")
    data = [X.value("Data", r, bw.DATA_COLS.index(q) + 1) for q in bw.QUARTERS]
    row = next(x for x in rows if x["key"] == "Master|Chainlink|master_revenue_q0")
    assert [row[q] for q in ("Q0", "Q-1", "Q-2", "Q-3", "Q-4")] == pytest.approx(data)
    assert row["T12M"] == pytest.approx(sum(data[:4])) and row["T12M basis"] == "sum"


def test_the_shift_maps_each_window_reference():
    import build_workbook as bw
    assert bw._shift_col("revenue_usd", "q0", 2) == "q2"
    assert bw._shift_col("revenue_usd", "q1_covered_days", 3) == "q4_covered_days"
    assert bw._shift_col("revenue_usd", "q0_events", 1) == "q1_events"
    assert bw._shift_col("revenue_usd", "q2", 3) is None                       # past Q-4
    assert bw._shift_col("price_usd", "now", 2) == "q2"                        # a stock's now -> the quarter's mean
    assert bw._shift_col("revenue_usd", "now", 2) is None                      # a flow's trailing 30 days has no shift
    assert bw._shift_col("revenue_usd", "y1", 1) is None
    assert bw._shift_col("revenue_usd", "latest_date", 1) == "latest_date"
    R = bw.Refs(10, 1, [])
    with bw.shifted(2):
        assert "NA()" == R.D(5, "revenue_usd", "now") and "$" in R.D(5, "revenue_usd", "q0")
    q1_share = R.C(5, bw.WINDOW_SHARE_COL["q1"])
    with bw.shifted(1):
        assert R.C(5, bw.WINDOW_SHARE_COL["q0"]) == q1_share          # the window's split moves with it
    with bw.shifted(4):
        assert R.C(5, bw.WINDOW_SHARE_COL["q0"]) == "NA()"            # no Q4 split column
    assert bw._SHIFT["k"] == 0


def test_the_statistics():
    import trajectory as t
    g, n, why = t.loglinear_growth([16.0, 8.0, 4.0, 2.0, 1.0])               # doubling each quarter
    assert n == 4 and why == "" and g == pytest.approx(2 ** 4 - 1)
    assert t.loglinear_growth([1.0, None, None, None])[0] is None
    assert "below 0" in t.loglinear_growth([1.0, -1.0, 2.0, 3.0])[2]
    assert t.improved([5, 4, 3, 2, 1], 1) == (4, 4) and t.improved([5, 4, 3, 2, 1], -1) == (0, 4)
    assert t.improved([5, 4], 0) == (None, 0)
    assert t.trough([5, 1, 3, 4, 0], 1) == "Q-1" and t.trough([5, 1, 3, 4, 0], -1) == "Q0"
    assert t.t12m([1, 2, 3, 4, 5], "sum") == (10, "sum") and t.t12m([1, 2, 3, 4], "mean") == (2.5, "mean")
    assert t.t12m([1, None, 3, 4], "sum")[0] is None
    assert t.qoq(110, 100) == pytest.approx(0.10) and t.qoq(-50, -100) == pytest.approx(0.5)
    y, why = t.years_to(1.0, 0.5, 1.0, 1)                                       # 0.5 doubling a year -> 1 year
    assert y == pytest.approx(1.0) and why == ""
    assert t.years_to(1.0, 2.0, -0.5, 1) == (0.0, "met")
    assert "points away" in t.years_to(1.0, 0.5, -0.1, 1)[1]
    assert t.years_to(0.71, 2.0, -0.5, -1)[0] == pytest.approx(math.log(0.71 / 2.0) / math.log(0.5))
    assert t.years_to(None, 1, 1, 1, "no tbill_3m_yield stored yet")[1] == "no tbill_3m_yield stored yet"


def test_benchmarks_and_parameters_live_in_config():
    import config
    import trajectory as t
    ids = {b["id"] for b in config.TRAJECTORY["benchmarks"]}
    # NRR is valuation HEADROOM, not a target (Jake's Phase 1 review 2026-10-12, item 7): no years-to-NRR benchmark;
    # A1 shows headroom = the Config global nrr_m2_gdp (0.71) ÷ NRR
    assert ids == {"retire_vs_dilution", "burn_vs_issuance", "subsidy_cover", "yield_vs_tbill"}
    assert t.benchmark_for("a1_nrr_settlement") is None and config.GLOBALS["nrr_m2_gdp"] == 0.71
    assert config.TRAJECTORY["rows"]["a1_nrr_settlement_headroom"] == ("Valuation", 1)
    # the T-bill benchmark sits on the stakers' TOTAL YIELD (Jake 2026-10-12, decision 9 on 6f8c86c)
    assert t.benchmark_for("a1_validator_yield") is None
    assert t.benchmark_target(t.benchmark_for("a1_total_yield_stakers"), {})[0] is None   # until Phase 3 stores it
    assert config.TRAJECTORY["history_days"] >= 5 * config.TRAJECTORY["period_days"]


def test_snapshot_carries_the_trajectory_columns(tmp_path, monkeypatch):
    import xlcalc
    import export_snapshot as es
    wb, _rows = _built(tmp_path, monkeypatch, "Pendle")
    X = xlcalc.Workbook(wb)
    got = es._trajectory(wb, lambda s, r, c: X.value(s, r, c))["Pendle"]
    assert any(k.startswith("traj | A3 | a3_circ_retirement | Q-1") for k in got)
    assert any(k.endswith("| Years to benchmark") for k in got)


def test_history_gaps_names_the_cheapest_fix_on_the_real_dump(tmp_path, monkeypatch):
    import config
    import history_gaps as hg
    wb, _rows = _built(tmp_path, monkeypatch, "Chainlink")
    rows = hg.gap_rows(wb, pd.Timestamp("2026-10-09"))
    by = {r["metric"]: r for r in rows}
    assert by["revenue_usd"]["gap"] == "ok"
    assert "CJ" in by["total_supply"]["cheapest_fix"] and "archive_backfill" in by["total_supply"]["cheapest_fix"]
    assert by["fdv_usd"]["fix_cost"] == "as its inputs"
    assert not [r for r in rows if r["project"] not in config.CREDIBILITY_PROJECTS]
    assert by["revenue_usd"]["Q-1"] == "full" and by["revenue_usd"]["Q-4 (bonus)"].startswith("partial")
    # judged on Q0..Q-3 (Jake's Phase 1 review, item 8): a short Q-4 alone is no gap
    assert all(r["gap"] == "ok" for r in rows if all(r[q] == "full" for q in ("Q0", "Q-1", "Q-2", "Q-3")))
    out = hg.write(rows, tmp_path / "history_gaps.csv")
    assert out.read_text(encoding="utf-8").splitlines()[0].split(",") == hg.FIELDS
    assert isinstance(hg.fix_commands(rows), list)
