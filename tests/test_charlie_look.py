"""Charlie edition LOOK (P2g-4): charlie_look.apply_look restyles the reader workbook and nothing else.

Hand-made workbook with the reader's tab names and header layout (row 3 bands, row 4 headers, project rows from 5),
so no network and no store build.
"""
from __future__ import annotations

import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import charlie_look as cl  # noqa: E402
import reader  # noqa: E402

GREEN = PatternFill("solid", fgColor="C6E0B4")
AMBER = PatternFill("solid", fgColor="FFE699")
RED = PatternFill("solid", fgColor="F4B6B6")

SUMMARY_HEADERS = ["Project", "Token", "Market cap", "Revenue growth, selected period", "Holder value yield",
                   "Net holder yield", "Net supply change, % a year", "Supply", "Staking yield minus T-bill"]
COMPARISON_HEADERS = ["Project", "Token", "Primary archetype", "Holder value yield", "Total dilution, % of float",
                      "Net holder yield", "Payout ratio", "Lock rate, % of circulating"]
BUSINESS_HEADERS = ["Project", "Primary archetype", "TVL", "TVL change, selected period", "Price"]
PROJECTS = [("Alpha", "AAA", False), ("Beta", "BBB", False), ("Gamma", "GGG", False), ("Delta", "DDD", True)]


def _tab(wb, name, headers, bands, formula_cols=(), token_cols=()):
    ws = wb.create_sheet(name)
    ws.cell(1, 1, f"{name} - title").font = reader.F_TITLE
    ws.cell(2, 1, "note").font = reader.F_SUB
    for c, label in bands.items():
        ws.cell(3, c, label).font = Font(name=reader.FONT, size=9, bold=True, color="1F3864")
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(4, c, h)
        cell.font, cell.fill = reader.F_HEAD, reader.FILL_HEAD
        cell.comment = Comment(f"definition of {h}", "t")
    for i, (proj, tok, parked) in enumerate(PROJECTS):
        r = 5 + i
        ws.cell(r, 1, proj).number_format = "@"
        second = tok if headers[1] == "Token" else "A1 Infrastructure"
        ws.cell(r, 2, second).number_format = "@"
        for c in range(3, len(headers) + 1):
            cell = ws.cell(r, c)
            cell.value = f"=A{r}" if False else (i + 1) * (c - 2) * (-1 if c % 2 else 1) / 100
            if c in formula_cols:
                cell.value = f"={openpyxl.utils.get_column_letter(c)}{r + 20}"      # a formula, same sheet
                ws.cell(r + 20, c, (i - 1) * 0.05)
            cell.number_format = "0.0%"
            if c in token_cols:
                cell.number_format = reader.FMT_TOK
                cell.value = 2_500_000 * (i + 1) if i != 2 else -3_000_000
        if parked:
            for c in range(1, len(headers) + 1):
                ws.cell(r, c).fill = reader.FILL_PARKED
    return ws


@pytest.fixture()
def wb():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in ("About", "Charts", "Definitions", "Business", "Comparison", reader.A4_TAB, reader.A3_TAB,
                 reader.A2_TAB, reader.A1_TAB, "Summary"):
        if name == "Summary":
            ws = _tab(wb, name, SUMMARY_HEADERS, {3: "Size", 5: "Rank", 8: "Supply"}, formula_cols=(4, 6),
                      token_cols=(8,))
            ws["A9"] = "A1 Infrastructure - section row"          # a section row: no column B
        elif name == "Comparison":
            ws = _tab(wb, name, COMPARISON_HEADERS, {4: "Holders", 7: "Holding"}, formula_cols=(6,))
        elif name == "Business":
            ws = _tab(wb, name, BUSINESS_HEADERS, {3: "Income"}, token_cols=(3,))
        elif name in (reader.A1_TAB, reader.A2_TAB, reader.A3_TAB, reader.A4_TAB):
            ws = _tab(wb, name, ["Project", "Token", "Key", "Quarter change"], {3: "(a) Rank key"})
        else:
            ws = wb.create_sheet(name)
            ws["A1"] = name
            ws["B2"] = 1_500_000
            ws["B2"].number_format = reader.FMT_TOK
    for n in ("_data", "_series"):
        wb.create_sheet(n).sheet_state = "hidden"
    wb["_data"]["A1"] = 1_000_000
    wb["_data"]["A1"].number_format = reader.FMT_TOK
    # codes, confidence fills, a text cell holding a number format '#,##0'
    ws = wb["Summary"]
    ws["C5"], ws["D6"], ws["E5"], ws["E7"] = "n/a", "blocked", "by design", "secondary"
    ws["C6"].fill, ws["C7"].fill, ws["D5"].fill = GREEN, AMBER, RED
    ws["E8"] = "parked"
    ws["E6"].comment = Comment("why: because", "t")
    ws["B6"].number_format = reader.FMT_TOK                # a TEXT cell in the token format: must not change
    ws = wb["Comparison"]
    ws["F6"] = "n/a"
    ws["D5"].fill = GREEN
    ws["D8"].font = Font(name=reader.FONT, size=10, bold=True, color="C00000", italic=True)
    return wb


def _snapshot(wb) -> dict:
    snap = {}
    for ws in wb.worksheets:
        for cell in ws._cells.values():
            snap[(ws.title, cell.coordinate)] = (cell.value, cell.number_format,
                                                 cell.comment.text if cell.comment else None,
                                                 cell.fill.fgColor.rgb if cell.fill.fill_type else None)
    return snap


def test_invariance_values_comments_formats_and_existing_fills(wb):
    before = _snapshot(wb)
    cl.apply_look(wb)
    after = _snapshot(wb)
    for key, (v, fmt, com, fill) in before.items():
        v2, fmt2, com2, fill2 = after[key]
        assert v2 == v, key
        assert com2 == com, key
        if fill is not None:
            assert fill2 == fill, key                                   # confidence / parked fills untouched
        if fmt != reader.FMT_TOK:
            assert fmt2 == fmt, key
        else:
            assert fmt2 in (reader.FMT_TOK, cl.TOKEN_FMT), key
    assert set(before) <= set(after)                                    # nothing removed; any new cell is empty
    assert all(after[k][0] is None and after[k][2] is None for k in set(after) - set(before))


def test_only_allowed_fills_were_added(wb):
    before = _snapshot(wb)
    cl.apply_look(wb)
    allowed = {cl.ZEBRA, *cl.BAND_BLUES}
    for (sheet, coord), (_, _, _, fill) in _snapshot(wb).items():
        if before.get((sheet, coord), (None,) * 4)[3] is None and fill is not None:
            assert fill[-6:] in allowed, (sheet, coord, fill)


def test_tab_order_colours_and_hidden_last(wb):
    out = cl.apply_look(wb)
    assert wb.sheetnames == ["Summary", reader.A1_TAB, reader.A2_TAB, reader.A3_TAB, reader.A4_TAB, "Comparison",
                             "Business", "Charts", "About", "Definitions", "_data", "_series"]
    assert out["tab_order"] == wb.sheetnames
    assert wb["_data"].sheet_state == wb["_series"].sheet_state == "hidden"
    col = {n: (wb[n].sheet_properties.tabColor.rgb or "")[-6:] for n in wb.sheetnames if n[0] != "_"}
    assert col["Summary"] == "14213D"
    assert {col[n] for n in (reader.A1_TAB, reader.A2_TAB, reader.A3_TAB, reader.A4_TAB)} == {"2F5597"}
    assert col["Comparison"] == col["Business"] == "5B6B7B"
    assert col["Charts"] == col["About"] == col["Definitions"] == "BFBFBF"
    assert sum(1 for ws in wb.worksheets if ws.sheet_view.tabSelected) == 1


def test_highlights_goes_first_and_gets_teal_without_content_change(wb):
    hl = wb.create_sheet("Highlights")
    hl["A1"] = "Token metrics - highlights"
    hl["B3"] = 123456
    hl["B3"].number_format = reader.FMT_TOK
    hl["A1"].font = Font(name="Arial", size=22)
    cl.apply_look(wb)
    assert wb.sheetnames[0] == "Highlights" and wb.sheetnames[1] == "Summary"
    assert hl.sheet_properties.tabColor.rgb[-6:] == "1F8A8A"
    assert hl.sheet_view.zoomScale == 90 and hl.sheet_view.showGridLines is False
    assert hl["A1"].value == "Token metrics - highlights" and hl["A1"].font.sz == 22      # layout left alone
    assert hl["B3"].value == 123456


def test_frozen_panes_gridlines_zoom(wb):
    cl.apply_look(wb)
    assert wb["Summary"].freeze_panes == "C5"                        # after Token (column B), under row 4
    assert wb["Comparison"].freeze_panes == "C5"
    for n in (reader.A1_TAB, reader.A2_TAB, reader.A3_TAB, reader.A4_TAB):
        assert wb[n].freeze_panes == "C5"
    assert wb["Business"].freeze_panes == "B5"                       # no Token column: after Project
    for ws in wb.worksheets:
        if ws.sheet_state == "visible":
            assert ws.sheet_view.showGridLines is False, ws.title
            assert ws.sheet_view.zoomScale == 90, ws.title


def test_header_title_bands_and_body_font(wb):
    cl.apply_look(wb)
    ws = wb["Summary"]
    t = ws["A1"].font
    assert (t.name, t.sz, t.b, t.color.rgb[-6:]) == ("Calibri", 18, True, "14213D")
    for c in range(1, len(SUMMARY_HEADERS) + 1):
        h = ws.cell(4, c)
        assert h.font.b and h.font.color.rgb[-6:] == "FFFFFF" and h.font.name == "Calibri"
        assert h.fill.fgColor.rgb[-6:] == "1F3864"                    # the reader's navy, kept
        assert h.alignment.wrap_text and h.alignment.vertical == "center"
    assert ws.row_dimensions[4].height in (32, 44)
    band = {c: ws.cell(3, c).fill.fgColor.rgb[-6:] for c in range(3, 10)}
    assert band[3] == band[4] == "2F5597" and band[5] == band[6] == band[7] == "5B8DB8" and band[8] == "2F5597"
    assert ws["C3"].font.b and ws["C3"].font.color.rgb[-6:] == "FFFFFF"
    assert ws["A3"].fill.fill_type is None                             # before the first block: unfilled
    assert ws["A5"].font.name == "Calibri" and ws["A5"].font.sz == 10
    assert ws["A5"].border.left.style == "thin" and ws["A5"].border.left.color.rgb[-6:] == "D9D9D9"
    assert ws["A9"].border.left.style is None                          # section row: no data border
    # zebra: odd project rows, only where there was no fill; existing confidence fills stay
    assert ws["F5"].fill.fill_type is None and ws["F6"].fill.fgColor.rgb[-6:] == "F7F8FA"
    assert ws["C6"].fill.fgColor.rgb[-6:] == "C6E0B4" and ws["C7"].fill.fgColor.rgb[-6:] == "FFE699"
    assert ws["A8"].fill.fgColor.rgb[-6:] == "EDEDED"                   # parked grey kept


def _cf(ws):
    return [(str(cf.sqref), r) for cf in ws.conditional_formatting for r in cf.rules]


def test_sign_colour_rules_on_matching_columns_only_and_no_fill(wb):
    out = cl.apply_look(wb)
    ws = wb["Summary"]
    cols = {sq.split(":")[0].rstrip("0123456789") for sq, r in _cf(ws) if r.type == "expression"}
    # header text -> letter: D growth, F net holder yield, G net supply change, I staking yield minus
    assert cols == {"D", "F", "G", "I"}
    assert sorted(out["sign_columns"]) == sorted(
        [f"Summary!{SUMMARY_HEADERS[i]}" for i in (3, 5, 6, 8)]
        + [f"Business!{BUSINESS_HEADERS[3]}", f"Comparison!{COMPARISON_HEADERS[5]}"]
        + [f"{n}!Quarter change" for n in (reader.A1_TAB, reader.A2_TAB, reader.A3_TAB, reader.A4_TAB)])
    for w in wb.worksheets:
        for sq, r in _cf(w):
            if r.type == "expression":
                assert r.dxf.fill is None, (w.title, sq)
                assert r.dxf.font.color.rgb[-6:] in ("1B7F5C", "B23A48")
                assert "ISNUMBER" in r.formula[0]                      # codes (text) never trigger
    colours = {r.dxf.font.color.rgb[-6:] for _, r in _cf(ws) if r.type == "expression"}
    assert colours == {"1B7F5C", "B23A48"}
    rng = next(sq for sq, r in _cf(ws) if r.type == "expression")
    assert rng == "D5:D8"                                                 # data rows only (row 9 is a section row)


def test_data_bars_only_on_the_three_comparison_columns(wb):
    out = cl.apply_look(wb)
    bars = [(w.title, sq) for w in wb.worksheets for sq, r in _cf(w) if r.type == "dataBar"]
    assert sorted(bars) == [("Comparison", "D5:D8"), ("Comparison", "E5:E8"), ("Comparison", "F5:F8")]
    assert out["bar_columns"] == ["Holder value yield", "Total dilution, % of float", "Net holder yield"]
    assert out["bar_columns_missing"] == []
    for w in wb.worksheets:
        for sq, r in _cf(w):
            if r.type == "dataBar":
                assert r.dataBar.color.rgb[-6:] == "8FD3CF"


def test_missing_bar_header_is_reported(wb):
    wb["Comparison"]["E4"] = "Dilution"
    out = cl.apply_look(wb)
    assert out["bar_columns_missing"] == ["Total dilution, % of float"]


def test_codes_grey_italic_centred_and_numbers_untouched(wb):
    out = cl.apply_look(wb)
    ws = wb["Summary"]
    for a in ("C5", "D6", "E5", "E7", "E8"):
        c = ws[a]
        assert c.font.i and c.font.color.rgb[-6:] == "7F7F7F" and c.alignment.horizontal == "center", a
    assert ws["C6"].font.i is not True                                     # a number stays as it was
    assert wb["Comparison"]["F6"].alignment.horizontal == "center"
    assert ws["E6"].comment.text == "why: because"
    assert wb["Comparison"]["D8"].font.b and wb["Comparison"]["D8"].font.i                 # kept bold/italic
    assert out["codes_styled"] == 6


def test_token_format_only_where_it_was_plain_thousands(wb):
    out = cl.apply_look(wb)
    assert wb["Summary"]["H5"].number_format == cl.TOKEN_FMT
    assert wb["Business"]["C6"].number_format == cl.TOKEN_FMT
    assert wb["About"]["B2"].number_format == cl.TOKEN_FMT
    assert wb["Summary"]["B6"].number_format == reader.FMT_TOK            # a text cell: unchanged
    assert wb["_data"]["A1"].number_format == reader.FMT_TOK              # hidden sheets: unchanged
    assert wb["Summary"]["D5"].number_format == "0.0%"
    assert out["token_format_cells"] == 4 + 4 + 3                          # Summary H, Business C, B2 of Charts/About/Definitions
    assert cl.TOKEN_FMT == reader.FMT_USD.replace("$", "")


def test_idempotent(wb):
    cl.apply_look(wb)
    first = _snapshot(wb)
    cf1 = [(w.title, sq, r.type) for w in wb.worksheets for sq, r in _cf(w)]
    styles1 = {(w.title, c.coordinate): (c.font.name, c.font.sz, c.font.b, c.font.i, c.border.left.style,
                                        c.alignment.horizontal) for w in wb.worksheets for c in w._cells.values()}
    out2 = cl.apply_look(wb)
    assert _snapshot(wb) == first
    assert [(w.title, sq, r.type) for w in wb.worksheets for sq, r in _cf(w)] == cf1
    assert {(w.title, c.coordinate): (c.font.name, c.font.sz, c.font.b, c.font.i, c.border.left.style,
                                     c.alignment.horizontal) for w in wb.worksheets for c in w._cells.values()} \
        == styles1
    assert out2["tab_order"] == wb.sheetnames


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_saved_file_recalcs_in_libreoffice_with_no_mismatch(wb, tmp_path):
    import json
    import workbook_values
    cl.apply_look(wb)
    p = tmp_path / "look.xlsx"
    wb.save(p)
    workbook_values.store_values(p, wb)
    res = subprocess.run([sys.executable, str(ROOT / "recalc_check.py"), str(p)], capture_output=True, text=True,
                         timeout=300)
    data = json.loads(res.stdout[res.stdout.index("{"):])
    assert data["formulas"] >= 8
    assert data["mismatches"] == 0 and data["errors"] == 0, data
