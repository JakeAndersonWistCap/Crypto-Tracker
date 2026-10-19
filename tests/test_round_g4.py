"""
tests/test_round_g4.py — Round G4 (Jake 2026-10-18 night), SPEC §19 "G4 — Round G4, old → new". One test (or more) per
item, each named for its row.
"""
from __future__ import annotations

import io
import json
import logging
import socket
import sys
import threading
import time
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import config  # noqa: E402
from fetch.base import FetchOutput  # noqa: E402


@pytest.fixture
def local_only(monkeypatch):
    for k in ("HTTP_PROXY", "http_proxy"):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("NO_PROXY", "127.0.0.1,localhost")
    monkeypatch.setenv("no_proxy", "127.0.0.1,localhost")


def _hanging_server():
    srv = socket.socket()
    srv.bind(("127.0.0.1", 0))
    srv.listen(64)
    conns = []

    def accept():
        while True:
            try:
                c, _ = srv.accept()
            except OSError:
                return
            conns.append(c)
    threading.Thread(target=accept, daemon=True).start()
    return srv, conns, f"http://127.0.0.1:{srv.getsockname()[1]}"


# ------------------------------------------------------------------------------------------ item 1: the timeouts
def test_g4_1a_a_paced_wait_is_registered_and_an_abandoned_tier_stops_before_reserving():
    from fetch import base
    pace = base.HostPace(2.0)                                         # 0.5 s between calls, real time.sleep
    seen = []

    def worker():
        pace.wait("api.etherscan.io")
        pace.wait("api.etherscan.io")                                 # this one sleeps ~0.5 s, registered
    t = threading.Thread(target=worker, name="explorer_l2:arbitrum")
    t.start()
    for _ in range(40):
        seen += [w for src, w, _ in base.inflight() if src == "explorer_l2:arbitrum"]
        time.sleep(0.02)
    t.join()
    assert any(w.startswith("paced for api.etherscan.io (2.00 req/s shared run-wide)") for w in seen), seen
    # an abandoned tier raises at the pacer, before it reserves a slot others would wait behind
    pace2 = base.HostPace(1.0)
    errs = []

    def abandoned():
        base.abandon_thread(threading.current_thread())
        try:
            pace2.wait("api.etherscan.io")
        except base.TierAbandoned as e:
            errs.append(e)
    t = threading.Thread(target=abandoned, name="explorer_l2:base")
    t.start()
    t.join()
    assert errs and pace2._next == 0.0, "no slot reserved by the abandoned tier"


def test_g4_1b_a_slow_http_host_is_stalled_after_two_timeouts_and_not_asked_again(local_only):
    from fetch import base
    srv, conns, url = _hanging_server()
    try:
        base.host_stall_reset()
        h = base.Http(timeout=0.3, retries=2)
        t0 = time.monotonic()
        with pytest.raises(base.HostStalled) as ei:
            h.get(url + "/api")
        assert "HOST STALLED this run: 127.0.0.1" in str(ei.value)
        assert len(conns) == 2, "a timeout is retried at most once; two in a row stall the host"
        assert time.monotonic() - t0 < 6
        t0 = time.monotonic()
        with pytest.raises(base.HostStalled):
            h.get(url + "/api")
        assert time.monotonic() - t0 < 0.1 and len(conns) == 2, "a stalled host is not asked again"
    finally:
        srv.close()
        base.host_stall_reset()


def test_g4_1c_each_scan_gets_its_share_of_the_tier_budget(monkeypatch):
    import fetch
    from fetch.logscan import TIER_MARGIN_S, LogScan
    monkeypatch.setitem(fetch.TIER_BUDGET_S, "explorer_l2:test", 100)
    monkeypatch.delenv("TOKEN_METRICS_BUDGET_EXPLORER_L2:TEST", raising=False)
    shares = []
    proj = {"name": "Pendle", "log_scans": [{"key": f"k{i}", "metric": f"m{i}", "chain": "arbitrum",
                                              "own_tier": "explorer_l2:test"} for i in range(3)]}
    ls = LogScan(own_tier="explorer_l2:test")
    ls._scan = lambda p, spec, w, out: shares.append(ls._share(10_000))
    t = threading.Thread(target=lambda: ls.run([proj], None, FetchOutput()), name="explorer_l2:test")
    t.start()
    t.join()
    want = (100 - TIER_MARGIN_S) / 3
    assert len(shares) == 3 and abs(shares[0] - want) < 2, shares
    # time an early scan did not use passes to the ones after it; none may run past the tier's own margin
    assert all(s <= 100 - TIER_MARGIN_S + 1 for s in shares), shares


def test_g4_1d_the_timeout_message_names_the_last_phase():
    import fetch
    from fetch.base import set_phase

    class A:
        planned = [("Morpho", "emissions_tokens_claims_base")]
    st = {"sub": FetchOutput(), "adapter": A(), "tier": 2, "t0": time.monotonic() - 240}
    t = threading.Thread(target=lambda: set_phase("scan merkl_morpho_out_base (base): reconciling each holder's "
                                                  "balanceOf to the wei"), name="explorer_l2:base")
    t.start()
    t.join()
    _, _, kept, _, _ = fetch._abandon("explorer_l2:base", st, 240)
    msg = " ".join(e.message for e in kept.log)
    assert "waiting on: nothing registered" in msg.replace("Still waiting on", "waiting on")
    assert "Last phase: scan merkl_morpho_out_base (base): reconciling each holder's balanceOf to the wei" in msg


def test_g4_1e_with_blockscout_and_katana_blocked_every_other_series_lands(local_only, monkeypatch, tmp_path):
    """Done-when of item 1: Blockscout (an explorer host) and rpc.katana.network both hang. Through the real dispatcher,
    the series served elsewhere land inside their budgets; the Run Log names each stalled source and the series it
    gapped."""
    import fetch
    from fetch import base
    from fetch import chain as ch
    from fetch.explorer import ExplorerLogs
    from fetch.logscan import LogScan
    srv, conns, url = _hanging_server()
    try:
        real = ch.rpc_endpoints
        monkeypatch.setattr(ch, "rpc_endpoints", lambda c: [url] if c == "katana" else real(c))
        monkeypatch.setenv("TOKEN_METRICS_RPC_READ_TIMEOUT", "0.3")
        monkeypatch.setenv("TOKEN_METRICS_LOGCACHE", str(tmp_path / "logscan"))
        monkeypatch.setenv("TOKEN_METRICS_DAILY_CHECKS", "always")
        monkeypatch.setenv("BLOCKSCOUT_API_KEY", "test-blockscout-key")
        monkeypatch.setitem(config.EXPLORERS["blockscout"], "pro_api", url + "/{chain_id}/api")
        monkeypatch.setattr(config, "EXPLORER_READ_TIMEOUT_S", 0.3)
        monkeypatch.setitem(config.EXPLORER_LOG_ROUTES, 10, ["blockscout"])
        ch.BREAKER.reset()
        base.host_stall_reset()
        morpho, pendle = config.PROJECT_BY_NAME["Morpho"], config.PROJECT_BY_NAME["Pendle"]
        katana = next(s for s in morpho["log_scans"] if s["key"] == "merkl_morpho_out_katana")
        op = {**next(s for s in pendle["log_scans"] if s["key"] == "gauge_pendle_out_optimism"),
              "own_tier": "explorer", "window_days": 0, "declared_zero": None}
        base_spec = {**katana, "key": "merkl_morpho_out_base_test", "chain": "base",
                     "metric": "emissions_tokens_claims_base", "own_tier": config.l2_tier("base"), "rpc_window": None}

        class BaseScan(LogScan):
            def _scan(self, p, spec, w, out):
                out.add(pd.DataFrame({"date": [pd.Timestamp("2026-10-17")], "project": p["name"],
                                      "metric": spec["metric"], "value": [7.0], "source": "explorer:x", "tier": 2}),
                        "explorer", p["name"], "stub", 2)
        projects = [{**morpho, "log_scans": [katana, base_spec]}, {**pendle, "log_scans": [op]}]
        todo = [("explorer", 2, lambda ctx: LogScan(own_tier="explorer", explorer=ExplorerLogs())),
                (config.l2_tier("katana"), 2, lambda ctx: LogScan(own_tier=config.l2_tier("katana"))),
                (config.l2_tier("base"), 2, lambda ctx: BaseScan(own_tier=config.l2_tier("base")))]
        t0 = time.monotonic()
        results = fetch._dispatch(todo, projects, None, {})
        assert time.monotonic() - t0 < 60
        by = {n: (sub, to) for n, _, sub, _, to in results}
        assert not any(to for _, to in by.values()), "no tier timed out"
        assert list(by[config.l2_tier("base")][0].frame().metric) == ["emissions_tokens_claims_base"]
        ex = [e.message for e in by["explorer"][0].log if e.status == "failed"]
        src = [m for m in ex if m.startswith("SOURCE STALLED THIS RUN: 127.0.0.1")]
        assert len(src) == 1 and "Pendle/emissions_tokens_gauge_optimism" in src[0], ex
        kat = [e.message for e in by[config.l2_tier("katana")][0].log if e.status == "failed"]
        assert any(m.startswith("CHAIN STALLED THIS RUN: katana") and "Morpho/emissions_tokens_claims_katana" in m
                   for m in kat), kat
    finally:
        srv.close()
        ch.BREAKER.reset()
        base.host_stall_reset()


def test_g4_1f_katana_reads_only_the_days_after_its_last_stored_row(tmp_path, monkeypatch):
    from fetch.logscan import LogScan
    from fetch.base import today
    monkeypatch.setenv("TOKEN_METRICS_DAILY_CHECKS", "always")
    morpho = config.PROJECT_BY_NAME["Morpho"]
    spec = next(s for s in morpho["log_scans"] if s["key"] == "merkl_morpho_out_katana")
    assert spec["incremental_after_seed"] is True
    yday = today().normalize() - pd.Timedelta(days=1)
    starts = []

    class LS(LogScan):
        def _block_at_ts(self, chain, cid, ts, to_block):
            starts.append(pd.Timestamp(ts, unit="s").normalize())
            raise RuntimeError("stop here")
    ls = LS(own_tier=config.l2_tier("katana"), last_dates={("Morpho", spec["metric"]): yday - pd.Timedelta(days=3)})
    ls.cache.root = tmp_path
    out = FetchOutput()
    ls._window_scan(morpho, spec, "katana", 747474, spec["token"], [h.lower() for h in spec["holders"]], 100, "x", out)
    assert starts == [yday - pd.Timedelta(days=3)], "the window starts at the last stored day, not 100 days back"
    ls2 = LS(own_tier=config.l2_tier("katana"), last_dates={("Morpho", spec["metric"]): yday})
    ls2.cache.root = tmp_path / "b"
    out = FetchOutput()
    assert ls2._window_scan(morpho, spec, "katana", 747474, spec["token"], [], 100, "x", out) is True
    assert "up to date" in out.log[-1].message and not starts[1:]


def test_g4_1g_a_density_narrowed_chunk_grows_back_a_range_limit_does_not():
    from fetch.chain import ChainReader
    r = ChainReader()
    calls = []

    def resilient(chain, base, lo, hi, chunk, only=None):
        span = min(chunk, hi - lo + 1)
        calls.append(span)
        if lo < 1_000 and span > 100:                               # one dense stretch at the start
            r.__dict__.setdefault("_dense", set()).add(chain)
            return [], lo + 99, 100
        return [], lo + span - 1, chunk
    r._get_logs_resilient = resilient
    r._rpc_traverse("ethereum", {}, 0, 50_000, 10_000)
    assert max(calls[2:]) == 10_000, f"grew back after the dense stretch: {calls[:8]}"
    assert len(calls) < 30, "the dense stretch costs calls, the rest of the walk does not"


# ------------------------------------------------------------------------------- item 2: market cap like with like
def _sanity_wb(mcap_date, price_date, circ_date):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master"
    heads = ["Project", "Price — spot ($)", "Market cap ($)", "Circulating supply (used)", "FDV ($)"]
    for c, h in enumerate(heads, start=1):
        ws.cell(4, c, h)
    for c, v in enumerate(["Aerodrome", 0.76949, 1_333_067_528.0, 1_516_808_772.0, 2e9], start=1):
        ws.cell(5, c, v)
    d = wb.create_sheet("Data")
    for c, h in enumerate(["Key", "Latest date", "Now (flow: trailing 30d sum · stock: latest)"], start=1):
        d.cell(1, c, h)
    for r, (k, dt) in enumerate((("Aerodrome|market_cap_usd", mcap_date), ("Aerodrome|price_usd", price_date),
                                 ("Aerodrome|circulating_supply_onchain", circ_date),
                                 ("Aerodrome|total_supply", "2026-10-18")), start=2):
        d.cell(r, 1, k)
        d.cell(r, 2, dt)
        d.cell(r, 3, 2_002_000_000.0 if k.endswith("total_supply") else None)
    return wb


def test_g4_2_the_mcap_rule_compares_like_with_like_and_names_a_stale_input():
    import sanity
    wb = _sanity_wb("2026-10-17", "2026-10-18", "2026-10-14")
    got = [h for h in sanity.hits(wb, lambda s, r, c: wb[s].cell(r, c).value) if h[0] == "Aerodrome"]
    rules = [r for _p, r, _d in got]
    assert "mcap_stale" in rules and "mcap_not_price_x_circ" not in rules, got
    d = next(d for _p, r, d in got if r == "mcap_stale")
    assert "dated 2026-10-17" in d and "spot price 0.76949 is dated 2026-10-18" in d
    assert "circulating_supply_onchain last stored 2026-10-14" in d and "-12.44%" in d
    # the same date on both sides and still apart: the rule fires as before, unchanged tolerance
    wb = _sanity_wb("2026-10-18", "2026-10-18", "2026-10-18")
    rules = [r for p, r, _d in sanity.hits(wb, lambda s, r, c: wb[s].cell(r, c).value) if p == "Aerodrome"]
    assert "mcap_not_price_x_circ" in rules and "mcap_stale" not in rules
    assert sanity.TOL_MCAP == 0.005


def test_g4_2_a_daily_mark_is_kept_only_when_its_series_reached_the_store(tmp_path, monkeypatch):
    import token_metrics as tm
    from fetch.logcache import DailyChecks
    monkeypatch.delenv("TOKEN_METRICS_DAILY_CHECKS", raising=False)
    d = DailyChecks(tmp_path)
    DailyChecks.defer()
    try:
        d.done("ve_managed:Aerodrome:filing_managed_lock_tokens", "2026-10-18")
        d.done("business:Near:near_intents_volume_usd:full-history", "2026-10-18")
        d.done("window:merkl_morpho_out_katana", "2026-10-18")
    finally:
        frame = pd.DataFrame({"date": [pd.Timestamp("2026-10-18")], "project": "Near",
                              "metric": "near_intents_volume_usd", "value": [1.0], "source": "x", "tier": 1})
        n = DailyChecks.commit(verify=tm._mark_stored(frame))
    assert n == 2 and DailyChecks.last_dropped == ["ve_managed:Aerodrome:filing_managed_lock_tokens"]
    assert set(json.loads(d.f.read_text())) == {"business:Near:near_intents_volume_usd:full-history",
                                                "window:merkl_morpho_out_katana"}


# ---------------------------------------------------------------------------------- item 3: pool release pairing
def test_g4_3_hyperliquid_pool_release_pairs_coingecko_with_coingecko():
    import fetch
    from fetch.base import point
    hl = config.PROJECT_BY_NAME["Hyperliquid"]
    out = FetchOutput()
    day = pd.Timestamp("2026-10-18")
    out.add(point("Hyperliquid", "circulating_supply", 340_000_000.0, "coingecko", 1, day), "coingecko", "Hyperliquid")
    out.add(point("Hyperliquid", "total_supply_coingecko", 999_000_000.0, "coingecko", 1, day), "coingecko",
            "Hyperliquid")
    prior = {("Hyperliquid", "circulating_supply"): 339_000_000.0,
             ("Hyperliquid", "total_supply_coingecko"): 999_000_000.0}
    dates = {k: pd.Timestamp("2026-10-17") for k in prior}
    fetch._derive_pool_release(out, [hl], prior, dates)
    f = out.frame()
    rel = f[f.metric == "pool_release_tokens"]
    assert len(rel) == 1 and rel.value.iloc[0] == 1_000_000.0
    assert "d(total_supply_coingecko) 0" in out.log[-1].message
    assert "missing total_supply" not in " ".join(e.message for e in out.log)


# ------------------------------------------------------------------------------------------- item 4: logging
def test_g4_4_a_line_with_a_minus_sign_is_never_lost_on_a_cp1252_console():
    import progress
    raw = io.TextIOWrapper(io.BytesIO(), encoding="cp1252")
    assert not progress.console_can_encode(raw)
    safe = progress.ConsoleSafe(raw)
    h = logging.StreamHandler(safe)
    lg = logging.getLogger("g4_console")
    lg.addHandler(h)
    lg.propagate = False
    lg.warning("net total = 998,862,059 − 48,007,021 → 950,546,499 ≠ Δ")
    raw.flush()
    assert raw.buffer.getvalue() == b"net total = 998,862,059 - 48,007,021 -> 950,546,499 != ?\n"


def test_g4_4_the_katana_seed_reports_progress_while_dating():
    from fetch.katana_seed import date_blocks
    said = []
    blocks = list(range(0, 3_000, 3))
    date_blocks(blocks, lambda b: 1_750_000_000 + b * 40, say=said.append, every=0.0)
    assert said and said[0].startswith("dating payouts: ") and "block(s) dated" in said[0]


# ----------------------------------------------------------------------------------- item 5: stale-series audit
def test_g4_5_stale_series_lists_old_rows_and_the_marked_but_empty_case(tmp_path, monkeypatch, capsys):
    import check_offline_items as coi
    import store
    from fetch.logcache import DailyChecks
    db = tmp_path / "m.db"
    st = store.Store(str(db))
    today = pd.Timestamp.now(tz="UTC").tz_localize(None).normalize()
    st.upsert(pd.DataFrame({"date": [today - pd.Timedelta(days=4), today - pd.Timedelta(days=1)],
                            "project": "Aerodrome", "metric": ["filing_managed_lock_tokens", "price_usd"],
                            "value": [444e6, 0.77], "source": ["ve_managed:0xeBf418Fe.weights", "coingecko"],
                            "tier": 2}))
    st.close()
    monkeypatch.setattr(store, "DB_PATH", db)
    monkeypatch.setenv("TOKEN_METRICS_LOGCACHE", str(tmp_path / "logscan"))
    monkeypatch.setattr(coi, "_portfolio_projects", lambda: [config.PROJECT_BY_NAME["Aerodrome"]])
    DailyChecks().f.parent.mkdir(parents=True, exist_ok=True)
    DailyChecks().f.write_text(json.dumps({"ve_managed:Aerodrome:filing_managed_lock_tokens": str(today.date())}))
    coi.stale_series()
    out = capsys.readouterr().out
    line = next(ln for ln in out.splitlines() if "filing_managed_lock_tokens" in ln)
    assert "4 day(s) old" in line and "MARKED READ TODAY" in line and "BACKFILL" in line
    assert "price_usd" not in out, "a daily series one day old is current"
    assert coi.stale_series in coi.CHECKS and coi.mcap_inputs in coi.CHECKS


# -------------------------------------------------------------------------------- item 6: seeds count in usage
def test_g4_6_seed_runs_report_and_add_their_alchemy_usage():
    src = (ROOT / "token_metrics.py").read_text()
    seed_block = src[src.index("if args.seed:"):src.index("# ===== BUILD ONLY. =====")]
    assert '_alchemy_report(log, f"--seed {args.seed}")' in seed_block
    assert "alchemy_usage.report" in (ROOT / "archive_backfill.py").read_text()


# ------------------------------------------------------------------------------------------------ item 7(a) Sky TMF
def test_g4_7a_the_sky_tmf_reference_is_the_msc13_figure_and_stays_broken_explained():
    import csv
    with open(ROOT / "independent_references.csv", encoding="utf-8") as f:
        row = next(r for r in csv.DictReader(f) if r["metric"] == "net_revenue_sep_2026")
    assert float(row["value"]) == 14_635_947.43 and row["tolerance"] == "2%"          # tolerance kept
    assert "settlement-cycle" in row["source"] and "summary.md" in row["source"]
    # the MSC #13 arithmetic the cause quotes: MSC leg + non-MSC leg, and each leg's lines
    assert round(20_154_616.00 - 6_050_154.00 - 177_113.78, 2) == 13_927_348.22
    assert round(14_938_630.49 - 14_230_031.28, 2) == 708_599.21
    assert round(13_927_348.22 + 708_599.21, 2) == 14_635_947.43
    assert 8_218_668 + 1_016_416 + 2_392_354 == 11_627_438                           # Spark mint incl. Jan–Aug fix
    assert round(5_930_611 + 1_041_399 - 165_013.90) == 6_806_996                    # Grove mint less its fix


# ------------------------------------------------------------------------------------------------ item 7(b) queue
def _circ_frame(name, rows):
    return pd.DataFrame([{"date": pd.Timestamp("2026-10-18"), "project": name, "metric": m, "value": float(v)}
                         for m, v in rows.items()])


def test_g4_7b_ether_fi_circulating_is_checked_on_coingeckos_basis_and_a_real_gap_still_flags():
    from fetch import circulating
    spec = config.CIRCULATING_ONCHAIN["Ether.fi"]
    assert spec["coingecko_counts"] == ("noncirculating_holding_tokens",)
    base = {"total_supply": 998_535_999, "treasury_holding_tokens": 0, "noncirculating_holding_tokens": 188_586_981,
            "total_supply_coingecko": 965_350_000}
    # ours 809,949,018 (the 10-18 row); + the wallets CoinGecko counts, on CoinGecko's own total = 965,350,000
    out = FetchOutput()
    circulating.check(out, _circ_frame("Ether.fi", {**base, "circulating_supply": 965_350_000}),
                      [config.PROJECT_BY_NAME["Ether.fi"]])
    assert not [r for r in out.review if r["reason"] == "onchain_vs_coingecko"]
    out = FetchOutput()                                            # a real 5% disagreement still raises the row
    circulating.check(out, _circ_frame("Ether.fi", {**base, "circulating_supply": 1_016_000_000}),
                      [config.PROJECT_BY_NAME["Ether.fi"]])
    hit = [r for r in out.review if r["reason"] == "onchain_vs_coingecko"]
    assert hit and "the wallets CoinGecko counts" in hit[0]["basis"] and "total_supply_coingecko" in hit[0]["basis"]
    assert config.CIRCULATING_TOLERANCE == 0.02


def test_g4_7b_answered_rows_leave_the_open_queue_and_say_the_decision():
    import openpyxl
    import build_workbook as bw
    rv = pd.DataFrame([
        {"project": "Ether.fi", "metric": "sethfi_reward_tokens_reconciled", "reason": "change_threshold",
         "action": "stored_flagged", "value": 6816.65, "prior_value": 24213.41, "date": "2026-10-17",
         "prior_date": "2026-10-10", "basis": "", "source": "explorer", "tier": 2},
        {"project": "Ether.fi", "metric": "sethfi_reward_tokens_reconciled", "reason": "change_threshold",
         "action": "stored_flagged", "value": 1.0, "prior_value": 9.0, "date": "2026-10-24",
         "prior_date": "2026-10-17", "basis": "", "source": "explorer", "tier": 2},
        {"project": "Ether.fi", "metric": "locked_tokens", "reason": "supply_partial", "action": "stored_flagged",
         "value": 87478043.18, "prior_value": 87478289.44, "date": "2026-10-18", "prior_date": None, "basis": "",
         "source": "chain", "tier": 2}])
    wb = openpyxl.Workbook()
    ws = wb.active
    bw.write_review_queue(ws, rv, "r")
    assert bw.BOARD_INPUTS["review_open"]["Ether.fi"] == 1          # only the 10-24 row is open
    text = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert any(t.startswith("ANSWERED — not open (2)") for t in text)
    assert any(t == "ANSWERED: ACCEPT" for t in text)
    # an allow-listed lost headline is answered under the sanity block, an unlisted one is open
    hits = [("Pendle", "headline_lost", "CP SUBSIDY COVER = revenue ÷ $ supply-side incentives (b): 12.578 in the "
                                        "previous run (2026-10-13), now BLOCKED — unequal coverage"),
            ("Pendle", "headline_lost", "CP SOMETHING ELSE: 1 in the previous run, now n/a")]
    bw._write_sanity(ws, hits)
    text = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert any("SANITY RULES — 1 hit(s)" in t for t in text)
    assert any(t == "headline_lost — ANSWERED" for t in text)
    assert bw.BOARD_INPUTS["review_open"]["Pendle"] == 1


# ------------------------------------------------------------------------------------------------ item 7(c) P6
def test_g4_7c_a_configured_route_that_stored_nothing_is_p4_and_an_empty_stub_stays_p6():
    from fetch.gaps import P_ACTIONABLE, P_ACTIONABLE_HEADLINE, P_UNCOVERED, _priority
    assert _priority("Aerodrome", "voter_rewards_onchain_fees_usd", "the voting-reward contracts' epoch read "
                     "(fetch/aero_voter.py) is configured but stored nothing this run") == P_ACTIONABLE
    assert _priority("Aerodrome", "actual_buyback_tokens", "the Foundation buyback window stored nothing this "
                     "run") == P_ACTIONABLE_HEADLINE
    assert _priority("Ether.fi", "cashback_paid_tokens", "TIER TIMED OUT this run — explorer (tier 2) timed out "
                     "after 300s") == P_ACTIONABLE
    assert _priority("Aerodrome", "max_supply", "CoinGecko returned no data for this series") == P_ACTIONABLE
    assert _priority("Aethir", "customer_revenue_usd", "entry disabled in sources.yaml — EMPTY STUB, no url on file "
                     "yet") == P_UNCOVERED


def test_g4_7c_gap_routes_groups_open_rows_by_project_and_route(tmp_path, monkeypatch, capsys):
    import sqlite3
    import check_offline_items as coi
    import store as sm
    db = tmp_path / "m.db"
    sm.Store(str(db)).close()
    con = sqlite3.connect(str(db))
    cols = [r[1] for r in con.execute("PRAGMA table_info(gap_report)")]
    rows = [("r1", "Aethir", "customer_revenue_usd", "entry disabled in sources.yaml — EMPTY STUB, no url on file yet",
             "Complete the entry", "P6 uncovered"),
            ("r1", "Aerodrome", "voter_rewards_onchain_fees_usd", "aero_voter read is configured but stored nothing "
             "this run", "python check_offline_items.py aerodrome_voter_rewards", "P4 fixable"),
            ("r1", "Maple", "x", "deliberately disabled", "", "P5 answered")]
    for run, proj, metric, reason, sugg, pl in rows:
        vals = {"run_id": run, "project": proj, "metric": metric, "reason": reason, "suggestion": sugg,
                "priority_label": pl, "priority": int(pl[1]), "tiers_attempted": "3",
                "timestamp": "2026-10-18T00:00:00", "ts": "2026-10-18T00:00:00"}
        keys = [k for k in vals if k in cols]
        con.execute(f"INSERT INTO gap_report ({', '.join(keys)}) VALUES ({', '.join('?' * len(keys))})",
                    [vals[k] for k in keys])
    con.commit()
    con.close()
    monkeypatch.setattr(sm, "DB_PATH", str(db))
    coi.gap_routes()
    out = capsys.readouterr().out
    assert "2 open row(s) in 2 project(s)" in out
    assert out.index("Aethir") < out.index("Aerodrome")                       # Jake's order: Aethir first
    assert "NO ROUTE — empty registry stub" in out and "ROUTE: aero_voter" in out and "Maple" not in out


# ------------------------------------------------------------------------------------------------ item 7(e) N/A
def test_g4_7e_na_reasons_flags_the_missing_routes(tmp_path, monkeypatch, capsys):
    import csv
    import check_offline_items as coi
    import coverage_audit as ca
    p = tmp_path / "coverage_audit.csv"
    rows = ([{"Project": "World Mobile", "Status": "N/A BY DESIGN",
              "Reason / route / error": "no cell written: the column does not apply to this project"}] * 3
            + [{"Project": "World Mobile", "Status": "N/A BY DESIGN", "Reason / route / error": "n/a — Data: no rows stored"}] * 2
            + [{"Project": "World Mobile", "Status": "LIVE-UNVERIFIED", "Reason / route / error": ""}] * 4)
    with open(p, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=ca.HEAD)
        w.writeheader()
        for r in rows:
            w.writerow({h: r.get(h, "") for h in ca.HEAD})
    monkeypatch.setenv("TOKEN_METRICS_COVERAGE_CSV", str(p))
    coi.na_reasons()
    out = capsys.readouterr().out
    assert "World Mobile: 5 N/A vs 4 value(s) — N/A EXCEEDS VALUES; 2 of them a MISSING ROUTE" in out
    assert "<< MISSING ROUTE" in out


# ------------------------------------------------------------------------------------------------ item 7(f) board
def test_g4_7f_the_scoreboard_counts_statuses_gaps_and_open_review_rows():
    import build_workbook as bw
    import coverage_audit as ca
    rows = [{"Project": "Sky", "Status": s} for s in ("VERIFIED", "LIVE-UNVERIFIED", "LIVE-UNVERIFIED", "BROKEN",
                                                       "N/A BY DESIGN", "NOT WIRED", "CHECK")]
    rows += [{"Project": "(all)", "Status": "NOT WIRED"}]
    board = ca.scoreboard(rows, {("Sky", 6): 3, ("Sky", 4): 2, ("Sky", 5): 7, ("Aethir", 6): 14}, {"Sky": 2})
    s = board["Sky"]
    assert (s["VERIFIED"], s["LIVE-UNVERIFIED"], s["N/A"], s["NOT WIRED"], s["BROKEN"], s["CHECK"]) == (1, 2, 1, 1, 1, 1)
    assert (s["Gap P4"], s["Gap P5"], s["Gap P6"], s["Review Queue"]) == (2, 7, 3, 2)
    assert s["Verified share"] == "1/5 = 20%"
    assert board["ALL"]["NOT WIRED"] == 2 and board["ALL"]["Gap P6"] == 17 and board["Aethir"]["Gap P6"] == 14
    line = bw.scoreboard_lines(board)[1]
    assert line.startswith("SCOREBOARD Aethir:") or line.startswith("SCOREBOARD Sky:")
    assert "Review Queue" in line and "Gap P4" in line


def test_g4_7b_geodnet_coingecko_still_on_its_documented_static_figure_is_recorded_not_queued():
    from fetch import circulating
    st = config.CIRCULATING_ONCHAIN["GEODNET"]["static_cross_check"]
    spec = config.CIRCULATING_ONCHAIN["GEODNET"]
    legs = {spec["total"]: 600_000_000, **{m: 0.0 for m in spec["subtract"]}}
    for cg, queued in ((float(st["figure"]), False), (float(st["figure"]) * 1.03, True)):
        out = FetchOutput()
        circulating.check(out, _circ_frame("GEODNET", {**legs, "circulating_supply": cg}),
                          [config.PROJECT_BY_NAME["GEODNET"]])
        hit = [r for r in out.review if r["reason"] == "onchain_vs_coingecko"]
        assert bool(hit) is queued, cg
        if not queued:
            assert any("documented static figure" in e.message for e in out.log)
