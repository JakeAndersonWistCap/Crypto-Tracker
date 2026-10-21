"""
tests/test_decisions_6f8c86c.py — Jake's decisions on 6f8c86c / 54a9437 (2026-10-12): Morpho on the issuer rule (1), Sky's
MKR backing, unlock stock and emission rate in force (2), NEAR's lockup.near + contributors.near (3), FDV history on one
chain basis (6), the trajectory basis rule (7), lumpy flows on the rolling T12M (8) and Total Yield (9). Real dumps
wherever they exist (fixtures/real_*.json).
"""
from __future__ import annotations

import sqlite3
import sys
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))


# ----------------------------------------------------------------------------------------------- 6 FDV history
def test_fdv_history_section_cj_moves_coingeckos_pre_handover_totals_on_the_real_dump(tmp_path, dumps_20261009):
    """Decision 6: CoinGecko's pre-handover total_supply rows move to total_supply_coingecko (values untouched), so the
    archive can fill total_supply from the chain at past blocks. The real Pendle dump holds 21 CoinGecko totals."""
    import builtins
    import run_sql
    from real_dump import load, to_store
    payload = load("Pendle")
    db = tmp_path / "metrics.db"
    to_store(payload, db)
    before = sorted((r["date"], r["value"]) for r in payload["metrics"] if r["metric"] == "total_supply")
    assert len(before) == 21
    # one day already holding CoinGecko's total under its own name: the rename leaves that row, the DELETE removes it
    con = sqlite3.connect(db)
    con.execute("INSERT INTO metrics (date, project, metric, value, source, tier, fetched_at) VALUES (?, 'Pendle', "
                "'total_supply_coingecko', ?, 'coingecko', 1, 'x')", before[0])
    con.commit()
    con.close()
    real_input = builtins.input
    builtins.input = lambda *a: "DELETE CJ"
    try:
        assert run_sql.main(["--delete", "CJ", "--db", str(db)]) == 0
    finally:
        builtins.input = real_input
    con = sqlite3.connect(db)
    left = con.execute("SELECT COUNT(*) FROM metrics WHERE project='Pendle' AND metric='total_supply'").fetchone()[0]
    moved = sorted(con.execute("SELECT date, value FROM metrics WHERE project='Pendle' "
                               "AND metric='total_supply_coingecko'").fetchall())
    con.close()
    assert left == 0
    assert moved == before


def test_fdv_history_archive_reads_on_chain_first_totals_labelled_pre_handover():
    import config
    from fetch import archive as ar
    from fetch.base import _measuring_point
    lab = config.TOTAL_SUPPLY_CHAIN_BACKFILL_LABEL
    assert lab == "chain backfill (pre-handover)"
    assert config.total_supply_chain_backfill_label("Pendle", "total_supply", "2026-10-10") == lab
    assert config.total_supply_chain_backfill_label("Pendle", "total_supply", "2026-10-11") == ""
    assert config.total_supply_chain_backfill_label("Uniswap", "total_supply", "2026-10-10") == ""
    assert config.archive_not_by_design("Pendle", "total_supply") == ""
    assert config.archive_not_by_design("Uniswap", "total_supply")
    a = ar.ArchiveBackfill(store=None, projects=[config.PROJECT_BY_NAME["Pendle"], config.PROJECT_BY_NAME["Uniswap"]],
                           reader=object())
    ok, no = a.targets({"ethereum"})
    assert "total_supply" in ok.get("Pendle", set())
    assert "total_supply" not in ok.get("Uniswap", set())
    # the label is an annotation: one measuring point with the live chain read
    assert _measuring_point(f"chain:ethereum:token[{lab}]:archive") == _measuring_point("chain:ethereum:token")


# ----------------------------------------------------------------------------------------------- 7 the basis rule
def test_a_change_of_basis_is_a_series_break_never_a_trend_on_the_real_dump(tmp_path, monkeypatch, dumps_20261009):
    """Decision 7: Aerodrome's circulating moved to the on-chain set on 2026-10-07 (circulating_supply_used, and the
    non-circulating wallet set 3 -> 7 that day). Every row reading it says SERIES BREAK 2026-10-07 and carries no QoQ,
    growth or direction; Q0 still equals the headline cell. A row reading no changed series has no break."""
    from test_trajectory import _built
    _wb, rows = _built(tmp_path, monkeypatch, "Aerodrome")
    by = {r["key"]: r for r in rows}
    circ = by["A3|Aerodrome|circulating_supply"]
    assert circ["series_break"].startswith("SERIES BREAK 2026-10-07")
    assert circ["QoQ"] is None and circ["growth_4q"] is None and circ["improved"] is None
    assert circ["trend"].startswith("SERIES BREAK 2026-10-07")
    assert circ["verdict"].startswith("PASS")
    clean = [r for r in rows if not r["series_break"]]
    assert clean, "every row broken — the rule must reach only rows reading a changed series"
    assert all(r["verdict"].startswith("PASS") for r in rows)


def test_clean_quarters_and_the_declared_handover_rule():
    import trajectory
    asof = pd.Timestamp("2026-10-09")
    assert trajectory.clean_quarters(asof, "2026-10-07", 90, 5) == 0          # inside Q0
    assert trajectory.clean_quarters(asof, "2026-07-12", 90, 5) == 1          # Q0 = 07-12..10-09, wholly after
    assert trajectory.clean_quarters(asof, "2025-01-01", 90, 5) == 5
    import config
    g = pd.DataFrame({"date": pd.to_datetime(["2026-09-01", "2026-09-02"]), "value": [1.0, 1.0],
                      "source": ["coingecko:mcap/price", "coingecko"]})
    # a declared handover with no composition change is one series (Pendle's CoinGecko history -> live)
    assert config.declared_handover("Pendle", "circulating_supply")
    assert trajectory.basis_breaks("Pendle", ["circulating_supply"], {("Pendle", "circulating_supply"): g},
                                   pd.Timestamp("2026-01-01")) == []
    # Ether.fi's total steps at its handover: a break until decision 6's chain backfill restates it
    g2 = pd.DataFrame({"date": pd.to_datetime(["2026-10-10", "2026-10-11"]), "value": [965.35e6, 998.5e6],
                       "source": ["coingecko", "chain:ethereum:token"]})
    br = trajectory.basis_breaks("Ether.fi", ["total_supply"], {("Ether.fi", "total_supply"): g2},
                                 pd.Timestamp("2026-01-01"))
    assert [b[0] for b in br] == [pd.Timestamp("2026-10-11")]


# ----------------------------------------------------------------------------------------------- 2 Sky
# THE CHAIN READ OF 2026-10-12 (fetch/sky_mkr.read, live, Ethereum mainnet): MKR.totalSupply 83,516.32, the Pause Proxy's
# MKR 49.56, converter rate 24,000 and fee 5%; vest 17 (2026-09-13..2026-12-12) 143,208,393 SKY
SKY_READ = {"outstanding": 83516.31879149908, "issuer": 49.55875015574528, "rate": 24000, "fee": 0.05,
            "backing": 1903042128.9426281,
            "vest": {"id": 17, "bgn": 1789303331, "fin": 1797079331, "tot": 143208393.0, "annual": 580789593.8333334}}


def test_sky_mkr_backing_is_circulating_and_the_vest_rate_is_stored():
    """Decision 2a: MKR-backed SKY held by non-issuers is circulating, added back after the subtraction; the issuer's
    MKR stays out. The vest in force is stored as the annual rate (580.8M/yr on the 10-12 read)."""
    import config
    from fetch.base import FetchOutput
    from fetch.circulating import series
    from fetch.sky_mkr import SkyMkr
    assert config.CIRCULATING_ONCHAIN["Sky"]["add"] == ("sky_mkr_nonissuer_backing_tokens",)
    assert config.unlock_stock_metrics("Sky") == ()                       # decision 2b
    out = FetchOutput()
    SkyMkr(reader=lambda spec: SKY_READ).run([config.PROJECT_BY_NAME["Sky"]], None, out)
    f = out.frame().set_index("metric")
    assert f.loc["sky_mkr_nonissuer_backing_tokens", "value"] == pytest.approx(
        (SKY_READ["outstanding"] - SKY_READ["issuer"]) * 24_000 * 0.95, rel=1e-12)
    assert f.loc["emissions_rate_current_annual_tokens", "value"] == pytest.approx(580_789_594, rel=1e-3)
    assert "vest 17" in f.loc["emissions_rate_current_annual_tokens", "source"]
    d = pd.Timestamp("2026-10-12")
    h = pd.DataFrame([{"date": d, "project": "Sky", "metric": m, "value": v} for m, v in (
        ("total_supply", 23_459_804_203.61), ("treasury_holding_tokens", 100.0), ("noncirculating_holding_tokens", 50.0),
        ("sky_mkr_nonissuer_backing_tokens", f.loc["sky_mkr_nonissuer_backing_tokens", "value"]))])
    assert series(h, "Sky").iloc[-1] == pytest.approx(23_459_804_203.61 - 150 + 1_903_042_128.0, rel=1e-9)
    # a day without the backing read computes nothing (never a figure on a basis missing a leg)
    assert series(h[h["metric"] != "sky_mkr_nonissuer_backing_tokens"], "Sky").empty


# ----------------------------------------------------------------------------------------------- 8 lumpy flows
def test_lumpy_rows_take_direction_from_the_rolling_t12m_on_the_real_dump(tmp_path, monkeypatch):
    """Decision 8: Maple's fees are a declared lumpy series; its fee rows read direction and growth from the rolling
    T12M (Q0..Q-3 vs Q-1..Q-4), with no single-quarter improved count."""
    import trajectory
    from test_trajectory import _built
    _wb, rows = _built(tmp_path, monkeypatch, "Maple")
    lumpy = [r for r in rows if r["direction_basis"].startswith("rolling T12M")]
    assert lumpy
    for r in lumpy:
        assert r["improved"] is None
        if r["T12M_prev"] is not None and r["T12M"] is not None and r["T12M_prev"] > 0 and r["T12M"] > 0:
            assert r["growth_4q"] == pytest.approx((r["T12M"] / r["T12M_prev"]) ** 4 - 1)
            assert "rolling T12M" in r["trend"]
    fees = next(r for r in rows if r["key"] == "Master|Maple|fees_q0" or r["hid"].startswith("fees_q0"))
    qs = [fees[k] for k in trajectory.LABELS]
    assert fees["T12M"] == pytest.approx(sum(qs[:4])) and fees["T12M_prev"] == pytest.approx(sum(qs[1:5]))
    assert all(r["direction_basis"].startswith("4 quarters") for r in rows if r not in lumpy)


def test_lumpy_reason_names_the_discrete_flow():
    import trajectory
    assert "unlock" in trajectory.lumpy_reason("Chainlink", {"unlock_release_tokens", "price_usd"}, {})
    assert "LUMPY_FLOWS" in trajectory.lumpy_reason("Maple", {"fees_usd"}, {})
    assert trajectory.lumpy_reason("Chainlink", {"price_usd"}, {}) == ""
    t0, t1, chg, g, why = trajectory.rolling_t12m([4, 3, 2, 1, 1], "sum")
    assert (t0, t1, why) == (10, 7, "") and chg == pytest.approx(10 / 7 - 1) and g == pytest.approx((10 / 7) ** 4 - 1)
    assert trajectory.rolling_t12m([4, 3, 2, 1, None], "sum")[2] is None


# ----------------------------------------------------------------------------------------------- 9 Total Yield
def test_total_yield_is_validator_yield_less_net_inflation_on_the_real_dump(tmp_path, monkeypatch):
    """Decision 9: stakers = validator staking yield − net inflation; non-stakers = − net inflation, with net inflation
    Master's own 'Net supply change, annualised % of circulating'. The T-bill benchmark sits on the stakers' figure."""
    import config
    from test_final_round import _build, _cell
    wb, value = _build(tmp_path, monkeypatch, "Ethereum")
    vy = _cell(wb, value, "A1", "Ethereum", "VALIDATOR STAKING YIELD")
    ni = _cell(wb, value, "Master", "Ethereum", "Net supply change, annualised % of circulating")
    st = _cell(wb, value, "A1", "Ethereum", "TOTAL YIELD — stakers")
    nh = _cell(wb, value, "A1", "Ethereum", "Total yield — non-stakers")
    assert isinstance(vy, float)
    if isinstance(ni, float):
        assert st == pytest.approx(vy - ni) and nh == pytest.approx(-ni)
    else:
        # alignment round B5: on the 2026-10-09 dump the issuance leg covers 10 of 90 days, so the net is BLOCKED —
        # and Total Yield carries the same BLOCKED text, never a figure on half a window
        assert str(ni).startswith("BLOCKED") and str(st).startswith("BLOCKED") and str(nh).startswith("BLOCKED")
    b = next(b for b in config.TRAJECTORY["benchmarks"] if b["id"] == "yield_vs_tbill")
    assert b["on"] == ("a1_total_yield_stakers",)


# ----------------------------------------------------------------------------------------------- Phase 1 review, item 1
def _pendle_build(tmp_path, monkeypatch, previous):
    import openpyxl
    import xlcalc
    import store as sm
    import build_workbook as bw
    from real_dump import load, to_store
    payload = load("Pendle")
    monkeypatch.setitem(sys.modules, "recalc", None)
    monkeypatch.setattr(bw, "_SCOPE", bw._SCOPE)
    db = tmp_path / "metrics.db"
    to_store(payload, db)
    monkeypatch.setattr(sm, "DB_PATH", str(db))
    st = sm.Store(str(db))
    try:
        bw.build_workbook(st, tmp_path / "wb.xlsx", asof=pd.Timestamp(payload["meta"]["asof"]), only=["Pendle"],
                          previous=previous)
    finally:
        st.close()
    wb = openpyxl.load_workbook(tmp_path / "wb.xlsx")
    X = xlcalc.Workbook(wb)
    return wb, (lambda s, r, c: X.value(s, r, c)), bw


def test_a_headline_lost_between_runs_is_a_check_and_a_sanity_hit_on_the_real_dump(tmp_path, monkeypatch):
    """Item 1 of Jake's Phase 1 review: a headline that read a value in the previous run's snapshot and reads n/a now
    is a CHECK on the Credibility tab (the project's continuity row) and a sanity hit naming it. The previous snapshot
    here is the real dump's own build, with one currently-empty cell recorded as having had a value."""
    import sanity
    (tmp_path / "a").mkdir()
    wb, value, bw = _pendle_build(tmp_path / "a", monkeypatch, previous=None)
    cells = [h for h in bw._HEADLINE_CELLS if h["project"] == "Pendle"]
    now = {}
    for hc in cells:
        c = wb[hc["sheet"]][hc["cell"]]
        v = value(hc["sheet"], c.row, c.column)
        now[f"{hc['sheet'][:2]} | {hc['header']}"] = float(v) if isinstance(v, (int, float)) else None
    held = sum(isinstance(v, float) for v in now.values())
    lost = next(k for k, v in now.items() if v is None)
    assert held >= 5
    # unchanged since the previous run: PASS, no hit
    (tmp_path / "b").mkdir()
    wb, value, bw = _pendle_build(tmp_path / "b", monkeypatch, previous=("2026-10-08", {"Pendle": dict(now)}))
    ws = wb["Credibility"]
    r = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 4).value or "").startswith("CONTINUITY"))
    assert value("Credibility", r, 11) == "PASS" and value("Credibility", r, 7) == held
    assert not [h for h in bw.SANITY_HITS if h[1] == "headline_lost"]
    # one headline read a value last run and none now: CHECK, and the sanity hit names it
    (tmp_path / "c").mkdir()
    prev = dict(now)
    prev[lost] = 0.0123
    wb, value, bw = _pendle_build(tmp_path / "c", monkeypatch, previous=("2026-10-08", {"Pendle": prev}))
    ws = wb["Credibility"]
    r = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 4).value or "").startswith("CONTINUITY"))
    assert str(value("Credibility", r, 11)).startswith("CHECK (HEADLINE LOST")
    assert value("Credibility", r, 7) - value("Credibility", r, 5) == 1
    hits = [h for h in bw.SANITY_HITS if h[1] == "headline_lost"]
    assert len(hits) == 1 and lost.split(" | ", 1)[1][:40] in hits[0][2] and "0.0123" in hits[0][2]
    assert sanity.lost_headlines(wb, {"Pendle": prev}, bw._HEADLINE_CELLS, "2026-10-08") == hits


def test_the_archive_never_interleaves_a_declared_handover():
    """Item 1: GEODNET's burn stock is a declared handover (Polygon only, then Polygon + Solana). A current-point
    archive reading dated before a stored Polygon-only row is refused — the build would see the legs overlap and
    withhold the series; one after the last Polygon-only row is written. The flow's Dune lead leg is resolved by month
    at read time and never blocks."""
    import config
    from fetch.archive import breaks_handover_order
    pts = config.declared_handover("GEODNET", "burn_address_balance")["ordered_points"]
    g = pd.DataFrame({"date": pd.to_datetime(["2026-09-22", "2026-09-27", "2026-09-28", "2026-10-01"]),
                      "source": [pts[0], pts[0], pts[1], pts[1]]})
    assert "interleave" in breaks_handover_order("GEODNET", "burn_address_balance", g, "2026-09-10",
                                                 pts[1] + ":archive")
    assert breaks_handover_order("GEODNET", "burn_address_balance", g, "2026-10-05", pts[1] + ":archive") == ""
    fp = config.declared_handover("GEODNET", "gross_burn_tokens")["ordered_points"]
    f = pd.DataFrame({"date": pd.to_datetime(["2025-08-01", "2026-09-28"]), "source": [fp[0], fp[2] + ":delta"]})
    assert breaks_handover_order("GEODNET", "gross_burn_tokens", f, "2025-09-15", fp[2] + ":archive:delta") == ""
    f2 = pd.concat([f, pd.DataFrame({"date": pd.to_datetime(["2026-09-25"]), "source": [fp[1] + ":delta"]})])
    assert "interleave" in breaks_handover_order("GEODNET", "gross_burn_tokens", f2, "2025-09-15",
                                                 fp[2] + ":archive:delta")


# ----------------------------------------------------------------------------------------------- Phase 1 review, item 8a
def test_a_memo_of_a_shorter_ask_never_reads_as_nothing_older(monkeypatch, tmp_path):
    """Item 8a: Aethir, GEODNET and Plume fees / revenue stayed at 2025-09-29 — 365 days before the year's backfill —
    after the window grew to 455 days: the memo of the 365-day ask ("first 2025-09-29") read as "the source starts
    there". A memo now carries the window it asked; one without it (the 365-day era) or shorter is re-asked."""
    import json
    import config
    from fetch import backfill as bf
    from fetch.base import BACKFILL_DAYS, today
    monkeypatch.setenv("TOKEN_METRICS_LOGCACHE", str(tmp_path))
    first = {("GEODNET", "fees_usd"): "2025-09-29"}
    geod = config.PROJECT_BY_NAME["GEODNET"]
    bf._depth_file().write_text(json.dumps({"GEODNET|fees_usd": {"first": "2025-09-29",
                                                                 "checked_on": str(today().date())}}))
    assert ("GEODNET", "fees_usd") in bf.plan([geod], first)[0]
    bf.record({("GEODNET", "fees_usd")}, first)
    assert json.loads(bf._depth_file().read_text())["GEODNET|fees_usd"]["asked_days"] == BACKFILL_DAYS
    assert ("GEODNET", "fees_usd") not in bf.plan([geod], first)[0]


# ----------------------------------------------------------------------------------------------- Phase 1 review, item 8b
def test_q4_prices_come_from_defillama_before_coingeckos_year_on_the_real_dump(tmp_path, monkeypatch):
    """Item 8b: CoinGecko's free chart serves 365 days; the fifth quarter's prices come from DefiLlama's coins chart
    (fixtures/llama_chart_link_2026-10-12.json, read live 2026-10-12) for the days before CoinGecko's first price — a
    declared older leg, so the series is one (no measuring-point block) and Q-4 has a price."""
    import json
    from fetch.xref import CrossRefs as XRef
    import build_workbook as bw
    from test_final_round import _build
    body = json.loads((ROOT / "fixtures" / "llama_chart_link_2026-10-12.json").read_text())["body"]
    key = "ethereum:0x514910771af9ca656af840dff83e8264ecf986ca"
    rows = XRef.parse_chart(body["coins"][key]["prices"])
    assert len(rows) >= 450
    extra = [{"date": str(d.date()), "metric": "price_usd_llama", "value": v, "source": f"defillama_coins:{key}",
              "tier": 1, "fetched_at": "2026-10-12T12:00:00Z"} for d, v in rows if d <= pd.Timestamp("2026-10-09")]
    wb, value = _build(tmp_path, monkeypatch, "Chainlink", extra_rows=extra)
    g = bw._LAST_GROUPS[("Chainlink", "price_usd")]
    first_cg = pd.Timestamp("2025-09-12")
    older = g[pd.to_datetime(g["date"]) < first_cg]
    assert len(older) >= 55 and older["source"].str.startswith("defillama_coins:history[").all()
    assert not (g[pd.to_datetime(g["date"]) >= first_cg]["source"] != "coingecko").any()
    ws = wb["Data"]
    r = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "Chainlink|price_usd")
    assert ws.cell(r, bw.DATA_COLS.index("status") + 1).value == "ok"
    assert isinstance(ws.cell(r, bw.DATA_COLS.index("q4") + 1).value, float)


# ----------------------------------------------------------------------------------------------- Phase 1 review, item 8c
def test_eth_validator_yield_has_history_from_the_spec_formula_on_the_real_dump(tmp_path, monkeypatch, dumps_20261009):
    """Item 8c: consensus rewards before Etherscan's first Eth2Staking read = the spec's 166.32 x sqrt(staked ETH) / 365
    on that day's staked ETH, scaled to the measured series over their overlap (k on every row). The validator yield
    then reads every quarter; Q0 is unchanged in kind (the measured days stay measured)."""
    import build_workbook as bw
    from test_trajectory import _built
    _wb, rows = _built(tmp_path, monkeypatch, "Ethereum")
    g = bw._LAST_GROUPS[("Ethereum", "consensus_rewards_tokens")]
    leg = g[g["source"].astype(str).str.startswith("derived:consensus_formula[")]
    assert len(leg) > 250 and "k=" in leg["source"].iloc[0]
    # G11 B6 (supersedes follow-up item 2's "never inside Q0", which left a hole between the leg and the first measured
    # read): the leg runs to the day before the measured series, and its days inside Q0 say so on the row
    meas = g[~g["source"].astype(str).str.startswith("derived:consensus_formula[")]
    assert pd.to_datetime(leg["date"]).max() == pd.to_datetime(meas["date"]).min() - pd.Timedelta(days=1)
    inside = leg[pd.to_datetime(leg["date"]) > pd.Timestamp("2026-10-09") - pd.Timedelta(days=90)]
    assert len(inside) and inside["source"].str.contains("INSIDE THE HEADLINE WINDOW").all()
    k = float(leg["source"].iloc[0].split("k=")[1].split(";")[0])
    assert 0.9 < k < 1.05                                  # participation and penalties: a little under the maximum
    vy = next(r for r in rows if r["hid"] == "a1_validator_yield")
    assert all(isinstance(vy[q], float) for q in ("Q0", "Q-1", "Q-2", "Q-3"))
    assert 0.02 < vy["Q-3"] < 0.05 and vy["verdict"].startswith("PASS")


def test_two_empty_trends_are_not_a_cross_tab_mismatch():
    """Item 9: PancakeSwap (parked) read as a trend mismatch with n/a on both sides."""
    import cross_tab
    import xlcalc
    assert cross_tab._trend_na("n/a — no trajectory row") and cross_tab._trend_na("– · 4Q n/a · QoQ n/a")
    assert cross_tab._trend_na(None) and cross_tab._trend_na(xlcalc.XLError("#N/A"))
    assert not cross_tab._trend_na("▲ · +12%/yr (4Q) · 3/4 improved · QoQ +2.0%")
    assert not cross_tab._trend_na("– · NEW since Q-2 · QoQ +5.0%")


# ----------------------------------------------------------------------------------------------- Phase 1 review, item 2
AERO_WEEKLY = ROOT / "fixtures" / "aero_onchain_weekly_2026-10-12.json"


def aerodrome_restated(payload: dict) -> dict:
    """The real Aerodrome dump as Jake's store reads after the restatement (RUNBOOK 11bq): SQL CJ (CoinGecko's
    pre-handover totals -> total_supply_coingecko) and CK (the 3-wallet rows go), then the archive reads at past
    blocks — here the REAL reads captured 2026-10-12 (fixtures/aero_onchain_weekly_2026-10-12.json: AERO totalSupply,
    the seven non-circulating wallets and the filing wallets' managed-veNFT deposits at the first block of every 7th
    day, public Base archive eth_call), with the sources the archive and the seed write."""
    import json
    import config
    reads = json.loads(AERO_WEEKLY.read_text())["reads"]
    keys = [k for k, v in config.PROJECT_BY_NAME["Aerodrome"]["contracts"].items()
            if v.get("metric_override") == "noncirculating_holding_tokens"]
    comp = "chain:sum(" + "+".join(f"base:{k}" for k in keys) + ")"
    rows = []
    for r in payload["metrics"]:
        if r["metric"] == "total_supply" and r["source"].startswith("coingecko") and r["date"] < "2026-10-11":
            r = {**r, "metric": "total_supply_coingecko"}                             # CJ
        if r["metric"] == "noncirculating_holding_tokens" and "noncirc_flight_school" not in r["source"]:
            continue                                                                  # CK
        rows.append(r)
    have = {(r["date"][:10], r["metric"]) for r in rows}
    for x in reads:
        for metric, value, src in (
                ("total_supply", x["total_supply"], f"chain:base:token[{config.TOTAL_SUPPLY_CHAIN_BACKFILL_LABEL}]:archive"),
                ("noncirculating_holding_tokens", sum(x["noncirc"][k] for k in keys), comp + ":archive"),
                ("filing_managed_lock_tokens", x["managed"], "ve_managed:0xeBf418Fe.weights:archive")):
            if (x["date"], metric) not in have:
                rows.append({"date": x["date"], "metric": metric, "value": value, "source": src, "tier": 2,
                             "fetched_at": "2026-10-12T14:00:00Z"})
    return {**payload, "metrics": rows}


def test_aerodrome_circulating_restated_on_the_onchain_basis_has_no_jump_on_the_real_dump(tmp_path, monkeypatch, dumps_20261009):
    """Item 2 of Jake's Phase 1 review: the workbook read CoinGecko's 1,004.5M (it excludes every veAERO lock) before
    2026-10-07 and the on-chain set's 1,516.8M after — a fake +51% at 10-07. CoinGecko never stands in now, and with
    the history restated on the on-chain basis (real archive reads) every quarter is on that basis: no SERIES BREAK,
    Q-1..Q-4 are the on-chain figures, and circulating rises smoothly into 10-07."""
    import config
    import build_workbook as bw
    from real_dump import load
    from test_final_round import _build
    if not AERO_WEEKLY.exists():
        pytest.skip("the captured reads are not on file")
    assert config.no_coingecko_fallback("Aerodrome")
    # without the restatement: no CoinGecko stand-in, so the days before the on-chain set are empty, never 1.0bn
    (tmp_path / "a").mkdir()
    _build(tmp_path / "a", monkeypatch, "Aerodrome", asof="2026-10-12")
    used = bw._LAST_GROUPS[("Aerodrome", "circulating_supply_used")]
    assert used["value"].min() > 1.4e9 and pd.to_datetime(used["date"]).min() >= pd.Timestamp("2026-10-07")
    # restated
    payload = aerodrome_restated(load("Aerodrome"))
    (tmp_path / "b").mkdir()
    import real_dump
    monkeypatch.setattr(real_dump, "load", lambda project: payload)
    _build(tmp_path / "b", monkeypatch, "Aerodrome", asof="2026-10-12")
    used = bw._LAST_GROUPS[("Aerodrome", "circulating_supply_used")].sort_values("date")
    assert (used["source"] == "derived:circulating_used[circulating_supply_onchain]").all()
    s = used.set_index(pd.to_datetime(used["date"]))["value"]
    assert 1.2e9 < s.iloc[0] < 1.3e9 and s.is_monotonic_increasing or s.diff().min() > -0.02 * s.max()
    before, after = s[s.index < "2026-10-07"].iloc[-1], s[s.index >= "2026-10-07"].iloc[0]
    assert abs(after / before - 1) < 0.05, (before, after)          # no fake +51% at 10-07
    rows = {r["hid"]: r for r in bw.TRAJECTORY_ROWS if r["project"] == "Aerodrome"}
    circ = rows["circulating_supply"]
    assert not circ["series_break"], circ["series_break"]
    assert all(isinstance(circ[q], float) and circ[q] > 1.2e9 for q in ("Q0", "Q-1", "Q-2", "Q-3"))


# ----------------------------------------------------------------------------------------------- follow-up 2 MODELLED
def test_the_spec_formula_leg_is_labelled_modelled_everywhere_on_the_real_dump(tmp_path, monkeypatch):
    """Jake 2026-10-12 (Phase 1 follow-up, item 2): the ETH consensus-rewards leg from the spec formula is MODELLED on the
    Data row's note, the Trajectory row (its quarters listed, excluded from any scoring; the snapshot carries them) and
    the Credibility note of the headline it sits under."""
    import config
    import build_workbook as bw
    import export_snapshot as es
    from test_trajectory import _built
    wb, rows = _built(tmp_path, monkeypatch, "Ethereum")
    assert config.is_modelled("derived:consensus_formula[MODELLED — x]") and not config.is_modelled(
        "etherscan:ethsupply2.Eth2Staking:delta")
    g = bw._LAST_GROUPS[("Ethereum", "consensus_rewards_tokens")]
    assert g[g["source"].str.startswith("derived:consensus_formula")]["source"].str.contains("MODELLED").all()
    ws = wb["Data"]
    r = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "Ethereum|consensus_rewards_tokens")
    assert "MODELLED leg" in str(ws.cell(r, bw.DATA_COLS.index("note") + 1).value)
    vy = next(x for x in rows if x["hid"] == "a1_validator_yield")
    # G11 B6: the leg now reaches the first measured read, so Q0 is MODELLED in part (said, excluded from scoring)
    assert vy["modelled_quarters"][:2] == ["Q0", "Q-1"]
    assert "MODELLED" in vy["trend"]
    assert not next(x for x in rows if x["hid"] == "master_price_spot")["modelled_quarters"]
    cw = wb["Credibility"]
    notes = [str(cw.cell(i, 12).value or "") for i in range(1, cw.max_row + 1)
             if cw.cell(i, 1).value == "Ethereum" and "VALIDATOR STAKING YIELD" in str(cw.cell(i, 4).value or "")]
    assert notes and notes[0].startswith("MODELLED history under this headline")
    assert "MODELLED quarters (excluded from any scoring)" in es.TRAJ_FIELDS


def test_near_stake_history_is_read_by_epoch_and_tolerates_the_epoch_age_limit():
    """Follow-up item 3: the archival RPC answers `validators` by epoch_id; a day whose epoch is past the node's
    retention (UNKNOWN_EPOCH) still gets its supply row, never a stake row."""
    import config
    from fetch import archive as ar
    assert config.HISTORY_FORWARD_ONLY["Near"]["metrics"] == ()

    class Fake(ar.NearArchive):
        def __init__(self, unknown):
            self.unknown = unknown

        def at(self, day):
            return {"total_supply": str(1_300_000_000 * 10**24), "epoch_id": "EPOCHabcdef123"}

        def rpc(self, method, params):
            assert method == "validators" and params == {"epoch_id": "EPOCHabcdef123"}
            if self.unknown:
                raise RuntimeError("archival-rpc.mainnet.fastnear.com: {'cause': {'name': 'UNKNOWN_EPOCH'}}")
            return {"epoch_start_height": 1, "current_validators": [{"stake": str(400 * 10**30)},
                                                                   {"stake": str(158 * 10**30)}]}
    p = config.PROJECT_BY_NAME["Near"]
    rows = Fake(False).rows(p, pd.Timestamp("2026-01-14"))
    st = next(r for r in rows if r["metric"] == "locked_tokens")
    assert st["value"] == 558_000_000 and st["source"].startswith("near_rpc:validators[epoch EPOCHabc")
    from fetch.base import _measuring_point
    assert _measuring_point(st["source"]) == "near_rpc:validators"
    rows = Fake(True).rows(p, pd.Timestamp("2025-07-14"))
    assert [r["metric"] for r in rows] == ["total_supply_protocol"]
