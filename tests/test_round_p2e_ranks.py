"""ROUND P2e (Jake 2026-10-28), reader items 1 and 5 — no definition changes.

P2e-1 RANKS AGREE: an archetype tab ranks ONLY the projects whose primary archetype it is, by the ranking Summary shows;
a secondary member keeps its key and reads "secondary" in the rank column, its comment naming its primary archetype.
P2e-5 ONE CODE PER PARKED ROW: every rank and key cell of a parked row that holds no number reads "parked", the
underlying state kept in its comment ("Behind it: …").

Built once on the committed real dumps (fixtures/real_*.json), as test_round_p1 does.
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import config  # noqa: E402
import reader  # noqa: E402


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    """The reader built on the committed real dumps (the build test_round_p1 makes)."""
    import build_workbook as bw
    import measuring_guard as mg
    import pandas as pd
    import store as sm
    from real_dump import to_store
    d = tmp_path_factory.mktemp("p2e")
    db = d / "metrics.db"
    asofs = []
    for pl in mg._payloads():
        to_store(pl, db)
        asofs.append(pl["meta"]["asof"])
    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(sm, "DB_PATH", str(db))
        mp.setitem(sys.modules, "recalc", None)
        mp.setenv("TOKEN_METRICS_STORE_VALUES", "1")
        st = sm.Store(str(db))
        try:
            bw.build_workbook(st, d / "audit.xlsx", asof=pd.Timestamp(max(asofs)), reader_path=d / "reader.xlsx")
        finally:
            st.close()
    assert "error" not in bw.READER_STATS, bw.READER_STATS
    assert bw.READER_STATS.get("blank_unexplained") == 0, bw.READER_STATS.get("blank_cells")
    return {"reader": d / "reader.xlsx"}


TABS = {1: reader.A1_TAB, 2: reader.A2_TAB, 3: reader.A3_TAB, 4: reader.A4_TAB}


def _col(ws, header: str) -> int:
    return next(c for c in range(1, ws.max_column + 1) if ws.cell(4, c).value == header)


def _rank_col(ws) -> int:
    return next(c for c in range(1, ws.max_column + 1) if str(ws.cell(4, c).value or "").startswith("Rank in"))


def _rows(ws) -> dict:
    return {ws.cell(r, 1).value: r for r in range(5, ws.max_row + 1)
            if ws.cell(r, 1).value and ws.cell(r, 2).value is not None}


def _arch_of_summary(ws) -> dict:
    """{project: primary archetype number} from Summary's section rows."""
    out = {}
    for label, lo, hi in reader._summary_sections(ws):
        a = int(label[1])
        for r in range(lo, hi + 1):
            out[ws.cell(r, 1).value] = a
    return out


def test_p2e_1_every_rank_is_the_same_on_summary_and_on_its_primary_tab(built):
    import xlcalc
    wb = openpyxl.load_workbook(built["reader"])
    X = xlcalc.Workbook(wb)
    s = wb["Summary"]
    srows, sr = _rows(s), _col(s, "Rank in archetype")
    prim = _arch_of_summary(s)
    checked = 0
    for name, arch in prim.items():
        ws = wb[TABS[arch]]
        rows = _rows(ws)
        if name not in rows:
            continue
        want = X.value("Summary", srows[name], sr)
        got = X.value(ws.title, rows[name], _rank_col(ws))
        assert got == want, (name, ws.title, got, want)
        checked += 1
    assert checked >= 20, checked
    # each tab's ranks are 1..n over its primaries, in order, and no rank is held by a non-primary
    for arch, tab in TABS.items():
        ws = wb[tab]
        rc = _rank_col(ws)
        nums = [X.value(tab, r, rc) for n, r in sorted(_rows(ws).items(), key=lambda kv: kv[1])]
        ranks = [v for v in nums if isinstance(v, (int, float))]
        assert ranks == [float(i) for i in range(1, len(ranks) + 1)], (tab, ranks)
        for n, r in _rows(ws).items():
            if isinstance(X.value(tab, r, rc), (int, float)):
                assert prim[n] == arch, (tab, n)


def test_p2e_1_a_secondary_member_reads_secondary_and_names_its_primary(built):
    wb = openpyxl.load_workbook(built["reader"])
    prim = _arch_of_summary(wb["Summary"])
    seen = 0
    for arch, tab in TABS.items():
        ws = wb[tab]
        rc = _rank_col(ws)
        kinds = []
        for n, r in sorted(_rows(ws).items(), key=lambda kv: kv[1]):
            v = ws.cell(r, rc).value
            if reader.parked_reason(n):
                assert v == "parked", (tab, n, v)
                kinds.append(2)
            elif prim[n] != arch:
                assert v == "secondary", (tab, n, v)
                c = ws.cell(r, rc)
                assert c.comment.text.startswith(f"Primary archetype: {reader.ARCHETYPE_NAMES[prim[n]]} — ranked "
                                                 "there; shown here for comparison, not ranked."), (tab, n)
                assert c.font.color.rgb.endswith(reader.CODES["secondary"])
                # the key value is kept: the key cell is the audit number or the key's own code, never "secondary"
                assert ws.cell(r, rc - 1).value != "secondary"
                kinds.append(1)
                seen += 1
            else:
                kinds.append(0)
        assert kinds == sorted(kinds), (tab, kinds)       # primaries, then secondaries, then parked
    assert seen > 0
    # Chainlink (primary A1) is a secondary member of A3 on the fixture build
    a3 = wb[reader.A3_TAB]
    assert a3.cell(_rows(a3)["Chainlink"], _rank_col(a3)).value == "secondary"
    # About and Definitions list the code
    about = " ".join(str(c.value or "") for c in wb["About"]["A"])
    assert "'secondary'" in about
    assert "secondary" in [c.value for c in wb["Definitions"]["A"]]


def test_p2e_5_a_parked_row_reads_parked_in_every_rank_and_key_cell(built):
    wb = openpyxl.load_workbook(built["reader"])
    parked = [n for n in _rows(wb["Summary"]) if reader.parked_reason(n)]
    assert parked
    checked = 0
    for tab in ("Summary",) + tuple(TABS.values()):
        ws = wb[tab]
        rows = _rows(ws)
        rc = _rank_col(ws)
        # Summary: the key and the rank; a tab: its whole (a) block (key, rank and the figures beside the key)
        block = [rc - 1, rc] if tab == "Summary" else [c for c in range(3, ws.max_column + 1)
                                                        if _block_of(ws, c) == "(a) Rank key"]
        assert len(block) >= 2, (tab, block)
        for n in parked:
            if n not in rows:
                continue
            for c in block:
                cell = ws.cell(rows[n], c)
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    continue                              # a number: parked figures are shown, not ranked
                assert v == "parked", (tab, n, ws.cell(4, c).value, v)
                assert cell.comment.text.startswith(reader.parked_reason(n)), (tab, n, ws.cell(4, c).value)
                assert "Behind it: " in cell.comment.text, (tab, n, ws.cell(4, c).value)
                checked += 1
    assert checked >= 4, checked


def _block_of(ws, c: int) -> str:
    """The block label over column c (row 3 holds it on the first column of each block)."""
    for j in range(c, 0, -1):
        v = ws.cell(3, j).value
        if v:
            return str(v)
    return ""


def test_p2e_5_a_parked_code_keeps_its_underlying_state(monkeypatch):
    wb = openpyxl.Workbook()
    ws = wb.active
    sh = reader.Sheet(ws)
    col = reader.Col("Retired ÷ emissions (x)", "RETIRED ÷ EMISSIONS (x).", reader.FMT_X)
    sh.cols += [reader.Col("Project", "The project."), col]
    reader.write_headers(sh)
    ws.cell(5, 1, "Aerodrome")
    ref = reader.put_code(sh, 5, col, "blocked", "unequal coverage: buyback covers 40 days, emissions 90.")
    ws.cell(6, 1, "Aerodrome")
    num = reader.put_formula(sh, 6, col, "=1.5")
    monkeypatch.setattr(reader, "parked_reason", lambda n: config.AERODROME_PARK_PLAIN if n == "Aerodrome" else "")
    reader.park_cell(sh, ref, "Aerodrome")
    reader.park_cell(sh, num, "Aerodrome")
    assert ws[ref].value == "parked" and ws[ref].font.color.rgb.endswith(reader.CODES["parked"])
    text = ws[ref].comment.text
    assert text.startswith(config.AERODROME_PARK_PLAIN)
    assert "Behind it: blocked — unequal coverage: buyback covers 40 days" in text
    assert ws[num].value == "=1.5"                         # a number stays: shown, not ranked
    reader.park_cell(sh, ref, "Aerodrome")                 # idempotent
    assert ws[ref].comment.text == text
    reader.park_cell(sh, ref, "Sky")                       # not parked: untouched
    assert ws[ref].value == "parked"


def test_p2e_1_the_ranking_is_one_helper(monkeypatch):
    monkeypatch.setattr(reader, "parked_reason", lambda n: "parked" if n == "P" else "")
    keys = {"A": (3, 0.5, None, ""), "B": (3, 2.0, None, ""), "C": (1, 9.0, None, ""), "D": (3, None, None, "x"),
            "P": (3, 5.0, None, "")}
    assert reader._arch_ranking(keys, 3) == [("B", 1, 2.0), ("A", 2, 0.5), ("D", None, None)]
    assert reader._arch_ranking(keys, 1) == [("C", 1, 9.0)]
    note = reader._secondary_note("C", keys)
    assert note.startswith("Primary archetype: A1 Infrastructure — ranked there; shown here for comparison, not "
                           "ranked.") and "1." in note
    assert "secondary" in reader.CODES


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
