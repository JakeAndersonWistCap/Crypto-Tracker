"""
tests/test_review_ef6dc86.py — Jake's review of the ef6dc86 run and his Phase 2 decisions (2026-10-12 evening).

Judged on the real dumps (fixtures/real_*.json) wherever one holds the shape. GEODNET and Sky have no dump in the repo
(the GEODNET dump named in the review was not attached), so their fixes are pinned on the exact shapes the run's log
shows, built here, and say so.
"""
from __future__ import annotations

import builtins
import sqlite3
import sys
from pathlib import Path

import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import config  # noqa: E402


def _rows(project, metric, dates, value, source, tier=2):
    return [{"date": pd.Timestamp(d), "project": project, "metric": metric, "value": float(v), "source": source,
             "tier": tier, "fetched_at": "2026-10-12T00:00:00", "is_manual": 0, "entered_on": None,
             "source_note": None} for d, v in zip(dates, value if isinstance(value, list) else [value] * len(dates))]


def _run_section(db, section):
    import run_sql
    real = builtins.input
    builtins.input = lambda *a: f"DELETE {section}"
    try:
        assert run_sql.main(["--delete", section, "--db", str(db)]) == 0
    finally:
        builtins.input = real


# ------------------------------------------------------------------------------------------- 7/12 GEODNET A4 burn
def test_geodnet_a_restated_current_leg_cannot_overlap_the_monthly_dune_leg(monkeypatch):
    """Items 7 / 12, the first hypothesis RULED OUT: the archive restates past days on the CURRENT leg (Polygon +
    Solana), but _monthly_leg_views moves every Dune month dated on or after the daily leg's first day to the cross-check
    metric, so the declared legs can never overlap — even with a missed day leaving a month not whole. The cause of
    GEODNET's n/a is elsewhere; sanity's headline_starved names it (the Data row's status and note) on the next run.
    Shape from the run's log (the GEODNET dump named in the review was not attached)."""
    import build_workbook as bw
    monkeypatch.setattr(bw, "_SCOPE", [config.PROJECT_BY_NAME["GEODNET"]])
    asof = pd.Timestamp("2026-10-09")
    summed = "chain:sum(polygon:burn_polygon+solana:burn_solana_token_account):archive:delta"
    months = pd.date_range("2025-01-01", "2026-08-01", freq="MS")
    daily = pd.date_range("2025-07-14", "2026-10-08").drop(pd.Timestamp("2026-03-15"))
    long = pd.DataFrame(_rows("GEODNET", "gross_burn_tokens", months, 1_500_000.0, "dune:8683175", 4)
                        + _rows("GEODNET", "gross_burn_tokens", daily, 50_000.0, summed))
    data = bw.aggregate(long, pd.DataFrame(), asof)
    row = data[(data.project == "GEODNET") & (data.metric == "gross_burn_tokens")].iloc[0]
    assert row["status"] != "measuring_point_changed", row["note"]
    k = row["q0"] / 50_000.0                                      # the daily leg alone: no Dune month inside Q0
    assert k == pytest.approx(round(k)) and 89 <= round(k) <= 90, row["q0"]
    g = bw._LAST_GROUPS[("GEODNET", "gross_burn_tokens")]
    assert pd.to_datetime(g[g["source"] == "dune:8683175"]["date"]).max() < daily.min()


def test_a_headline_that_is_na_while_its_input_is_full_is_a_sanity_hit():
    """Item 7's guard: the continuity guard compares with the previous snapshot, so a headline that was already n/a
    there passed. headline_starved names it with the input row's status and note."""
    import openpyxl
    import sanity
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "A4 Permanent Burn"
    ws["B5"] = "n/a"
    ws["C5"] = 12.0
    hcs = [{"sheet": "A4 Permanent Burn", "cell": "B5", "project": "GEODNET", "header": "GROSS BURN Q0 (tokens)",
            "input_metric": "gross_burn_tokens"},
           {"sheet": "A4 Permanent Burn", "cell": "C5", "project": "GEODNET", "header": "a figure",
            "input_metric": "gross_burn_tokens"}]
    data = {"GEODNET|gross_burn_tokens": {"first_date": "2025-07-14", "latest_date": "2026-10-08",
                                          "status": "measuring_point_changed", "note": "NOT REPORTED — overlap"}}
    hits = sanity.starved_headlines(wb, hcs, data, "2026-10-09")
    assert [h[1] for h in hits] == ["headline_starved"]
    assert "measuring_point_changed" in hits[0][2] and "2025-07-14..2026-10-08" in hits[0][2]
    # an input that does not cover Q0 is not "full": no hit
    data["GEODNET|gross_burn_tokens"]["first_date"] = "2026-09-28"
    assert not sanity.starved_headlines(wb, hcs, data, "2026-10-09")


# ------------------------------------------------------------------------------------------- 13 restatement
def test_cm_moves_every_coingecko_total_and_cn_restates_fluid_and_maple_on_the_real_dumps(tmp_path):
    from real_dump import load, to_store
    pen = load("Pendle")
    db = tmp_path / "metrics.db"
    to_store(pen, db)
    con = sqlite3.connect(db)
    # the 2026-10-11 row CJ left (dated ON the handover day): CM moves it too
    con.execute("INSERT INTO metrics (date, project, metric, value, source, tier, fetched_at) VALUES "
                "('2026-10-11', 'Pendle', 'total_supply', 281527448.0, 'coingecko', 1, 'x')")
    con.commit()
    con.close()
    _run_section(db, "CM")
    con = sqlite3.connect(db)
    assert con.execute("SELECT COUNT(*) FROM metrics WHERE project='Pendle' AND metric='total_supply' "
                       "AND source LIKE 'coingecko%'").fetchone()[0] == 0
    assert con.execute("SELECT COUNT(*) FROM metrics WHERE project='Pendle' AND metric='total_supply_coingecko' "
                       "AND date='2026-10-11'").fetchone()[0] == 1
    con.close()
    mp = load("Maple")
    db2 = tmp_path / "maple.db"
    to_store(mp, db2)
    held = [r for r in mp["metrics"] if r["metric"] == "noncirculating_holding_tokens"]
    assert held and not any("noncirc_ssf_safe" in r["source"] for r in held)
    _run_section(db2, "CN")
    con = sqlite3.connect(db2)
    assert con.execute("SELECT COUNT(*) FROM metrics WHERE project='Maple' "
                       "AND metric='noncirculating_holding_tokens'").fetchone()[0] == 0
    con.close()


def test_the_archive_says_how_many_runs_are_left_and_continues_where_it_stopped():
    """13c: the budget line names the days still to read, the measured pace and the runs it takes; stored days are
    skipped, so a re-run continues by itself. The Solana note tells 'cached' from 'nothing read'."""
    src = (Path(__file__).resolve().parents[1] / "fetch" / "archive.py").read_text(encoding="utf-8")
    assert "more run(s) at this budget" in src and "--budget-min" in src
    assert "already cached (complete)" in src


# ------------------------------------------------------------------------------------------- 9 Sky circulating
def test_sky_onchain_circulating_takes_its_added_back_leg(monkeypatch):
    """Item 9: _circulating_views handed fetch/circulating.series the total and the subtracted legs but not `add`, so
    Sky's MKR-backed SKY was an empty series and every day dropped (circulating_used fell back to CoinGecko's
    23,433.6M). Shape from the run's log (no Sky dump in the repo)."""
    import build_workbook as bw
    monkeypatch.setattr(bw, "_SCOPE", [config.PROJECT_BY_NAME["Sky"]])
    spec = config.circulating_onchain("Sky")
    assert spec["add"] == ["sky_mkr_nonissuer_backing_tokens"] or tuple(spec["add"]) == ("sky_mkr_nonissuer_backing_tokens",)
    days = pd.date_range("2026-10-05", "2026-10-11")
    groups = {}
    for m, v in (("total_supply", 23_459_804_204.0), ("treasury_holding_tokens", 27_996_757.0),
                 ("noncirculating_holding_tokens", 2_043_352_267.0),
                 ("sky_mkr_nonissuer_backing_tokens", 1_903_000_000.0)):
        groups[("Sky", m)] = pd.DataFrame(_rows("Sky", m, days, v, f"chain:ethereum:{m}"))
    bw._circulating_views(groups)
    g = groups[("Sky", "circulating_supply_onchain")]
    assert len(g) == len(days)
    assert g["value"].iloc[-1] == pytest.approx(23_459_804_204 - 27_996_757 - 2_043_352_267 + 1_903_000_000)


# ------------------------------------------------------------------------------------------- 10 / 11 trajectory
def test_a_series_that_begins_inside_its_oldest_quarter_reads_new_since_not_a_rate():
    import trajectory
    # NEAR's retired tokens: the Intents buyback began 2026-03-19, inside Q-2 of an asof 2026-10-09 window
    assert trajectory.partial_start_quarter("2026-03-19", "2026-10-09", 90, 5) == 2
    assert trajectory.partial_start_quarter("2025-01-01", "2026-10-09", 90, 5) is None
    g, n, label, why = trajectory.growth_rule([900.0, 700.0, 40.0, None], None, starts_in=2)
    assert g is None and label == "NEW since Q-2" and "part-quarter" in why
    g, *_ = trajectory.growth_rule([900.0, 700.0, 600.0, None], None, starts_in=None)
    assert g is not None


def test_divergence_headline_is_revenue_minus_price_with_capture_apart_on_the_real_dump(tmp_path, monkeypatch):
    import build_workbook as bw
    from test_trajectory import _built
    _built(tmp_path, monkeypatch, "Maple")
    row = next(r for r in bw.DIVERGENCE_ROWS if r["Project"] == "Maple")
    rev, px = row["Revenue 4Q growth (business, USD)"], row["Price 4Q growth"]
    gap = row["REVENUE − PRICE (pp) — the headline gap"]
    if isinstance(rev, float) and isinstance(px, float):
        assert gap == pytest.approx(rev - px)
    assert any(h.startswith("Capture: tokens retired") for h in row)
    # alignment round C10: NEAR's revenue line is gas + Intents now, and the note says so
    assert "TOTAL = gas revenue + NEAR Intents" in config.TRAJECTORY["divergence"]["revenue_notes"]["Near"]


# ------------------------------------------------------------------------------------------- decisions 1-6
def test_maple_switches_to_the_issuer_rule_with_the_safe_out():
    from fetch import archive as ar
    assert config.circulating_primary_metric("Maple") == "circulating_supply_onchain"
    keys = [k for k, _ in ar.serving_contracts(config.PROJECT_BY_NAME["Maple"], "noncirculating_holding_tokens")]
    assert "noncirc_ssf_safe" in keys
    safe = config.PROJECT_BY_NAME["Maple"]["contracts"]["noncirc_ssf_safe"]
    assert safe["address"] == "0x58BE00490D6a5D298d61f24226d2397983Ee39A1" and safe["verified"] == "2026-10-12"
    assert "etherscan.io/address/0x58BE" in safe["source_url"]


def test_etherfi_lp_withdrawal_is_excluded_and_the_2025_burn_recorded():
    scan = next(s for s in config.PROJECT_BY_NAME["Ether.fi"]["log_scans"] if s["key"] == "buyback_wallet_inflow")
    assert "0x3b52e19b49fca8a4cd417b3e2549a761e64cb9f03fd22d6848df8c635261112b" in scan["exclude_txs"]
    assert "lp-withdrawal-removed" in config.SOURCE_MARKERS
    ev = config.ONE_OFF_BURN_EVENTS["Ether.fi"][0]
    assert ev["tokens"] == 1_464_000 and ev["funding"] == "treasury" and ev["from"] == "2025-08-08"
    note = config.PROJECT_BY_NAME["Ether.fi"]["burn_split"]["note"]
    assert "ONE BURN WAS EXECUTED" in note and "NO EXECUTED BURN" not in note


def test_aerodrome_credits_liquid_bribes_with_fees_only_beside_on_the_real_dump(tmp_path, monkeypatch):
    """Decision 4: bribes paid to direct emissions are holder income — the liquid ones (the illiquid-token rule still
    applies to fees and bribes). Credited − fees-only over the dump = Σ (on-chain bribes − illiquid bribes)."""
    import build_workbook as bw
    from test_trajectory import _built
    _built(tmp_path, monkeypatch, "Aerodrome")
    g = bw._LAST_GROUPS
    cred = g[("Aerodrome", "holder_value_revenue_usd")].set_index("date")["value"]
    fees = g[("Aerodrome", "holder_value_fees_only_usd")].set_index("date")["value"]
    br = g[("Aerodrome", "voter_rewards_onchain_bribes_usd")].set_index("date")["value"]
    ill = g[("Aerodrome", "voter_rewards_illiquid_bribes_usd")].set_index("date")["value"]
    want = sum(max(0.0, v - ill.get(d, 0.0)) for d, v in br.items())
    assert (cred.sum() - fees.sum()) == pytest.approx(want, rel=1e-9)
    assert want < br.sum()                                    # the LAPTOP bribe's illiquid part is out
    src = g[("Aerodrome", "holder_value_revenue_usd")]
    assert set(src[src["date"] >= "2026-07-02"]["source"].str.split("[").str[0]) == {
        "derived:holder_value_with_liquid_bribes"}


def test_uniswap_one_off_is_its_own_event_and_chainlink_reserve_is_held_at_discretion():
    assert "Uniswap" not in config.HOLDER_VALUE
    assert "2025-12-27..2025-12-27: 100,000,000" in config.one_off_events_text("Uniswap")
    assert config.PHASE2_ATTRIBUTES["Chainlink"]["destination"]["value"] == "held-discretionary"


def test_the_rebase_reference_keeps_its_total_after_cj_on_the_real_dump(tmp_path, monkeypatch):
    """Item 14: after CJ moved the pre-handover totals, the rebase reference had no T and Aerodrome's in_emissions read
    CHECK (no reference) — a3_net_absorption followed it. With CoinGecko's total standing in where the chain total is
    not yet restated, the row is judged again."""
    from real_dump import evaluate_real, load
    payload = load("Aerodrome")
    for r in payload["metrics"]:
        if r["metric"] == "total_supply" and str(r["source"]).startswith("coingecko"):
            r["metric"] = "total_supply_coingecko"
    o = evaluate_real(tmp_path, monkeypatch, "Aerodrome", payload=payload)
    assert o["in_emissions"]["verdict"] != "CHECK (no reference)", o["in_emissions"]


def test_the_daily_run_line_names_the_bigquery_fix():
    """Item 15: a failed check carries its exact fix in the closing line (adc_check.py's FIX line, else the table)."""
    ps = (Path(__file__).resolve().parents[1] / "windows" / "daily_run.ps1").read_text(encoding="utf-8")
    assert "gcloud auth application-default login" in ps.split("$AdcFix")[1]
    assert 'Where-Object { $_ -like "FIX: *" }' in ps and '$_[0] -like "BigQuery check*"' in ps


# ------------------------------------------------------------------------------------------- 16 BigQuery on user ADC
def test_the_daily_run_checks_bigquery_first_logs_in_by_hand_and_skips_when_scheduled():
    """Item 16 (replaced 2026-10-12 evening): user ADC stays — a real one-row query before any step; run by hand, a
    failure runs the login and the quota-project command itself and re-tests; scheduled, the BigQuery reads are skipped
    and the summary says so FIRST."""
    root = Path(__file__).resolve().parents[1]
    ps = (root / "windows" / "daily_run.ps1").read_text(encoding="utf-8")
    steps = [ln.strip() for ln in ps.splitlines() if "Invoke-Step " in ln and not ln.lstrip().startswith(("#", "function"))]
    assert steps[0].startswith('$adc = Invoke-Step "BigQuery check" @("adc_check.py")'), steps[:2]
    login = ps.index("& gcloud auth application-default login")
    assert ps.index("if ((Test-Interactive) -and $gcloud)") < login < ps.index(
        "& gcloud auth application-default set-quota-project near-data-510309") < ps.index(
        'Invoke-Step "BigQuery check (after login)"')
    assert '$env:TOKEN_METRICS_SKIP_NEAR_BQ = "1"' in ps and 'if ($BqSkipped) { $summary = "$BqSkipped | $summary" }' in ps
    assert "[switch]$Scheduled" in ps and "if ($Scheduled) { return $false }" in ps and "svchost.exe" in ps
    assert "-Scheduled" in (root / "windows" / "register_task.ps1").read_text(encoding="utf-8")
    src = (root / "adc_check.py").read_text(encoding="utf-8")
    assert "crypto_near_mainnet_us.blocks" in src and "LIMIT 1" in src and 'print(f"FIX: {fix}")' in src


def test_a_skipped_run_reads_no_bigquery_and_the_summary_leads_with_it(tmp_path, monkeypatch):
    import daily_summary as ds
    from fetch import near_bigquery as nbq
    from fetch.base import FetchOutput

    class Boom:
        def query(self, *a, **k):
            raise AssertionError("no query may run when the daily run skipped BigQuery")
    monkeypatch.setenv(nbq.SKIP_ENV, "1")
    a = nbq.NearBigQuery(client=Boom(), bq=object(), cache_file=tmp_path / "c.json", csv_dir=tmp_path / "none")
    out = FetchOutput()
    a._project("Near", config.PROJECT_BY_NAME["Near"]["near_bigquery"], out, False)
    msgs = [e.message for e in out.log]
    assert any(m.startswith("NEAR BigQuery skipped — log in: gcloud auth application-default login") for m in msgs)
    lines = ds.summarise({}, None, None, "[06:30:02] NEAR BigQuery skipped — log in: gcloud auth "
                                          "application-default login\n")
    assert lines[1] == "NEAR BigQuery skipped — log in: gcloud auth application-default login"


def test_a_bigquery_lapse_is_caught_up_from_each_series_last_stored_day(tmp_path, monkeypatch):
    """Item 16: after a lapse the next good run reads from the day after the last one stored — one query for the supply
    flows (burn / issuance), the activity split's top-up covers every missing newer day, p2p's top-up up to a chunk."""
    from types import SimpleNamespace
    import copy
    from fetch import near_bigquery as nbq
    from fetch.base import FetchOutput
    monkeypatch.setattr(nbq, "today", lambda: pd.Timestamp("2026-10-12"))

    class Bq:
        class QueryJobConfig:
            def __init__(self, **kw):
                self.__dict__.update(kw)
                self.dry_run = kw.get("dry_run", False)
        ScalarQueryParameter = staticmethod(lambda k, t, v: (k, t, v))
        ArrayQueryParameter = staticmethod(lambda k, t, v: (k, t, tuple(v)))

    class Client:
        windows: list = []

        def query(self, sql, job_config=None):
            p = {k: v for k, _, v in job_config.query_parameters}
            if job_config.dry_run:
                return SimpleNamespace(total_bytes_processed=3_000_000)
            Client.windows.append((str(p["d0"]), str(p["d1"])))
            days = pd.date_range(p["d0"], p["d1"])
            rows = [{"day": d, "n_blocks": 1, "n_epoch_starts": 0, "burn_near": 1.0, "issuance_near": 0.0, "n": 5,
                     "kaiching_n": 1} for d in days]
            return SimpleNamespace(result=lambda: rows, total_bytes_billed=10 * 1024 ** 2)
    spec = copy.deepcopy(config.PROJECT_BY_NAME["Near"]["near_bigquery"])
    # stored through 2026-10-05; the credential lapsed 10-06..10-11
    st = {"ledger": {}, "topup_bytes": {"bytes": 3e6, "on": "2026-10-12", "source": "t"},
          "supply_flows_through": "2026-10-05",
          "activity": {str(d.date()): [5, 1] for d in pd.date_range("2025-10-13", "2026-10-05")}}
    a = nbq.NearBigQuery(client=Client(), bq=Bq, cache_file=tmp_path / "c.json", csv_dir=tmp_path / "none")
    a._run_bytes = a._backfill_bytes = 0
    a._supply_flows(a._client_obj, spec, st, FetchOutput(), "Near")
    assert Client.windows[-1] == ("2026-10-06", "2026-10-11") and st["supply_flows_through"] == "2026-10-11"
    n = len(Client.windows)
    a._activity_ex(a._client_obj, spec, st, FetchOutput(), "Near")
    assert Client.windows[n] == ("2026-10-06", "2026-10-11")              # the top-up: the whole lapse, first
    assert all(str(d.date()) in st["activity"] for d in pd.date_range("2026-10-06", "2026-10-11"))
