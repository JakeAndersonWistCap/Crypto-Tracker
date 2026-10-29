"""P2g (Jake 2026-10-29): the Charlie edition - `--edition charlie`, the Highlights first tab. Plain asserts, no network.

Two kinds of test. The SYNTHETIC ones build a small reader-shaped workbook (the tabs, header rows and project rows the
cards quote, with numbers chosen to meet each condition) and run charlie_highlights on it in a second. The FIXTURE ones
build the committed real dumps (fixtures/real_*.json) twice - without and with the flag - and compare. The fixture build
holds other numbers than Jake's live store, so on it only some cards render (C5, C7 and C8); the fixture integration
tests lower the floor to 3 (charlie_edition.MIN_CARDS) and say so. The production floor stays 5.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import charlie_edition  # noqa: E402
import charlie_highlights as H  # noqa: E402
import reader  # noqa: E402

FIXTURE_CARDS = 3          # C5, C7, C8 render on the committed dumps (the others need Jake's live numbers)


# ======================================================================================================= conditions
def test_c1_key_at_least_one_and_strictly_decreasing_earlier_quarters():
    assert H.c1_holds(1.06, 0.86, 0.79, 0.73)
    assert H.c1_holds(1.0, 0.9, 0.8, 0.7)                          # the key may equal 1.0
    assert not H.c1_holds(0.999, 0.86, 0.79, 0.73)                 # below the crossover
    assert not H.c1_holds(1.06, 1.06, 0.79, 0.73)                  # not strictly above -1q
    assert not H.c1_holds(1.06, 0.86, 0.86, 0.73)                  # earlier quarters not strictly decreasing
    assert not H.c1_holds(1.06, 0.86, 0.79, 0.80)                  # -3q above -2q
    assert not H.c1_holds(1.06, "blocked", 0.79, 0.73)             # a code is no number
    assert not H.c1_holds(1.06, None, 0.79, 0.73)


def test_c2_at_least_ten_and_at_least_one_point_three_times_three_quarters_back():
    assert H.c2_holds(22.95, 14.25)
    assert H.c2_holds(10.0, 7.0)
    assert not H.c2_holds(9.99, 1.0)
    assert not H.c2_holds(20.0, 16.0)                              # 20 < 1.3 x 16
    assert H.c2_holds(13.0, 10.0)                                  # exactly 1.3x passes
    assert not H.c2_holds("blocked", 1.0)
    assert not H.c2_holds(20.0, "n/a")


def test_c3_highest_net_holder_yield_above_zero_whoever_it_is():
    assert H.c3_pick({"Sky": 0.032, "Near": -0.01, "Fluid": -0.009}) == ("Sky", 1)
    assert H.c3_pick({"Sky": 0.032, "Pendle": 0.05, "Near": -0.01}) == ("Pendle", 2)      # another project is highest
    assert H.c3_pick({"Sky": -0.001, "Near": -0.01}) == (None, 0)                        # none above 0
    assert H.c3_pick({"Sky": 0.0, "Near": -0.01}) == (None, 0)                           # 0 is not above 0
    assert H.c3_pick({"Sky": "blocked", "Near": 0.002, "X": None}) == ("Near", 1)        # codes are skipped
    assert H.c3_pick({}) == (None, 0)


def test_c4_rank_one_in_a3_and_key_at_least_point_nine():
    assert H.c4_holds(1.0, True, 1.02)
    assert H.c4_holds(1, True, 0.9)
    assert not H.c4_holds(1.0, True, 0.899)
    assert not H.c4_holds(2.0, True, 1.5)
    assert not H.c4_holds(1.0, False, 1.5)                         # rank 1 in another archetype
    assert not H.c4_holds("unranked", True, 1.5)
    assert not H.c4_holds(1.0, True, "blocked")


def test_c4_largest_revenue_base_clause_needs_the_single_largest():
    assert H.is_largest("Hyperliquid", {"Hyperliquid": 849e6, "Sky": 162e6, "Aave": None, "X": "n/a"})
    assert not H.is_largest("Hyperliquid", {"Hyperliquid": 100.0, "Sky": 162e6})
    assert not H.is_largest("Hyperliquid", {"Hyperliquid": 100.0, "Sky": 100.0})           # a tie is not the largest
    assert not H.is_largest("Hyperliquid", {"Sky": 1.0})


def test_c5_growth_above_half_and_supply_falling():
    assert H.c5_holds(1.406, -0.036)
    assert not H.c5_holds(0.5, -0.036)                             # strictly above 0.5
    assert H.c5_holds(0.51, -0.0001)
    assert not H.c5_holds(1.4, 0.0)                                # 0 is not falling
    assert not H.c5_holds(1.4, 0.01)
    assert not H.c5_holds("n/a", -0.1)


def test_c6_latest_zero_and_an_earlier_quarter_above_zero():
    assert H.c6_holds(0.0, [0.0, 0.65, 1.23])
    assert H.c6_holds(0, [0, 0, 0.01])
    assert not H.c6_holds(0.0, [0.0, 0.0, 0.0])                    # never above zero
    assert not H.c6_holds(0.1, [0.5, 0.65, 1.23])                  # latest is not zero
    assert not H.c6_holds("blocked", [0.5])
    assert H.c6_holds(0.0, ["blocked", "n/a", 0.4])                # codes among the earlier quarters are skipped
    assert not H.c6_holds(0.0, ["blocked", "n/a", "n/a"])


def test_c6_second_best_clause_is_true_only_when_exactly_one_is_better():
    assert H.is_second_best("Fluid", {"Sky": 0.03, "Fluid": -0.009, "Near": -0.01, "X": "blocked"})
    assert not H.is_second_best("Fluid", {"Sky": 0.03, "Pendle": 0.02, "Fluid": -0.009})
    assert not H.is_second_best("Fluid", {"Fluid": 0.5, "Sky": 0.03})                       # best, not second
    assert not H.is_second_best("Fluid", {"Sky": 0.03})


def test_c7_highest_revenue_growth_strictly():
    assert H.c7_holds("Near", {"Near": 1.63, "Uniswap": 1.4, "Sky": -0.02, "Y": "n/a"})
    assert not H.c7_holds("Near", {"Near": 1.4, "Uniswap": 1.4})
    assert not H.c7_holds("Near", {"Near": 1.0, "Uniswap": 1.4})
    assert not H.c7_holds("Near", {"Near": "n/a", "Uniswap": 1.4})


def test_c8_three_lowest_and_three_highest_with_a_value():
    vals = {"A": 0.30, "B": 0.01, "C": 0.20, "D": 0.04, "E": "n/a", "F": 0.10, "G": 0.05, "H": None}
    low, high = H.c8_split(vals)
    assert low == [("B", 0.01), ("D", 0.04), ("G", 0.05)]
    assert high == [("A", 0.30), ("C", 0.20), ("F", 0.10)]
    assert H.c8_split({"A": 1.0, "B": 2.0, "C": 3.0, "D": 4.0, "E": 5.0}) is None         # fewer than six: no card
    assert H.c8_split({"A": 1.0, "B": 2.0, "C": 3.0, "D": 4.0, "E": 5.0, "F": "n/a"}) is None
    low, high = H.c8_split({n: float(i) for i, n in enumerate("ABCDEF")})
    assert [n for n, _ in low] == list("ABC") and [n for n, _ in high] == list("FED")


def test_number_display_follows_the_source_format():
    f = H.format_value
    assert f(0.0679, "0.0%") == "6.8%"
    assert f(-0.0092, "0.0%") == "-0.9%"
    assert f(1.0612, '0.0"x"') == "1.1x"
    assert f(0.0, '0.0"x"') == "0.0x"
    assert f(849e6, reader.FMT_USD) == "$849.0m"
    assert f(1.27e9, reader.FMT_USD) == "$1.3bn"
    assert f(0.0, reader.FMT_USD) == "$0"
    assert f(1234567.0, reader.FMT_TOK) == "1,234,567"
    assert f(0.0, '0.0"x";-0.0"x";"0.0x — halted"') == "0.0x — halted"
    assert f(-3.2, '0.0"x";-0.0"x"') == "-3.2x"
    assert f("blocked", "General") == "blocked"
    assert f(0.5, "General") == "0.5"


# ============================================================================================ the synthetic reader
ABOUT_LINES = [
    "Token metrics",
    "Data as of 2026-10-29. Every figure is checked in token_metrics_audit.xlsx, built in the same run.",
    "",
    "What this file measures",
    "It measures whether a token's holders are paid by the business behind it, or diluted to keep it going. For each "
    "project it sets what the business earns against the new tokens it hands out.",
    "",
    "How to read a cell",
    "Fill is confidence: GREEN verified and usable · AMBER a real number with a caveat (hover it) · RED not "
    "usable · no fill: computed here from the cells it names.",
    "",
]
PARKED = {"World Mobile"}
# (project, summary section, key, rank, revenue annualised, revenue growth, net supply change, NHY, HVY, dilution,
#  lock rate, subsidy cover, tvl change, price change, burn yield)
BASE = {
    "GEODNET": dict(sec="A2", key=1.06, rank=1, rev=3e7, gro=0.2, sup=0.01, nhy=-0.045, hvy=0.068, dil=0.12, lock=0.01,
                    cover=1.0, tvl=0.1, price=0.1),
    "Morpho": dict(sec="A2", key=22.95, rank=2, rev=0.0, gro=0.1, sup=0.0, nhy=-0.02, hvy=0.0, dil=0.07, lock=0.3,
                   cover=2.0, tvl=0.459, price=0.15),
    "Sky": dict(sec="A3", key=0.5, rank=2, rev=162e6, gro=-0.02, sup=-0.002, nhy=0.032, hvy=0.039, dil=0.062,
                lock=0.857, cover=3.0, tvl=0.05, price=0.0),
    "Fluid": dict(sec="A3", key=0.0, rank=3, rev=6e6, gro=0.3, sup=0.0, nhy=-0.009, hvy=0.0, dil=0.2, lock=0.1,
                  cover=56.4, tvl=0.2, price=0.1),
    "Hyperliquid": dict(sec="A3", key=1.02, rank=1, rev=849e6, gro=0.486, sup=-0.04, nhy=-0.03, hvy=0.05, dil=0.227,
                        lock=0.2, cover=4.0, tvl=0.3, price=0.2),
    "Uniswap": dict(sec="A4", key="n/a", rank="unranked", rev=126.6e6, gro=1.406, sup=-0.036, nhy=-0.011, hvy=0.028,
                    dil=0.036, lock=0.0, cover=5.0, tvl=0.0, price=0.0, burn=0.036),
    "Near": dict(sec="A1", key="n/a", rank="unranked", rev=16e6, gro=1.63, sup=0.024, nhy=-0.013, hvy=0.002, dil=0.041,
                 lock=0.45, cover=6.0, tvl=0.994, price=1.6),
    "Ethereum": dict(sec="A1", key="n/a", rank="unranked", rev=5e8, gro=0.1, sup=0.0, nhy=-0.02, hvy=0.01, dil=0.014,
                     lock=0.2, cover=7.0, tvl=0.1, price=0.1),
    "Pendle": dict(sec="A3", key=0.2, rank=4, rev=2e7, gro=0.2, sup=0.01, nhy=-0.02, hvy=0.01, dil=0.039, lock=0.5,
                   cover=8.0, tvl=0.1, price=0.1),
    "Aethir": dict(sec="A2", key=0.4, rank=3, rev=1e7, gro=0.05, sup=0.0, nhy=-0.3, hvy=0.01, dil=0.288, lock=0.1,
                   cover=1.0, tvl=0.1, price=0.1),
    "Plume": dict(sec="A1", key="n/a", rank="unranked", rev=1e6, gro=0.1, sup=0.0, nhy=-0.2, hvy=0.0, dil=0.208,
                  lock=0.1, cover=1.0, tvl=0.1, price=0.1),
    # parked: figures that would take over C3 (highest NHY), C4's revenue clause, C7 and C8 if they were not excluded
    "World Mobile": dict(sec="A2", key="parked", rank="parked", rev=9e12, gro=9.99, sup=-0.5, nhy=0.9, hvy=0.9,
                         dil=0.99, lock=0.9, cover=99.0, tvl=9.0, price=9.0),
}


def make_reader(over: dict | None = None, drop: tuple = ()) -> openpyxl.Workbook:
    """A reader-shaped workbook: About, Summary (section rows), A2/A3/A4, Comparison, Business. `over` maps
    project -> {field: value} to change; `drop` names projects left out."""
    data = {n: dict(v, **(over or {}).get(n, {})) for n, v in BASE.items() if n not in drop}
    import build_workbook as bw                                        # noqa: F401  (key_headers reads its constants)
    wb = openpyxl.Workbook()
    ab = wb.active
    ab.title = "About"
    for i, t in enumerate(ABOUT_LINES, start=1):
        ab.cell(i, 1, t or None)
    r0 = len(ABOUT_LINES) + 2
    ab.cell(r0, 1, "Honest coverage — counted from the Summary tab (the figures follow the data)")
    for j, h in enumerate(("Archetype", "Projects", "Ranked", "Unranked", "Parked"), start=1):
        ab.cell(r0 + 1, j, h)
    for k, (lab, n, rk, un, pk) in enumerate((("A1 Infrastructure", 3, 0, 3, 0), ("A2 Coordination", 3, 3, 0, 1),
                                              ("All archetypes", 6, 3, 3, 1))):
        ab.cell(r0 + 2 + k, 1, lab)
        for j, v in enumerate((n, rk, un, pk), start=2):
            ab.cell(r0 + 2 + k, j, f"=COUNTA(Summary!A5:A6)+{v}-2" if k == 0 and j == 2 else v).number_format = "0"
    ab.cell(r0 + 6, 1, "Why a project is unranked, by the code in its key cell")

    def sheet(name, heads, rows):
        ws = wb.create_sheet(name)
        ws.cell(1, 1, name)
        for j, h in enumerate(heads, start=1):
            ws.cell(4, j, h)
        for i, vals in enumerate(rows, start=5):
            for j, v in enumerate(vals, start=1):
                if v is not None:
                    ws.cell(i, j, v)
        return ws

    # Summary: section rows (the name alone), then project rows
    sh = ["Project", "Token", "Archetypes", "Protocol revenue, annualised", "Revenue growth, selected period",
          "Archetype key", "Rank in archetype", "Net supply change, % a year"]
    rows = []
    for sec, label in (("A1", "A1 Infrastructure — ranked by x"), ("A2", "A2 Coordination — ranked by x"),
                       ("A3", "A3 Revenue Buyback — ranked by x"), ("A4", "A4 Permanent Burn — ranked by x")):
        rows.append((label,))
        for n, d in data.items():
            if d["sec"] == sec:
                rows.append((n, n[:3].upper(), sec, d["rev"], d["gro"], d["key"], d["rank"], d["sup"]))
    S = sheet("Summary", sh, rows)
    for r in range(5, S.max_row + 1):
        if S.cell(r, 2).value is None:
            continue
        S.cell(r, 4).number_format = reader.FMT_USD
        S.cell(r, 5).number_format = reader.FMT_PCT
        S.cell(r, 8).number_format = reader.FMT_PCT
        if isinstance(S.cell(r, 6).value, float):
            S.cell(r, 6).number_format = reader.FMT_X
        name = S.cell(r, 1).value
        for c in (6, 7):                                               # the key / rank cell's comment: the unranked reason
            if isinstance(S.cell(r, c).value, str):
                from openpyxl.comments import Comment
                S.cell(r, c).comment = Comment(
                    "No measured key.\n\nn/a — customer_revenue_usd (q0): no rows stored. A second sentence.\n\n"
                    "Working: audit file, A2 Coordination!BN10." if name != "World Mobile" else
                    "PARKED — no further work on World Mobile; its row is shown with its numbers.\n\nBehind it: x",
                    "t")

    import charlie_highlights as ch
    head_a2, early_a2 = ch.key_headers(reader.A2_TAB)
    head_a3, early_a3 = ch.key_headers(reader.A3_TAB)
    q = {"GEODNET": [0.86, 0.79, 0.73], "Morpho": [18.0, 16.0, 14.25], "Fluid": [0.0, 0.65, 1.23],
         "Hyperliquid": [1.01, 1.75, 2.57], "Sky": [0.4, 0.3, 0.2]}
    q.update({k: v for k, v in (over or {}).get("_q", {}).items()})
    for tab, head, early, secs in ((reader.A2_TAB, head_a2, early_a2, ("A2",)), (reader.A3_TAB, head_a3, early_a3, ("A3",))):
        rows = []
        for n, d in data.items():
            if d["sec"] in secs or n in ("GEODNET", "Sky", "Hyperliquid", "Fluid", "Morpho"):
                if (tab == reader.A2_TAB) != (n in ("GEODNET", "Morpho", "Aethir", "World Mobile")):
                    continue
                rows.append((n, n[:3].upper(), d["key"]) + tuple(q.get(n, ["n/a"] * 3)))
        ws = sheet(tab, ["Project", "Token", head] + early, rows)
        for r in range(5, ws.max_row + 1):
            for c in range(3, 7):
                if isinstance(ws.cell(r, c).value, (int, float)):
                    ws.cell(r, c).number_format = reader.FMT_X
    sheet(reader.A4_TAB, ["Project", "Token", "Burn ÷ issuance (x)", "Permanent burn yield"],
          [(n, n[:3].upper(), "n/a", d["burn"]) for n, d in data.items() if "burn" in d])
    ws = wb[reader.A4_TAB]
    ws["D5"].number_format = reader.FMT_PCT
    cm = ["Project", "Token", "Primary archetype", "Holder value yield", "Total dilution, % of float", "Net holder yield",
          "Subsidy cover (x)", "Lock rate, % of circulating"]
    C = sheet(reader.COMPARISON, cm, [(n, n[:3].upper(), d["sec"], d["hvy"], d["dil"], d["nhy"], d["cover"], d["lock"])
                                      for n, d in data.items()])
    for r in range(5, C.max_row + 1):
        for c, f in ((4, reader.FMT_PCT), (5, reader.FMT_PCT), (6, reader.FMT_PCT), (7, reader.FMT_X), (8, reader.FMT_PCT)):
            C.cell(r, c).number_format = f
    B = sheet(reader.BUSINESS, ["Project", "Primary archetype", "TVL change, selected period",
                                "Price change, selected period"],
              [(n, d["sec"], d["tvl"], d["price"]) for n, d in data.items()])
    for r in range(5, B.max_row + 1):
        B.cell(r, 3).number_format = reader.FMT_PCT
        B.cell(r, 4).number_format = reader.FMT_PCT
    return wb


def run(over=None, drop=(), min_cards=5, wb=None):
    wb = wb or make_reader(over, drop)
    res = H.add_highlights(wb, None, min_cards=min_cards, parked=lambda n: n in PARKED)
    return wb, res


def test_synthetic_all_eight_cards_render_in_order_and_the_tab_is_first():
    wb, res = run()
    assert res["cards"] == ["C1", "C2", "C3", "C4", "C5", "C6", "C7", "C8"], res["omitted"]
    assert wb.sheetnames[0] == "Highlights" and wb.sheetnames[1] == "About"
    ws = wb["Highlights"]
    assert ws.sheet_properties.tabColor.rgb.endswith("1F8A8A")
    assert ws.page_setup.orientation == "landscape" and ws.page_setup.paperSize == 9
    assert ws.page_setup.fitToWidth == 1 and ws.sheet_properties.pageSetUpPr.fitToPage
    assert ws.sheet_view.showGridLines is False and ws.sheet_view.zoomScale == 90
    assert ws["A1"].value == "Token metrics — highlights"
    assert ws["A1"].font.bold and ws["A1"].font.size == 18
    assert sum(1 for w in wb.worksheets if w.sheet_view.tabSelected) == 1
    # the as-of and the two About lines are formulas on About's cells
    assert ws["A2"].value.startswith("=LEFT(About!A2,") and ws["A3"].value.startswith("=LEFT(About!A5,")
    assert ws["A4"].value == "=About!A8"
    text = "\n".join(res["lines"])
    for needle in ("GEODNET", "Morpho", "Sky", "Hyperliquid", "Uniswap", "Fluid", "Near", "Dilution league"):
        assert needle in text
    last = [c.value for row in ws.iter_rows() for c in row if c.value is not None][-1]
    assert last == "Confidence is shown by cell fill throughout; hover any number for its source." == H.FOOTER


def test_synthetic_the_printed_card_shows_label_and_the_value_as_displayed():
    _, res = run()
    lines = res["lines"]
    i = next(k for k, ln in enumerate(lines) if "[C1]" in ln)
    block = lines[i:i + 8]
    assert "GEODNET" in block[0]
    assert any(ln.strip() == "Customer revenue ÷ emissions (x): 1.1x" for ln in block), block
    assert any(ln.strip() == "Holder value yield: 6.8%" for ln in block)
    j = next(k for k, ln in enumerate(lines) if "[C4]" in ln)
    assert any("Protocol revenue, annualised: $849.0m" in ln for ln in lines[j:j + 9])
    assert any("Reading:" in ln for ln in block) and any("Caveat:" in ln for ln in block)


def test_synthetic_a_failed_card_is_omitted_and_the_rest_reflow():
    wb, res = run(over={"GEODNET": {"key": 0.98}}, min_cards=5)
    assert res["cards"] == ["C2", "C3", "C4", "C5", "C6", "C7", "C8"]
    assert any(o.startswith("C1 GEODNET") for o in res["omitted"])
    ws = wb["Highlights"]
    heads = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and " — " in c.value
             and c.font.size == 12 and c.font.bold]
    assert not any(h.startswith("GEODNET") for h in heads) and heads[0].startswith("Morpho")
    # the first card sits in the top-left, the second beside it: reflow
    assert ws.cell(6, 1).value.startswith("Morpho") and ws.cell(6, 10).value.startswith("Sky")


def test_synthetic_card_three_is_about_the_highest_project_whoever_it_is_and_says_so():
    wb, res = run(over={"Pendle": {"nhy": 0.05, "lock": 0.2}})
    assert "C3" in res["cards"]
    ws = wb["Highlights"]
    heads = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.font.size == 12
             and c.font.bold]
    assert any(h.startswith("Pendle — Holders are paid more") for h in heads)
    reading = next(c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)
                   and c.value.startswith("Reading: Measured returns"))
    assert "the highest in the set" in reading and "the only project" not in reading
    assert "locked" not in reading                                     # its lock rate is not above 0.5
    cav = next(c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)
               and c.value.startswith("Caveat: Net holder yield is the measured"))
    assert "buyback" not in cav                                        # the Sky caveat is not said of another project
    # and Sky itself: the only one, locked clause (lock rate 0.857 > 0.5)
    wb2, _ = run()
    r2 = next(c.value for row in wb2["Highlights"].iter_rows() for c in row if isinstance(c.value, str)
              and c.value.startswith("Reading: Measured returns"))
    assert "the only project in the set where that holds" in r2 and "most of the circulating supply is locked" in r2
    assert any(isinstance(c.value, str) and c.value.startswith("Caveat: The buyback retires well below")
               for row in wb2["Highlights"].iter_rows() for c in row)


def test_synthetic_conditional_clauses_follow_the_quoted_cells():
    # C4: the revenue clause only when Hyperliquid has the largest base among non-parked
    _, a = run()
    assert any("on the largest revenue base in the set" in ln for ln in a["lines"])
    _, b = run(over={"Sky": {"rev": 9e9}})
    assert not any("largest revenue base" in ln for ln in b["lines"])
    # C6: second best only when true
    assert any("second best in the set" in ln for ln in a["lines"])
    _, c = run(over={"Near": {"nhy": 0.02}})
    assert not any("second best" in ln for ln in c["lines"])
    # C3: the locked clause only above 0.5
    _, d = run(over={"Sky": {"lock": 0.5}})
    assert not any("is locked" in ln for ln in d["lines"])


def test_synthetic_card_conditions_decide_each_card():
    cases = {"C1": {"GEODNET": {"key": 0.5}}, "C2": {"Morpho": {"key": 9.0}}, "C3": {"Sky": {"nhy": -0.1}},
             "C4": {"Hyperliquid": {"key": 0.5}}, "C5": {"Uniswap": {"gro": 0.4}}, "C6": {"Fluid": {"key": 0.2}},
             "C7": {"Hyperliquid": {"gro": 2.0}}}
    for cid, over in cases.items():
        _, res = run(over=over, min_cards=1)
        assert cid not in res["cards"], (cid, res["cards"])
        assert any(o.startswith(cid) for o in res["omitted"]), (cid, res["omitted"])
        assert len(res["cards"]) == 7, (cid, res["cards"], res["omitted"])


def test_synthetic_fewer_than_five_cards_raises_listing_each_failed_condition_with_values():
    over = {"GEODNET": {"key": 0.5}, "Morpho": {"key": 9.0}, "Sky": {"nhy": -0.1}, "Hyperliquid": {"key": 0.5}}
    with pytest.raises(H.HighlightsError) as e:
        run(over=over)                                                   # 4 render: C5, C6, C7, C8 -> short of 5
    msg = str(e.value)
    assert "only 4 of 8" in msg and "5 are needed" in msg
    for cid in ("C1 GEODNET", "C2 Morpho", "C3", "C4 Hyperliquid"):
        assert cid in msg, msg
    assert "key=0.5" in msg and "needs key >= 1.0" in msg and "needs the Summary" not in msg
    assert "Cards that did: C5, C6, C7, C8" in msg
    run(over=over, min_cards=4)                                          # the floor is a parameter; 4 render


def test_synthetic_a_missing_cell_omits_its_card_with_the_reason():
    _, res = run(drop=("Fluid",), min_cards=1)
    assert "C6" not in res["cards"]
    assert any(o.startswith("C6") and "Fluid" in o for o in res["omitted"]), res["omitted"]


def _refs(ws):
    """[(cell, sheet, coord)] for every formula on Highlights that quotes another tab (not About)."""
    out = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                m = re.fullmatch(r"=('?)([^'!]+)\1!([A-Z]+\d+)", c.value)
                if m and m.group(2) != "About":
                    out.append((c, m.group(2), m.group(3)))
    return out


def test_synthetic_every_number_is_a_formula_with_the_source_cells_format_and_no_typed_number():
    wb, res = run()
    ws = wb["Highlights"]
    refs = _refs(ws)
    assert len(refs) >= 30
    for c, sheet, coord in refs:
        assert wb[sheet][coord].number_format == c.number_format, (c.coordinate, c.value)
    for row in ws.iter_rows():
        for c in row:
            assert not isinstance(c.value, (int, float)) or isinstance(c.value, bool) and False, (c.coordinate, c.value)
    # no number inside a sentence: strip tab names (A1..A4) and tab-less words, then no digit remains
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and not c.value.startswith("=") and (
                    c.value.startswith(("Reading:", "Caveat:")) or (c.font.size == 12 and c.font.bold)):
                t = re.sub(r"\bA[1-4]\b", "", c.value)
                assert not re.search(r"\d", t), c.value


def test_synthetic_found_by_header_and_row_not_by_address():
    """Move the columns and rows of the source tabs: the chips follow."""
    wb = make_reader()
    ws = wb["Comparison"]
    ws.insert_cols(3)                                                   # every Comparison header moves one column right
    ws.insert_rows(5)
    ws.cell(5, 1, "Zzz")
    ws.cell(5, 2, "ZZZ")
    wb2, res = run(wb=wb)
    assert "C3" in res["cards"] and "C8" in res["cards"]
    for c, sheet, coord in _refs(wb2["Highlights"]):
        assert wb2[sheet][coord].number_format == c.number_format
    sky = next(c for c, s, k in _refs(wb2["Highlights"]) if s == "Comparison" and wb2[s][k].value == 0.032)
    assert wb2["Comparison"].cell(wb2["Comparison"][sky.value.split("!")[1]].row, 1).value == "Sky"


def test_synthetic_the_only_fills_are_the_source_cells_confidence_fills():
    from openpyxl.styles import PatternFill
    wb = make_reader()
    wb["A2 Coordination"]["C5"].fill = PatternFill("solid", fgColor="FFC6EFCE")
    wb["Comparison"]["D5"].fill = PatternFill("solid", fgColor="FFFFEB9C")
    wb, res = run(wb=wb)
    ws = wb["Highlights"]
    filled = {}
    for c, sheet, coord in _refs(ws):
        src = wb[sheet][coord].fill
        if c.fill.fill_type == "solid":
            filled[c.coordinate] = c.fill.fgColor.rgb
            assert src.fill_type == "solid" and src.fgColor.rgb == c.fill.fgColor.rgb, (c.coordinate, coord)
        else:
            assert src.fill_type != "solid", (c.coordinate, sheet, coord)         # a confidence fill is never dropped
    assert filled, "the mirrored confidence fill is carried to the chip"
    for row in ws.iter_rows():
        for c in row:
            if c.fill.fill_type == "solid":                              # a chip's fill, or its merged second cell
                assert c.fill.fgColor.rgb in [v for k, v in filled.items() if int(re.sub(r"\D", "", k)) == c.row], \
                    f"decorative fill at {c.coordinate}"


def test_synthetic_each_chip_carries_a_comment_naming_its_source_and_the_card_a_link():
    wb, res = run()
    ws = wb["Highlights"]
    cov = next(c.row for row in ws.iter_rows() for c in row if c.value == "Coverage")
    cards = [(c, s, k) for c, s, k in _refs(ws) if c.row < cov]
    commented = [c for c, s, k in cards if c.comment is not None]
    assert all(c.comment is not None for c, s, k in cards if not re.search(r"!A\d+$", c.value))     # numbers
    assert all(c.comment is None for c, s, k in cards if re.search(r"!A\d+$", c.value))            # project names
    assert any("Source: A2 Coordination, GEODNET" in c.comment.text for c in commented)
    links = [c.hyperlink for row in ws.iter_rows() for c in row if c.hyperlink is not None]
    assert len(links) == 8 and all(h.location and "!A" in h.location for h in links)
    assert {h.display for h in links} >= {"Open A2 Coordination", "Open A3 Revenue Buyback", "Open Summary",
                                          "Open Comparison"}
    assert all(h.display == f"Open {h.location.split('!')[0].strip(chr(39))}" for h in links)


FORBIDDEN_WORDS = ("still", "now", "newly", "improved since")


def test_synthetic_no_string_on_the_tab_dates_it_or_refers_to_an_earlier_build():
    wb, res = run()
    strings = H.all_strings(wb["Highlights"])
    assert len(strings) > 60
    for s in strings:
        low = s.lower()
        for w in FORBIDDEN_WORDS:
            assert not re.search(rf"\b{w}\b", low), (w, s)
        assert not re.search(r"earlier build|previous build|last build|yesterday|last run", low), s
        assert not H.FORBIDDEN.search(s), s
    # the scanner itself catches them
    for bad in ("it is still negative", "Now positive", "newly ranked", "improved since the last build"):
        assert H.FORBIDDEN.search(bad), bad


def test_synthetic_parked_projects_are_in_no_card_and_no_table_only_in_the_coverage_block():
    wb, res = run()
    ws = wb["Highlights"]
    assert "World Mobile" not in "\n".join(res["lines"])
    cov_row = next(c.row for row in ws.iter_rows() for c in row if c.value == "Coverage")
    for c, sheet, coord in _refs(ws):
        name = wb[sheet].cell(wb[sheet][coord].row, 1).value
        if sheet != "Summary" or c.row < cov_row:
            assert name not in PARKED, (c.coordinate, c.value, name)          # no chip, no table row on a parked row
    for row in ws.iter_rows(max_row=cov_row - 1):
        for c in row:
            if isinstance(c.value, str):
                assert "World Mobile" not in c.value, c.value
    # ... and they ARE listed in the coverage block, with the code and the reason
    rows = [r for r in ws.iter_rows(min_row=cov_row) if any(isinstance(c.value, str) and c.value.startswith("=Summary!A")
                                                            for c in r)]
    parked = [r for r in rows if wb["Summary"][r[0].value[1:].split("!")[1]].value == "World Mobile"]
    assert len(parked) == 1
    cells = [c for c in parked[0] if c.value is not None]
    assert any(c.value.startswith("=Summary!F") for c in cells)               # the code is a formula on the Summary cell
    assert any(isinstance(c.value, str) and c.value.startswith("PARKED — no further work on World Mobile")
               for c in cells)
    # every unranked project is listed with its reason (first sentence of the comment), none ranked
    listed = {wb["Summary"][r[0].value[1:].split("!")[1]].value for r in rows}
    assert listed == {"Uniswap", "Near", "Ethereum", "Plume", "World Mobile"}
    why = next(c.value for r in rows for c in r if isinstance(c.value, str) and c.value.startswith("n/a — customer"))
    assert why == "n/a — customer_revenue_usd (q0): no rows stored."


def test_synthetic_coverage_block_copies_abouts_table_by_formula():
    wb, res = run()
    ws, ab = wb["Highlights"], wb["About"]
    cov_row = next(c.row for row in ws.iter_rows() for c in row if c.value == "Coverage")
    title = next(c for c in ws[cov_row + 1] if c.value)
    assert title.value == "=About!A11"
    got = {}
    for row in ws.iter_rows(min_row=cov_row + 2):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=About!"):
                got[c.value] = c
    for r in range(12, 16):
        assert f"=About!A{r}" in got, (r, sorted(got))
    assert "=About!B13" in got and "=About!E14" in got
    assert got["=About!B13"].number_format == ab["B13"].number_format == "0"
    wb2, _ = run(over=None)
    X = H.xlcalc.Workbook(wb2)
    assert X.value("Highlights", got["=About!A15"].row, got["=About!A15"].column) == "All archetypes"


def test_synthetic_c8_sorts_at_build_time_and_leaves_parked_out():
    wb, res = run()
    ws = wb["Highlights"]
    X = H.xlcalc.Workbook(wb)
    row = next(c.row for r in ws.iter_rows() for c in r if c.value == "Dilution league")
    names = [[X.value("Highlights", row + 2 + i, c) for i in range(3)] for c in (1, 10)]
    vals = [[X.value("Highlights", row + 2 + i, c) for i in range(3)] for c in (6, 15)]
    assert names[0] == ["Ethereum", "Uniswap", "Pendle"] and vals[0] == [0.014, 0.036, 0.039]
    assert names[1] == ["Aethir", "Hyperliquid", "Plume"] and vals[1] == [0.288, 0.227, 0.208]
    assert "World Mobile" not in names[0] + names[1]
    # the values are formulas with the source format; the names formulas on the Comparison project cells
    for c in (6, 15):
        for i in range(3):
            cell = ws.cell(row + 2 + i, c)
            assert cell.value.startswith("=Comparison!E") and cell.number_format == "0.0%"
    assert all(ws.cell(row + 2 + i, 1).value.startswith("=Comparison!A") for i in range(3))


def test_synthetic_a_chip_that_evaluates_to_an_error_or_blank_fails_the_build():
    wb = make_reader()
    wb["A4 Permanent Burn"]["D5"] = "=1/0"                               # Uniswap's permanent burn yield: #DIV/0!
    with pytest.raises(H.HighlightsError) as e:
        run(wb=wb)
    assert "#DIV/0!" in str(e.value) and "C5 Uniswap" in str(e.value)
    wb = make_reader()
    wb["A4 Permanent Burn"]["D5"] = '=""'                                # a blank
    with pytest.raises(H.HighlightsError) as e:
        run(wb=wb)
    assert "C5 Uniswap" in str(e.value) and "evaluates to ''" in str(e.value)
    wb = make_reader()
    # a chip of C7 (Near's TVL change) that is #NAME?
    r = next(r for r in range(5, wb["Business"].max_row + 1) if wb["Business"].cell(r, 1).value == "Near")
    wb["Business"].cell(r, 3).value = "=NOSUCHFN(1)"
    with pytest.raises(H.HighlightsError):
        run(wb=wb)
    # a chip that quotes an EMPTY source cell is not rendered at all (the card is omitted, never a blank chip)
    wb = make_reader()
    r = next(r for r in range(5, wb["A4 Permanent Burn"].max_row + 1) if wb["A4 Permanent Burn"].cell(r, 1).value == "Uniswap")
    wb["A4 Permanent Burn"].cell(r, 4).value = None
    _, res = run(wb=wb, min_cards=1)
    assert "C5" not in res["cards"] and any(o.startswith("C5") and "empty" in o for o in res["omitted"])


def test_synthetic_a_code_in_a_chip_is_shown_as_the_code_not_as_a_number():
    wb, res = run(over={"Morpho": {"tvl": "n/a"}})
    assert "C2" in res["cards"]
    assert any(ln.strip() == "TVL change, selected period: n/a" for ln in res["lines"])


# ================================================================================================= the hooks
def test_the_flag_is_in_the_command_line_and_defaults_to_none():
    import token_metrics as tm
    assert tm.parse_args(["--no-fetch"]).edition is None
    assert tm.parse_args(["--no-fetch", "--edition", "charlie"]).edition == "charlie"
    with pytest.raises(SystemExit):
        tm.parse_args(["--edition", "dave"])


def test_the_look_is_called_when_charlie_look_exists_and_skipped_with_a_line_when_it_does_not(monkeypatch):
    import types
    wb = make_reader()
    monkeypatch.setattr(reader, "parked_reason", lambda n: "parked" if n in PARKED else "")
    monkeypatch.setitem(sys.modules, "charlie_look", None)                 # import raises ImportError
    out = charlie_edition.apply_edition(wb, None, "charlie")
    assert "not present" in out["look"] and wb.sheetnames[0] == "Highlights"
    assert any("not present" in ln for ln in charlie_edition.report_lines("x"))
    calls = []
    fake = types.SimpleNamespace(apply_look=lambda w: calls.append(w.sheetnames[0]))
    monkeypatch.setitem(sys.modules, "charlie_look", fake)
    wb2 = make_reader()
    out2 = charlie_edition.apply_edition(wb2, None, "charlie")
    assert calls == ["Highlights"] and "applied" in out2["look"]       # called after the tab is added
    with pytest.raises(ValueError):
        charlie_edition.apply_edition(make_reader(), None, "dave")


def test_a_second_highlights_sheet_is_refused():
    wb, _ = run()
    with pytest.raises(H.HighlightsError):
        H.add_highlights(wb, None, parked=lambda n: False)


# ============================================================================================ the fixture builds
def _build(d: Path, edition):
    """The reader (and, with `edition`, the edition) built on the committed real dumps, as test_round_p1 builds it."""
    import build_workbook as bw
    import config
    import measuring_guard as mg
    import store as sm
    from real_dump import to_store
    getattr(config, "RESTATEMENT_ACTIVE", {}).clear()
    db = d / "metrics.db"
    asofs = []
    for pl in mg._payloads():
        to_store(pl, db)
        asofs.append(pl["meta"]["asof"])
    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(sm, "DB_PATH", str(db))
        mp.setitem(sys.modules, "recalc", None)
        mp.setenv("TOKEN_METRICS_STORE_VALUES", "0")
        mp.setattr(charlie_edition, "MIN_CARDS", FIXTURE_CARDS)
        mp.delitem(sys.modules, "charlie_look", raising=False)
        st = sm.Store(str(db))
        try:
            bw.build_workbook(st, d / "audit.xlsx", asof=pd.Timestamp(max(asofs)), reader_path=d / "reader.xlsx",
                              edition=edition)
        finally:
            st.close()
    assert "error" not in bw.READER_STATS, bw.READER_STATS
    getattr(config, "RESTATEMENT_ACTIVE", {}).clear()
    return {"dir": d, "stats": dict(bw.READER_STATS), "locations": dict(reader.HEADLINE_LOCATIONS),
            "column_map": len(reader.COLUMN_MAP), "last": dict(charlie_edition.LAST)}


@pytest.fixture(scope="module")
def plain(tmp_path_factory):
    d = tmp_path_factory.mktemp("charlie_plain")
    charlie_edition.LAST.clear()
    return _build(d, None)


@pytest.fixture(scope="module")
def flagged(tmp_path_factory):
    d = tmp_path_factory.mktemp("charlie_flag")
    return _build(d, "charlie")


def _cells(path: Path, data_only: bool = False) -> dict:
    wb = openpyxl.load_workbook(path, data_only=data_only)
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None or c.comment is not None:
                    out[(ws.title, c.coordinate)] = (c.value, c.number_format, c.comment.text if c.comment else None)
        out[(ws.title, "#dims")] = (ws.max_row, ws.max_column, ws.sheet_state)
    return out, wb.sheetnames


def test_a_invariance_the_flag_changes_nothing_on_the_normal_files(plain, flagged):
    """With and without --edition: every cell value / formula / number format / comment of every reader tab is
    identical, the normal reader has no Highlights tab, and the audit file's tabs are the same."""
    a, names_a = _cells(plain["dir"] / "reader.xlsx")
    b, names_b = _cells(flagged["dir"] / "reader.xlsx")
    assert names_a == names_b and "Highlights" not in names_b
    assert a.keys() == b.keys()
    diff = [k for k in a if a[k] != b[k]]
    assert not diff, diff[:10]
    av, _ = _cells(plain["dir"] / "reader.xlsx", data_only=True)
    bv, _ = _cells(flagged["dir"] / "reader.xlsx", data_only=True)
    assert [k for k in av if av[k] != bv[k]] == []
    assert len(a) > 20000
    # the reader's module state and stats are the normal build's, not the edition's second build's
    assert plain["locations"] == flagged["locations"] and plain["column_map"] == flagged["column_map"]
    assert plain["stats"]["tabs"] == flagged["stats"]["tabs"] and "Highlights" not in flagged["stats"]["tabs"]
    assert plain["stats"]["numbers"] == flagged["stats"]["numbers"]
    pa = openpyxl.load_workbook(plain["dir"] / "audit.xlsx", read_only=True).sheetnames
    pb = openpyxl.load_workbook(flagged["dir"] / "audit.xlsx", read_only=True).sheetnames
    assert pa == pb
    assert (plain["dir"] / "reader_column_map.csv").read_bytes() == (flagged["dir"] / "reader_column_map.csv").read_bytes()


def test_b_the_flag_writes_the_edition_with_highlights_first_and_no_second_audit_file(plain, flagged):
    assert not (plain["dir"] / charlie_edition.CHARLIE_NAME).exists()
    edition = flagged["dir"] / charlie_edition.CHARLIE_NAME
    assert edition.exists()
    assert sorted(p.name for p in flagged["dir"].glob("*.xlsx")) == sorted(["audit.xlsx", "reader.xlsx", edition.name])
    wb = openpyxl.load_workbook(edition)
    normal = openpyxl.load_workbook(flagged["dir"] / "reader.xlsx")
    assert wb.sheetnames[0] == "Highlights" and wb.sheetnames[1:] == normal.sheetnames
    assert wb["Highlights"].sheet_properties.tabColor.rgb.endswith("1F8A8A")
    # charts are made, not round-tripped: the Charts tab of the edition holds them
    assert len(wb["Charts"]._charts) == len(normal["Charts"]._charts) > 0
    # the edition's existing tabs are the normal reader's, cell for cell
    a, _ = _cells(flagged["dir"] / "reader.xlsx")
    b, _ = _cells(edition)
    assert {k: v for k, v in b.items() if k[0] != "Highlights"} == a
    last = flagged["last"]
    assert last["cards"] == ["C5", "C7", "C8"], last["cards"]        # the fixture holds other numbers than Jake's store
    why = {o.split(" ")[0] for o in last["omitted"]}
    assert why == {"C1", "C2", "C3", "C4", "C6"}, last["omitted"]


def test_fixture_cards_every_number_is_a_formula_and_the_text_is_clean(flagged):
    edition = flagged["dir"] / charlie_edition.CHARLIE_NAME
    wb = openpyxl.load_workbook(edition)
    ws = wb["Highlights"]
    refs = _refs(ws)
    assert len(refs) >= 16
    for c, sheet, coord in refs:
        if not re.search(r"!A\d+$", c.value):                        # a project's name is text: General, not '@'
            assert wb[sheet][coord].number_format == c.number_format, (c.coordinate, c.value)
    for row in ws.iter_rows():
        for c in row:
            assert not isinstance(c.value, (int, float)), (c.coordinate, c.value)
    strings = H.all_strings(ws)
    for s in strings:
        assert not H.FORBIDDEN.search(s), s
    # the cached values are in the file and none is an error
    wv = openpyxl.load_workbook(edition, data_only=True)["Highlights"]
    for c, sheet, coord in refs:
        v = wv[c.coordinate].value
        assert v is not None and v != "" and not (isinstance(v, str) and v.startswith("#")), (c.coordinate, v)
    # parked projects only in the coverage block
    cov = next(c.row for row in ws.iter_rows() for c in row if c.value == "Coverage")
    parked = [n for n in ("Aerodrome", "World Mobile") if reader.parked_reason(n) or n == "World Mobile"]
    for row in ws.iter_rows(max_row=cov - 1):
        for c in row:
            if isinstance(c.value, str):
                assert not any(p in c.value for p in parked), c.value
    for c, sheet, coord in refs:
        if c.row < cov:
            assert wb[sheet].cell(wb[sheet][coord].row, 1).value not in parked, (c.coordinate, c.value)
    listed = {wv.cell(c.row, c.column).value for row in ws.iter_rows(min_row=cov) for c in row
              if isinstance(c.value, str) and c.value.startswith("=Summary!A")}
    assert {"Aerodrome", "World Mobile"} <= listed


def test_d_the_production_floor_of_five_fails_on_the_fixture_numbers_and_lists_the_conditions(flagged):
    wb = openpyxl.load_workbook(flagged["dir"] / "reader.xlsx")            # the normal reader: no Highlights yet
    with pytest.raises(H.HighlightsError) as e:
        H.add_highlights(wb, None)                                          # min_cards = 5
    msg = str(e.value)
    assert "only 3 of 8" in msg and "5 are needed" in msg
    for cid in ("C1 GEODNET", "C2 Morpho", "C3", "C4 Hyperliquid", "C6 Fluid"):
        assert cid in msg, msg
