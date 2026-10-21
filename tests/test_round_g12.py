"""
tests/test_round_g12.py — Round G12, "wiring. nothing new is fetched" (Jake 2026-10-21), SPEC §19 "G12". The wiring
audit (A1) and the empty-cell rows (A2), the series wired to a cell, a cross-check or a declared reason (B), and the
items carried over from G11 (C1–C8).
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import build_workbook as bw  # noqa: E402
import config  # noqa: E402
import wiring  # noqa: E402


def _g(project, metric, days, values, source="x"):
    return pd.DataFrame({"date": pd.to_datetime(list(days)), "project": project, "metric": metric,
                         "value": [float(v) for v in values], "tier": 1, "source": source, "is_manual": False,
                         "entered_on": ""})


def _row(project, metric, now=1.0, **kw):
    return {"project": project, "metric": metric, "now": now, "label": metric, "latest_date": "2026-10-13", **kw}


# ================================================================================================== A1 the audit
def test_a1_reads_and_writes_are_attributed_to_the_view_and_its_helpers():
    wiring.reset()
    g = wiring.TrackingDict({("P", "a"): 1, ("P", "b"): 2})

    def _helper(groups):
        return groups.get(("P", "a"))

    def _some_view(groups):
        _helper(groups)
        groups[("P", "c")] = 3                       # written in the view's body, read in a helper
    _some_view(g)
    edges = wiring.view_edges()
    assert ("a", "_some_view") in edges[("P", "c")]


def test_a1_classes_shown_cross_input_not_shown_unwired(monkeypatch):
    wiring.reset()
    g = wiring.TrackingDict({("P", "parent"): 1})

    def _builder(groups):
        groups.get(("P", "parent"))
        groups[("P", "child")] = 2
    _builder(g)
    d = wiring.TrackingDict({"P|child": {}, "P|crossed": {}, "P|styled": {}})
    with wiring.phase("shown"):
        def _cell():
            return d.get("P|child")
        _cell()

        def _write_table():                          # a styling read: no cell shows it
            return d.get("P|styled")
        _write_table()
    with wiring.phase("cross"):
        def _cred():
            return d.get("P|crossed")
        _cred()
    monkeypatch.setitem(config.NOT_SHOWN, ("P", "declared"), "a reason")
    rows = [_row("P", m) for m in ("child", "parent", "crossed", "declared", "styled", "orphan")]
    rows.append(_row("P", "empty", now=None))         # no number in any window: not audited
    out = wiring.audit(rows)
    cls = {m: v["class"] for (_p, m), v in out.items()}
    assert cls == {"child": "SHOWN", "parent": "INPUT", "crossed": "CROSS-CHECK", "declared": "NOT SHOWN",
                   "styled": "UNWIRED", "orphan": "UNWIRED"}
    assert "child (_builder)" in out[("P", "parent")]["why"]
    assert wiring.summary(out)["unwired"] == 2


def test_a1_a_formula_on_a_display_tab_and_on_credibility_counts():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "A1 Infrastructure"
    ws["A5"] = "Near"
    ws["C5"] = '=INDEX(Data!$J$2:$J$9,MATCH($A5&"|tx_count",Data!$A$2:$A$9,0))'
    cr = wb.create_sheet("Credibility")
    cr["A7"] = "Near"
    cr["E7"] = '=IFERROR(INDEX(Data!$J$2:$J$9,MATCH($A7&"|fees_usd",Data!$A$2:$A$9,0)),NA())'
    shown, cross = wiring.formula_refs(wb)
    assert set(shown) == {("Near", "tx_count")} and set(cross) == {("Near", "fees_usd")}
    assert shown[("Near", "tx_count")] == {"A1 Infrastructure"}


def test_a1_a_series_named_only_in_prose_is_not_read():
    rows = [{"project": "P", "ref": {"metric": "used", "why": "mentions unused in a sentence", "note": "unused"}}]
    assert wiring.credibility_refs(rows, {("P", "used"), ("P", "unused")}) == {("P", "used")}


def test_a1_fetch_inputs_are_edges_and_each_module_really_derives_the_child():
    for child, parents, module, _projects in config.FETCH_INPUTS:
        src = (ROOT / module).read_text(encoding="utf-8")
        assert child in src and all(p in src for p in parents), (child, parents, module)
    wiring.reset()
    out = wiring.audit([_row("Ethereum", "gross_burn_tokens"), _row("Ethereum", "burn_cumulative_tokens")])
    # gross_burn_tokens is not shown in this synthetic audit, so its parent is not an input either
    assert out[("Ethereum", "burn_cumulative_tokens")]["class"] == "UNWIRED"
    d = wiring.TrackingDict({"Ethereum|gross_burn_tokens": {}})
    with wiring.phase("shown"):
        def _cell():
            return d.get("Ethereum|gross_burn_tokens")
        _cell()
    out = wiring.audit([_row("Ethereum", "gross_burn_tokens"), _row("Ethereum", "burn_cumulative_tokens")])
    assert out[("Ethereum", "burn_cumulative_tokens")]["class"] == "INPUT"
    assert "fetch: fetch/etherscan_supply.py" in out[("Ethereum", "burn_cumulative_tokens")]["why"]


def test_a1_not_shown_rules_follow_config():
    assert "GUESSED" in config.not_shown_reason("Aethir", "price_usd_coinbase")
    assert config.not_shown_reason("Ethereum", "price_usd_coinbase") == ""          # a declared product
    assert "on-chain-first" in config.not_shown_reason("Chainlink", "circulating_supply_implied")
    assert config.not_shown_reason("World Mobile", "circulating_supply_implied") == ""
    assert "circulating_supply_implied" in config.fetch_inputs("World Mobile")["circulating_supply"][0]


def test_a1_check_and_run_log_line(tmp_path):
    out = {("Near", "tx_count_kaiching"): {"class": "UNWIRED", "why": "", "label": "Kai-Ching tx", "latest": 696799.0,
                                          "latest_date": "2026-10-12"},
           ("Near", "tx_count"): {"class": "SHOWN", "why": "", "label": "tx", "latest": 1.0, "latest_date": ""}}
    wiring.write(out, tmp_path / "wiring.json")
    text = wiring.report(tmp_path / "wiring.json", unwired_only=True)
    assert "UNWIRED 1" in text and "tx_count_kaiching" in text and "696,799" in text and "  SHOWN" not in text
    line = " ".join(bw.wiring_lines(out))
    assert "UNWIRED 1" in line
    import check_offline_items as coi
    assert coi.wiring in coi.CHECKS


# ================================================================================================== A2 empty cells
def test_a2_a_row_on_a_cell_empty_by_design_says_so():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CP Comparability"
    ws["CW6"] = None
    ws["CT6"] = "=1"
    assert bw._empty_by_design(wb, "'CP Comparability'!CW6")
    assert not bw._empty_by_design(wb, "'CP Comparability'!CT6")
    assert not bw._empty_by_design(wb, "'Nope'!A1")
    import inspect
    src = inspect.getsource(bw.write_credibility)
    assert '"no figure by design"' in src and "N/A (no figure by design)" in src


# ================================================================================================== B wiring
JAKE_NAMED = (
    "volume_usd fdv_usd_coingecko total_supply_coingecko own_chain_tx_count own_chain_active_addresses "
    "own_chain_fees_native_tokens tx_count_ex_kaiching tx_count_kaiching future_emissions_tokens hypercore_users_total "
    "hypercore_new_users tx_count_successful open_interest_usd_asxn latent_supply_tokens vependle_voting_supply_tokens "
    "compute_service_fee_tokens compute_service_fee_cumulative_tokens compute_rewards_tokens compute_power_total "
    "supply_locations supply_units_checker_licences capacity_monthly_total utilisation_cumulative_pct "
    "utilisation_containers_pct arr_tile_usd customer_revenue_monthly_usd customer_revenue_cumulative_usd "
    "issuance_feed_polygon_tokens issuance_feed_solana_tokens ccip_token_transfers ccip_v2_fees_usd ccip_v2_premium_usd "
    "ssf_holdings_tokens ssf_net_outflow_tokens liquidation_bonus_usd_day_morpho_api mkr_outstanding_tokens "
    "sky_mkr_nonissuer_backing_tokens net_revenue_msc_usd usds_farm_reward_rate_usds_per_s sethfi_topup_tokens "
    "sethfi_aps_rebuilt holders_revenue_usd_defillama programmatic_buyback_tokens fees_usd_defillama fees_native_tokens "
    "near_intents_volume_usd_nearblocks mev_relay_block_share nft_volume_usd p2p_transfer_volume_usd "
    "tvl_maturity_le_30d_usd tvl_maturity_31_90d_usd tvl_maturity_91_180d_usd tvl_maturity_gt_180d_usd "
    "emissions_tokens_gauge_rate_ethereum emissions_tokens_gauge_rate_hyperevm").split()


def _wired_names() -> set:
    import inspect

    import credibility
    names: set = set()
    for d in config.DETAIL_COLUMNS:
        names |= {m for m, _w in d["cells"].values()}
    wiring._strings(config.CREDIBILITY_WIRING, names)
    src = inspect.getsource(credibility.wiring_inputs) + inspect.getsource(bw.write_master)
    names |= {m for m in config.METRICS if f'"{m}"' in src}
    for _c, parents, _m, _p in config.FETCH_INPUTS:
        names |= set(parents)
    return names


def test_b_every_series_jake_named_is_wired_by_code():
    missing = [m for m in JAKE_NAMED if m not in _wired_names()]
    assert not missing, missing


def test_b_detail_columns_are_well_formed_and_map_to_a_spec_concept():
    import spec
    for d in config.DETAIL_COLUMNS:
        assert d["header"].startswith("Detail — ") and spec.concept_for(d["tab"], d["header"]) == "detail"
        assert d["tab"] in ("A1", "A2", "A3", "A4", "CP")
        for p, (m, w) in d["cells"].items():
            assert p in config.PROJECT_BY_NAME and m in config.METRICS and w in ("now", "q0", "m1"), (p, m, w)
    assert spec.concept_for("Master", "Trading volume, 30d ($) — CoinGecko spot, all venues") == "liquidity"
    assert spec.concept_for("Master", "Turnover = 30d volume ÷ market cap (x)") == "liquidity"
    assert all("Aerodrome" not in d["cells"] for t in ("A1", "A2", "A3", "A4", "CP") for d in config.detail_columns(t))


def test_b_a_detail_column_sits_beside_its_anchor_and_reads_its_project_only():
    R = bw.Refs(10, 0, [])
    specs = [("Project", None, None, "text"), ("Transactions Q0", None, None, "pull"), ("Fees Q0 ($)", None, None, "pull"),
             ("Notes", None, None, "text")]
    out = bw._with_details("A1", specs, R)
    heads = [s[0] for s in out]
    i = heads.index("Transactions Q0")
    assert heads[i + 1].startswith("Detail — ") and heads[-1] == "Notes"
    organic = next(s for s in out if s[0].startswith("Detail — Transactions Q0, ORGANIC"))
    assert "tx_count_ex_kaiching" in organic[1](5, {"name": "Near"})
    assert organic[1](5, {"name": "Ethereum"}) == ""
    overhang = next(s for s in out if s[0].startswith("Detail — Supply overhang, % of total supply"))
    f = overhang[1](5, {"name": "Hyperliquid"})
    assert "future_emissions_tokens" in f and "total_supply" in f and f.startswith("=IFERROR(")


def test_b2_world_mobile_fees_reach_the_cell_master_references():
    R = bw.Refs(10, 0, [])
    head, build, _fmt, _k, _b, meta = bw._fees_q0_col(R)
    assert head == "Fees Q0 ($)" and "own_chain_fees_usd" in build(5, {"name": "World Mobile"})
    assert meta["fmt_fn"]("World Mobile").endswith('" (own-chain fees)"') and meta["fmt_fn"]("Near") is None
    import inspect
    for fn in (bw.write_a1, bw.write_a3, bw.write_a4, bw.write_master):
        assert "_fees_q0_col(R)" in inspect.getsource(fn), fn.__name__


def test_b_cross_check_rows_only_where_the_second_source_holds_a_number():
    import credibility as cred
    rows = {"Hyperliquid|open_interest_usd_asxn": {"now": 1.5e10}, "Hyperliquid|total_supply_coingecko": {"now": 9.5e8},
            "Hyperliquid|actual_buyback_usd": {"q0": 1.8e8}, "Hyperliquid|holders_revenue_usd": {"q0": 1.7e8}}
    out = cred.wiring_inputs("Hyperliquid", rows)
    assert "in_oi_asxn" in out and "in_total_supply_coingecko" in out
    assert "in_fdv_coingecko" not in out                     # CoinGecko's FDV holds no number: nothing to compare
    assert "in_buyback_asxn" not in out                      # its reference holds no number here
    assert out["in_total_supply_coingecko"]["ref"]["tol"] == 1.0
    assert out["in_oi_asxn"]["ref"]["force_verdict"] == "N/A (recorded)"
    assert "in_holders_revenue_llama" in out
    # Aerodrome is parked: only B1's two CoinGecko rows
    aero = cred.wiring_inputs("Aerodrome", {"Aerodrome|fdv_usd_coingecko": {"now": 1.8e9},
                                            "Aerodrome|holders_revenue_usd": {"q0": 2.4e7}})
    assert set(aero) == {"in_fdv_coingecko"}


def test_b_a_recorded_reference_shows_its_value_and_never_judges():
    import credibility as cred
    spec = config._c_rec("windows differ", metric="arr_tile_usd", window="now")
    ref = cred.reference("Aethir", spec, {"Aethir|arr_tile_usd": {"now": 62.49e6, "latest_date": "2026-10-13"}},
                         None, pd.Timestamp("2026-10-13"))
    assert ref["value"] == pytest.approx(62.49e6) and ref["verdict"] == "N/A (recorded)" and ref["mode"] == "static"
    assert "recorded, not judged" in ref["note"]


def test_b_style_reads_never_count_as_shown():
    assert {"_write_table", "_flags", "_confidence_tally", "_reasoned_na"} <= wiring.STYLE_FUNCS


# ================================================================================================== C carried over
def test_c1_guard_compares_the_previous_snapshot_key_for_key(tmp_path, monkeypatch):
    """END TO END (Jake's 10-21 store: 329 HEADLINE LOST, "Ethereum Master | Price — spot ($)" among them): a real-dump
    store, the previous snapshot written as the daily run writes it (export_snapshot.snapshot_rows -> write_csv), then
    measuring_guard.live_check on the same store — every Master and A-tab figure the snapshot holds is found."""
    import openpyxl
    import xlcalc

    import export_snapshot as es
    import measuring_guard as mg
    import store as sm
    from real_dump import load, to_store
    payload = load("Chainlink")
    db = tmp_path / "metrics.db"
    to_store(payload, db)
    monkeypatch.setitem(sys.modules, "recalc", None)
    monkeypatch.setattr(bw, "_SCOPE", bw._SCOPE)
    monkeypatch.setattr(config, "read_portfolio", lambda _path: (["Chainlink"], []))
    import sqlite3
    con = sqlite3.connect(str(db))
    today = pd.Timestamp.now("UTC").tz_localize(None).normalize()
    newest = con.execute("SELECT MAX(date) FROM metrics WHERE date <= ?", (str(today.date()) + "T23:59:59",)).fetchone()[0]
    con.close()
    asof = min(pd.Timestamp(newest).normalize(), today)
    monkeypatch.setattr(sm, "DB_PATH", str(db))
    st = sm.Store(str(db))
    try:
        bw.build_workbook(st, tmp_path / "snap.xlsx", asof=asof, only=["Chainlink"])
    finally:
        st.close()
    wb = openpyxl.load_workbook(tmp_path / "snap.xlsx")
    X = xlcalc.Workbook(wb)
    rows = es.snapshot_rows(wb, lambda s, r, c: xlcalc.display(X.value(s, r, c)), [], [], list(bw._HEADLINE_CELLS),
                            asof)
    snaps = tmp_path / "snapshots"
    es.write_csv(rows, snaps / f"snapshot_{(asof - pd.Timedelta(days=1)).date()}.csv")
    monkeypatch.setattr(es, "SNAPSHOT_DIR", snaps)
    _d, prev = es.load_previous(asof - pd.Timedelta(days=1))
    held = {k for k, v in prev["Chainlink"].items() if isinstance(v, float)}
    assert any(k.startswith("Master | Price") for k in held) and any(k.startswith("A3 | ") for k in held)
    bad, snap = mg.live_check(str(db), {})
    lost = [b for b in bad if b.startswith("HEADLINE LOST")]
    assert not lost, lost[:5]
    assert held <= set(snap["headlines"]["Chainlink"])


def test_c2_a_declared_zero_holds_the_headline():
    import credibility as cred
    cells = [{"project": "Plume", "sheet": "A4 Permanent Burn", "cell": "AS9", "header": "GROSS ISSUANCE Q0 (tokens)"}]
    row = cred.continuity_row("Plume", cells, {"Plume": {"A4 | GROSS ISSUANCE Q0 (tokens)": 0.0}}, "2026-10-20")
    assert 'LEFT(\'A4 Permanent Burn\'!AS9,3)="0 —"' in row["ours"]["cell"]
    import inspect
    assert "sanity.stated_number(now) is None" in inspect.getsource(bw.write_credibility)
    import sanity
    assert sanity.stated_number("0 — fixed supply: 1bn minted at genesis") == 0.0


def test_c3_the_daily_twin_is_compared_only_where_both_series_cover():
    ends = pd.date_range("2025-09-30", "2026-09-30", freq="ME")
    daily_ends = pd.date_range("2024-11-30", "2026-09-30", freq="ME")    # the by-day read reaches back further
    groups = {("Sky", "buyback_spending_usd_ba"): _g("Sky", "buyback_spending_usd_ba", ends, [10.0] * len(ends)),
              ("Sky", "buyback_spending_usd_ba_daily"): _g("Sky", "buyback_spending_usd_ba_daily", daily_ends,
                                                           [10.0] * len(daily_ends))}
    bw._daily_twin_views(groups, pd.Timestamp("2026-10-21"))
    assert ("Sky", "buyback_spending_usd_ba_daily") not in groups
    # a REAL daily series (one row a day) is never dropped
    days = pd.date_range("2025-09-01", "2026-10-20")
    groups = {("Sky", "buyback_spending_usd_ba"): _g("Sky", "buyback_spending_usd_ba", ends, [300.0] * len(ends)),
              ("Sky", "buyback_spending_usd_ba_daily"): _g("Sky", "buyback_spending_usd_ba_daily", days,
                                                           [7.0] * len(days))}
    bw._daily_twin_views(groups, pd.Timestamp("2026-10-21"))
    assert ("Sky", "buyback_spending_usd_ba_daily") in groups


def test_c4_a_365_day_cell_with_a_held_hole_is_not_a_stored_but_unshown_hit(tmp_path, monkeypatch, capsys):
    import csv
    import sqlite3

    import check_offline_items as coi
    import store as sm
    db = tmp_path / "metrics.db"
    con = sqlite3.connect(str(db))
    con.execute("CREATE TABLE metrics (date TEXT, project TEXT, metric TEXT, value REAL, source TEXT, tier INT, "
                "fetched_at TEXT)")
    days = list(pd.date_range("2025-10-21", "2026-10-21"))
    held = days[:100] + days[-103:]                         # 203 of 366 days, spanning both ends
    con.executemany("INSERT INTO metrics VALUES (?, 'Near', 'dex_volume_usd', 1.0, 'defillama', 1, '')",
                    [(str(d.date()),) for d in held])
    con.commit()
    con.close()
    path = tmp_path / "coverage_audit.csv"
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["Project", "Tab", "Column", "Status", "Reason / route / error", "Value now"])
        w.writeheader()
        w.writerow({"Project": "Near", "Tab": "A1", "Column": "Settlement volume, trailing 365d ($)",
                    "Status": "MISSING", "Value now": "",
                    "Reason / route / error": "MISSING [D] — dex_volume_usd — the route stored nothing"})
    monkeypatch.setenv("TOKEN_METRICS_COVERAGE_CSV", str(path))
    monkeypatch.setattr(sm, "DB_PATH", str(db))
    coi.stored_but_unshown()
    assert "0 hit(s)" in capsys.readouterr().out


def test_c5_a_day_without_its_eth_row_is_asked_again_and_an_incomplete_price_is_not_cached():
    from fetch.eth_settlement import EthP2PBigQuery
    e = {"days": {"2026-09-19": {"ETH": "1", "0xa0b8": "2"}, "2026-09-20": {"0xa0b8": "2"},
                  "2026-09-21": {}}, "prices": {}}
    assert EthP2PBigQuery.held_days(e) == {"2026-09-19"}
    p2p = {"validation_extend_from": "2026-09-14", "validation_from": "2026-09-14"}
    held = {str(d.date()) for d in pd.date_range("2026-09-14", "2026-10-18")} - {"2026-09-20", "2026-09-21",
                                                                                  "2026-09-22", "2026-09-23"}
    weeks = EthP2PBigQuery.pieces(held, p2p, pd.Timestamp("2026-10-21"))
    assert (pd.Timestamp("2026-09-14"), pd.Timestamp("2026-09-20")) in weeks
    assert (pd.Timestamp("2026-09-21"), pd.Timestamp("2026-09-27")) in weeks and len(weeks) == 2
    import inspect
    src = inspect.getsource(EthP2PBigQuery._value_eth)
    assert 'if all((px.get(f"ethereum:{t}") or (None,))[0] is not None for t in stables):' in src


def test_c6_the_buyback_stitch_is_a_series_break(monkeypatch):
    import trajectory
    days_h = pd.date_range("2026-06-01", "2026-09-21")
    days_l = pd.date_range("2026-09-22", "2026-10-13")
    g = pd.concat([_g("Hyperliquid", "actual_buyback_tokens", days_h, [1.0] * len(days_h),
                      "defillama:holders_revenue_usd/price"),
                   _g("Hyperliquid", "actual_buyback_tokens", days_l, [1.0] * len(days_l),
                      "hypercore_info:spotClearinghouseState:delta:as-buyback")], ignore_index=True)
    monkeypatch.setattr(bw, "_LAST_GROUPS", {("Hyperliquid", "actual_buyback_tokens"): g})
    br = bw._stitch_breaks("Hyperliquid", "actual_buyback_usd")
    assert br and br[0][0] == "2026-09-22" and "SERIES BREAK" in br[0][1]
    breaks = trajectory._own_breaks("Hyperliquid", "actual_buyback_tokens", {("Hyperliquid", "actual_buyback_tokens"): g},
                                    pd.Timestamp("2026-01-01"))
    assert [str(b[0].date()) for b in breaks] == ["2026-09-22"]
    # a benign handover ("none …") stays one series
    assert bw._stitch_breaks("Hyperliquid", "circulating_supply") == []
    # the $ series and the retired tokens are DERIVED from the tokens: they carry its declared step
    usd = g.assign(metric="actual_buyback_usd", source="derived:tokens*price")
    groups = {("Hyperliquid", "actual_buyback_tokens"): g, ("Hyperliquid", "actual_buyback_usd"): usd}
    assert trajectory._stepping_parents("Hyperliquid", "actual_buyback_usd", groups) == ("actual_buyback_tokens",)
    br = trajectory.basis_breaks("Hyperliquid", ["actual_buyback_usd"], groups, pd.Timestamp("2026-01-01"))
    assert [str(b[0].date()) for b in br] == ["2026-09-22"]
    # an undeclared parent is not followed (no new break reported for it)
    assert trajectory._stepping_parents("Uniswap", "actual_buyback_usd",
                                        {("Uniswap", "actual_buyback_usd"): usd.assign(project="Uniswap")}) == ()


def test_c7_the_chain_tier_retries_a_429_like_the_explorer_tier(monkeypatch, caplog):
    import requests

    from fetch import chain as ch
    key = "testkey0000000000000000000000000"
    prov = ch._guarded_provider_class(__import__("web3").HTTPProvider)
    p = prov(f"https://eth-mainnet.g.alchemy.com/v2/{key}", exception_retry_configuration=None)
    p.chain = "ethereum"
    calls, waits = [], []

    def post(*_a, **_k):
        calls.append(1)
        if len(calls) <= 3:
            r = requests.Response()
            r.status_code = 429
            r.headers["Retry-After"] = "7" if len(calls) == 2 else None or ""
            raise requests.exceptions.HTTPError("429 Client Error: Too Many Requests", response=r)
        return b'{"jsonrpc":"2.0","id":1,"result":"0x1"}'
    monkeypatch.setattr(p._request_session_manager, "make_post_request", post)
    monkeypatch.setattr(ch, "_sleep", waits.append)
    monkeypatch.setattr(ch.BREAKER, "check", lambda *a: None)
    with caplog.at_level("WARNING", logger="token_metrics.fetch.chain"):
        assert p._make_request("eth_blockNumber", b"{}") == b'{"jsonrpc":"2.0","id":1,"result":"0x1"}'
    assert len(calls) == 4 and waits == [2.0, 7.0, 8.0]      # 2 s, Retry-After 7 s (above 4), then 8 s
    assert key not in caplog.text and "eth-mainnet.g.alchemy.com" in caplog.text
    assert ch.HTTP_RETRIES == 4


def test_c8_a_series_younger_than_its_window_says_needs_history(monkeypatch):
    monkeypatch.setitem(bw._WINDOWS, "q0", ("2026-07-15", "2026-10-13"))
    monkeypatch.setitem(bw._WINDOWS, "q1", ("2026-04-16", "2026-07-15"))
    R = bw.Refs(10, 0, [])
    f = f"=IFERROR({R.D(5, 'supply_units', 'now')}/{R.D(5, 'supply_units', 'q1')}-1,{bw.NA})"
    rows = {"GEODNET|supply_units": {"now": 20000.0, "q1": None, "status": "ok", "first_date": "2026-08-14"}}
    out = bw._reasoned_na(f, {"name": "GEODNET"}, rows)
    assert "NEEDS HISTORY: supply_units is held from 2026-08-14" in out and "valid from 2027-02-10" in out
    import coverage_audit as ca
    assert any(rx.search("n/a — NEEDS HISTORY: supply_units is held from 2026-08-14") and st == "MISSING C"
               for rx, st in ca._HONEST_RULES)
