"""Item 17 of Jake's review of 97b29ae (2026-10-11): permanent sanity rules over the built numbers."""
from __future__ import annotations

import openpyxl

import sanity

HEADS = ["Project", "Price — spot ($)", "Market cap ($)", "Circulating supply (used)", "FDV ($)", "Revenue Q0 ($)",
         "Actual buyback Q0 ($) — observed", "Buyback status", "Burn status", "Gross burn Q0 (tokens)",
         "NET SUPPLY CHANGE Q0 (tokens) = issuance − burn", "Net supply change, annualised % of circulating"]


def _wb(rows: list, totals: dict):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master"
    for c, h in enumerate(HEADS, start=1):
        ws.cell(4, c, h)
    for r, row in enumerate(rows, start=5):
        for c, h in enumerate(HEADS, start=1):
            ws.cell(r, c, row.get(h))
    d = wb.create_sheet("Data")
    d.cell(1, 1, "Key")
    d.cell(1, 2, "Now (flow: trailing)")
    for i, (p, v) in enumerate(totals.items(), start=2):
        d.cell(i, 1, f"{p}|total_supply")
        d.cell(i, 2, v)
    return wb


def test_every_rule_fires_on_its_case_and_a_clean_row_passes():
    clean = {"Project": "Clean", "Price — spot ($)": 2.0, "Market cap ($)": 200.0, "Circulating supply (used)": 100.0,
             "FDV ($)": 400.0, "Revenue Q0 ($)": 10.0, "Actual buyback Q0 ($) — observed": 10.2,
             "Buyback status": "active (measured Q0)", "Burn status": "n/a", "Gross burn Q0 (tokens)": None,
             "NET SUPPLY CHANGE Q0 (tokens) = issuance − burn": 5.0,
             "Net supply change, annualised % of circulating": 0.01}
    bad = {**clean, "Project": "Bad", "FDV ($)": 150.0, "Circulating supply (used)": 120.0,
           "Actual buyback Q0 ($) — observed": 33.0, "Revenue Q0 ($)": 29.7, "Burn status": "unconfirmed",
           "Gross burn Q0 (tokens)": 5.886e6,
           "NET SUPPLY CHANGE Q0 (tokens) = issuance − burn": "BLOCKED — unequal coverage: issuance covers 12 days",
           "Net supply change, annualised % of circulating": 0.0087}
    wb = _wb([clean, bad], {"Clean": 1000.0, "Bad": 110.0})
    got = sanity.hits(wb, lambda s, r, c: wb[s].cell(r, c).value)
    rules = sorted(rule for p, rule, _d in got if p == "Bad")
    assert not [h for h in got if h[0] == "Clean"], got
    assert rules == sorted(["fdv_below_mcap", "circ_above_total", "mcap_not_price_x_circ", "buyback_above_revenue",
                            "status_contradicts", "mixed_windows"]), got


def test_an_active_label_with_no_flow_and_a_documented_funding_source(monkeypatch):
    import config
    row = {"Project": "Aerodrome", "Price — spot ($)": 1.0, "Market cap ($)": 10.0, "Circulating supply (used)": 10.0,
           "FDV ($)": 20.0, "Revenue Q0 ($)": 1.0, "Actual buyback Q0 ($) — observed": 5.0,
           "Buyback status": "active (measured Q0)", "Burn status": "active (declared; no burn series measured)"}
    wb = _wb([row], {"Aerodrome": 100.0})
    got = sanity.hits(wb, lambda s, r, c: wb[s].cell(r, c).value)
    # Aerodrome's buyback is funded from the Foundation's voting share (config foundation_buyback.funding): no flag
    assert (config.PROJECT_BY_NAME["Aerodrome"].get("foundation_buyback") or {}).get("funding")
    assert [h[1] for h in got] == ["status_contradicts"] and "no burn measured" in got[0][2], got
