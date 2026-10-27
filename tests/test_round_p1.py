"""ROUND P1 STAGE 1 (Jake 2026-10-27): the reader workbook — no data changes.

token_metrics_audit.xlsx is the full build; token_metrics.xlsx (reader.py) shows About, Summary, A3 and Definitions,
every number a formula on a hidden `_data` sheet. Built once here on the committed real dumps (fixtures/real_*.json).
"""
from __future__ import annotations

import csv
import json
import re
import shutil
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import config  # noqa: E402
import reader  # noqa: E402

CODES = set(reader.CODES)


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    import build_workbook as bw
    import measuring_guard as mg
    import store as sm
    from real_dump import to_store
    d = tmp_path_factory.mktemp("p1")
    db = d / "metrics.db"
    asofs = []
    for pl in mg._payloads():
        to_store(pl, db)
        asofs.append(pl["meta"]["asof"])
    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(sm, "DB_PATH", str(db))
        mp.setitem(sys.modules, "recalc", None)
        mp.setenv("TOKEN_METRICS_STORE_VALUES", "1")
        st = sm.Store(str(db))
        try:
            bw.build_workbook(st, d / "audit.xlsx", asof=pd.Timestamp(max(asofs)), reader_path=d / "reader.xlsx")
        finally:
            st.close()
    assert "error" not in bw.READER_STATS, bw.READER_STATS
    return {"dir": d, "audit": d / "audit.xlsx", "reader": d / "reader.xlsx", "stats": dict(bw.READER_STATS),
            "headlines": list(bw._HEADLINE_CELLS), "locations": dict(reader.HEADLINE_LOCATIONS),
            "asof": pd.Timestamp(max(asofs))}


# ================================================================================================== P1-1 / P1-3
def test_p1_1_two_files_the_reader_tabs_in_order_and_the_audit_file_whole(built):
    r = openpyxl.load_workbook(built["reader"])
    assert r.sheetnames == ["About", "Summary", "A3 Revenue Buyback", "Definitions", "_data"]
    assert r["_data"].sheet_state == "hidden"
    a = openpyxl.load_workbook(built["audit"], read_only=True)
    for tab in ("Master", "A1 Infrastructure", "A2 Coordination", "A3 Revenue Buyback", "A4 Permanent Burn",
                "CP Comparability", "Flows", "Credibility", "Trajectory", "Charts", "Config & Sources", "Gap Report",
                "Review Queue", "Staging", "Run Log", "Data", "Monthly", "Coverage"):
        assert tab in a.sheetnames, tab
    assert built["stats"]["failed"] == 0


def test_p1_1_a_reader_that_cannot_be_built_says_so_and_is_never_left_stale(tmp_path, monkeypatch):
    import build_workbook as bw
    monkeypatch.setattr(reader, "build_reader", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")))
    bw.READER_STATS.clear()
    out = bw.write_reader(openpyxl.Workbook(), tmp_path / "r.xlsx", tmp_path / "audit.xlsx", pd.Timestamp("2026-10-27"))
    assert out["error"] == "RuntimeError: boom"
    ws = openpyxl.load_workbook(tmp_path / "r.xlsx").active
    assert ws.title == "READER NOT BUILT" and "every figure is on audit.xlsx" in ws["A1"].value


def test_p1_1_the_run_writes_the_reader_and_the_audit_file():
    import token_metrics as tm
    assert tm.WORKBOOK.name == "token_metrics.xlsx" and tm.AUDIT_WORKBOOK.name == "token_metrics_audit.xlsx"
    src = (ROOT / "token_metrics.py").read_text(encoding="utf-8")
    assert src.count("build_workbook(st, AUDIT_WORKBOOK") == 2 and src.count("reader_path=WORKBOOK") == 2


# ================================================================================================== P1-2
def _value_cells(ws, first_row=5):
    """(cell, header) for every body cell of a reader tab (group rows skipped)."""
    heads = {c: ws.cell(4, c).value for c in range(1, ws.max_column + 1)}
    for r in range(first_row, ws.max_row + 1):
        if ws.cell(r, 2).value is None:                  # an archetype group row
            continue
        for c in range(3, ws.max_column + 1):
            yield ws.cell(r, c), heads[c]


def test_p1_2_headers_are_short_plain_and_defined(built):
    r = openpyxl.load_workbook(built["reader"])
    defs = {(row[0].value, row[2].value) for row in r["Definitions"].iter_rows(min_row=4) if row[0].value}
    for tab in ("Summary", "A3 Revenue Buyback"):
        ws = r[tab]
        for c in range(1, ws.max_column + 1):
            h = ws.cell(4, c).value
            assert h and len(h) <= 35, (tab, h)
            assert not re.search(r"\bQ-?\d\b|\b[GDP]\d+\b|config|_usd|_tokens", h), (tab, h)
            assert ws.cell(4, c).comment is not None, (tab, h)
            assert (tab, h) in defs, (tab, h)


def test_p1_2_a_value_cell_is_a_number_or_one_short_code(built):
    r = openpyxl.load_workbook(built["reader"])
    text_cols = {"What changed"}
    for tab in ("Summary", "A3 Revenue Buyback"):
        ws = r[tab]
        for c, h in _value_cells(ws):
            v = c.value
            if v is None or h in text_cols or c.number_format == "@":
                continue
            if isinstance(v, str) and v.startswith("="):
                assert "_data!" in v, (tab, c.coordinate, v)
                continue
            if isinstance(v, (int, float)):
                continue
            assert v in CODES, (tab, c.coordinate, h, v)
            assert c.comment is not None and c.comment.text, (tab, c.coordinate)
            assert c.font.color.rgb.endswith(reader.CODES[v]), (tab, c.coordinate)


def test_p1_2_codes_and_plain_words():
    assert reader.code_of("BLOCKED — NEAR'S LOCKUP RULE CANNOT FORM") == "blocked"
    assert reader.code_of("MISSING [D — ROUTE NOT FILLED]") == "not wired"
    assert reader.code_of("none available") == "by design"
    assert reader.code_of("n/a") == "n/a" and reader.code_of("") == "n/a"
    s = reader.plain("Revenue Q0 (G13 6, Jake 2026-10-21) beside (config.CURRENT_EMISSION_RATE) the Q-1 figure")
    assert "Q0" not in s and "G13" not in s and "config" not in s
    assert "last 90 days" in s and "the 90 days before" in s


# ================================================================================================== P1-5
def test_p1_5_periods_are_like_for_like_windows_and_a_short_history_is_said():
    d = pd.date_range("2026-01-01", "2026-10-26")
    s = pd.Series(1.0, index=d)
    s.loc["2026-10-20":] = 2.0
    v, why = reader.revenue_growth(s, 7, pd.Timestamp("2026-10-27"))
    assert v == pytest.approx(14 / 7 - 1) and "2026-10-20..2026-10-26 against 2026-10-13..2026-10-19" in why
    v, why = reader.revenue_growth(s, 365, pd.Timestamp("2026-10-27"))
    assert v is None and "730 days" in why and "held from 2026-01-01" in why
    p = pd.Series([1.0, 1.5], index=pd.to_datetime(["2026-09-27", "2026-10-27"]))
    assert reader.price_change(p, 30, pd.Timestamp("2026-10-27"))[0] == pytest.approx(0.5)
    assert reader.price_change(p, 7, pd.Timestamp("2026-10-27"))[0] is None


def test_p1_5_the_selector_picks_the_period_precomputed_on_data(built):
    import xlcalc
    r = openpyxl.load_workbook(built["reader"])
    ws = r["Summary"]
    assert ws["B2"].value == "90D" and ws.data_validations.dataValidation
    gcol = next(c for c in range(1, ws.max_column + 1) if ws.cell(4, c).value == "Price change, selected period")
    rows = [rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 2).value]
    seen = 0
    for period in ("30D", "3Y"):
        ws["B2"] = period
        X = xlcalc.Workbook(r)
        for rr in rows:
            v = X.value("Summary", rr, gcol)
            com = ws.cell(rr, gcol).comment.text
            line = next(x for x in com.splitlines() if x.startswith(f"{period}:"))
            if isinstance(v, float):
                assert f"{v:+.1%}" in line, (period, line)
                seen += 1
            else:
                assert v == "—" and "— —" in line, (period, v, line)
    assert seen > 0


def test_p1_6_the_a1_key_needs_two_years_and_says_so():
    d = pd.date_range("2025-10-27", "2026-10-27")
    g = {("X", "settlement_volume_365d_usd"): pd.DataFrame({"date": d, "value": 100.0}),
         ("X", "circulating_supply_used"): pd.DataFrame({"date": d, "value": 10.0})}
    v, why = reader.settlement_key(g, "X")
    assert v == pytest.approx(0.0) and "2025-10-27 to 2026-10-27" in why
    g[("X", "settlement_volume_365d_usd")] = pd.DataFrame({"date": d[100:], "value": 100.0})
    v, why = reader.settlement_key(g, "X")
    assert v is None and "two years of history" in why


# ================================================================================================== P1-6
def test_p1_6_each_group_is_ranked_by_its_key_and_the_unmeasured_rank_last(built):
    import xlcalc
    r = openpyxl.load_workbook(built["reader"])
    ws = r["Summary"]
    X = xlcalc.Workbook(r)
    heads = {ws.cell(4, c).value: c for c in range(1, ws.max_column + 1)}
    kc, rc = heads["Archetype key"], heads["Rank in archetype"]
    groups, cur = [], None
    for rr in range(5, ws.max_row + 1):
        if ws.cell(rr, 2).value is None and ws.cell(rr, 1).value:
            cur = []
            groups.append(cur)
            continue
        if ws.cell(rr, 2).value:
            cur.append((X.value("Summary", rr, kc), X.value("Summary", rr, rc), ws.cell(rr, rc).comment.text))
    assert len(groups) == 4
    for g in groups:
        assert [x[1] for x in g] == list(range(1, len(g) + 1))
        keys = [x[0] for x in g]
        nums = [k for k in keys if isinstance(k, float)]
        assert nums == sorted(nums, reverse=True)
        first_text = next((i for i, k in enumerate(keys) if not isinstance(k, float)), len(keys))
        assert all(not isinstance(k, float) for k in keys[first_text:]), "the unmeasured rank last"
        assert all("Ranked last" in x[2] for x in g[first_text:])


def test_p1_6_the_a3_key_is_the_dilution_ratio_without_its_unlock_leg(built):
    import build_workbook as bw
    import xlcalc
    a = openpyxl.load_workbook(built["audit"])
    X = xlcalc.Workbook(a)
    ws = a["A3 Revenue Buyback"]
    hdr = {ws.cell(4, c).value: c for c in range(1, ws.max_column + 1)}
    kc, dc, nc = hdr[bw.RETIRE_VS_EMISSIONS_HEAD], hdr[bw.RETIRE_VS_DILUTION_HEAD], hdr[bw.RETIRE_VS_DILUTION_NOW_HEAD]
    checked = 0
    for rr in range(5, ws.max_row + 1):
        name = ws.cell(rr, 1).value
        if name not in config.PROJECT_BY_NAME:
            continue
        k = X.value("A3 Revenue Buyback", rr, kc)
        ref = X.value("A3 Revenue Buyback", rr, nc if name in config.CURRENT_EMISSION_RATE else dc)
        if not isinstance(k, float) or not isinstance(ref, float):
            continue
        if config.unlock_stock_metrics(name):
            assert k >= ref - 1e-12, (name, k, ref)        # unlocks only add to the denominator the key drops
        else:
            assert k == pytest.approx(ref, rel=1e-9), name
        checked += 1
    assert checked >= 3


# ================================================================================================== P1-9
def test_p1_9_every_headline_shown_on_the_reader_equals_its_audit_value(built):
    import sanity
    import xlcalc
    a, r = openpyxl.load_workbook(built["audit"]), openpyxl.load_workbook(built["reader"])
    XA, XR = xlcalc.Workbook(a), xlcalc.Workbook(r)
    a3 = [h for h in built["headlines"] if h["sheet"] == "A3 Revenue Buyback"]
    assert a3 and all((h["sheet"], h["project"], h["header"]) in built["locations"] for h in a3)
    numbers = 0
    for h in a3:
        sheet, ref = built["locations"][(h["sheet"], h["project"], h["header"])]
        ca, cr = a[h["sheet"]][h["cell"]], r[sheet][ref]
        va = sanity.stated_number(XA.value(h["sheet"], ca.row, ca.column))
        vr = XR.value(sheet, cr.row, cr.column)
        if va is None:
            blank = not str(XA.value(h["sheet"], ca.row, ca.column) or "").strip()
            assert (vr is None) if blank else (vr in CODES), (h["project"], h["header"], vr)
        else:
            assert vr == pytest.approx(float(va), rel=1e-12), (h["project"], h["header"])
            numbers += 1
    assert numbers >= 20
    # every other headline stays on the audit file at its own cell
    rest = [h for h in built["headlines"] if h["sheet"] != "A3 Revenue Buyback"]
    assert rest
    for h in rest:
        c = a[h["sheet"]][h["cell"]]
        assert a[h["sheet"]].cell(4, c.column).value == h["header"] and a[h["sheet"]].cell(c.row, 1).value == h["project"]


def test_p1_9_the_column_map_names_every_master_and_a3_column(built):
    import build_workbook as bw
    rows = list(csv.DictReader(open(built["dir"] / "reader_column_map.csv", encoding="utf-8")))
    a = openpyxl.load_workbook(built["audit"], read_only=True)
    for tab in ("Master", "A3 Revenue Buyback"):
        ws = a[tab]
        heads = [c.value for c in next(ws.iter_rows(min_row=4, max_row=4)) if isinstance(c.value, str)]
        mapped = {x["audit_header"] for x in rows if x["audit_tab"] == tab}
        assert set(heads) <= mapped, set(heads) - mapped
    key = [x for x in rows if x["audit_header"] == bw.RETIRE_VS_EMISSIONS_HEAD]
    assert key and key[0]["reader_tab"] == "A3 Revenue Buyback" and key[0]["reader_header"] == "Retired ÷ emissions (x)"
    others = {x["audit_tab"] for x in rows if x["audit_column"] == "(all)"}
    assert {"Data", "Run Log", "Coverage", "Credibility"} <= others


def test_p1_9_the_guard_reads_a_headline_at_its_new_location(tmp_path, monkeypatch):
    import measuring_guard as mg
    import xlcalc
    a, r = openpyxl.Workbook(), openpyxl.Workbook()
    a.active.title = r.active.title = "A3 Revenue Buyback"
    a.active["C5"], r.active["F9"] = 1.0, 2.0
    r.save(tmp_path / "reader.xlsx")
    monkeypatch.setattr(reader, "HEADLINE_LOCATIONS", {("A3 Revenue Buyback", "P", "H"): ("A3 Revenue Buyback", "F9")})
    value = mg.headline_reader(a, xlcalc.Workbook(a), tmp_path / "reader.xlsx")
    assert value({"sheet": "A3 Revenue Buyback", "project": "P", "header": "H", "cell": "C5"}) == 2.0
    assert value({"sheet": "A3 Revenue Buyback", "project": "Q", "header": "H", "cell": "C5"}) == 1.0


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_p1_1_libreoffice_agrees_with_every_stored_reader_value(built, tmp_path):
    import recalc_check
    shutil.copy(built["reader"], tmp_path / "reader.xlsx")
    out = recalc_check.check(tmp_path / "reader.xlsx", 300)
    assert out.get("formulas", 0) > 100 and out["mismatches"] == 0 and out["errors"] == 0, json.dumps(out)[:800]
