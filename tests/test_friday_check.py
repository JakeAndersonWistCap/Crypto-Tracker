"""P2f-2: friday_check.py — the Friday checks, read-only, one PASS/FAIL line each.

Unit cases on small fake workbooks (openpyxl); an integration run on the 16-project fixture build (the build
test_round_p2e_ranks makes, from the committed real dumps) through main(); and a negative test per check, each
corrupting ONE cell of a copy of the clean build. No network.
"""
from __future__ import annotations

import re
import sys
from datetime import date
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import friday_check as fc  # noqa: E402
import reader  # noqa: E402


# ============================================================================================ the fixture build
@pytest.fixture(scope="module")
def built(tmp_path_factory):
    """The two workbooks built on the committed real dumps (the build test_round_p2e_ranks makes)."""
    import build_workbook as bw
    import measuring_guard as mg
    import pandas as pd
    import store as sm
    from real_dump import to_store
    d = tmp_path_factory.mktemp("friday")
    db = d / "metrics.db"
    asofs = []
    for pl in mg._payloads():
        to_store(pl, db)
        asofs.append(pl["meta"]["asof"])
    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(sm, "DB_PATH", str(db))
        mp.setitem(sys.modules, "recalc", None)
        mp.setenv("TOKEN_METRICS_STORE_VALUES", "1")
        st = sm.Store(str(db))
        try:
            bw.build_workbook(st, d / "token_metrics_audit.xlsx", asof=pd.Timestamp(max(asofs)),
                              reader_path=d / "token_metrics.xlsx")
        finally:
            st.close()
    assert "error" not in bw.READER_STATS, bw.READER_STATS
    return {"dir": d, "reader": d / "token_metrics.xlsx", "audit": d / "token_metrics_audit.xlsx",
            "asof": pd.Timestamp(max(asofs)).date()}


@pytest.fixture()
def rwb(built):
    """A fresh in-memory copy of the clean reader; a test corrupts it, the file on disk is never touched."""
    return openpyxl.load_workbook(built["reader"])


def _summary_row(wb, name: str) -> int:
    return fc._project_rows(wb[reader.SUMMARY])[name]


# ============================================================================================ integration
def test_integration_fixture_build_runs_every_check(built, capsys):
    """main() on the clean fixture build, the guard on the committed fixtures (--compare): eight checks PASS. The
    ninth, the scoreboard, is a FINDING on this build (see the report): it FAILs on Ethereum's BROKEN reference (the
    fixture store predates orphan_cleanup.sql); the standing NOT WIRED list is printed under it, informational."""
    code = fc.main(["--reader", str(built["reader"]), "--audit", str(built["audit"]), "--compare",
                    "--asof", built["asof"].isoformat()])
    out = capsys.readouterr().out
    lines = [ln for ln in out.splitlines() if ln.startswith(("PASS", "FAIL"))]
    assert len(lines) == 9, out
    fails = [ln for ln in lines if ln.startswith("FAIL")]
    assert len(fails) == 1 and fails[0].startswith("FAIL  scoreboard"), out
    assert re.search(r"— 1 BROKEN: Ethereum: 1$", fails[0]), fails[0]
    assert "NOT WIRED (declared, informational): " in out and "Sky: 5" in out
    assert "FRIDAY CHECK: 8 PASS, 1 FAIL" in out and code == 1
    assert "      ran on: " + str(fc.GUARD_BASELINE) in out        # the guard says what it ran on
    assert re.search(r"PASS  guard: 0 violations", out)


def test_integration_all_nine_pass_with_a_clean_scoreboard_exit_code_0(built, monkeypatch, capsys):
    """With the guard and the scoreboard stood in as clean, the other seven checks on the real build give exit 0."""
    monkeypatch.setattr(fc, "check_guard", lambda mode, store: ([], ["ran on: x"]))
    monkeypatch.setattr(fc, "check_scoreboard", lambda wb: ([], ["BROKEN 0, NOT WIRED 0"]))
    code = fc.main(["--reader", str(built["reader"]), "--audit", str(built["audit"]), "--compare",
                    "--asof", built["asof"].isoformat()])
    out = capsys.readouterr().out
    assert code == 0 and "FRIDAY CHECK: 9 PASS, 0 FAIL" in out, out


def test_the_clean_build_passes_each_reader_check_directly(built, rwb):
    assert fc.check_ranks(rwb) == []
    assert fc.check_about_counts(rwb) == []
    assert fc.check_headers(rwb) == []
    assert fc.check_no_blank(rwb) == []
    assert fc.check_parked(rwb, built["asof"]) == []
    assert fc.check_comments(rwb) == []
    assert fc.parked_projects(list(fc._project_rows(rwb[reader.SUMMARY])), built["asof"])   # not vacuous


def test_it_writes_nothing(built, monkeypatch):
    before = {p.name: (p.stat().st_mtime_ns, p.stat().st_size) for p in built["dir"].iterdir()}
    monkeypatch.setattr(fc, "check_guard", lambda mode, store: ([], []))
    fc.main(["--reader", str(built["reader"]), "--audit", str(built["audit"]), "--asof", built["asof"].isoformat()])
    assert {p.name: (p.stat().st_mtime_ns, p.stat().st_size) for p in built["dir"].iterdir()} == before


# ============================================================================================ negatives (one cell each)
def test_1_a_wrong_rank_on_the_primary_tab_is_caught(rwb):
    ws = rwb[reader.A3_TAB]
    rows, rc = fc._project_rows(ws), fc._rank_col(ws)
    assert ws.cell(rows["Sky"], rc).value is not None
    ws.cell(rows["Sky"], rc).value = 7                      # Summary says Sky is 1
    bad = fc.check_ranks(rwb)
    assert any(b.startswith("Sky: rank 7.0 on A3 Revenue Buyback") for b in bad), bad


def test_1_a_wrong_rank_on_summary_is_caught(rwb):
    s = rwb[reader.SUMMARY]
    rk = fc._col(s, "Rank in archetype")
    s.cell(_summary_row(rwb, "Aethir"), rk).value = 5
    bad = fc.check_ranks(rwb)
    assert any(b.startswith("Aethir:") and "Summary" in b for b in bad), bad


def test_1_a_secondary_cell_on_the_primary_tab_is_caught(rwb):
    ws = rwb[reader.A3_TAB]
    ws.cell(fc._project_rows(ws)["Sky"], fc._rank_col(ws)).value = "secondary"
    assert any("Sky: reads 'secondary' on A3 Revenue Buyback" in b for b in fc.check_ranks(rwb))


def test_2_an_about_count_that_differs_from_summary_is_caught(rwb):
    a = rwb["About"]
    top = next(r for r in range(1, a.max_row + 1) if str(a.cell(r, 1).value or "").startswith("Honest coverage"))
    a.cell(top + 2, 2).value = 99                           # A1's Projects
    bad = fc.check_about_counts(rwb)
    assert any(b.startswith("A1 Infrastructure Projects: About says 99.0, Summary has ") for b in bad), bad


def test_3_a_40_character_header_is_caught(rwb):
    ws = rwb[reader.A2_TAB]
    ws.cell(4, 5).value = "H" * 40
    assert fc.check_headers(rwb) == [f"A2 Coordination!E4: 40 characters ('{'H' * 40}')"]


def test_4_a_blank_cell_without_a_code_is_caught(rwb):
    s = rwb[reader.SUMMARY]
    r = _summary_row(rwb, "Ethereum")
    head = s.cell(4, 6).value
    s.cell(r, 6).value = None
    assert fc.check_no_blank(rwb) == [f"Summary / Ethereum / {head}: blank, no code"]


def test_4_a_code_without_its_comment_is_caught(rwb):
    s = rwb[reader.SUMMARY]
    r = _summary_row(rwb, "Bitcoin")
    c = next(c for c in range(3, s.max_column + 1) if s.cell(r, c).value in reader.CODES and s.cell(r, c).comment)
    s.cell(r, c).comment = None
    bad = fc.check_no_blank(rwb)
    assert len(bad) == 1 and "without a comment" in bad[0] and "/ Bitcoin /" in bad[0], bad


def test_5_a_non_parked_code_on_a_parked_row_is_caught(built, rwb):
    s = rwb[reader.SUMMARY]
    r = _summary_row(rwb, "Aerodrome")
    rk = fc._col(s, "Rank in archetype")
    assert s.cell(r, rk).value == "parked"
    s.cell(r, rk).value = "blocked"
    assert fc.check_parked(rwb, built["asof"]) == ["Summary / Aerodrome / Rank in archetype: reads 'blocked', not "
                                                   "'parked'"]


def test_5_a_parked_key_cell_that_reads_blocked_is_caught(built, rwb):
    s = rwb[reader.SUMMARY]
    r = _summary_row(rwb, "Aerodrome")
    kc = fc._col(s, "Archetype key")
    s.cell(r, kc).value = "blocked"
    bad = fc.check_parked(rwb, built["asof"])
    assert bad and "Aerodrome / Archetype key: reads 'blocked'" in bad[0], bad


def test_6_a_comment_that_starts_with_working_is_caught(rwb):
    from openpyxl.comments import Comment
    s = rwb[reader.SUMMARY]
    r = _summary_row(rwb, "Ethereum")
    s.cell(r, 4).comment = Comment("Working: audit file, Master!B3.", "x")
    bad = fc.check_comments(rwb)
    assert len(bad) == 1 and bad[0].startswith(f"Summary!D{r} (Ethereum /") and "starts with 'Working:'" in bad[0]


def test_6_the_rule_on_small_cases():
    assert fc.comment_problem("Working: audit file, Master!B3.")
    assert fc.comment_problem("  Working: x")
    assert fc.comment_problem("=Master!B3*2")
    assert fc.comment_problem("Master!B3 is the source.")
    assert fc.comment_problem("B3 = Master!C4")
    assert fc.comment_problem("$B$3:$C$9 holds it")
    assert fc.comment_problem("'Config & Sources'!A4 says so")
    assert fc.comment_problem("Market cap for Aave: price × circulating supply, as of 2026-10-13.\n\nWorking: x") == ""
    assert fc.comment_problem("A3 Revenue Buyback ranks by retired ÷ emissions.") == ""
    assert fc.comment_problem("Bitcoin has no market cap here.") == ""
    assert fc.comment_problem("") == ""


def test_9_a_changed_as_of_is_caught_by_main_and_names_both_dates(built, tmp_path, monkeypatch, capsys):
    wb = openpyxl.load_workbook(built["reader"])
    a = wb["About"]
    r = next(r for r in range(1, 30) if str(a.cell(r, 1).value or "").startswith("Data as of"))
    a.cell(r, 1).value = a.cell(r, 1).value.replace(built["asof"].isoformat(), "2026-10-12", 1)
    p = tmp_path / "token_metrics.xlsx"
    wb.save(p)
    monkeypatch.setattr(fc, "check_guard", lambda mode, store: ([], []))
    monkeypatch.setattr(fc, "check_scoreboard", lambda wb: ([], []))
    code = fc.main(["--reader", str(p), "--audit", str(built["audit"]), "--asof", built["asof"].isoformat()])
    out = capsys.readouterr().out
    assert code == 1
    assert f"FAIL  as-of date is today — reader (About): as-of 2026-10-12, today {built['asof']}" in out, out
    assert "FRIDAY CHECK: 8 PASS, 1 FAIL" in out


# ============================================================================================ unit cases, fake inputs
def test_headers_on_a_fake_workbook():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for tab in fc.VALUE_TABS:
        ws = wb.create_sheet(tab)
        ws.cell(4, 1, "Project")
        ws.cell(4, 2, "x" * reader.MAX_HEADER)               # exactly the limit passes
    assert fc.check_headers(wb) == []
    wb[reader.BUSINESS].cell(4, 3, "x" * (reader.MAX_HEADER + 1))
    assert fc.check_headers(wb) == [f"Business!C4: {reader.MAX_HEADER + 1} characters ('{'x' * (reader.MAX_HEADER + 1)}')"]


def test_blank_and_code_on_a_fake_tab():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(reader.BUSINESS)
    for c, h in enumerate(("Project", "Token", "Fees", "Revenue", "TVL"), start=1):
        ws.cell(4, c, h)
    ws.cell(5, 1, "Aave")
    ws.cell(5, 2, "AAVE")
    ws.cell(5, 3, "=_data!B2")
    ws.cell(5, 4, "n/a")
    ws.cell(5, 5, "   ")
    bad = fc.check_no_blank(wb)
    assert bad == ["Business / Aave / Revenue: code 'n/a' without a comment", "Business / Aave / TVL: blank, no code"]


def test_format_line_caps_reasons_at_ten_and_counts_the_rest():
    assert fc.format_line("PASS", "x", []) == "PASS  x"
    assert fc.format_line("FAIL", "x", ["a", "b"]) == "FAIL  x — a; b"
    line = fc.format_line("FAIL", "x", [f"r{i}" for i in range(13)])
    assert line == "FAIL  x — " + "; ".join(f"r{i}" for i in range(10)) + "; and 3 more"
    assert "r10" not in line


def test_the_guard_verdict_needs_zero_violations_and_lists_accepted():
    ok = ("MEASURING-POINT GUARD — 0 violation(s)\nACCEPTED LOSSES (config.REGRESSION_ALLOWLIST) — 1, listed, not "
          "counted\n  ACCEPTED: Hyperliquid A4 BURN — reason (accepted 2026-10-28)\nSANITY lost_headlines — 0 hit(s)\n"
          "  ACCEPTED: Hyperliquid A4 BURN — reason (accepted 2026-10-28)\n")
    assert fc.judge_guard(0, ok) == ([], ["ACCEPTED: Hyperliquid A4 BURN — reason (accepted 2026-10-28)"])
    ten = "MEASURING-POINT GUARD — 10 violation(s)\n  HEADLINE LOST Aave X: 1 before, no value after\n"
    reasons, _ = fc.judge_guard(1, ten)
    assert reasons[0] == "10 violation(s)" and "HEADLINE LOST Aave X" in reasons[1]      # "10" is not "0 violation(s)"
    live = "MEASURING-POINT GUARD ON THE LIVE STORE — 0 violation(s)\n"
    assert fc.judge_guard(0, live)[0] == []
    assert fc.judge_guard(1, live)[0]                                                     # exit code 1 is not a PASS
    assert fc.judge_guard(1, "Traceback (most recent call last):\nValueError: x\n")[0]    # a crash is a FAIL
    assert fc.judge_guard(0, "")[0]


def test_the_guard_runs_as_the_repos_entry_point_and_a_crash_is_a_fail(tmp_path):
    store = tmp_path / "metrics.db"
    store.write_bytes(b"")
    argv, ran_on = fc.guard_command("newer", str(store))
    assert argv[1:] == [str(ROOT / "measuring_guard.py"), "--newer", str(store)] and ran_on == str(store)
    argv, ran_on = fc.guard_command("compare", "ignored")
    assert argv[-2:] == ["--compare", str(ROOT / "tests" / "fixtures" / "guard_baseline.json")]
    calls = []

    class Done:
        returncode = 0
        stdout = ("MEASURING-POINT GUARD ON THE LIVE STORE — 0 violation(s)\nACCEPTED LOSSES — 1\n"
                  "  ACCEPTED: P H — why (accepted 2026-10-28)\n").encode()
        stderr = b""

    def fake_run(argv, **kw):
        calls.append((argv, kw["cwd"]))
        return Done()

    reasons, notes = fc.check_guard("newer", str(store), run=fake_run)
    assert reasons == [] and notes == [f"ran on: {store}", "ACCEPTED LOSSES — 1",
                              "ACCEPTED: P H — why (accepted 2026-10-28)"]
    assert calls[0][0][-2:] == ["--newer", str(store)]

    def boom(argv, **kw):
        raise OSError("no python")
    reasons, notes = fc.check_guard("newer", str(store), run=boom)
    assert reasons and "did not run" in reasons[0] and notes == [f"ran on: {store}"]
    reasons, _ = fc.check_guard("newer", str(tmp_path / "missing.db"), run=fake_run)
    assert reasons == [f"store not found: {tmp_path / 'missing.db'}"]


def _coverage_wb(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Coverage"
    ws.cell(1, 1, "Coverage — every column against SPEC.md")
    ws.cell(3, 1, "SCOREBOARD")
    for j, h in enumerate(("Project", "VERIFIED", "NOT WIRED", "BROKEN"), start=1):
        ws.cell(4, j, h)
    for i, row in enumerate(rows, start=5):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j, v)
    ws.cell(5 + len(rows) + 1, 1, "Project")                  # the lower table's own header: not part of the board
    return wb


def test_scoreboard_pass_when_both_are_zero_and_fail_names_the_projects():
    ok = fc.check_scoreboard(_coverage_wb([("ALL", 9, 0, 0), ("Aave", 4, 0, 0), ("Sky", 5, 0, 0)]))
    assert ok == ([], ["BROKEN 0, NOT WIRED 0"])
    reasons, notes = fc.check_scoreboard(_coverage_wb([("ALL", 9, 3, 2), ("Aave", 4, 1, 0), ("Sky", 5, 1, 2)]))
    assert reasons == ["2 BROKEN: Sky: 2"]
    assert notes == ["BROKEN 2, NOT WIRED 3",
                     "NOT WIRED (declared, informational): Aave: 1, Sky: 1, (no project: SPEC list): 1"]
    # NOT WIRED alone is printed, never a FAIL
    only = fc.check_scoreboard(_coverage_wb([("ALL", 9, 3, 0), ("Aave", 4, 3, 0)]))
    assert only == ([], ["BROKEN 0, NOT WIRED 3", "NOT WIRED (declared, informational): Aave: 3"])
    assert fc.check_scoreboard(openpyxl.Workbook())[0]        # no Coverage tab: a FAIL, said


def test_scoreboard_counts_on_parked_projects_are_printed_never_failed():
    """P2g-4: a parked project's BROKEN and NOT WIRED are information; another project's BROKEN still fails."""
    board = fc.read_scoreboard(_coverage_wb([("ALL", 9, 3, 2), ("Aerodrome", 4, 3, 1), ("Sky", 5, 0, 1)]))
    reasons, notes = fc.scoreboard_verdict(board, parked={"Aerodrome"})
    assert reasons == ["1 BROKEN: Sky: 1"]
    assert "BROKEN on parked projects (informational, not judged): Aerodrome: 1" in notes
    assert "NOT WIRED on parked projects (informational, not judged): Aerodrome: 3" in notes
    only = fc.scoreboard_verdict(fc.read_scoreboard(_coverage_wb([("ALL", 4, 3, 1), ("Aerodrome", 4, 3, 1)])),
                                 parked={"Aerodrome"})
    assert only[0] == []                                              # nothing but a parked project: no FAIL


def test_asof_equals_today_and_says_both_dates():
    d = date(2026, 10, 30)
    assert fc.check_asof(d, d, d) == []
    assert fc.check_asof(date(2026, 10, 29), d, d) == ["reader (About): as-of 2026-10-29, today 2026-10-30"]
    assert fc.check_asof(d, date(2026, 10, 28), d) == ["audit (Run Log): as-of 2026-10-28, today 2026-10-30"]
    assert len(fc.check_asof(None, None, d)) == 2


def test_a_missing_workbook_fails_the_checks_and_does_not_crash(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(fc, "check_guard", lambda mode, store: ([], []))
    code = fc.main(["--reader", str(tmp_path / "no.xlsx"), "--audit", str(tmp_path / "no2.xlsx"), "--asof",
                    "2026-10-30"])
    out = capsys.readouterr().out
    assert code == 1 and "FAIL  ranks identical" in out and "cannot open" in out
    assert out.strip().splitlines()[-1] == "FRIDAY CHECK: 1 PASS, 8 FAIL"


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
