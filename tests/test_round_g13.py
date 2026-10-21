"""
tests/test_round_g13.py — Round G13, "finish the wiring. nothing new is fetched" (Jake 2026-10-21), SPEC §19 "G13".
The last UNWIRED series (1), one Fees definition per project (2), the P2P cache-to-store drop (3), the 365-day cells
(4), attribute columns in the guard (5), Aethir's subsidy cover (6), cached zeros (7), Hyperliquid's burn break (8),
wiring.json's reasons (10).
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import build_workbook as bw  # noqa: E402
import config  # noqa: E402
import wiring  # noqa: E402


def _g(project, metric, days, values, source="x"):
    return pd.DataFrame({"date": pd.to_datetime(list(days)), "project": project, "metric": metric,
                         "value": [float(v) for v in values], "tier": 1, "source": source, "is_manual": False,
                         "entered_on": ""})


# ================================================================================================== 1 the last UNWIRED
def test_1a_fetch_time_components_are_inputs_from_config_and_code():
    fi = config.fetch_inputs("Pendle")["noncirculating_holding_tokens"]
    assert set(fi[0]) == {"noncirc_ecosystem_fund_tokens", "noncirc_governance_multisig_tokens",
                          "noncirc_team_multisig_tokens"}
    src = (ROOT / "fetch" / "chain.py").read_text(encoding="utf-8")
    assert "component_metric" in src and "PER-WALLET ROWS" in src       # the per-wallet rows are the sum's own reads
    wm = config.fetch_inputs("World Mobile")["locked_tokens"][0]
    assert set(wm) == {"locked_tokens_base", "locked_tokens_cardano"}
    s = (ROOT / "fetch" / "wm_supply.py").read_text(encoding="utf-8")
    assert 'legs["locked_tokens_base"] + legs["locked_tokens_cardano"]' in s


def test_1a_the_settlement_and_stake_views_read_back_what_they_back_fill():
    import inspect
    src = inspect.getsource(bw._settlement_rebuild_views)
    assert 'rebuilt = _daily(groups[(name, "settlement_rebuilt_usd")])' in src
    assert src.index('rebuilt = _daily(groups[(name, "settlement_rebuilt_usd")])') < src.index("after = rebuilt[")
    hl = inspect.getsource(bw._hl_model_views)
    assert 'stk = _daily(groups[(name, spec["stake_modelled"])])' in hl
    # the graph sees the read: a view reading r and writing w gives the INPUT edge r -> w
    wiring.reset()
    g = wiring.TrackingDict({("Ethereum", "settlement_rebuilt_usd"): 1})

    def _settlement_rebuild_views(groups):
        groups[("Ethereum", "settlement_rebuilt_usd")] = 2
        groups.get(("Ethereum", "settlement_rebuilt_usd"))
        groups[("Ethereum", "settlement_volume_usd")] = 3
    _settlement_rebuild_views(g)
    assert ("settlement_rebuilt_usd", "_settlement_rebuild_views") in wiring.view_edges()[("Ethereum",
                                                                                          "settlement_volume_usd")]


def test_1a_not_shown_reasons_and_detail_columns():
    assert "MODELLED BACKWARD" in config.not_shown_reason("Hyperliquid", "future_emissions_modelled_tokens")
    assert "never in a total" in config.not_shown_reason("World Mobile", "wmt_legacy_supply_tokens")
    cells = {(d["tab"], p, m) for d in config.DETAIL_COLUMNS for p, (m, _w) in d["cells"].items()}
    for want in (("A2", "Aethir", "compute_purchases_cumulative_tokens"),
                 ("A2", "Aethir", "compute_service_fee_cumulative_tokens"),
                 ("A2", "World Mobile", "wmtx_cardano_supply_tokens"),
                 ("A1", "Chainlink", "staker_yield_community_pool"),
                 ("CP", "Chainlink", "staking_funding_release_tokens")):
        assert want in cells, want
    heads = {m: d["header"] for d in config.DETAIL_COLUMNS for _p, (m, _w) in d["cells"].items()}
    assert "Onchain Compute Purchases" in heads["compute_purchases_cumulative_tokens"]
    assert "totalServiceFee" in heads["compute_service_fee_cumulative_tokens"]


def test_1b_cross_check_rows_for_morpho_ccip_and_hyperliquid():
    import credibility as cred
    m = config.CREDIBILITY_WIRING["Morpho"]["in_merkl_off_ethereum"]
    assert "INCLUDES the off-Ethereum claims" in m["ref"]["why"] and m["ref"]["metric"] == "merkl_campaigns_non_ethereum_tokens"
    c = config.CREDIBILITY_WIRING["Chainlink"]["in_ccip_v2_arbitrum_overlap"]
    assert c["ref"]["tol"] == 100.0 * config.CCIP_FEE_OVERLAP["tolerance"] and "verdict" not in c["ref"]
    rows = {"Hyperliquid|issuer_release_replay_tokens": {"q0": 5.0}}
    assert "in_issuer_release_replay" in cred.wiring_inputs("Hyperliquid", rows)


def test_1b_world_mobile_fdv_is_recorded_while_its_credibility_is_parked():
    import credibility as cred
    assert "World Mobile" not in config.CREDIBILITY_PROJECTS
    rows = {"World Mobile|fdv_usd_coingecko": {"now": 1.0e8, "latest_date": "2026-10-20"},
            "World Mobile|fdv_usd": {"now": 1.1e8, "latest_date": "2026-10-20"}}
    out = cred.build_rows([], rows, pd.DataFrame(), pd.Timestamp("2026-10-21"), projects={"World Mobile"})
    wm = [r for r in out if r["project"] == "World Mobile"]
    assert len(wm) == 1 and wm[0]["id"] == "in_fdv_coingecko"
    assert wm[0]["verdict"] == "N/A (recorded)" and wm[0]["mode"] == "static" and wm[0]["value"] == 1.0e8


def test_1c_the_monthly_twin_is_found_whatever_day_it_is_dated():
    ends = pd.date_range("2025-01-31", "2026-09-30", freq="ME")
    monthly = pd.Series([1.0e6 + i for i in range(len(ends))], index=ends)
    assert "Credibility reference" in config.not_shown_reason("Sky", "buyback_spending_usd_ba_daily")
    starts = ends - pd.offsets.MonthBegin(1)                        # the same totals dated on the 1st
    assert bw.daily_twin(pd.Series(monthly.values, index=starts), monthly)[0]
    days = pd.date_range("2024-11-14", "2026-10-20")              # zero-filled days, the month total on its last day
    zf = pd.Series(0.0, index=days)
    zf.loc[ends] = monthly.values
    assert bw.daily_twin(zf, monthly)[0]
    real = pd.Series(1.0e6 / 30, index=pd.date_range("2025-01-01", "2026-10-20"))   # a real daily series
    assert not bw.daily_twin(real, monthly)[0]
    from fetch.sky_accounting import looks_daily
    assert not looks_daily([(d, 1.0) for d in ends[-6:]])        # month-END dates are a monthly answer
    assert looks_daily([(d, 1.0) for d in pd.date_range("2026-09-01", "2026-09-10")])


# ================================================================================================== 2 one Fees definition
def test_2_no_world_mobile_cell_reads_fees_usd(tmp_path, monkeypatch):
    import openpyxl
    import store as sm
    monkeypatch.setitem(sys.modules, "recalc", None)
    st = sm.Store(tmp_path / "empty.db")
    out = tmp_path / "wb.xlsx"
    bw.build_workbook(st, out)
    st.close()
    wb = openpyxl.load_workbook(out)
    bad, seen = [], 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("=") and '|fees_usd"' in v:
                    proj = ws.cell(c.row, 1).value
                    if proj == "World Mobile":
                        bad.append(f"{ws.title}!{c.coordinate}")
                if isinstance(v, str) and '|own_chain_fees_usd"' in v and ws.cell(c.row, 1).value == "World Mobile":
                    seen += 1
    assert not bad, bad[:10]
    assert seen >= 4                                               # Fees Q0, ÷ issuance, ÷ FDV, the Δ columns


# ================================================================================================== 3 P2P cache -> store
def test_3_a_cached_day_keyed_by_symbol_with_an_incomplete_cached_price_is_stored(monkeypatch):
    from fetch.base import FetchOutput
    from fetch.eth_settlement import EthP2PBigQuery
    p2p = config.SETTLEMENT_REBUILD["ethereum_rebuild"]["p2p"]
    addr = {t[0]: t[1] for t in p2p["tokens"]}
    day = "2026-09-21"
    e = {"days": {day: {"ETH": str(10 ** 18), "DAI": str(5 * 10 ** 18), "USDC": str(7 * 10 ** 6),
                        "WETH": str(10 ** 18), "USDT": str(3 * 10 ** 6), addr["USDS"]: str(2 * 10 ** 18)}},
         "prices": {day: {"ethereum:DAI": [None, None]}}}         # an old answer cached for the symbol keys
    prices = {(day, f"ethereum:{addr[s]}"): (1.0, d) for s, a, d in [(t[0], t[1], t[2]) for t in p2p["tokens"]]}
    stored = pd.DataFrame({"project": ["Ethereum"], "metric": ["price_usd"], "date": [pd.Timestamp(day)],
                           "value": [2000.0]})
    a = EthP2PBigQuery(stored_long=stored, prices=prices)
    monkeypatch.setattr(a, "_save", lambda st: None)
    out = FetchOutput()
    a._value_eth(p2p, {}, e, out, "Ethereum", "p2p_transfer_volume_usd")
    f = out.frame()
    assert len(f) == 1 and str(f["date"].iloc[0])[:10] == day
    # 1 ETH + 1 WETH at $2,000, 5 DAI + 7 USDC + 3 USDT + 2 USDS at $1
    assert abs(float(f["value"].iloc[0]) - (4000 + 17)) < 1e-6
    assert set(e["days"][day]) == {"ETH", *addr.values()}            # the cache is rewritten on the SQL's keys


def test_3_every_cached_day_not_stored_says_why(monkeypatch):
    from fetch.base import FetchOutput
    from fetch.eth_settlement import EthP2PBigQuery
    p2p = config.SETTLEMENT_REBUILD["ethereum_rebuild"]["p2p"]
    e = {"days": {"2026-09-20": {"ETH": "1"}, "2026-09-22": {"ETH": "1", "XYZ": "5"}}, "prices": {}}
    stored = pd.DataFrame({"project": ["Ethereum"], "metric": ["price_usd"], "date": [pd.Timestamp("2026-09-22")],
                           "value": [2000.0]})
    a = EthP2PBigQuery(stored_long=stored, prices={})
    monkeypatch.setattr(a, "_save", lambda st: None)
    out = FetchOutput()
    a._value_eth(p2p, {}, e, out, "Ethereum", "p2p_transfer_volume_usd")
    text = " ".join(x.message for x in out.log)
    assert "2 cached day(s) NOT stored" in text
    assert any(x.status == "warn" and "NOT stored" in x.message for x in out.log)   # a WARN on the run summary
    assert "2026-09-20: no stored Ethereum price_usd that day" in text and "2026-09-22: token key(s)" in text


# ================================================================================================== 4 the 365-day cells
def test_4_a_365_day_cell_tagged_with_its_inputs_q0_is_judged_on_365_days(tmp_path, monkeypatch, capsys):
    import csv
    import sqlite3

    import check_offline_items as coi
    import store as sm
    db = tmp_path / "metrics.db"
    con = sqlite3.connect(str(db))
    con.execute("CREATE TABLE metrics (date TEXT, project TEXT, metric TEXT, value REAL, source TEXT, tier INT, "
                "fetched_at TEXT)")
    days = list(pd.date_range("2025-10-21", "2026-10-20"))
    held = days[:114] + days[-89:]                         # 203 days; the last 89 cover Q0 (07-24..10-20)
    for m in ("dex_volume_usd", "p2p_transfer_volume_usd"):
        con.executemany(f"INSERT INTO metrics VALUES (?, 'Near', '{m}', 1.0, 'x', 1, '')",
                        [(str(d.date()),) for d in (days if m == "dex_volume_usd" else held)])
    con.commit()
    con.close()
    path = tmp_path / "coverage_audit.csv"
    text = ("BLOCKED — unequal coverage: dex (dex_volume_usd, q0) covers 90 days of a 90-day window; p2p "
            "(p2p_transfer_volume_usd, q0) covers 89 days of a 90-day window")
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["Project", "Tab", "Column", "Status", "Reason / route / error", "Value now"])
        w.writeheader()
        for col in ("Settlement volume, trailing 365d ($)", "NRR ON SETTLEMENT VOLUME — market cap ÷ trailing-365d "
                                                            "settlement volume"):
            w.writerow({"Project": "Near", "Tab": "A1", "Column": col, "Status": "MISSING", "Value now": "",
                        "Reason / route / error": text})
        # a cell BUILT ON a 365-day figure, its header naming no window (the fixture's Near "NRR HEADROOM (settlement)")
        w.writerow({"Project": "Near", "Tab": "A1", "Column": "NRR HEADROOM (settlement) = M2/GDP ÷ NRR",
                    "Status": "MISSING", "Reason / route / error": "MISSING [C — NEEDS HISTORY] — p2p_transfer_volume_usd "
                    "is stored from 2026-04-01; the 365-day window needs it from 2025-10-21",
                    "Value now": "n/a — network_reserve_ratio (now): BLOCKED — the settlement REBUILD (dex_volume_usd + "
                                 "p2p_transfer_volume_usd (now)) holds 203 of the 365 days"})
    monkeypatch.setenv("TOKEN_METRICS_COVERAGE_CSV", str(path))
    monkeypatch.setattr(sm, "DB_PATH", str(db))
    coi.stored_but_unshown()
    assert "0 hit(s)" in capsys.readouterr().out


# ================================================================================================== 5 attribute columns
def test_5_attribute_columns_are_not_headlines():
    import inspect

    import export_snapshot as es
    import measuring_guard as mg
    assert es.is_attribute("Master | Archetypes (primary first)") and es.is_attribute("Master | Held")
    assert not es.is_attribute("Master | Price — spot ($)") and not es.is_attribute("A3 | Fees Q0 ($)")
    assert "_es.is_attribute(head)" in inspect.getsource(mg.live_check)


# ================================================================================================== 6 subsidy cover
def test_6_subsidy_cover_over_the_days_every_leg_covers(monkeypatch):
    asof = pd.Timestamp("2026-10-20")
    days = pd.date_range("2026-07-01", "2026-10-20")
    b = pd.concat([_g("Aethir", "supplier_rewards_b_tokens", days[days < "2026-10-05"], [1.0] * 96, "schedule:config"),
                   _g("Aethir", "supplier_rewards_b_tokens", days[days >= "2026-10-05"], [1.0] * 16,
                      "aethir_page:supplier_rewards")])
    groups = {("Aethir", "supplier_rewards_b_tokens"): b,
              ("Aethir", "customer_revenue_usd"): _g("Aethir", "customer_revenue_usd", days, [100.0] * len(days)),
              ("Aethir", "outflow_b_usd"): _g("Aethir", "outflow_b_usd", days, [10.0] * len(days))}
    monkeypatch.setitem(bw.GLOBALS, "period_days", 90)
    bw._subsidy_common_views(groups, asof)
    info = bw.SUBSIDY_COMMON["Aethir"]
    assert info == {"from": "2026-10-05", "days": 16, "first": "2026-10-05", "last": "2026-10-20"}
    assert float(groups[("Aethir", "subsidy_cover_num_usd")]["value"].sum()) == 1600.0
    assert "2026-10-05..2026-10-20 (16 common days" in bw._subsidy_dates_fmt("Aethir")
    # a WEEKLY revenue leg (Aethir's weeklyNetworkRevenue: dated the week's first day, finished weeks only) is spread
    # over its seven days — never a week of revenue against one day of incentives
    weeks = pd.date_range("2026-07-06", "2026-10-12", freq="7D")
    groups[("Aethir", "customer_revenue_usd")] = _g("Aethir", "customer_revenue_usd", weeks, [700.0] * len(weeks))
    bw._subsidy_common_views(groups, asof)
    info = bw.SUBSIDY_COMMON["Aethir"]
    assert info == {"from": "2026-10-05", "days": 14, "first": "2026-10-05", "last": "2026-10-18"}
    num = float(groups[("Aethir", "subsidy_cover_num_usd")]["value"].sum())
    den = float(groups[("Aethir", "subsidy_cover_den_usd")]["value"].sum())
    assert (num, den) == (1400.0, 140.0) and num / den == 10.0
    # under RATE_MIN_DAYS common days the two series are not written (nothing would read them: the cell says NEEDS
    # HISTORY), and a held (c) leaves the (b) view to the (b)-only cell
    groups[("Aethir", "customer_revenue_usd")] = _g("Aethir", "customer_revenue_usd", weeks[weeks < "2026-10-12"],
                                                    [700.0] * int((weeks < "2026-10-12").sum()))
    for k in (("Aethir", "subsidy_cover_num_usd"), ("Aethir", "subsidy_cover_den_usd")):
        groups.pop(k)
    bw._subsidy_common_views(groups, asof)
    assert bw.SUBSIDY_COMMON["Aethir"]["days"] == 7 and ("Aethir", "subsidy_cover_num_usd") in groups
    groups[("Aethir", "customer_revenue_usd")] = _g("Aethir", "customer_revenue_usd", weeks[weeks < "2026-10-05"],
                                                    [700.0] * int((weeks < "2026-10-05").sum()))
    for k in (("Aethir", "subsidy_cover_num_usd"), ("Aethir", "subsidy_cover_den_usd")):
        groups.pop(k)
    bw._subsidy_common_views(groups, asof)
    assert bw.SUBSIDY_COMMON["Aethir"]["days"] == 0 and ("Aethir", "subsidy_cover_num_usd") not in groups
    import inspect
    assert 'got == ["b"]' in inspect.getsource(bw._cp_cols)
    # under RATE_MIN_DAYS common days: NEEDS HISTORY with the date
    monkeypatch.setitem(bw.SUBSIDY_COMMON, "Aethir", {"from": "2026-10-18", "days": 3, "first": "2026-10-18",
                                                      "last": "2026-10-20"})
    assert bw._subsidy_dates_fmt("Aethir") is None
    import inspect
    assert "n/a — NEEDS HISTORY: the (b) denominator is complete from" in inspect.getsource(bw._cp_cols)
    import coverage_audit as ca                     # graded MISSING C (status honesty), never LIVE-UNVERIFIED
    assert any(st == "MISSING C" and rx.search("n/a — NEEDS HISTORY: the (b) denominator is complete from 2026-10-05")
               for rx, st in ca._HONEST_RULES)


# ================================================================================================== 7 cached zeros
def test_7_a_zero_is_saved_as_zero_and_no_credibility_gap_is_an_empty_string(tmp_path):
    import openpyxl

    import workbook_values
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Credibility"
    ws["A1"], ws["A2"], ws["B1"] = "x", "y", "=COUNTIF(A1:A2,\"z\")"
    ws["B2"], ws["B3"] = "=1-1", "=B1+B2"
    p = tmp_path / "v.xlsx"
    wb.save(p)
    st = workbook_values.store_values(p)
    assert st["failed"] == 0
    wv = openpyxl.load_workbook(p, data_only=True)["Credibility"]
    assert [wv["B1"].value, wv["B2"].value, wv["B3"].value] == [0, 0, 0]
    import numpy as np
    assert workbook_values._cached(np.int64(0)) == ("", "0.0") and workbook_values._cached(0) == ("", "0")
    import inspect
    src = inspect.getsource(bw.write_credibility)
    assert '"no reference value"' in src and "ABS(G{r})),\"\")" not in src


# ================================================================================================== 8 Hyperliquid's burn
def test_8_hyperliquid_burn_trends_break_at_the_buyback_stitch(monkeypatch):
    for m in ("gross_burn_tokens", "total_burn_tokens", "burn_ex_chain_fee_tokens"):
        assert "SERIES BREAK" in config.declared_handover("Hyperliquid", m)["composition_change"]
    days_h, days_l = pd.date_range("2026-06-01", "2026-09-11"), pd.date_range("2026-09-12", "2026-10-13")
    g = pd.concat([_g("Hyperliquid", "gross_burn_tokens", days_h, [1.0] * len(days_h),
                      "defillama:holders_revenue_usd/price"),
                   _g("Hyperliquid", "gross_burn_tokens", days_l, [1.0] * len(days_l),
                      "hypercore_info:spotClearinghouseState:delta")], ignore_index=True)
    monkeypatch.setattr(bw, "_LAST_GROUPS", {("Hyperliquid", "gross_burn_tokens"): g})
    assert bw._stitch_breaks("Hyperliquid", "gross_burn_tokens")[0][0] == "2026-09-12"
    # the total / ex-chain-fee burn: the Core burns join at the first day both legs exist (a second stitch); the trend
    # breaks at both, and the row names the buyback's stitch as the earlier one
    import trajectory
    days_a, days_b = pd.date_range("2026-06-01", "2026-09-24"), pd.date_range("2026-09-25", "2026-10-13")
    tot = pd.concat([_g("Hyperliquid", "total_burn_tokens", days_a, [1.0] * len(days_a),
                        "defillama:holders_revenue_usd/price"),
                     _g("Hyperliquid", "total_burn_tokens", days_b, [1.0] * len(days_b),
                        "derived:burn_total(gross_burn_tokens+core_burn_tokens)")], ignore_index=True)
    groups = {("Hyperliquid", "gross_burn_tokens"): g, ("Hyperliquid", "total_burn_tokens"): tot}
    found = trajectory.basis_breaks("Hyperliquid", ["total_burn_tokens"], groups, pd.Timestamp("2025-10-13"))
    assert [str(x[0].date()) for x in found] == ["2026-09-12", "2026-09-25"]
    import inspect
    assert "earlier: " in inspect.getsource(bw) and 'f"{x[0].date()} {x[1]}" for x in found[-4:-1]' in \
        inspect.getsource(bw)


# ================================================================================================== 10 wiring.json
def test_10_every_shown_and_cross_checked_series_names_where_it_is_read(tmp_path):
    import openpyxl
    wiring.reset()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master"
    ws["A5"] = "Near"
    ws["C5"] = '=INDEX(Data!$J$2:$J$9,MATCH($A5&"|price_usd",Data!$A$2:$A$9,0))'
    out = wiring.audit([{"project": "Near", "metric": "price_usd", "now": 2.0}], wb=wb)
    assert out[("Near", "price_usd")]["class"] == "SHOWN" and out[("Near", "price_usd")]["why"] == "read on Master"
    wiring.write(out, tmp_path / "wiring.json")
    import json
    rows = json.loads((tmp_path / "wiring.json").read_text())["rows"]
    assert rows[0]["why"] == "read on Master" and rows[0]["class"] == "SHOWN"
