"""P2e-3 / P2e-4 (Jake 2026-10-28).

P2e-3: Summary's revenue column reads PROTOCOL revenue (config.revenue_headline_metric) on every row; the header says so
and each cell's comment names its series. Morpho reads 0 — a measured 0 (DefiLlama reports 0 on every stored day), and
the comment gives the reason only from Morpho's own record (its fee split, with source and date). An unknown is never 0.

P2e-4: one list (config.REGRESSION_ALLOWLIST) for measuring_guard and sanity_hits.json. An accepted loss is printed by
the guard as "ACCEPTED: <project> <headline> — <reason> (accepted <date>)" and carried under sanity_hits.json's
"accepted", never counted; a loss not on the list — or on an expired entry — is a hit in both.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))
import config  # noqa: E402
import measuring_guard as mg  # noqa: E402
import reader  # noqa: E402
import sanity  # noqa: E402

ON = "2026-10-28"
HL = "A4 | BURN ÷ ISSUANCE (x) — crossover, on the basis to the left"
AERO = "A3 | Net absorption Q0 (tokens) = RETIRED buyback (burned + held) − emissions"
NOT_LISTED = ("Pendle", "Master | Fees Δ6m")


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    """The fixture build (fixtures/real_*.json), as tests/test_round_p1.py builds it: the audit file and the reader."""
    import pandas as pd

    import build_workbook as bw
    import store as sm
    from real_dump import to_store
    d = tmp_path_factory.mktemp("p2e")
    db = d / "metrics.db"
    asofs = []
    for pl in mg._payloads():
        to_store(pl, db)
        asofs.append(pl["meta"]["asof"])
    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(sm, "DB_PATH", str(db))
        mp.setitem(sys.modules, "recalc", None)
        mp.setenv("TOKEN_METRICS_STORE_VALUES", "1")
        st = sm.Store(str(db))
        try:
            bw.build_workbook(st, d / "audit.xlsx", asof=pd.Timestamp(max(asofs)), reader_path=d / "reader.xlsx")
        finally:
            st.close()
    assert "error" not in bw.READER_STATS, bw.READER_STATS
    return {"audit": d / "audit.xlsx", "reader": d / "reader.xlsx"}


# ============================================================================================== P2e-3 Summary revenue
def _summary_col(ws, header):
    return next(c for c in range(1, ws.max_column + 1) if ws.cell(4, c).value == header)


def _summary_row(ws, project):
    return next(r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value == project)


def test_p2e3_the_header_names_the_revenue_it_reads_on_every_row(built):
    assert len(reader.REV_HEADER) <= reader.MAX_HEADER
    assert reader.REV_HEADER == "Protocol revenue, annualised"
    ws = openpyxl.load_workbook(built["reader"])["Summary"]
    c = _summary_col(ws, reader.REV_HEADER)
    assert "Revenue, annualised" not in [ws.cell(4, k).value for k in range(1, ws.max_column + 1)]
    # every project's series is protocol revenue: revenue_usd, or the sum of its protocol revenue streams — never the
    # customer revenue series
    for p in config.PROJECTS:
        m = config.revenue_headline_metric(p["name"])
        assert m in ("revenue_usd", "revenue_total_usd"), (p["name"], m)
        assert "customer_revenue_usd" not in (config.REVENUE_TOTAL.get(p["name"]) or ())
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 2).value is None:                   # an archetype group row
            continue
        name = ws.cell(r, 1).value
        com = ws.cell(r, c).comment
        assert com is not None and f"{name}|{config.revenue_headline_metric(name)}" in com.text, (name, com)


def test_p2e3_morphos_zero_is_the_sources_zero_with_the_series_and_the_reason(built):
    import xlcalc
    wb = openpyxl.load_workbook(built["reader"])
    ws = wb["Summary"]
    c, r = _summary_col(ws, reader.REV_HEADER), _summary_row(ws, "Morpho")
    v = xlcalc.Workbook(wb).value("Summary", r, c)
    assert v == 0, v                                       # a number, not a code: the source reports 0
    text = ws.cell(r, c).comment.text
    assert "Morpho|revenue_usd 'Protocol revenue' from DefiLlama morpho: dailyRevenue" in text, text
    assert "reads 0 because the source reports 0 — a measured 0" in text, text
    assert "a stored row on 90 of the 90 days, 0 of them non-zero" in text, text
    # the reason only from Morpho's own record, with its source and date
    fs = config.PROJECT_BY_NAME["Morpho"]["fee_split"]
    assert "FEE SWITCH IS OFF" in text and fs["source_url"] in text and fs["source_date"] in text, text
    assert "Customer revenue (what users pay — the A2 tab's figure) is a different series" in text, text
    # the stored series itself: every Morpho revenue_usd row in the dump is a stored 0 from DefiLlama
    dump = json.loads((ROOT / "fixtures" / "real_morpho.json").read_text(encoding="utf-8"))
    rows = [x for x in dump["metrics"] if x["metric"] == "revenue_usd"]
    assert rows and {x["source"] for x in rows} == {"defillama"} and all(x["value"] == 0 for x in rows)


def test_p2e3_an_unknown_revenue_never_shows_0(built):
    """A Summary revenue cell holds a number only where its series covers a day of the last 90; else a code."""
    import xlcalc
    wb = openpyxl.load_workbook(built["reader"])
    ws = wb["Summary"]
    audit = reader.Audit(openpyxl.load_workbook(built["audit"])) if hasattr(reader, "Audit") else None
    X = xlcalc.Workbook(wb)
    c = _summary_col(ws, reader.REV_HEADER)
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 2).value is None:
            continue
        name, raw = ws.cell(r, 1).value, ws.cell(r, c).value
        if isinstance(raw, str) and not raw.startswith("="):
            assert raw in reader.CODES and ws.cell(r, c).comment is not None, (name, raw)
            continue
        assert isinstance(X.value("Summary", r, c), (int, float)), (name, raw)
        if audit is not None:
            days = audit.data_field(f"{name}|{config.revenue_headline_metric(name)}",
                                    "Q0: days of the 90 the series existed for")
            assert isinstance(days, (int, float)) and days > 0, (name, days)


class _A:
    def __init__(self, fields):
        self.fields = fields

    def data_field(self, key, header):
        return self.fields.get((key, header))


class _Ctx:
    def __init__(self, fields):
        self.config, self.A = config, _A(fields)


def _fields(name, metric, src, days, events):
    return {(f"{name}|{metric}", "Label"): "Protocol revenue", (f"{name}|{metric}", "Source (of latest point)"): src,
            (f"{name}|{metric}", "Q0: days of the 90 the series existed for"): days,
            (f"{name}|{metric}", "Q0: days with a non-zero value (events)"): events}


def test_p2e3_a_zero_with_no_published_reason_says_so_and_an_unknown_says_no_zero():
    p = {"name": "Testproj", "defillama_fees_slug": "testproj", "archetypes": [3]}
    ctx = _Ctx(_fields("Testproj", "revenue_usd", "defillama", 90, 0))
    note = reader.revenue_series_note(ctx, p, "revenue_usd", 0.0)
    assert "Testproj|revenue_usd 'Protocol revenue' from DefiLlama testproj: dailyRevenue" in note
    assert "reads 0 because the source reports 0" in note and "The source publishes no reason for the 0." in note
    # a fee split that is not 'off' gives no reason either — nothing is inferred
    note = reader.revenue_series_note(ctx, dict(p, fee_split={"status": "active", "note": "x"}), "revenue_usd", 0.0)
    assert "The source publishes no reason for the 0." in note
    # an unknown (no figure) never says "reads 0"
    assert "reads 0" not in reader.revenue_series_note(ctx, p, "revenue_usd", None)


def test_p2e3_a_summed_revenue_names_every_leg():
    p = config.PROJECT_BY_NAME["Near"]
    f = {**_fields("Near", "revenue_usd", "defillama", 90, 90),
         **_fields("Near", "intents_revenue_usd", "defillama:near-intents", 90, 90)}
    note = reader.revenue_series_note(_Ctx(f), p, "revenue_total_usd", 1.0)
    assert "Near|revenue_usd" in note and "Near|intents_revenue_usd" in note and "DefiLlama near-intents" in note


# ============================================================================================== P2e-4 accepted losses
@pytest.fixture
def parked(monkeypatch):
    monkeypatch.setattr(config, "aerodrome_park_reason", lambda state=None, day=None: "PARKED — migration: test")


def test_p2e4_every_allowlist_entry_has_a_reason_and_an_accepted_date():
    for a in config.REGRESSION_ALLOWLIST:
        assert a.get("reason", "").strip(), a
        assert len(str(a.get("accepted") or "")) == 10, a     # YYYY-MM-DD
        assert not a.get("until") or a["accepted"] <= a["until"], a


def test_p2e4_one_match_for_the_guards_key_and_sanitys_detail(parked):
    # the guard keys "A4 | BURN ÷ ISSUANCE …"; sanity.lost_headlines writes "A4 BURN ÷ ISSUANCE …: …"
    for proj, head in (("Hyperliquid", HL), ("Aerodrome", AERO)):
        guard_key = head
        sanity_detail = f"{head[:2]} {head[5:][:70]}: 1.5 in the previous run (2026-10-27), now n/a"
        a, b = config.regression_allowance(proj, guard_key, on=ON), config.regression_allowance(proj, sanity_detail,
                                                                                                on=ON)
        assert a is not None and a is b, (proj, head)


def test_p2e4_the_guard_prints_an_accepted_loss_and_counts_only_the_unlisted(parked, capsys):
    before = {"headlines": {"Hyperliquid": {HL: 1.5}, "Aerodrome": {AERO: 2.0e6}, NOT_LISTED[0]: {NOT_LISTED[1]: 0.1}},
              "bases": {}}
    after = {"headlines": {}, "bases": {}}
    accepted: list = []
    bad = mg.compare(before, after, accepted, on=ON)
    assert bad == [f"HEADLINE LOST {NOT_LISTED[0]} {NOT_LISTED[1]}: 0.1 before, no value after"]
    hl = config.regression_allowance("Hyperliquid", HL, on=ON)
    assert f"ACCEPTED: Hyperliquid {HL} — {hl['reason']} (accepted 2026-10-28)" in accepted
    assert any(x.startswith(f"ACCEPTED: Aerodrome {AERO} — parked (P2c-4)") and x.endswith("(accepted 2026-10-28)")
               for x in accepted), accepted
    mg.print_accepted(accepted)
    out = capsys.readouterr().out
    assert "ACCEPTED LOSSES (config.REGRESSION_ALLOWLIST) — 2, listed, not counted" in out
    assert f"  ACCEPTED: Hyperliquid {HL} — " in out
    # the live check (--newer) by the same rule
    bad, acc = mg.lost_against_previous(before["headlines"], {}, None, "2026-10-27", on=ON)
    assert len(bad) == 1 and bad[0].startswith(f"HEADLINE LOST {NOT_LISTED[0]} ") and len(acc) == 2


def _hits():
    return [("Hyperliquid", "headline_lost", f"A4 {HL[5:]}: 1.5 in the previous run (2026-10-27), now blocked — rule 9"),
            ("Pendle", "headline_lost", "Ma Fees Δ6m: 0.1 in the previous run (2026-10-27), now n/a"),
            ("Chainlink", "fdv_below_mcap", "FDV 1 < market cap 2")]


def test_p2e4_sanity_hits_json_carries_the_accepted_loss_apart(tmp_path):
    body = sanity.write_hits(tmp_path / "sanity_hits.json", "wb.xlsx", ON, _hits(), on=ON)
    assert body == json.loads((tmp_path / "sanity_hits.json").read_text(encoding="utf-8"))
    assert [h[0] for h in body["hits"]] == ["Pendle", "Chainlink"]          # the unlisted loss is a hit
    (acc,) = body["accepted"]
    hl = config.regression_allowance("Hyperliquid", HL, on=ON)
    assert acc["project"] == "Hyperliquid" and acc["rule"] == "headline_lost"
    assert acc["reason"] == hl["reason"] and acc["accepted"] == "2026-10-28" and acc["until"] == "2027-01-31"
    assert acc["line"] == f"ACCEPTED: Hyperliquid {HL} — {hl['reason']} (accepted 2026-10-28)"
    # sanity.py --workbook reads the file back: the accepted loss is answered, not open
    found = sanity.with_build_hits([], tmp_path / "wb.xlsx")
    open_, answered = sanity.split_answered(found, on=ON)   # judged on the test's day, not the clock's
    assert [h[0] for h in answered] == ["Hyperliquid"] and {h[0] for h in open_} == {"Pendle", "Chainlink"}


def test_p2e4_an_expired_entry_is_a_hit_in_both(monkeypatch):
    expired = [dict(a, until="2026-10-01") if a["project"] == "Hyperliquid" else a for a in config.REGRESSION_ALLOWLIST]
    monkeypatch.setattr(config, "REGRESSION_ALLOWLIST", expired)
    accepted: list = []
    bad = mg.compare({"headlines": {"Hyperliquid": {HL: 1.5}}, "bases": {}}, {"headlines": {}, "bases": {}},
                     accepted, on=ON)
    assert bad == [f"HEADLINE LOST Hyperliquid {HL}: 1.5 before, no value after"] and accepted == []
    body = sanity.hits_payload("wb.xlsx", ON, _hits()[:1], on=ON)
    assert body["accepted"] == [] and body["hits"][0][0] == "Hyperliquid"
