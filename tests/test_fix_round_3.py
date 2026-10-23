"""
tests/test_fix_round_3.py — fix round 3 (Jake 2026-10-13): each mechanism on REAL data — the dumps (fixtures/real_*.json)
plus the real API responses captured on 2026-10-13 for the series the dumps predate (fixtures/api_hl_ledgers.json,
api_sky_pnl.json, api_chainlink_staking_path.json, api_chainlink_reward_vault.json, api_merkl_morpho.json,
api_boros_market2.csv, api_etherfi_cashback.json, api_llama_pendle_tokens.json).

  B1   Hyperliquid's issuer-side wallets replayed from their own ledgers: end levels = live state; the team's releases;
       the build: (e) unlocks, NHY and Total Yield numeric, the modelled rewards history labelled MODELLED
  Q2   Sky's payout-type ratios on revenue BEFORE distributions; the implied buyback stays on NPS
  B7   the corrected reference rows grade VERIFIED
  B2   Chainlink: the staking part of each release, FIFO through 0xD50f074E…, out of (e); the community-pool rate
  B3   Morpho: Merkl campaigns PAID off Ethereum (the Arc campaign pays on Ethereum and is not counted twice)
  B4   Aerodrome: the weekly rows dated by the flip that paid them; the team emission on its own flip's rate
  B5   NEAR: the two bridged outflows named on their days
  B6   W1 TVL at one date's prices; W2 Boros $ volume; W3 the cashback its purchases did not cover; W4–W7 config
"""
from __future__ import annotations

import csv
import datetime as dt
import json
import sys
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import config  # noqa: E402
from test_fix_round_2 import _groups  # noqa: E402

FX = ROOT / "fixtures"


def _fx(name: str):
    return json.loads((FX / name).read_text(encoding="utf-8"))


def _rows(metric: str, pts, source: str, tier: int = 2) -> list[dict]:
    return [{"date": str(pd.Timestamp(d).date()), "metric": metric, "value": float(v), "source": source, "tier": tier,
             "fetched_at": "2026-10-13T20:00:00"} for d, v in pts]


# --------------------------------------------------------------------------------------------------------- B1 replay
def _hl_replays():
    from fetch import hl_ledger as h
    d = _fx("api_hl_ledgers.json")
    spec = config.PROJECT_BY_NAME["Hyperliquid"]["hl_ledger_replay"]
    pairs, internal = h.hype_pairs(d["spotMeta_hype"]), set(spec["wallets"].values())
    reps = {n: h.replay_wallet(w["address"], w["userNonFundingLedgerUpdates"], w["delegatorHistory"],
                               w["delegatorRewards"], w["fills"], pairs, internal) for n, w in d["wallets"].items()}
    return d, spec, reps


def test_hl_replay_end_levels_equal_the_live_state():
    d, _spec, reps = _hl_replays()
    for n, w in d["wallets"].items():
        spot = sum(float(b.get("total") or 0) for b in w["spotClearinghouseState"].get("balances") or ()
                   if b.get("coin") == "HYPE")
        ds = w["delegatorSummary"]
        live = spot + sum(float(ds.get(k) or 0) for k in ("delegated", "undelegated", "totalPendingWithdrawal"))
        assert abs(float(sum(reps[n]["end"])) - live) < 1e-6, n


def _hl_rows():
    from fetch import hl_ledger as h
    _d, spec, reps = _hl_replays()
    fr = h.daily(reps, dt.date(2025, 7, 10), dt.date(2026, 10, 13))
    out = []
    for m, pts in h.series(fr, spec).items():
        out += _rows(m, pts, h.SOURCE, 1)
    return out


def test_hl_team_releases_are_the_ledgers_own():
    rows = pd.DataFrame(_hl_rows())
    rel = rows[rows.metric == "issuer_release_replay_tokens"].set_index("date")["value"]
    assert rel["2026-10-07"] == pytest.approx(3_750_000, abs=1)
    assert rel["2025-11-29"] == pytest.approx(1_745_746, abs=1)
    team = rows[rows.metric == "team_wallet_tokens_replay"].set_index("date")["value"]
    assert team["2026-10-12"] == pytest.approx(237_915_860.467, abs=0.01)


def test_hl_build_unblocks_unlocks_and_labels_the_model():
    g = _groups("Hyperliquid", _hl_rows())
    unl = g[("Hyperliquid", "unlock_release_tokens")]
    q0 = unl[(pd.to_datetime(unl.date) > "2026-07-15") & (pd.to_datetime(unl.date) <= "2026-10-13")]
    assert q0.value.sum() > 4_000_000
    em = g[("Hyperliquid", "emissions_tokens")]
    assert em.source.astype(str).str.startswith("modelled:hl_reference_yield").any()
    assert config.is_modelled("modelled:hl_reference_yield")
    assert ("Hyperliquid", "future_emissions_modelled_tokens") in g


def test_hl_burn_total_keeps_the_funds_own_reads_and_covers_q0():
    """Round 3 review H1 (Jake 2026-10-14): the total's history leg kept only the DefiLlama rows (to 2026-09-11), so the
    Fund's own reads 2026-09-12..09-28 fell out — Q0 2,146,959 HYPE over 18 'covered' days read -18.25% net supply and
    put non-stakers' Total Yield at +9.0% against NHY -6.40%. The leg is now continuous and covers all 90 days."""
    import build_workbook as bw
    g = _groups("Hyperliquid", _hl_rows())
    tb = g[("Hyperliquid", "total_burn_tokens")].drop_duplicates("date", keep="last")
    s = tb.set_index(pd.to_datetime(tb.date))["value"].astype(float)
    asof = pd.Timestamp("2026-10-13")
    hole = s[(s.index >= "2026-09-12") & (s.index <= "2026-09-28")]
    assert len(hole) >= 10 and hole.sum() > 300_000         # the Fund's spotClearinghouseState reads, kept
    q0 = s[(s.index > asof - pd.Timedelta(days=90)) & (s.index <= asof)].sum()
    assert q0 == pytest.approx(2_467_817.2, abs=1)
    s_cov = bw._after_leg_gap(tb, s)
    assert bw._window_coverage(s_cov, asof - pd.Timedelta(days=90), asof)[0] == 90
    # the retirement above DefiLlama's revenue_usd it exposes (114%) is the Fund's documented intake, not a gap
    f = config.PROJECT_BY_NAME["Hyperliquid"]["buyback_funding"]
    assert "hyperliquid-docs/trading/fees" in f["source_url"] and f["source_date"] == "2026-10-14"


# ------------------------------------------------------------------------------------------------------------ Q2 Sky
def _sky_rows():
    from fetch.sky_accounting import nps_candidates
    d = _fx("api_sky_pnl.json")
    c = nps_candidates(d["pnl"]["data"], d["sm"]["data"])["revenue-expense"]
    return _rows("net_revenue_pre_distribution_usd",
                 [(pd.Period(m, "M").to_timestamp("M"), v) for m, v in c.items() if m < "2026-10"],
                 "sky_accounting:pnl.revenue-expense", 3)


def test_sky_ratio_base_is_before_distributions_and_the_implied_buyback_stays_on_nps():
    assert config.revenue_ratio_base_metric("Sky") == "net_revenue_pre_distribution_usd"
    assert config.revenue_base_metric("Sky") == "net_protocol_surplus_usd"
    assert config.revenue_ratio_base_metric("Near") == config.revenue_base_metric("Near") == "revenue_total_usd"
    sep = {r["date"]: r["value"] for r in _sky_rows()}["2026-09-30"]
    assert sep == pytest.approx(15_379_185, abs=1)          # revenue − expense; NPS (after distribution) 7,153,241


def test_sky_reference_reads_the_msc_series_and_ba_is_a_labelled_basis_difference():
    # G5 D1: the TMF months are checked against our own read of Soter's report; Block Analitica's month is not graded
    import references as rf
    msc = _rows("net_revenue_msc_usd", [(pd.Timestamp("2026-09-30"), 14_635_947.43)],
                "sky_msc:soterlabs/settlement-cycle sky_total[accrual]", 3)
    g = _groups("Sky", _sky_rows() + msc)
    r = {x["metric"]: x for x in rf.check(g)}["net_revenue_sep_2026"]
    assert r["ours"] == pytest.approx(14_635_947.43, abs=0.01) and r["status"] == "VERIFIED"
    b = {x["metric"]: x for x in rf.basis_rows(g)}["net_revenue_ba_vs_msc_2026-09"]
    assert b["status"] == "N/A BY DESIGN" and b["status"] != "BROKEN"
    assert b["ours"] == pytest.approx(15_379_185, abs=1)
    assert "KNOWN ACCOUNTING-BASIS DIFFERENCE" in b["detail"] and "accrual" in b["detail"]
    assert "15,379,185" in b["detail"] and "14,635,947" in b["detail"] and "+5.08%" in b["detail"]


# ------------------------------------------------------------------------------------------------------- B7 references
def test_corrected_reference_rows_verify():
    import references as rf
    uni = {r["metric"]: r for r in rf.check(_groups("Uniswap"))}
    assert uni["franchiser_recall_inflow"]["status"] == "VERIFIED"
    cl = {r["metric"]: r for r in rf.check(_groups("Chainlink"))}
    assert cl["reserve_holdings_2026-04-15"]["status"] == cl["reserve_holdings_end_q1_2026"]["status"] == "VERIFIED"


# ------------------------------------------------------------------------------------------------------- B2 Chainlink
def _cl_path_rows():
    P = config.CL_STAKING_PATH
    t = _fx("api_chainlink_staking_path.json")["transfers"]
    lump, fwd = P["lump_wallet"].lower(), P["forwarder"].lower()
    hops = {h.lower() for h in P["staking_hops"]}
    nc = {c["address"].lower() for c in config.PROJECT_BY_NAME["Chainlink"]["contracts"].values()
          if c.get("metric_override") == "noncirculating_holding_tokens" and c.get("chain") == "ethereum"}
    acc = {m: {} for m in ("staking_path_in_tokens", "staking_path_out_tokens", "staking_funding_tokens")}

    def add(m, ts, v):
        acc[m][ts[:10]] = acc[m].get(ts[:10], 0.0) + v
    for x in t[P["lump_wallet"]]:
        v = int(x["value"]) / 1e18
        if x["to"] == lump and x["from"] in nc:
            add("staking_path_in_tokens", x["ts"], v)
        elif x["from"] == lump:
            add("staking_path_out_tokens", x["ts"], v)
    for x in t[P["forwarder"]]:
        if x["from"] == fwd and x["to"] in hops:
            add("staking_funding_tokens", x["ts"], int(x["value"]) / 1e18)
    out = []
    for m, a in acc.items():
        out += _rows(m, sorted(a.items()), "logscan:" + m)
    return out


def test_chainlink_staking_part_of_each_release_comes_out_of_e_fifo():
    rows = _cl_path_rows()
    fund = sum(r["value"] for r in rows if r["metric"] == "staking_funding_tokens")
    assert fund == pytest.approx(6_682_500 + 1_445_000, abs=1)       # the two staking hops, all time
    base, g = _groups("Chainlink"), _groups("Chainlink", rows)
    split = g[("Chainlink", "staking_funding_release_tokens")].set_index("date")["value"]
    assert split.sum() == pytest.approx(327_500, abs=1)               # FIFO: the 2025-12-19 lump, within the dump
    b, a = (x[("Chainlink", "unlock_release_tokens")].value.sum() for x in (base, g))
    assert b - a == pytest.approx(327_500, abs=1)
    assert "staking_funding_release_tokens" in str(g[("Chainlink", "unlock_release_tokens")].source.iloc[-1])
    # P1 (round 3 review): FIFO is labelled an attribution convention, ACCEPTED, on both series
    for m in ("unlock_release_tokens", "staking_funding_release_tokens"):
        assert "ATTRIBUTION CONVENTION: FIFO" in str(g[("Chainlink", m)].source.iloc[-1]) and "ACCEPTED" in str(
            g[("Chainlink", m)].source.iloc[-1])


def test_chainlink_community_pool_rate_from_the_vault():
    from fetch.reward_vault import community_rate, decode_buckets
    d = _fx("api_chainlink_reward_vault.json")
    assert d["getStakingPools"].lower().endswith(config.PROJECT_BY_NAME["Chainlink"]["reward_vault_rates"]
                                                 ["community_pool"].lower()[2:])
    rate = community_rate(decode_buckets(d["getRewardBuckets"]), int(d["community_getTotalPrincipal"], 16),
                          d["timestamp"])
    assert rate == pytest.approx(0.0432, abs=0.0005)


# ------------------------------------------------------------------------------------------------------ B3 Morpho
def _morpho_rows():
    from fetch import merkl_campaigns as mc
    fx = _fx("api_merkl_morpho.json")
    spec = config.PROJECT_BY_NAME["Morpho"]["merkl_campaigns"]
    tot, seen = {}, set()
    for cid, ch in spec["chains"].items():
        per, _ = mc.daily_amounts(fx["chains"][str(cid)], ch["token"], dt.date(2026, 10, 13), int(cid), seen)
        for d, v in per.items():
            tot[d] = tot.get(d, 0.0) + v
    d0, pts = min(tot), []
    while d0 < dt.date(2026, 10, 13):
        pts.append((d0, tot.get(d0, 0.0)))
        d0 += dt.timedelta(days=1)
    return _rows(spec["metric"], pts, "merkl_campaigns", 1)


def test_merkl_counts_where_a_campaign_pays_not_where_it_computes():
    from fetch import merkl_campaigns as mc
    fx = _fx("api_merkl_morpho.json")
    arc = fx["arc_not_declared"]
    assert arc and all(c["distributionChainId"] == 1 for c in arc)
    per, skipped = mc.daily_amounts(arc, "0x58D97B57BB95320F9a05dC918Aef65434969c2B2", dt.date(2026, 10, 13), 5042)
    assert not per and "distributed on chain 1" in skipped[0][1]


def test_morpho_b_against_the_merkl_reference():
    import references as rf
    g = _groups("Morpho", _morpho_rows())
    r = {x["metric"]: x for x in rf.check(g)}["merkl_rewards_rate"]
    # F5 (2026-10-14): the reference reads the same Merkl API our series does — same-source, so a match is
    # LIVE-UNVERIFIED (it checks the arithmetic, not the figure), never VERIFIED
    assert r["status"] == "LIVE-UNVERIFIED" and 6_500 < r["ours"] < 8_500 and "same-source" in r["detail"]


# --------------------------------------------------------------------------------------------------- B4 Aerodrome
def test_aerodrome_weekly_rows_sit_on_the_flip_that_paid_them():
    g = _groups("Aerodrome")
    gi = g[("Aerodrome", "gross_issuance_tokens")].set_index("date")["value"]
    gi.index = pd.to_datetime(gi.index)
    assert gi[pd.Timestamp("2026-10-08")] == pytest.approx(4_184_974.45, abs=1)   # the Voter transfer of 10-08
    assert pd.Timestamp("2026-10-15") not in gi.index                              # the epoch still in progress
    im = g[("Aerodrome", "issuer_mint_tokens")].set_index("date")["value"]
    im.index = pd.to_datetime(im.index)
    assert im[pd.Timestamp("2026-10-08")] == pytest.approx(220_689.5, rel=1e-4)    # on-chain team transfer
    assert im[pd.Timestamp("2025-10-09")] == 0                                     # teamRate 0 from the 10-09 flip
    assert "flip_offset_days" not in config.ISSUER_MINTS["Aerodrome"]


# ------------------------------------------------------------------------------------------------------------ B5 NEAR
def test_near_bridged_outflows_are_recorded_and_named():
    by_day = {}
    for o in config.NEAR_REVENUE_OUTFLOWS:
        by_day[o["date"]] = by_day.get(o["date"], 0) + o["tokens"]
    assert by_day["2026-04-22"] == 45_010 and by_day["2026-07-10"] == 12_000      # 45,000 + 10 on 04-22 (run Y2)
    assert all("not retired" in o["class"] for o in config.NEAR_REVENUE_OUTFLOWS)
    g = _groups("Near")
    b = g[("Near", "actual_buyback_tokens")]
    named = b[b.source.astype(str).str.contains("bridged out")]
    assert all(v < 0 for v in named.value)


# --------------------------------------------------------------------------------------------------------------- B6
def test_w1_tvl_at_one_dates_prices():
    from fetch.llama import constant_price_tvl
    j = _fx("api_llama_pendle_tokens.json")
    rows, note = constant_price_tvl(j)
    last_usd = sum(v for v in j["tokensInUsd"][-1]["tokens"].values() if v)
    assert rows[-1][1] == pytest.approx(last_usd, rel=1e-9)                         # the latest day is itself
    assert "priced at" in note
    assert "protocol_tvl_constant_price_usd" in config.BUSINESS_SLOTS["Pendle"]
    assert "eth_staked_etherfi_tokens" in config.BUSINESS_SLOTS["Ether.fi"]


def test_w2_boros_units_are_the_collateral():
    from fetch.boros import collateral_pricing
    d = _fx("api_boros_meta.json")
    assets = {int(a["tokenId"]): a for a in d["assets"] if a.get("tokenId") is not None}
    by_id = {m["marketId"]: m for m in d["markets"]}
    assert collateral_pricing(by_id[2], assets) == "row"        # WBTC collateral, BTC underlying: Boros's own price
    assert collateral_pricing(by_id[161], assets) == "par"      # USD₮0 collateral on a BTC market: units are dollars


def test_w2_boros_dollar_volume_is_yield_units_times_borors_own_price():
    from fetch.boros import daily_usd
    with open(FX / "api_boros_market2.csv", newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    per, unpriced = daily_usd(rows, dt.date(2025, 8, 13), dt.date(2025, 8, 13))
    assert per[dt.date(2025, 8, 13)] == pytest.approx(129.44898787684116 * 120134.09, rel=1e-9)
    assert unpriced == 0


def test_w3_only_the_cashback_purchases_did_not_cover_is_dilution():
    d = _fx("api_etherfi_cashback.json")
    paid = _rows("cashback_paid_tokens", d["paid_daily"].items(), "logscan:cashback_out_optimism")
    bought = _rows("programmatic_buyback_tokens", [(x["date"], x["ethfi_bought"]) for x in d["cow_buys_daily"]],
                   "logscan:programmatic_buyback_inflow")
    days = pd.date_range("2026-09-01", "2026-10-12")
    have = {r["date"] for r in bought}
    bought += _rows("programmatic_buyback_tokens", [(x, 0.0) for x in days if str(x.date()) not in have],
                    "logscan:programmatic_buyback_inflow")
    have = {r["date"] for r in paid}
    paid += _rows("cashback_paid_tokens", [(x, 0.0) for x in days if str(x.date()) not in have],
                  "logscan:cashback_out_optimism")
    g = _groups("Ether.fi", paid + bought)
    unf = g[("Ether.fi", "cashback_unfunded_tokens")]
    assert sum(r["value"] for r in paid) == pytest.approx(1_378_243.74, abs=1)
    assert 0 < unf.value.sum() < 1_000                                              # 404 ETHFI ahead of the buys
    assert ("Ether.fi", "outflow_c_tokens") in g


def test_w4_to_w7_are_wired():
    assert config.outflow_metrics("Maple", "b") == ("drips_rewards_tokens",)
    assert any(s["key"] == "drips_out" for s in config.PROJECT_BY_NAME["Maple"]["log_scans"])
    assert config.LEAKAGE_III["Sky"]["spend"] == "buyback_spending_usd_ba"
    for n in ("Ether.fi", "Maple", "Pendle"):
        assert n not in config.INSIDER_LOCK                      # review W6: no insider identified -> NOT WIRED
        assert config.INSIDER_SHARE_MONITOR[n]["numerator"] == "insider_stake_shares"
        assert any(k.startswith("insider_stake_") for k in config.PROJECT_BY_NAME[n]["contracts"])
    assert config.RELEASE_LESS["Pendle"] == ("issuer_to_gauges_tokens",)


def test_w7_pendle_gauge_transfers_come_out_of_e():
    rows = _rows("issuer_to_gauges_tokens", [("2025-12-23", 100_000), ("2026-02-06", 50_000)],
                 "logscan:issuer_to_gauges")
    base, g = _groups("Pendle"), _groups("Pendle", rows)
    b = base[("Pendle", "unlock_release_tokens")].set_index("date")["value"]
    a = g[("Pendle", "unlock_release_tokens")].set_index("date")["value"]
    # each transfer comes out of its day's release (or the next stored day's), never below 0: on 2025-12-23 the stock
    # fell 99,924.66 net of an inflow the same day, so 75.34 of the 100,000 had no release left to come out of
    assert 149_900 < b.sum() - a.sum() <= 150_000
    assert (a >= 0).all()
    assert "issuer_to_gauges_tokens" in str(g[("Pendle", "unlock_release_tokens")].source.iloc[-1])


# ------------------------------------------------------------------------------------- BUG (Jake's run): llama 6840 pts
def test_llama_price_history_is_asked_in_calls_under_the_500_point_cap(tmp_path):
    """The first read of the 456-day window used to ask every coin in one URL: 15 x 456 = 6,840 points, HTTP 400 for all
    15 projects (the price cross-check and the Q-4 price history were blind). Each call now stays at or under 500."""
    from fetch import xref
    from fetch.base import FetchOutput
    from fetch.logcache import DailyChecks
    calls = []

    class Fake:
        def get(self, url, params=None, **_):
            keys = url.rsplit("/", 1)[1].split(",")
            calls.append((keys, params["span"]))
            start = int(params["start"])
            return {"coins": {k: {"prices": [{"timestamp": start + 86_400 * i, "price": 1.0 + i}
                                             for i in range(params["span"])]} for k in keys}}
    cr = xref.CrossRefs(http=Fake(), daily=DailyChecks(tmp_path), robots=lambda u: (True, "ok"))
    out = FetchOutput()
    cr._llama(list(config.PROJECTS), {"days": 30}, "2026-10-13", out)
    assert calls and all(len(k) * span <= xref.LLAMA_MAX_POINTS for k, span in calls)
    assert sorted(k for ks, _ in calls for k in ks) == sorted({cr.llama_key(p)[0] for p in config.PROJECTS})
    got = pd.concat(out.frames)
    assert set(got.project) == {p["name"] for p in config.PROJECTS} and set(got.metric) == {"price_usd_llama"}
    assert not [e for e in out.log if e.status == "failed"]


def test_w4_maple_drips_form_outflow_b_on_the_real_series():
    d = _fx("api_maple_drips_daily.json")["daily"]
    paid = {r["date"]: r["syrup_out"] for r in d}
    days = pd.date_range("2025-07-01", "2026-10-12")
    rows = _rows("drips_rewards_tokens", [(x, paid.get(str(x.date()), 0.0)) for x in days], "logscan:drips_out")
    g = _groups("Maple", rows)
    b = g[("Maple", "outflow_b_tokens")]
    assert b.value.sum() == pytest.approx(4_097_369.65, abs=1)
    assert b[pd.to_datetime(b.date) > "2026-02-18"].value.sum() == 0          # Drips ended: the last payout 02-18


# ---------------------------------------------------------------------------------------------------- V2 the guard
def test_a_stated_zero_is_a_value_not_a_lost_headline():
    """Jake's round-2 run: Uniswap A4 and Plume A1 gross issuance Q0 flagged lost ("0 last run, now 0") — the cell
    reads "0 — fixed supply: …", which is a value with its reason."""
    import openpyxl
    import sanity
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "A4 Permanent Burn"
    ws["A1"] = "0 — fixed supply: UNI mints nothing (the 2% inflation switch has never been used)"
    ws["A2"] = "n/a — gross_issuance_tokens: no rows stored"
    cells = [{"project": "Uniswap", "sheet": "A4 Permanent Burn", "header": "Gross issuance Q0", "cell": c}
             for c in ("A1", "A2")]
    prev = {"Uniswap": {"A4 | Gross issuance Q0": 0.0}}
    assert sanity.stated_number(ws["A1"].value) == 0.0 and sanity.stated_number("1,250 — x") == 1250.0
    assert sanity.lost_headlines(wb, prev, cells[:1], "2026-10-12") == []
    hits = sanity.lost_headlines(wb, prev, cells[1:], "2026-10-12")
    assert len(hits) == 1 and hits[0][1] == "headline_lost"


# ------------------------------------------------------------------------------------------------- V1 unequal coverage
@pytest.mark.parametrize("project", ["Fluid", "Pendle"])
def test_subsidy_cover_on_an_incomplete_denominator_is_blocked(tmp_path, monkeypatch, project):
    import xlcalc
    from test_overnight_ab5f2a9 import _get, _row
    from test_trajectory import _built
    wb, _ = _built(tmp_path, monkeypatch, project)
    o = _row(wb, xlcalc.Workbook(wb), "CP Comparability", project)
    cover = _get(o, "SUBSIDY COVER")
    assert str(cover).startswith("BLOCKED — unequal coverage: the incentive denominator is incomplete"), cover
    assert config.OUTFLOW_INCOMPLETE[(project, "b")][:20] in str(cover)


@pytest.mark.parametrize("project,sym,n", [("Ether.fi", "sETHFI", 10), ("Maple", "stSYRUP", 8), ("Pendle", "sPENDLE", 3)])
def test_w6_insider_share_is_not_wired_when_no_checked_wallet_holds_shares(tmp_path, monkeypatch, project, sym, n):
    """Round 3 review W6 (Jake 2026-10-14): 'If none are identified, the cell is NOT WIRED with the route, not 0.'"""
    import xlcalc
    from test_overnight_ab5f2a9 import _get, _row
    from test_trajectory import _built
    wb, _ = _built(tmp_path, monkeypatch, project)
    o = _row(wb, xlcalc.Workbook(wb), "CP Comparability", project)
    cell = str(_get(o, "Insider share of staked / locked supply"))
    # G17-2a: still never 0 — now MISSING E, naming what Jake must supply (the issuer wallet list)
    assert cell.startswith(f"{config.INSIDER_LIST_NEEDED}no insider wallet identified among {sym} holders"), cell
    assert len(cell) <= 255                                                   # an Excel string constant
    what = str(_get(o, "Insider share — what is divided"))
    assert f"Wallets checked ({n}, 0 {sym} each on 2026-10-14)" in what     # E6: re-read 2026-10-14 (E1 added two)
    leak = _get(o, "Leakage (i)")
    assert leak in ("", None) or str(leak).startswith(config.INSIDER_LIST_NEEDED), leak     # G17-2a: the same reason


def test_only_declared_gaps_block():
    assert config.outflow_incomplete("Morpho", ["b"]) == []         # B3: off-Ethereum campaigns now read
    assert config.outflow_incomplete("Aethir", ["b"]) and config.outflow_incomplete("Fluid", ["b", "c"])


# ------------------------------------------------------------------------------------- V3–V6 on Jake's 10-13 dumps
def test_v5_chainlink_quarterly_release_reads_the_flows_line(monkeypatch):
    import build_workbook as bw
    import references as rf
    from real_dump import load
    monkeypatch.setattr(bw, "_SCOPE", [config.PROJECT_BY_NAME["Chainlink"]])
    g = _groups("Chainlink")
    flows = bw.flow_statements(g, pd.Timestamp(load("Chainlink")["meta"]["asof"]))
    r = {x["metric"]: x for x in rf.check(g, flows=flows)}["quarterly_release"]
    assert r["ours"] == pytest.approx(18_750_000) and "Flows: Token flow" in r["route"]
    assert r["status"] == "LIVE-UNVERIFIED"                      # a press figure: a cross-check, never VERIFIED


def test_v3_maple_dao_safe_move_is_named_and_only_the_eoa_leg_is_a_release():
    import references as rf
    g = _groups("Maple")
    r = {x["metric"]: x for x in rf.check(g)}["dao_safe_balance"]
    # D5 (run review 2026-10-14): the reference is now read AFTER the move (block 26,185,275) — VERIFIED, the move kept
    # in the row's note
    assert r["status"] == "VERIFIED" and r["ours"] == pytest.approx(42_290_111.3257, abs=0.001)
    row = next(x for x in csv.DictReader(open(ROOT / "independent_references.csv", encoding="utf-8"))
               if x["project"] == "Maple" and x["metric"] == "dao_safe_balance")
    assert "0xd628a262" in row["note"] and "4,000,000" in row["note"]
    rel = g[("Maple", "unlock_release_tokens")]
    assert float(rel[pd.to_datetime(rel.date) == pd.Timestamp("2026-10-13")].value.iloc[0]) == pytest.approx(4e6, abs=1)


def test_e5_eth_burn_on_one_measured_basis_after_cq(monkeypatch, tmp_path):
    """E5 (i), was V4: Etherscan's BurntFees is the previous day's closing snapshot, so each stored delta measured the
    day BEFORE its date and 2026-09-29 was counted twice at the seam. Section CQ, run on Jake's real dump, moves the
    rows and removes the doubled derived day; the 30 days to 2026-10-12 then compare with the execution-layer burn
    measured block by block (2,458.9064 ETH) — the Etherscan leg equal to it, the DefiLlama-derived leg ~22 ETH low."""
    import builtins
    import contextlib
    import io
    import sqlite3

    import real_dump
    import references as rf
    import run_sql
    payload = real_dump.load("Ethereum")
    db = tmp_path / "metrics.db"
    real_dump.to_store(payload, db)
    sec = run_sql.parse_sections(run_sql.SQL_FILE.read_text(encoding="utf-8"))["CQ"]
    conn = sqlite3.connect(db)
    monkeypatch.setattr(builtins, "input", lambda *a: "DELETE CQ")
    with contextlib.redirect_stdout(io.StringIO()):
        assert run_sql.run_delete(conn, sec["text"], "CQ", db, sec["authored"]) == 0
        assert run_sql.run_delete(conn, sec["text"], "CQ", db, sec["authored"]) == 1      # a second run: nothing
    q = lambda sql: conn.execute(sql).fetchall()  # noqa: E731
    es = "project = 'Ethereum' AND source LIKE 'etherscan:ethsupply2.BurntFees%'"
    assert q(f"SELECT MIN(date), MAX(date), COUNT(*) FROM metrics WHERE {es} AND metric = 'gross_burn_tokens'") \
        == [("2026-09-29", "2026-10-12", 11)]
    assert q(f"SELECT COUNT(*) FROM metrics WHERE {es} AND substr(fetched_at, 1, 10) = date") == [(0,)]
    day = dict(q("SELECT source, value FROM metrics WHERE project = 'Ethereum' AND metric = 'gross_burn_tokens' "
                 "AND date = '2026-09-29'"))
    assert list(day) == ["etherscan:ethsupply2.BurntFees:delta"] and day[list(day)[0]] == pytest.approx(230.751861, abs=1e-6)
    cols = ("date", "metric", "value", "source", "tier", "fetched_at")
    rows = [dict(zip(cols, r)) for r in q(f"SELECT {', '.join(cols)} FROM metrics WHERE project = 'Ethereum'")]
    conn.close()
    monkeypatch.setattr(real_dump, "load", lambda project, vintage=None: {**payload, "metrics": rows})
    r = {x["metric"]: x for x in rf.check(_groups("Ethereum", []))}["burn_30d"]
    assert r["status"] == "VERIFIED", r
    assert r["ours"] == pytest.approx(2_437.0675, abs=0.001) and "-0.89%" in r["detail"]


def test_v6_a_short_leg_is_named_a_coverage_gap(monkeypatch):
    """GEODNET Q-4 on Jake's 10-13 run read +2.71%: the opening was stored (451,925,750.34, his workbook) while the
    mining release starts 2025-09-29 — 20 of 90 days. The status names a COVERAGE gap."""
    import build_workbook as bw
    from real_dump import load
    monkeypatch.setattr(bw, "_SCOPE", [config.PROJECT_BY_NAME["GEODNET"]])
    g = dict(_groups("GEODNET"))
    asof = pd.Timestamp(load("GEODNET")["meta"]["asof"])
    start, _end = bw._flow_window(asof, 4)
    c = g[("GEODNET", "circulating_supply_used")]
    first = c.iloc[[0]].assign(date=start, value=451_925_750.3365579)
    g[("GEODNET", "circulating_supply_used")] = pd.concat([first, c], ignore_index=True)
    st = next(x for x in bw.flow_statements(g, asof) if x["Project"] == "GEODNET" and x["Line"] == "Status")
    assert "COVERAGE gap" in st["Q-4"] and "pool_release_tokens covers 20/90 days" in st["Q-4"], st["Q-4"]


# --------------------------------------------------------------------------- round 3 run results (Jake 2026-10-14)
def test_x3_near_probe_is_inconclusive_on_429_and_on_empty_reads(monkeypatch, capsys):
    """X3: 'A 429 or empty read must give INCONCLUSIVE, never a pass.' Two of three wallets answered HTTP 429 and the
    verdict said every outflow was on the list. The probe now waits a 429 out (paced by fetch.nearblocks' own pacer)
    and gives no verdict on a partial or empty read."""
    import check_offline_items as coi
    from fetch import nearblocks as nb
    clock = [0.0]
    monkeypatch.setattr(nb, "_clock", lambda: clock[0])
    monkeypatch.setattr(nb, "_sleep", lambda s: clock.__setitem__(0, clock[0] + s))
    calls = []

    class _Limited:
        status_code = 429

        def json(self):
            return {"message": "rate limited"}

    monkeypatch.setattr(coi.requests, "get", lambda url, **k: calls.append(url) or _Limited())
    coi.near_revenue_wallet_flows()
    out = capsys.readouterr().out
    assert "VERDICT: INCONCLUSIVE" in out and "every outflow of 1,000 NEAR or more is on" not in out
    assert len(calls) == 6 * 4                         # each of the 6 reads: the call + 3 retries after the minute
    assert clock[0] >= 3 * 6 * nb.RATE_LIMIT_WAIT      # it waited the limit out, it did not hammer

    ok = [("a receipts", "complete", 12), ("a ft-txns", "complete", None)]
    assert coi.near_flows_verdict(ok, [], []).startswith("every outflow")
    empty = ok + [("b receipts", "complete", 0)]
    assert coi.near_flows_verdict(empty, [], []).startswith("INCONCLUSIVE") and "b receipts: 0 rows" in \
        coi.near_flows_verdict(empty, [], [])
    assert coi.near_flows_verdict(ok + [("c receipts", "HTTP 429", 1000)], [], []).startswith("INCONCLUSIVE")


def test_x4_nearblocks_budget_covers_a_day_of_free_plan_pacing():
    """X4: the tier timed out at 60s on free-plan pacing (6 credits a minute) with 5 of its calls read. The budget now
    covers every pacing window the day's calls need, plus one window of slack."""
    import math

    import fetch
    from fetch import nearblocks as nb
    calls = fetch.nearblocks_daily_calls()
    assert calls >= 9
    windows = math.ceil(calls / nb.CREDITS_PER_MINUTE)
    assert fetch.TIER_BUDGET_S["nearblocks"] >= 60 * windows + 60       # the windows + one of slack


def test_x2_boros_first_read_resumes_across_runs_and_stores_only_when_complete(tmp_path, monkeypatch):
    """X2: the tier timed out at 120s with 0 rows. The adapter now keeps each market's days in .cache/boros/state.json,
    stops itself at its own budget, reads the stalest markets first and stores nothing until every market is read
    through a day. Real export rows (fixtures/api_boros_market2.csv) for every market."""
    import fetch.boros as fb
    from fetch.base import FetchOutput
    monkeypatch.setattr(fb, "today", lambda: pd.Timestamp("2026-10-13"))
    rows = list(csv.DictReader((FX / "api_boros_market2.csv").open(newline="", encoding="utf-8")))
    meta = _fx("api_boros_meta.json")
    real = [m for m in meta["markets"] if fb.collateral_pricing(m, {int(a["tokenId"]): a for a in meta["assets"]})]
    assert real
    mk = [dict(real[i % len(real)], marketId=100 + i) for i in range(5)]   # the real records, five market ids
    clock, calls = [0.0], []

    class _Http:
        def get(self, url, params=None, **k):
            if url.endswith("/v1/markets"):
                return {"results": mk}
            return {"results": meta["assets"]}

        def get_text(self, url, params=None, **k):
            calls.append(params["marketId"])
            clock[0] += 100.0                                # each export "takes" 100s of the budget
            return "\n".join([",".join(rows[0].keys())] + [",".join(str(r[c]) for c in rows[0]) for r in rows])

    class _Daily:
        def __init__(self):
            self.done_keys = set()

        def due(self, k, day):
            return k not in self.done_keys

        def done(self, k, day):
            self.done_keys.add(k)

    daily = _Daily()
    spec = config.PROJECT_BY_NAME["Pendle"]["boros_volume"]
    outs = []
    for _run in range(4):
        b = fb.BorosVolume(http=_Http(), daily_checks=daily, cache_dir=tmp_path, max_seconds=150.0,
                           clock=lambda: clock[0])
        out = FetchOutput()
        b._project("Pendle", spec, out)
        outs.append(out)
    # run 1 and 2: two markets each (cut by the budget), nothing stored; run 3: the fifth, then everything stored
    stored = [sum(len(f) for f in o.frames) if getattr(o, "frames", None) else 0 for o in outs]
    assert stored[0] == stored[1] == 0 and stored[2] > 0, stored
    assert len(calls) == 5 and len(set(calls)) == 5                         # each market read once, none twice
    assert ("boros_volume:Pendle" in daily.done_keys)                       # complete for yesterday: done for today
    assert stored[3] == 0                                                   # run 4: skipped as current
    st = json.loads((tmp_path / "state.json").read_text())
    assert all(v["through"] == "2026-10-12" for v in st["markets"].values())


def _hole_fakes(monkeypatch):
    """The X1 explorer hole, as fakes over real data: (fixture, logs, explorer missing the hole block, reader)."""
    import fetch.logscan as ls
    fx = _fx("api_etherfi_cashback_hole.json")
    logs = [ls._norm_log(x) for x in fx["logs"]]
    hole, h = fx["explorer_missing_block"], fx["holder"].lower()
    bal0 = int(fx["balance_before"])
    monkeypatch.setattr(ls, "today", lambda: pd.Timestamp("2026-10-13"))

    def match(e, topics, lo, hi):
        return lo <= e["blockNumber"] <= hi and all(t is None or e["topics"][i] == t.lower()
                                                    for i, t in enumerate(topics))

    class _Explorer:
        def get_logs(self, chain_id, token, topics, lo, hi):
            return [e for e in logs if match(e, topics, lo, hi) and e["blockNumber"] != hole], {"explorer": "blockscout"}

        def block_at(self, chain_id, ts):
            return fx["from_block"]

    class _Reader:
        balance_reads = 0

        def erc20_balance_at(self, chain, token, holder, block):
            _Reader.balance_reads += 1
            sign = lambda e: (1 if ls.topic_address(e["topics"][2]) == h else -1)  # noqa: E731
            return bal0 + sum(sign(e) * ls._amount(e) for e in logs if e["blockNumber"] <= block)

        def scan_logs(self, chain, base, lo, hi, rpc_only=False):
            assert rpc_only, "the hole must be read from the chain, not the explorer that lacks it"
            return [e for e in logs if match(e, base["topics"], lo, hi)], 1, 1

        def checksum(self, a):
            return a

        def block_timestamp(self, chain, b):
            # OP: 2 s blocks; block 158,047,200 is 2026-10-11 04:32:57 UTC (eth_getBlockByNumber, read 2026-10-14)
            return int(pd.Timestamp("2026-10-11 04:32:57", tz="UTC").timestamp()) + (b - hole) * 2

        def erc20(self, chain, token):
            class _F:
                class functions:
                    @staticmethod
                    def decimals():
                        class _C:
                            @staticmethod
                            def call():
                                return 18
                        return _C
            return _F

    return fx, logs, _Explorer(), _Reader()


def test_x1_an_explorer_hole_is_read_from_the_chain_not_tolerated(tmp_path, monkeypatch):
    """X1: Ether.fi's OP cashback window was 19.808964 ETHFI short of balanceOf. Cause: Blockscout has NO logs for OP
    block 158,047,200 (both hosts, every API), while the chain has the distributor's Transfer of exactly that amount.
    Real data: the RPC's Transfer logs in and out of the distributor around the hole and the real balanceOf at both
    ends (fixtures/api_etherfi_cashback_hole.json). The explorer fake serves them WITHOUT the hole block, as Blockscout
    does; the scan must find the block, read it from the chain and reconcile to the wei — no tolerance."""
    import fetch.logscan as ls
    from fetch.base import FetchOutput
    from fetch.logcache import LogCache
    fx, logs, explorer, reader = _hole_fakes(monkeypatch)
    hole, h = fx["explorer_missing_block"], fx["holder"].lower()
    scan = ls.LogScan(explorer=explorer, reader=reader, cache=LogCache(tmp_path))
    spec = {"key": "cashback_out_optimism", "metric": "cashback_paid_tokens", "window_days": 100}
    out = FetchOutput()
    assert scan._window_scan({"name": "Ether.fi"}, spec, "optimism", 10, fx["token"], [h], fx["to_block"], "test", out)
    msg = " ".join(e.message for e in out.log)
    assert f"block {hole:,} (1 Transfer(s) the explorer does not index)" in msg, msg[:400]
    assert "RECONCILED to the wei" in msg
    paid = sum(f["value"].sum() for f in out.frames)
    true_out = sum(ls._amount(e) for e in logs if ls.topic_address(e["topics"][1]) == h) / 1e18
    assert paid == pytest.approx(true_out, abs=1e-9)                     # the hole's 19.808964 ETHFI is in
    assert type(reader).balance_reads < 60                                    # one bisection, not a block-by-block walk


def test_x1_scan_logs_rpc_only_never_asks_the_explorer(monkeypatch):
    """The hole read must bypass the explorer: ChainReader.scan_logs(..., rpc_only=True) goes straight to eth_getLogs."""
    from fetch.chain import ChainReader
    r = ChainReader()

    class _NoExplorer:
        def configured(self, cid):
            return True

        def get_logs(self, *a, **k):
            raise AssertionError("the explorer was asked for a block it does not index")

        def start_budget(self, s):
            pass

    r.explorer = _NoExplorer()
    monkeypatch.setattr(r, "_get_logs_resilient", lambda chain, base, lo, hi, used: (["log"], hi, used))
    assert r.scan_logs("optimism", {"address": "0x0", "topics": []}, 5, 5, rpc_only=True)[0] == ["log"]


def test_x1_followup_the_direct_count_reconciles_and_patches_the_hole(tmp_path, monkeypatch):
    """X1 follow-up: Pendle's direct counts read the outflows only, with no balance check — the same hole would have
    lowered them silently. They now read the inflows too and reconcile against balanceOf at both ends, patching an
    explorer hole from the chain; a count that cannot reconcile stores nothing."""
    import fetch.logscan as ls
    from fetch.base import FetchOutput
    from fetch.logcache import LogCache
    fx, logs, explorer, reader = _hole_fakes(monkeypatch)
    h = fx["holder"].lower()
    scan = ls.LogScan(explorer=explorer, reader=reader, cache=LogCache(tmp_path))
    spec = {"key": "cashback_out_optimism", "metric": "cashback_paid_tokens", "chain": "optimism",
            "direct_metric": "cashback_paid_direct", "direct_days": 100}
    out = FetchOutput()
    scan._direct_count("Ether.fi", spec, 10, fx["token"], [h], set(), fx["to_block"], 10 ** 18, out)
    msg = " ".join(e.message for e in out.log)
    assert "RECONCILED to the wei" in msg and f"block {fx['explorer_missing_block']:,}" in msg, msg[:300]
    true_out = sum(ls._amount(e) for e in logs if ls.topic_address(e["topics"][1]) == h) / 1e18
    assert sum(f["value"].sum() for f in out.frames) == pytest.approx(true_out, abs=1e-9)

    # without the RPC (the hole cannot be read): nothing stored, and the failure names the difference
    class _NoRpc(type(reader)):
        def scan_logs(self, *a, **k):
            return [], 1, 1
    scan2 = ls.LogScan(explorer=explorer, reader=_NoRpc(), cache=LogCache(tmp_path / "b"))
    out2 = FetchOutput()
    scan2._direct_count("Ether.fi", spec, 10, fx["token"], [h], set(), fx["to_block"], 10 ** 18, out2)
    assert not out2.frames and "does NOT RECONCILE" in " ".join(e.message for e in out2.log)


def test_x1_followup_sky_burn_scan_reconciles_against_total_supply(monkeypatch):
    """X1 follow-up: Sky's burns go to address(0), so no balance can check them — but SKY's totalSupply moves only by
    mint and burn, so totalSupply(end) − totalSupply(start) = Σ mints − Σ burns, to the wei. Real chain slice
    (fixtures/api_sky_supply_slice.json: the 7,372,288 SKY Stage 2 burn at block 26,176,859 and totalSupply at both
    ends). The explorer side drops that block, as Blockscout dropped OP 158,047,200: the reconciliation finds it,
    reads it from the chain and closes; a block it cannot recover leaves the scan refused."""
    from fetch.chain import Chain
    fx = _fx("api_sky_supply_slice.json")
    ev = lambda x: {"from": "0x" + x["topics"][1][-40:], "to": "0x" + x["topics"][2][-40:],  # noqa: E731
                    "value": int(x["data"], 16), "block": int(x["blockNumber"], 16)}
    true_m, true_b = [ev(x) for x in fx["mints"]], [ev(x) for x in fx["burns"]]
    hole = true_b[0]["block"]
    s0 = int(fx["supply_before"])

    class _Reader:
        def __init__(self, rpc=True):
            self.rpc = rpc

        def burn_transfer_events(self, chain, token, burn_to, from_block, chunk=10_000, max_blocks=None, side="to",
                                 head=None):
            evs = true_m if side == "from" else [e for e in true_b if e["block"] != hole]   # the explorer's view
            return list(evs), from_block, head, 1

        def erc20_supply_at(self, chain, token, block):
            return s0 + sum(e["value"] for e in true_m if e["block"] <= block) - \
                sum(e["value"] for e in true_b if e["block"] <= block)

        def scan_logs(self, chain, base, lo, hi, rpc_only=False):
            assert rpc_only
            if not self.rpc:
                return [], 1, 1
            src = fx["mints"] if base["topics"][1] is not None else fx["burns"]
            return [x for x in src if lo <= int(x["blockNumber"], 16) <= hi], 1, 1

        def checksum(self, a):
            return a

    c = Chain.__new__(Chain)
    c.reader = _Reader()
    burns, verdict = c._supply_reconcile("ethereum", fx["token"], "0x" + "0" * 40, fx["from_block"], fx["to_block"],
                                         [e for e in true_b if e["block"] != hole], {}, 10 ** 18)
    assert verdict.startswith("RECONCILED to the wei against totalSupply"), verdict
    assert f"block {hole:,}" in verdict and sum(e["value"] for e in burns) == 7_372_288 * 10 ** 18
    c.reader = _Reader(rpc=False)
    _b, verdict = c._supply_reconcile("ethereum", fx["token"], "0x" + "0" * 40, fx["from_block"], fx["to_block"],
                                      [e for e in true_b if e["block"] != hole], {}, 10 ** 18)
    assert verdict.startswith("DOES NOT RECONCILE") and "-7,372,288.0000 tokens" in verdict


def test_y2_near_probe_passes_offline_on_the_saved_responses(monkeypatch, capsys):
    """Y2: the probe flagged 2026-04-22 1csfundsadmin -> fdd59b04… 45,000 NEAR (then bridge-mng.near) as not on the
    list — the list held one 45,010 entry (45,000 + 10 summed). One entry per transfer now. The probe's own logic, run
    offline on the saved NearBlocks responses (fixtures/api_near_revenue_wallets.json: every page of the three wallets
    read 2026-10-13, the one-time accounts' receipts 2026-10-14): the verdict passes; with the old summed entry it
    flags exactly the 45,000."""
    import check_offline_items as coi
    from fetch import nearblocks as nb
    saved = _fx("api_near_revenue_wallets.json")["responses"]
    clock = [0.0]
    monkeypatch.setattr(nb, "_clock", lambda: clock[0])
    monkeypatch.setattr(nb, "_sleep", lambda s: clock.__setitem__(0, clock[0] + s))

    class _R:
        status_code = 200

        def __init__(self, rows):
            self.rows = rows

        def json(self):
            return {"data": self.rows, "meta": {"next_page": None}}

    monkeypatch.setattr(coi.requests, "get", lambda url, **k: _R(saved[url.replace("https://api.nearblocks.io", "")]))
    coi.near_revenue_wallet_flows()
    out = capsys.readouterr().out
    assert "VERDICT: every outflow of 1,000 NEAR or more is on config.NEAR_REVENUE_OUTFLOWS" in out, out[-600:]
    assert "45,000.00 NEAR" in out and "12,000.00 NEAR" in out and out.count("BRIDGED OUT — not retired") >= 2
    import re
    assert re.search(r"native IN\s+1,978,728\.95", out) and re.search(r"native OUT\s+12,025\.00", out)  # buybacks wallet

    summed = [dict(config.NEAR_REVENUE_OUTFLOWS[0], tokens=45_010.0)] + list(config.NEAR_REVENUE_OUTFLOWS[2:])
    monkeypatch.setattr(config, "NEAR_REVENUE_OUTFLOWS", tuple(summed))
    coi.near_revenue_wallet_flows()
    out = capsys.readouterr().out
    assert "1 outflow(s) of 1,000+ NEAR NOT on config.NEAR_REVENUE_OUTFLOWS" in out and \
        "2026-04-22 1csfundsadmin.sputnik-dao.near 45,000" in out


def test_y1_logs_in_block_falls_through_endpoints_and_receipts_and_names_every_try(monkeypatch):
    """Y1: on Jake's machine optimism-rpc.publicnode.com refuses every past-block read ("Archive requests require a
    personal token") and mainnet.optimism.io rate-limits. The hole's block is read over that ONE block on every
    endpoint, an empty answer does not end the search, then from the block's receipts; every try is named. Real log
    and receipt shapes: the X1 hole (fixtures/api_etherfi_cashback_hole.json, tx 0xb09bfbe3…)."""
    from fetch.chain import ChainReader
    fx = _fx("api_etherfi_cashback_hole.json")
    hole = fx["explorer_missing_block"]
    real = [x for x in fx["logs"] if int(x["blockNumber"], 16) == hole]
    assert len(real) == 1
    monkeypatch.setenv("RPC_OPTIMISM", "https://refuses.example,https://empty.example,https://receipts.example")
    monkeypatch.delenv("OPTIMISM_RPC_URL", raising=False)

    class _Eth:
        def __init__(self, url):
            self.url = url

        def get_logs(self, flt):
            assert flt["fromBlock"] == flt["toBlock"] == hole                      # ONE block, nothing to refuse
            if "refuses" in self.url:
                raise RuntimeError("403 Client Error: Forbidden — Archive requests require a personal token")
            if "empty" in self.url:
                return []                                                          # a pruned log index answers []
            raise RuntimeError("429 Too Many Requests")

        def get_block(self, n, full_transactions=False):
            if "receipts" not in self.url:
                raise RuntimeError("403")
            return {"transactions": [real[0]["transactionHash"], "0x" + "ab" * 32]}

        def get_transaction_receipt(self, h):
            return {"logs": real if h == real[0]["transactionHash"] else []}

    class _W3:
        def __init__(self, url):
            self.eth = _Eth(url)

    r = ChainReader()
    monkeypatch.setattr(ChainReader, "make_web3", staticmethod(lambda chain, url, t=30: _W3(url)))
    topics = [real[0]["topics"][0], real[0]["topics"][1], None]
    got, tried = r.logs_in_block("optimism", fx["token"], topics, hole)
    assert len(got) == 1 and int(got[0]["data"], 16) == 19_808_964_390_546_887_139
    assert any("refuses.example eth_getLogs" in t and "Archive requests" in t for t in tried)
    assert any("empty.example eth_getLogs" in t and "0 log(s)" in t for t in tried)
    assert any("receipts.example receipts of block" in t and "1 log(s)" in t for t in tried)

    monkeypatch.setenv("RPC_OPTIMISM", "https://refuses.example")
    with pytest.raises(RuntimeError) as e:
        r.logs_in_block("optimism", fx["token"], topics, hole)
    assert "could not be read from the chain" in str(e.value) and "refuses.example" in str(e.value)


def test_y1_the_failure_line_names_the_hole_search(tmp_path, monkeypatch):
    """Y1: when the block cannot be read, the failure line says what was tried — not the old message."""
    import fetch.logscan as ls
    from fetch.base import FetchOutput
    from fetch.logcache import LogCache
    fx, logs, explorer, reader = _hole_fakes(monkeypatch)

    def refuse(chain, address, topics, block):
        raise RuntimeError(f"block {block:,} could not be read from the chain: optimism-rpc.publicnode.com "
                           f"eth_getLogs -> 403 Archive requests require a personal token")
    reader.logs_in_block = refuse
    scan = ls.LogScan(explorer=explorer, reader=reader, cache=LogCache(tmp_path))
    out = FetchOutput()
    spec = {"key": "cashback_out_optimism", "metric": "cashback_paid_tokens", "window_days": 100}
    assert not scan._window_scan({"name": "Ether.fi"}, spec, "optimism", 10, fx["token"], [fx["holder"].lower()],
                                 fx["to_block"], "test", out)
    msg = " ".join(e.message for e in out.log)
    assert "bisected on balanceOf" in msg and f"to block {fx['explorer_missing_block']:,}" in msg
    assert "EXPLORER-HOLE SEARCH" in msg and "Archive requests require a personal token" in msg
    assert "does not reconcile either" not in msg


# ------------------------------------------------------------------------------------------- run review D (2026-10-14)
def test_d2_subsidy_cover_blocks_are_allowlisted_only_while_the_gap_is_declared(monkeypatch):
    """D2: the three subsidy covers BLOCKED on an incomplete denominator are correct blocks; they are allow-listed with a
    reason, a re-check date AND a condition — the allowance lapses the moment OUTFLOW_INCOMPLETE stops declaring the gap."""
    import coverage_audit as ca
    detail = ("CP SUBSIDY COVER = revenue ÷ $ supply-side incentives (b), per covered da: 12.578 in the previous run "
              "(2026-10-13), now BLOCKED — unequal coverage: the incentive denominator is incomplete")
    rows = {r["Project"]: r for r in ca.regression_rows([(p, "headline_lost", detail)
                                                         for p in ("Pendle", "Aethir", "Fluid", "Morpho")])}
    assert all(rows[p]["Status"] == "N/A BY DESIGN" for p in ("Pendle", "Aethir", "Fluid"))
    assert "re-check" in rows["Pendle"]["Reason / route / error"] and rows["Morpho"]["Status"] == "BROKEN"
    gaps = {k: v for k, v in config.OUTFLOW_INCOMPLETE.items() if k != ("Pendle", "b")}
    monkeypatch.setattr(config, "OUTFLOW_INCOMPLETE", gaps)               # the Pendle scans wired: the allowance lapses
    assert ca.regression_rows([("Pendle", "headline_lost", detail)])[0]["Status"] == "BROKEN"


def _arb_gauge_world(drop_every: int):
    """Real dated Transfer logs of Pendle's Arbitrum gauge (fixtures/api_pendle_arb_gauge_slice.json, tenderly, with the
    real balanceOf at both ends — they close to the wei); the 'explorer' lacks every `drop_every`-th outflow block."""
    import fetch.logscan as ls
    fx = _fx("api_pendle_arb_gauge_slice.json")
    true_in = [dict(ls._norm_log(x), timeStamp=int(x["blockTimestamp"], 16)) for x in fx["ins"]]
    true_out = [dict(ls._norm_log(x), timeStamp=int(x["blockTimestamp"], 16)) for x in fx["outs"]]
    blocks = sorted({e["blockNumber"] for e in true_out})
    dropped = set(blocks[::drop_every])
    b0 = int(fx["balance_before"])
    raw = {(int(x["blockNumber"], 16), x["topics"][1][-40:], x["topics"][2][-40:]): x for x in fx["ins"] + fx["outs"]}

    class _Reader:
        balance_reads = 0

        def erc20_balance_at(self, chain, token, holder, block):
            _Reader.balance_reads += 1
            if block < fx["from_block"]:
                return b0
            return b0 + sum(ls._amount(e) for e in true_in if e["blockNumber"] <= block) - \
                sum(ls._amount(e) for e in true_out if e["blockNumber"] <= block)

        def logs_in_block(self, chain, address, topics, block):
            return [x for x in fx["ins"] + fx["outs"] if int(x["blockNumber"], 16) == block], \
                [f"tenderly eth_getLogs block {block:,} -> logs"]

        def dated_logs(self, chain, address, topics, lo, hi, chunk, trace, deadline=None):
            upper = min(hi, lo + chunk - 1)
            src = fx["outs"] if topics[1] is not None else fx["ins"]
            trace.append(f"tenderly eth_getLogs {lo:,}..{upper:,}")
            return [x for x in src if lo <= int(x["blockNumber"], 16) <= upper], upper

        def block_timestamp(self, chain, b):
            return next(int(x["blockTimestamp"], 16) for k, x in raw.items() if k[0] == b)

        def checksum(self, a):
            return a

    explorer_out = [e for e in true_out if e["blockNumber"] not in dropped]
    return fx, true_in, true_out, explorer_out, dropped, _Reader()


@pytest.mark.parametrize("drop_every,expect_span", [(9, False), (2, True)])
def test_d1_full_history_repair_reconciles_to_the_wei(tmp_path, drop_every, expect_span):
    """D1: Pendle's Arbitrum gauge scan never reconciled (Etherscan's completed seeds: in − out 812,025.60 PENDLE vs a
    balance near 78,400; the chain's 936,270 logs close to the wei). The repair: bisect balanceOf to the first block
    where it parts from the scan, read it from the chain; past EXPLORER_HOLE_MAX holes, re-read the span from the first
    divergence from the chain, dated, resumable — and the result reconciles TO THE WEI. Real dated logs."""
    import fetch.logscan as ls
    from fetch.logcache import LogCache
    fx, true_in, true_out, explorer_out, dropped, reader = _arb_gauge_world(drop_every)
    scan = ls.LogScan(explorer=object(), reader=reader, cache=LogCache(tmp_path))
    spec = {"key": "gauge_pendle_out_arbitrum", "repair_chunk_blocks": 100_000}
    base = (fx["from_block"] - 1, int(fx["balance_before"]))     # the slice starts from the real balance before it
    ins, outs, status, trace = scan._repair_history(spec, "arbitrum", 42161, fx["token"], fx["holder"], list(true_in),
                                                    explorer_out, fx["to_block"], base)
    assert status == "repaired", trace
    assert sum(ls._amount(e) for e in outs) == sum(ls._amount(e) for e in true_out)            # to the wei
    assert ("re-reading blocks" in " ".join(trace)) is expect_span, trace
    assert all(int(e["timeStamp"]) > 0 for e in outs)                                          # every outflow dated
    assert any("bisected on balanceOf" in t for t in trace)


def test_d1_span_repair_resumes_across_runs(tmp_path):
    """The span re-read stops at its budget, keeps its state in the log cache and resumes next run."""
    import fetch.logscan as ls
    from fetch.logcache import LogCache
    fx, true_in, true_out, explorer_out, dropped, reader = _arb_gauge_world(2)
    spec = {"key": "gauge_pendle_out_arbitrum", "repair_chunk_blocks": 50_000, "repair_budget_s": 1e-9}
    scan = ls.LogScan(explorer=object(), reader=reader, cache=LogCache(tmp_path))
    statuses = []
    for _ in range(40):
        ins, outs, status, trace = scan._repair_history(spec, "arbitrum", 42161, fx["token"], fx["holder"],
                                                        list(true_in), explorer_out, fx["to_block"],
                                                        (fx["from_block"] - 1, int(fx["balance_before"])))
        statuses.append(status)
        if status != "in progress":
            break
    assert statuses[0] == "in progress" and statuses[-1] == "repaired" and len(statuses) > 2
    assert not list(tmp_path.glob("repair-*"))                                        # state removed once complete


def test_d1_dated_logs_passes_over_an_undated_endpoint(monkeypatch):
    """arb1.arbitrum.io answers eth_getLogs with blockTimestamp 0x0 on every log; an undated outflow cannot be put on a
    day, so that answer is passed over for the next endpoint (tenderly dates them). Every answer is named."""
    from fetch.chain import ChainReader
    fx = _fx("api_pendle_arb_gauge_slice.json")
    monkeypatch.setenv("RPC_ARBITRUM", "https://arb1.example,https://tenderly.example")
    monkeypatch.delenv("ARBITRUM_RPC_URL", raising=False)

    class _Eth:
        def __init__(self, url):
            self.url = url

        def get_logs(self, flt):
            got = [dict(x) for x in fx["outs"] if flt["fromBlock"] <= int(x["blockNumber"], 16) <= flt["toBlock"]]
            if "arb1" in self.url:
                for x in got:
                    x["blockTimestamp"] = "0x0"
            return got

    monkeypatch.setattr(ChainReader, "make_web3", staticmethod(lambda chain, url, t=30: type("W", (), {"eth": _Eth(url)})()))
    trace = []
    got, upper = ChainReader().dated_logs("arbitrum", fx["token"], [], fx["from_block"], fx["to_block"],
                                          fx["to_block"] - fx["from_block"] + 1, trace)
    assert len(got) == len(fx["outs"]) and upper == fx["to_block"]
    assert any("tenderly.example" in t and "dated" in t for t in trace)


def test_d7_near_status_falls_through_when_the_primary_drops_tls():
    """D7: on Jake's 10-14 run rpc.mainnet.near.org dropped TLS (SSLEOFError — so did rpc.plume.org and Alchemy in the
    same minute), near.lava.build answers 410 (discontinued) and 1rpc.io a usage limit. The list now carries endpoints
    that serve `status` (read 2026-10-14), and a dropped primary falls through to them."""
    import fetch.near as fn
    api = config.PROJECT_BY_NAME["Near"]["node_api"]
    eps = api["endpoints"]
    assert "https://near.lava.build" not in eps and "https://1rpc.io/near" not in eps
    assert len([u for u in eps if u != "https://rpc.mainnet.near.org"]) >= 2

    class _Http:
        calls = []

        def post(self, url, json_body=None):
            self.calls.append(url)
            if url == eps[0]:
                raise ConnectionError("SSLError(SSLEOFError(8, '[SSL: UNEXPECTED_EOF_WHILE_READING] EOF occurred'))")
            return {"result": {"protocol_version": 86, "sync_info": {"latest_block_height": 219_936_884}}}

    obj = fn.NearNode.__new__(fn.NearNode)
    obj.http = _Http()
    payload, detail = obj._call(api, {"jsonrpc": "2.0", "id": "t", "method": "status", "params": []})
    assert payload["result"]["protocol_version"] == 86 and detail == eps[1] and obj.http.calls[:2] == eps[:2]


# =================================================================================================== run review E1–E9
def test_e1_etherfi_docs_treasury_wallets_are_subtracted_and_the_policy_says_what_the_ratios_use():
    """E1 (2026-10-14): ether.fi's own allocations page lists four Treasury wallets; two were missing from the set. The
    policy text of Ether.fi, Fluid and GEODNET still read CoinGecko / the 462M plan while the ratios used the on-chain
    set — corrected to what is used."""
    cs = config.PROJECT_BY_NAME["Ether.fi"]["contracts"]
    for k, a in (("noncirc_treasury_docs_3", "0x5f0E7A424d306e9E310be4f5Bb347216e473Ae55"),
                 ("noncirc_treasury_docs_4", "0xD022d6bb8B6C1C357ec77D930Dc6A0aD40FFC90b")):
        c = cs[k]
        assert c["address"] == a and c["metric_override"] == "noncirculating_holding_tokens"
        assert c["source_url"] == "https://etherfi.gitbook.io/gov/ethfi-allocations" and c["verified"] == "2026-10-14"
    for n in ("Ether.fi", "Fluid", "GEODNET"):
        assert config.circulating_primary_metric(n) == "circulating_supply_onchain"
        assert config.CIRCULATING_POLICY[n]["kind"] == "onchain", n
        assert "CoinGecko —" not in config.CIRCULATING_POLICY[n]["source"], n


def test_e2_aethir_supply_is_read_on_ethereum_and_flagged_for_solana():
    """E2: Arbitrum's ATH (3.18bn) sits in the Ethereum escrow, so the supply read is Ethereum's canonical 42bn, PARTIAL
    while Solana's 18.44bn is not reconciled; Arbitrum is never summed."""
    cs = config.PROJECT_BY_NAME["Aethir"]["contracts"]
    sup = [k for k, c in cs.items() if c.get("kind") == "erc20_total_supply"]
    assert sup == ["token_supply_ethereum"]
    s = cs["token_supply_ethereum"]
    assert s["chain"] == "ethereum" and s["address"] == "0xbe0Ed4138121EcFC5c0E56B40517da27E6c5226B"
    assert s["supply_is_partial"] and "SOLANA NOT RECONCILED" in s["partial_reason"]
    assert cs["token_arbitrum"]["kind"] == "bridged_representation"
    # F3 (Jake 2026-10-14): on-chain first anyway, the PARTIAL flag kept; CoinGecko the cross-check
    assert config.TOTAL_SUPPLY_ONCHAIN_FIRST["Aethir"]["decided_on"] == "2026-10-14"
    assert "Aethir" in config.METRICS["total_supply_coingecko"]["only_projects"]
    assert config.PROJECT_BY_NAME["Aethir"]["series_handover"]["total_supply"]["ordered_points"] == (
        "coingecko", "chain:ethereum:token_supply_ethereum")


def test_e3_world_mobile_band_is_not_moved_and_the_partial_names_what_is_missing():
    """E3: the post-deprecation sum (269,379,304.50) is not the aggregate; the 1.42bn floor refuses it, correctly."""
    wm = config.PROJECT_BY_NAME["World Mobile"]
    band = next(v for k, v in wm.items() if isinstance(v, dict) and "total_supply_gross" in v
                and isinstance(v["total_supply_gross"], dict) and "min" in v["total_supply_gross"])
    assert band["total_supply_gross"] == {"min": 1_420_000_000, "max": 2_000_000_000}
    assert "650,787,233.25" in config.WMTX_CARDANO_PARTIAL and "269,379,304.50" in config.WMTX_CARDANO_PARTIAL


def test_e5_hyperliquid_team_unlock_references_are_the_ledgers_and_verify():
    """E5 (ii): the team wallet's own ledger (HL info API, read 2026-10-14): 1,125,766 + 1 sold on 2026-01-06 and
    1,745,745 + 1 sold on 2025-11-29. The secondary 1.2M was the unstake; 1.75M the net unstake rounded."""
    import references as rf
    r = {x["metric"]: x for x in rf.check(_groups("Hyperliquid", _hl_rows()))}   # the replay, from the ledgers fixture
    for m, v in (("team_unlock_2026-01-06", 1_125_767), ("team_unlock_2025-11-29", 1_745_746)):
        assert r[m]["status"] == "VERIFIED", r[m]
        assert r[m]["ours"] == pytest.approx(v, abs=1)


def test_e5_near_intents_is_labelled_since_its_first_day():
    import references as rf
    r = {x["metric"]: x for x in rf.check(_groups("Near"))}["intents_cumulative_volume"]
    assert "since" in (r.get("detail") or "").lower() or "SINCE" in (r.get("detail") or ""), r
    row = next(x for x in csv.DictReader(open(ROOT / "independent_references.csv", encoding="utf-8"))
               if x["metric"] == "intents_cumulative_volume")
    assert "SINCE 2025-07-15" in row["note"] and "30685539216" in row["note"]


def test_e9_chainlink_staking_rate_on_both_bases():
    """E9: the vault read (2026-10-14) gives 4.323% on the community pool and 4.766% blended; the reference row the A1
    validator yield is compared with is on the blended basis, the community rate is stored and labelled apart."""
    import references as rf
    r = {x["metric"]: x for x in rf.check(_groups("Chainlink"))}
    assert r["staking_blended_rate"]["status"] == "VERIFIED", r["staking_blended_rate"]
    assert r["staking_blended_rate"]["ours"] == pytest.approx(0.04766, abs=0.0002)
    rows = {x["metric"]: x for x in csv.DictReader(open(ROOT / "independent_references.csv", encoding="utf-8"))
            if x["project"] == "Chainlink"}
    assert rows["staking_community_rate"]["note"].startswith("BASIS: COMMUNITY POOL ONLY")
    assert rows["staking_blended_rate"]["note"].startswith("BASIS: BLENDED")
    man = config.CREDIBILITY["Chainlink"]["a1_validator_yield"]
    assert man["manual"]["value"] == 0.04766 and man["note"].startswith("BASIS: BLENDED")
    assert config.VALIDATOR_YIELD["Chainlink"]["note"].startswith("BASIS: BLENDED")


def test_e1_cr_restates_etherfi_noncirculating_on_the_current_set(monkeypatch, tmp_path):
    """E1: section CR, run on Jake's real Ether.fi dump, removes exactly the old-set noncirculating_holding_tokens rows
    (none carries the newest wallet) and nothing else; a second run matches nothing."""
    import builtins
    import contextlib
    import io
    import sqlite3

    import real_dump
    import run_sql
    payload = real_dump.load("Ether.fi")
    db = tmp_path / "metrics.db"
    real_dump.to_store(payload, db)
    conn = sqlite3.connect(db)
    before = conn.execute("SELECT COUNT(*) FROM metrics").fetchone()[0]
    old = conn.execute("SELECT COUNT(*) FROM metrics WHERE project = 'Ether.fi' AND "
                       "metric = 'noncirculating_holding_tokens'").fetchone()[0]
    assert old > 300
    sec = run_sql.parse_sections(run_sql.SQL_FILE.read_text(encoding="utf-8"))["CR"]
    monkeypatch.setattr(builtins, "input", lambda *a: "DELETE CR")
    with contextlib.redirect_stdout(io.StringIO()):
        assert run_sql.run_delete(conn, sec["text"], "CR", db, sec["authored"]) == 0
        assert run_sql.run_delete(conn, sec["text"], "CR", db, sec["authored"]) == 1
    assert conn.execute("SELECT COUNT(*) FROM metrics").fetchone()[0] == before - old
    assert "noncirc_treasury_docs_4" in str(config.PROJECT_BY_NAME["Ether.fi"]["contracts"])
