#!/usr/bin/env python3
"""
sanity.py — rules no published figure may break (Jake's review of the 97b29ae run, 2026-10-11, item 17). PERMANENT.

Each rule reads the EVALUATED workbook (the numbers a reader sees) and returns (project, rule, detail):
  fdv_below_mcap        FDV < market cap
  circ_above_total      the circulating supply used > total supply
  mcap_not_price_x_circ market cap != spot price x the circulating supply shown (beyond 0.5%)
  buyback_above_revenue actual buyback $ > revenue $ by more than 5% with no documented funding source
  status_contradicts    a buyback / burn status label contradicts the data (a measured flow > 0 under n/a /
                        unconfirmed / none; or an 'active' label with no flow measured in the window)
  mixed_windows         a ratio printed beside a total BLOCKED for unequal coverage (the same two series, one row)
  cross_tab             the same figure differs between Master and an A-tab, or between two A-tabs (cross_tab.py)
  trajectory_q0_mismatch a Trajectory row's Q0 is not its headline cell (the shifted formula changed what Q0 reads)
  headline_starved      a headline reads n/a while the series it reads is stored across the whole Q0 window (the
                        input row is withheld or its window blanked) — the continuity guard misses it when the
                        previous run was n/a too (Jake's review of the ef6dc86 run, item 7: GEODNET's A4 gross burn)
  float_outside_circ    CP's tradeable float < 0 or > the circulating supply it is taken from (Phase 2, 2026-10-12)
  discretionary_above_total CP's discretionary supply > total supply

Only the projects in portfolio.txt are checked (parked ones are skipped; with no portfolio.txt, every project).

build_workbook runs it on every build and writes the hits under the Review Queue tab; credibility_report prints
them; `python sanity.py` builds today's workbook from metrics.db and lists them by project.
"""
from __future__ import annotations

import argparse
import sys

import config

TOL_MCAP = 0.005
TOL_BUYBACK = 0.05


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def stated_number(v):
    """A cell's value as a number, counting a STATED number too (fix round 3, V2: the fixed-supply convention writes
    "0 — fixed supply: <why>" — a value of 0 with its reason, not an n/a). A text whose first token is a number
    followed by " — " is that number; any other text is None."""
    n = _num(v)
    if n is not None or not isinstance(v, str):
        return n
    head = v.strip().split(" — ", 1)
    if len(head) < 2:
        return None
    try:
        return float(head[0].replace(",", ""))
    except ValueError:
        return None


def _rows(wb, sheet: str) -> dict:
    """{project: {header: (row, col)}} for the first table of a Master / A-tab sheet."""
    ws = wb[sheet]
    hr = next((r for r in range(1, 12) if ws.cell(r, 1).value == "Project"), None)
    if hr is None:
        return {}
    heads = {c: ws.cell(hr, c).value for c in range(1, ws.max_column + 1) if ws.cell(hr, c).value}
    out = {}
    for r in range(hr + 1, ws.max_row + 1):
        p = ws.cell(r, 1).value
        if not p:
            break
        out[str(p)] = {str(h): (r, c) for c, h in heads.items()}
    return out


def _data_now(wb, value) -> dict:
    """{(project, metric): the Data tab's 'now' value}."""
    ws = wb["Data"]
    col = next(c for c in range(1, ws.max_column + 1) if str(ws.cell(1, c).value or "").startswith("Now"))
    out = {}
    for r in range(2, ws.max_row + 1):
        k = str(ws.cell(r, 1).value or "")
        if "|" in k:
            out[tuple(k.split("|", 1))] = value("Data", r, col)
    return out


def hits(wb, value) -> list[tuple[str, str, str]]:
    import config
    import cross_tab
    found = []
    master = _rows(wb, "Master")
    data = _data_now(wb, value)

    def mv(p, h):
        loc = master.get(p, {}).get(h)
        return value("Master", *loc) if loc else None
    for p in master:
        mcap, fdv = _num(mv(p, "Market cap ($)")), _num(mv(p, "FDV ($)"))
        circ, px = _num(mv(p, "Circulating supply (used)")), _num(mv(p, "Price — spot ($)"))
        total = _num(data.get((p, "total_supply")))
        if mcap and fdv is not None and fdv < mcap * (1 - 1e-9):
            found.append((p, "fdv_below_mcap", f"FDV ${fdv:,.0f} < market cap ${mcap:,.0f}"))
        if circ and total and circ > total * (1 + 1e-9):
            found.append((p, "circ_above_total", f"circulating used {circ:,.0f} > total supply {total:,.0f}"))
        if mcap and circ and px and abs(mcap - px * circ) / mcap > TOL_MCAP:
            found.append((p, "mcap_not_price_x_circ", f"market cap ${mcap:,.0f} vs price {px:,.6g} x circulating shown "
                                                      f"{circ:,.0f} = ${px * circ:,.0f} ({px * circ / mcap - 1:+.2%})"))
        bb, rev = _num(mv(p, "Actual buyback Q0 ($) — observed")), _num(mv(p, "Revenue Q0 ($)"))
        proj = config.PROJECT_BY_NAME.get(p) or {}
        funded = proj.get("buyback_funding") or (proj.get("foundation_buyback") or {}).get("funding")
        if bb and rev and bb > rev * (1 + TOL_BUYBACK) and not funded:
            found.append((p, "buyback_above_revenue", f"actual buyback Q0 ${bb:,.0f} = {bb / rev:.0%} of revenue Q0 "
                                                      f"${rev:,.0f}, no documented funding source (config "
                                                      f"buyback_funding)"))
        for label_h, flow_h, what in (("Buyback status", "Actual buyback Q0 ($) — observed", "buyback"),
                                      ("Burn status", "Gross burn Q0 (tokens)", "burn")):
            st, flow = str(mv(p, label_h) or ""), _num(mv(p, flow_h))
            low = st.lower()
            if flow and flow > 0 and low.startswith(("n/a", "unconfirmed", "none", "halted")):
                found.append((p, "status_contradicts", f"{what} status '{st}' but {flow:,.0f} measured in Q0"))
            elif low.startswith("active") and (flow == 0 or "no " + what + " series measured" in low):
                found.append((p, "status_contradicts", f"{what} status '{st}' but no {what} measured in the window"))
        net = mv(p, "NET SUPPLY CHANGE Q0 (tokens) = issuance − burn")
        ann = _num(mv(p, "Net supply change, annualised % of circulating"))
        if isinstance(net, str) and net.startswith("BLOCKED") and ann is not None:
            found.append((p, "mixed_windows", f"net supply change is {net[:90]}, yet its annualised % reads {ann:.3%}"))
    # CP (Phase 2, 2026-10-12): float is a part of circulating; discretionary supply a part of total supply
    if "CP Comparability" in wb.sheetnames:
        cp = _rows(wb, "CP Comparability")
        for p, cols in cp.items():
            def cv(prefix):
                loc = next((v for h, v in cols.items() if h.startswith(prefix)), None)
                return _num(value("CP Comparability", *loc)) if loc else None
            flt, circ, tot = cv("TRADEABLE FLOAT"), cv("Circulating supply (tokens)"), cv("Total supply (tokens)")
            disc = cv("Discretionary supply (tokens)")
            if flt is not None and circ is not None and (flt < 0 or flt > circ):
                found.append((p, "float_outside_circ", f"tradeable float {flt:,.0f} outside 0..circulating {circ:,.0f}"))
            if disc is not None and tot is not None and disc > tot:
                found.append((p, "discretionary_above_total", f"discretionary supply {disc:,.0f} > total supply "
                                                              f"{tot:,.0f}"))
    # TRAJECTORY (theory layer, Phase 1): every row's Q0 is its headline cell's own formula — a row that disagrees means
    # the shift changed what Q0 reads (build_workbook.write_trajectory writes the verdict)
    if "Trajectory" in wb.sheetnames:
        import trajectory                                  # noqa: PLC0415
        tw = wb["Trajectory"]
        hr = next((r for r in range(1, 12) if tw.cell(r, 1).value == "Project"), None)
        vc, kc = trajectory.COL["Credibility (Q0 = headline cell)"], trajectory.COL["Key (tab|project|headline id)"]
        for r in range(hr + 1, tw.max_row + 1) if hr else ():
            v = str(tw.cell(r, vc).value or "")
            if v.startswith("CHECK"):
                found.append((str(tw.cell(r, 1).value), "trajectory_q0_mismatch", f"{tw.cell(r, kc).value}: {v}"))
    for m in cross_tab.mismatches(wb, value):
        found.append((m["project"], "cross_tab", cross_tab.describe(m).splitlines()[0][len(m["project"]) + 3:]))
    # PARKED PROJECTS ARE NOT CHECKED (Jake's run on e4e50b0, 2026-10-12 item 5: Injective, PancakeSwap and Render hits
    # were noise): only the names in portfolio.txt, as completeness_report scopes; no file, every project
    held, _unknown = config.read_portfolio(config.PORTFOLIO_FILE)
    return [h for h in found if h[0] in held] if held else found


def lost_headlines(wb, previous: dict, headline_cells: list, prev_date: str = "") -> list[tuple[str, str, str]]:
    """headline_lost (Jake's Phase 1 review 2026-10-12, item 1): a headline cell that read a VALUE in the previous run's
    snapshot and reads n/a / BLOCKED / text now. One hit per cell, with both values."""
    import xlcalc                                          # noqa: PLC0415
    X = xlcalc.Workbook(wb)
    out = []
    for hc in headline_cells:
        was = ((previous or {}).get(hc["project"]) or {}).get(f"{hc['sheet'][:2]} | {hc['header']}")
        if not isinstance(was, float):
            continue
        cell = wb[hc["sheet"]][hc["cell"]]
        try:
            now = X.value(hc["sheet"], cell.row, cell.column)
        except Exception as e:  # noqa: BLE001
            now = f"unevaluable ({type(e).__name__})"
        if stated_number(now) is None:
            out.append((hc["project"], "headline_lost",
                        f"{hc['sheet'][:2]} {hc['header'][:70]}: {was:,.6g} in the previous run ({prev_date}), now "
                        f"{str(xlcalc.display(now) if not isinstance(now, str) else now)[:80]}"))
    return out


def starved_headlines(wb, headline_cells: list, data_by_key: dict, asof, period_days: int = 90,
                      fresh_days: int = 3) -> list[tuple[str, str, str]]:
    """headline_starved: a headline cell that evaluates to n/a while its input series (the cell's metric) holds stored
    rows from before the Q0 window to within `fresh_days` of asof. One hit per cell, naming the input row's status and
    note — the reason the figure is not shown (a withheld series, a blanked window)."""
    import pandas as pd                                    # noqa: PLC0415
    import xlcalc                                          # noqa: PLC0415
    X = xlcalc.Workbook(wb)
    asof = pd.Timestamp(asof).normalize()
    q0_start = asof - pd.Timedelta(days=period_days)
    out = []
    for hc in headline_cells:
        m = hc.get("input_metric")
        row = (data_by_key or {}).get(f"{hc['project']}|{m}") if m else None
        if not row:
            continue
        first, last = row.get("first_date"), row.get("latest_date")
        try:
            full = (first and last and pd.Timestamp(first) <= q0_start
                    and pd.Timestamp(last) >= asof - pd.Timedelta(days=fresh_days))
        except (TypeError, ValueError):
            full = False
        if not full:
            continue
        cell = wb[hc["sheet"]][hc["cell"]]
        try:
            v = X.value(hc["sheet"], cell.row, cell.column)
        except Exception as e:  # noqa: BLE001
            v = f"unevaluable ({type(e).__name__})"
        shown = v if isinstance(v, str) else xlcalc.display(v)
        # a bare n/a (the formula's own fallback) or an error — never a cell that states its reason ("n/a — …")
        txt = str(shown).strip().lower()
        # a pull of a missing Data value carries the Data row's own reason ("n/a — data: …", alignment round) — still a
        # starved headline when the input is stored
        if _num(v) is None and (txt == "n/a" or txt.startswith(("#", "unevaluable", "n/a — data:"))):
            out.append((hc["project"], "headline_starved",
                        f"{hc['sheet'][:2]} {hc['header'][:60]} reads {str(shown)[:20]} while {m} is stored "
                        f"{str(first)[:10]}..{str(last)[:10]} (Data status '{row.get('status')}': "
                        f"{str(row.get('note') or '')[:160]})"))
    return out


def check_workbook(path_or_wb) -> list[tuple[str, str, str]]:
    import openpyxl
    import xlcalc
    wb = openpyxl.load_workbook(path_or_wb) if not hasattr(path_or_wb, "worksheets") else path_or_wb
    X = xlcalc.Workbook(wb)
    return hits(wb, lambda s, r, c: xlcalc.display(X.value(s, r, c)))


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workbook")
    a = ap.parse_args(argv)
    path = a.workbook
    if not path:
        import tempfile
        from pathlib import Path
        import store as store_mod
        from build_workbook import build_workbook
        path = str(Path(tempfile.mkdtemp()) / "sanity.xlsx")
        build_workbook(store_mod.Store(store_mod.DB_PATH), path)
        # the build's own hits: every rule, including those that need the build's state (headline_starved, the
        # continuity guard) — the workbook alone cannot answer them
        import build_workbook as bw                        # noqa: PLC0415
        found = list(bw.SANITY_HITS)
    else:
        found = check_workbook(path)
    # PARKED PROJECTS ARE SHOWN APART AND DO NOT SET THE EXIT STATUS (Jake's run on 3500042, item 26: World Mobile,
    # parked 2026-10-05 — config.CREDIBILITY_PARKED — kept the run red): live hits first, parked ones listed below
    live, parked = split_parked(found)
    print(f"SANITY RULES — {len(live)} hit(s) on live projects")
    _print_hits(live)
    if parked:
        print(f"\nPARKED PROJECTS — {len(parked)} hit(s), shown for the record, not counted "
              f"({', '.join(f'{n}: {config.CREDIBILITY_PARKED[n]}' for n in sorted({h[0] for h in parked}))})")
        _print_hits(parked)
    return 1 if live else 0


def split_parked(found: list) -> tuple[list, list]:
    """(live hits, hits on parked projects — config.CREDIBILITY_PARKED)."""
    return ([h for h in found if h[0] not in config.CREDIBILITY_PARKED],
            [h for h in found if h[0] in config.CREDIBILITY_PARKED])


def _print_hits(found: list) -> None:
    for p in sorted({h[0] for h in found}):
        print(f"  {p}")
        for _p, rule, detail in (h for h in found if h[0] == p):
            print(f"    [{rule}] {detail}".encode("ascii", "replace").decode("ascii"))


if __name__ == "__main__":
    import progress                                    # stdout line-buffered when piped (2026-10-11)
    progress.line_buffered()
    sys.exit(main())
