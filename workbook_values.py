"""
workbook_values.py — THE WORKBOOK CARRIES ITS CALCULATED VALUES (Round G11 A1, Jake 2026-10-21).

openpyxl writes formulas with an empty cached value (`<f>…</f><v/>`), so anything that reads the file without
recalculating — openpyxl data_only, measuring_guard, Credibility's CONTINUITY check, a viewer that does not recalculate —
saw blanks. After the build saves, every formula cell is evaluated in Python (xlcalc, the evaluator the sanity rules
and the guard use; it agreed with LibreOffice on all 11,151 formulas of the 2026-10-13 fixture build) and its value is
written into the file as the formula's cached value. Excel and LibreOffice still recalculate on open; nothing else in
the file changes. No LibreOffice needed.
"""
from __future__ import annotations

import os
import re

import tempfile
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

_CELL = re.compile(r'<c r="([A-Z]+[0-9]+)"([^>]*)><f([^>]*)>(.*?)</f><v\s*/>', re.S)
_CELL_V = re.compile(r'<c r="([A-Z]+[0-9]+)"([^>]*)><f([^>]*)>(.*?)</f><v>[^<]*</v>', re.S)


def _sheet_paths(z: zipfile.ZipFile) -> dict:
    """{sheet name: its part path} from workbook.xml and its relationships."""
    wbx = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    target = {m.group(1): m.group(2) for m in re.finditer(r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rels)}
    target.update({m.group(2): m.group(1) for m in re.finditer(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="([^"]+)"',
                                                                rels)})
    out = {}
    for m in re.finditer(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbx):
        name = (m.group(1).replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"')
                .replace("&apos;", "'"))
        t = target.get(m.group(2), "")
        out[name] = t.lstrip("/") if t.startswith("/") else "xl/" + t
    return out


def _cached(v) -> tuple[str, str] | None:
    """(type attribute, cached text) for a computed value; None leaves the cell's value empty."""
    import xlcalc                                          # noqa: PLC0415
    if isinstance(v, xlcalc.XLError):
        return ' t="e"', escape(v.code)
    if isinstance(v, bool):
        return ' t="b"', "1" if v else "0"
    if isinstance(v, (int, float)):
        if v != v or v in (float("inf"), float("-inf")):
            return ' t="e"', "#NUM!"
        return "", repr(float(v)) if isinstance(v, float) else str(v)
    if v is None:
        return ' t="str"', ""
    return ' t="str"', escape(str(v))


def store_values(path, wb=None) -> dict:
    """Evaluate every formula of the saved workbook at `path` (or of `wb`, the same workbook in memory) and write each
    value into the file as the formula's cached value. Returns {"formulas", "valued", "failed", "failed_cells"}."""
    import openpyxl                                        # noqa: PLC0415
    import xlcalc                                          # noqa: PLC0415
    path = Path(path)
    wb = wb if wb is not None else openpyxl.load_workbook(path)
    X = xlcalc.Workbook(wb)
    stats = {"formulas": 0, "valued": 0, "failed": 0, "failed_cells": []}
    with zipfile.ZipFile(path) as z:
        parts = _sheet_paths(z)
        names = z.namelist()
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx", dir=str(path.parent))
        tmp.close()
        with zipfile.ZipFile(tmp.name, "w", zipfile.ZIP_DEFLATED) as out:
            by_part = {v: k for k, v in parts.items()}
            for n in names:
                data = z.read(n)
                sheet = by_part.get(n)
                if sheet is not None and sheet in wb.sheetnames:
                    text = data.decode("utf-8")

                    def put(m, sheet=sheet):
                        coord, attrs, fattrs, ftext = m.group(1), m.group(2), m.group(3), m.group(4)
                        stats["formulas"] += 1
                        cell = wb[sheet][coord]
                        try:
                            v = X.value(sheet, cell.row, cell.column)
                        except Exception:  # noqa: BLE001 — an unevaluable cell keeps an empty cache, counted
                            stats["failed"] += 1
                            if len(stats["failed_cells"]) < 20:
                                stats["failed_cells"].append(f"{sheet}!{coord}")
                            return m.group(0)
                        c = _cached(v)
                        stats["valued"] += 1
                        attrs = re.sub(r'\s+t="[^"]*"', "", attrs) + c[0]
                        return f'<c r="{coord}"{attrs}><f{fattrs}>{ftext}</f><v>{c[1]}</v>'
                    text = _CELL.sub(put, text)
                    data = text.encode("utf-8")
                info = z.getinfo(n)
                out.writestr(info, data)
    os.replace(tmp.name, path)
    return stats


def missing_values(path, sheets) -> list[str]:
    """'sheet!cell' for every formula cell of `sheets` whose cached value reads None (data_only) — the A1 test."""
    import openpyxl                                        # noqa: PLC0415
    wf = openpyxl.load_workbook(path)
    wv = openpyxl.load_workbook(path, data_only=True)
    out = []
    for s in sheets:
        for row in wf[s].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and wv[s][c.coordinate].value is None:
                    out.append(f"{s}!{c.coordinate}")
    return out


__all__ = ["store_values", "missing_values"]

