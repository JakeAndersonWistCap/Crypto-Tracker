"""recalc_check.py — THE STORED VALUES AGAINST LIBREOFFICE (Round P1, Jake 2026-10-27; SPEC §19 P1-1).

Every formula's cached value (written by workbook_values from the Python evaluator) is compared with the value
LibreOffice computes when it recalculates a copy of the file. Prints, as JSON: formulas, mismatches, errors (an Excel
error value in LibreOffice's result) and the first locations of each. Both files of a run must give 0 and 0:

    python recalc_check.py token_metrics.xlsx
    python recalc_check.py token_metrics_audit.xlsx 900

Writes nothing beside the file (the copy is recalculated in a temporary folder). Needs LibreOffice.
"""
from __future__ import annotations

import json
import shutil
import sys
import tempfile
from pathlib import Path

ERRORS = ("#N/A", "#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!")
SHOW = 20


def _same(a, b) -> bool:
    if a is None or a == "":
        return b is None or b == ""
    if isinstance(a, bool) or isinstance(b, bool):
        return bool(a) == bool(b)
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= 1e-9 * max(1.0, abs(a), abs(b))
    return str(a) == str(b)


def check(path, timeout: int = 300) -> dict:
    """{"formulas", "mismatches", "errors", "mismatch_at", "error_at"} — or {"error"} when LibreOffice cannot run."""
    import openpyxl
    import recalc
    path = Path(path)
    with tempfile.TemporaryDirectory() as tmp:
        copy = Path(tmp) / path.name
        shutil.copy(path, copy)
        res = recalc.recalc(str(copy), timeout)
        if "error" in res:
            return {"error": res["error"]}
        f = openpyxl.load_workbook(path)
        ours = openpyxl.load_workbook(path, data_only=True)
        lo = openpyxl.load_workbook(copy, data_only=True)
        n, bad, errs = 0, [], []
        for ws in f.worksheets:
            a, b = ours[ws.title], lo[ws.title]
            for row in ws.iter_rows():
                for c in row:
                    if not (isinstance(c.value, str) and c.value.startswith("=")):
                        continue
                    n += 1
                    va, vb = a[c.coordinate].value, b[c.coordinate].value
                    if isinstance(vb, str) and vb in ERRORS:
                        errs.append(f"{ws.title}!{c.coordinate} {vb}")
                    if not _same(va, vb):
                        bad.append(f"{ws.title}!{c.coordinate} ours {va!r} libreoffice {vb!r}")
    return {"file": path.name, "formulas": n, "mismatches": len(bad), "errors": len(errs),
            "mismatch_at": bad[:SHOW], "error_at": errs[:SHOW]}


def main(argv=None) -> int:
    args = list(sys.argv[1:] if argv is None else argv)
    if not args:
        print(__doc__)
        return 1
    timeout = int(args[1]) if len(args) > 1 else 300
    out = check(args[0], timeout)
    print(json.dumps(out, indent=2, ensure_ascii=False))
    return 0 if out.get("mismatches") == 0 and out.get("errors") == 0 else 1


if __name__ == "__main__":
    import progress
    progress.line_buffered()
    sys.exit(main())
