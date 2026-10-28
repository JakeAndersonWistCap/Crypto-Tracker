"""reader.py — THE READER WORKBOOK (Round P1, Jake 2026-10-27; Round P2a, the Charlie build, Jake 2026-10-27).

build_workbook writes the full workbook — every tab, every working column — which the daily run saves as
token_metrics_audit.xlsx. This module derives token_metrics.xlsx from the same build, in memory: About, Summary, the
four archetype tabs, Comparison, Business, Charts, Definitions, a hidden `_data` sheet holding every number the reader
shows (each value cell is a formula on it), and a hidden `_series` sheet holding the daily series the period columns
read, so a growth or price change follows the period selector — any period, custom included — and recalculates.

Cell rules (P1-2): a value cell holds a number, or one short code — n/a, blocked, not wired, by design, parked (and
"unranked" in a rank column) — in the code's colour with the reason in its comment. Headers are at most 35 characters,
plain English; the full definition is the header's comment and a Definitions row. Confidence is the fill of the audit
cell each number comes from. Every number's comment names its audit cell.

Nothing is lost (P1-9): reader_column_map.csv names every audit column's place in the reader (or "audit only"), and
HEADLINE_LOCATIONS maps every headline the reader shows to its new cell (measuring_guard reads them there).
"""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormatObject, FormulaRule, IconSet, Rule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
DATA_SHEET = "_data"
SERIES_SHEET = "_series"
ABOUT, SUMMARY, DEFS = "About", "Summary", "Definitions"
A1_TAB, A2_TAB, A3_TAB, A4_TAB = "A1 Infrastructure", "A2 Coordination", "A3 Revenue Buyback", "A4 Permanent Burn"
COMPARISON, BUSINESS, CHARTS = "Comparison", "Business", "Charts"
CP_TAB = "CP Comparability"
AUDIT_NAME = "token_metrics_audit.xlsx"
READER_NAME = "token_metrics.xlsx"
STAGE = 2
MAX_HEADER = 35

FMT_USD = '[>=1000000000]$#,##0.0,,,"bn";[>=1000000]$#,##0.0,,"m";$#,##0'
FMT_PRICE = "$#,##0.00##"
FMT_PCT = "0.0%"
FMT_X = '0.0"x"'
FMT_TOK = "#,##0"
FMT_RANK = "0"
TEXT_FMT = "@"

# THE CODES (P1-2): the cell's text, its colour; the reason goes in the comment
CODES = {"n/a": "808080", "blocked": "C00000", "not wired": "C65911", "by design": "44546A", "parked": "7F7F7F",
         "unranked": "595959"}
PERIODS = (("1D", 1), ("7D", 7), ("30D", 30), ("90D", 90), ("6M", 182), ("1Y", 365), ("3Y", 1095), ("custom", None))
DEFAULT_PERIOD = "90D"
DEFAULT_CUSTOM_DAYS = 60
SERIES_YEARS_BACK = 2200                   # days of daily rows on _series: both windows of a 3-year period
SELECT_CELL, CUSTOM_CELL = f"{SUMMARY}!$B$2", f"{SUMMARY}!$D$2"
DAYS_REF = f"{SERIES_SHEET}!$B$1"          # the selected period in days — every period formula reads this one cell
MOVE_THRESHOLD = 0.10                      # "what changed" lists moves above 10% over the selected period
ARCHETYPE_NAMES = {1: "A1 Infrastructure", 2: "A2 Coordination", 3: "A3 Revenue Buyback", 4: "A4 Permanent Burn"}
ARCHETYPE_KEYS = {
    1: ("Settlement per token, 1-year change",
        "SETTLEMENT VOLUME PER CIRCULATING TOKEN, 1-YEAR CHANGE = (trailing-365-day settlement volume ÷ circulating "
        "supply, now) ÷ (the same a year earlier) − 1. Settlement volume is Artemis's DEX + NFT + P2P transfers, the "
        "series A1's NRR reads; it needs two years of history. Higher is better."),
    2: ("Customer revenue ÷ emissions (x)",
        "CUSTOMER REVENUE ÷ EMISSIONS $ (x) — customer revenue over the last 90 days ÷ the tokens emitted to suppliers "
        "over the same days at their 90-day average price (A2's column). Higher is better; 1.0x is the crossover."),
    3: ("Retired ÷ emissions (x)",
        "RETIRED BUYBACK ÷ EMISSIONS (x) — tokens bought back and burned or held, per covered day, ÷ tokens emitted per "
        "covered day, unlocks EXCLUDED, at the emission rate in force (Sky's scheduled rate where declared, else the "
        "last 90 days' measured rate). Higher is better; 1.0x is the crossover."),
    4: ("Burn ÷ issuance (x)",
        "BURN ÷ ISSUANCE (x) — tokens burned ÷ tokens issued over the last 90 days, each per covered day; for a "
        "pre-minted supply, burn ÷ pool release, so the direction matches — the RECIPROCAL of the audit file's "
        "'Release ÷ burn (x)'. The last 90 days' measured issuance is the rate in force (no A4 project declares a "
        "scheduled rate). Higher is better; 1.0x is the crossover."),
}

# PARKED (P2a A8): shown greyed, outside the ranks, numbers visible; left out of the charts
READER_PARKED = {"World Mobile": "PARKED — no further work on World Mobile until Jake reopens it (parked 2026-09-29); "
                                 "its row is shown with its numbers, outside the ranks and the charts."}
_AERO_PARK: dict = {}                       # this build's Aerodrome park reason (build_reader resets it with the as-of day)

# STATE OF THE LAST BUILD — read by measuring_guard and the tests
HEADLINE_LOCATIONS: dict = {}               # (audit sheet, project, audit header) -> (reader sheet, cell)
STATS: dict = {}
COLUMN_MAP: list = []


def parked_reason(name: str) -> str:
    """Why a project is parked on the reader, or ''. Aerodrome is parked while config.aerodrome_park_reason gives a
    reason on the as-of day (P2a D2). P2c-3: the reader says it in one plain sentence (config.AERODROME_PARK_PLAIN);
    the full statement, with its figures, is on the audit file (Notes)."""
    if name in READER_PARKED:
        return READER_PARKED[name]
    if name == "Aerodrome":
        if "why" not in _AERO_PARK:
            import config                                      # noqa: PLC0415
            fn = getattr(config, "aerodrome_park_reason", None)
            why = fn(day=_AERO_PARK.get("day")) if fn is not None else ""
            _AERO_PARK["why"] = config.AERODROME_PARK_PLAIN if why else ""
        return _AERO_PARK["why"]
    return ""


# ============================================================================================ reading the audit build
@dataclass
class Src:
    """One audit cell, evaluated: its number (or None), its text, its fill and comment, and where it is."""
    sheet: str
    coord: str
    value: object = None
    num: float | None = None
    text: str = ""
    fill: PatternFill | None = None
    comment: str = ""
    header: str = ""

    @property
    def where(self) -> str:
        return f"audit file, {self.sheet}!{self.coord}"


class Audit:
    """The built workbook (openpyxl, in memory) and its values (xlcalc): cells found by tab, header and project."""

    def __init__(self, wb):
        import xlcalc                                      # noqa: PLC0415
        self.wb, self.X = wb, xlcalc.Workbook(wb)
        self._cols: dict = {}
        self._rows: dict = {}

    def cols(self, sheet: str, header_row: int = 4) -> dict:
        if sheet not in self._cols:
            ws = self.wb[sheet] if sheet in self.wb.sheetnames else None
            out = {}
            if ws is not None:
                for c in range(1, ws.max_column + 1):
                    h = ws.cell(header_row, c).value
                    if isinstance(h, str) and h not in out:
                        out[h] = c
            self._cols[sheet] = out
        return self._cols[sheet]

    def col(self, sheet: str, header: str) -> int | None:
        """The column of `header` on `sheet` — exact, else the one header starting with it."""
        cols = self.cols(sheet)
        if header in cols:
            return cols[header]
        hits = [c for h, c in cols.items() if h.startswith(header)]
        return hits[0] if len(hits) == 1 else None

    def header(self, sheet: str, col: int) -> str:
        return str(self.wb[sheet].cell(4, col).value or "")

    def row(self, sheet: str, project: str) -> int | None:
        return self.rows(sheet).get(project)

    def rows(self, sheet: str) -> dict:
        if sheet not in self._rows:
            ws = self.wb[sheet] if sheet in self.wb.sheetnames else None
            out = {}
            if ws is not None:
                for r in range(5, ws.max_row + 1):
                    v = ws.cell(r, 1).value
                    if v == "Project":                     # a second table below the first (A1, A2): the first wins
                        break
                    if isinstance(v, str) and v and ws.cell(r, 2).value and v not in out:
                        out[v] = r
            self._rows[sheet] = out
        return self._rows[sheet]

    def at(self, sheet: str, row: int, col: int) -> Src:
        import sanity                                      # noqa: PLC0415
        ws = self.wb[sheet]
        c = ws.cell(row, col)
        try:
            v = self.X.value(sheet, row, col)
        except Exception as e:  # noqa: BLE001 — an unevaluable cell is no number, said in its reason
            v = f"#ERROR {type(e).__name__}"
        n = sanity.stated_number(v)
        fill = c.fill if c.fill is not None and c.fill.fill_type == "solid" else None
        return Src(sheet, c.coordinate, v, None if n is None else float(n),
                   "" if v is None else str(v), fill, c.comment.text if c.comment else "", self.header(sheet, col))

    def cell(self, sheet: str, header: str, project: str) -> Src | None:
        r, c = self.row(sheet, project), self.col(sheet, header)
        if r is None or c is None:
            return None
        return self.at(sheet, r, c)

    def data_row(self, key: str) -> int | None:
        if "Data" not in self._rows:
            ws = self.wb["Data"]
            self._rows["Data"] = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
            self._cols["Data"] = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
        return self._rows["Data"].get(key)

    def data_field(self, key: str, header: str):
        """A field of the Data tab's row for `project|metric`."""
        r = self.data_row(key)
        c = self._cols["Data"].get(header)
        if r is None or c is None:
            return None
        return self.X.value("Data", r, c)


# ============================================================================================ codes and plain words
def code_of(text: str) -> str:
    """One of the codes for a cell that holds no number (P1-2)."""
    t = str(text or "").strip().lower()
    if t.startswith("blocked") or "blocked —" in t or "cannot form" in t:
        return "blocked"
    if "not wired" in t or "unwired" in t or "route not filled" in t or "route pending" in t:
        return "not wired"
    if ("none available" in t or "by design" in t or "suppressed" in t or "closed" in t or "structural" in t
            or "not applicable" in t or "programme halted" in t):
        return "by design"
    return "n/a"


_ROUND_TAG = re.compile(r"\s*\((?:[^()]*?\b(?:Round [A-Z]\d*\w*|[GDP]\d{1,2}(?:[- ]?\d+[a-z]?)*|config\.?\s?\w*|"
                        r"decision \d+ on \w+)\b[^()]*)\)")


def audit_notes(text) -> str:
    """An audit Notes / flags cell for the reader: the park statement line the audit file carries (P2c-3,
    build_workbook._with_park_note) is dropped — the reader gives its own one-sentence park reason."""
    lines = [x for x in str(text or "").split("\n") if not x.startswith("PARKED — migration:")]
    return plain("\n".join(lines))


def plain(text: str) -> str:
    """Internal codes out of reader text (P1-2): the round / config parentheticals go, Q0 reads "last 90 days"."""
    s = _ROUND_TAG.sub("", str(text or ""))
    s = re.sub(r"\bconfig\.(\w+)", r"\1", s)
    s = re.sub(r"\bQ0\b", "last 90 days", s)
    s = re.sub(r"\bQ-1\b", "the 90 days before", s)
    return re.sub(r"[ \t]{2,}", " ", s).strip()


def _code_reason(src: Src | None, why: str = "") -> str:
    if src is None:
        return why or "no audit cell for this project"
    parts = [why] if why else []
    if src.text:
        parts.append(plain(src.text))
    if src.comment and plain(src.comment) not in parts:
        parts.append(plain(src.comment))
    parts.append(f"Working: {src.where}.")
    return "\n\n".join(p for p in parts if p)[:2400]


def reader_format(audit_fmt: str | None) -> str:
    f = str(audit_fmt or "")
    if "%" in f:
        return FMT_PCT
    if '"x"' in f or f.endswith("x"):
        return FMT_X
    if "$" in f:
        return FMT_USD
    return FMT_TOK


def _zero_format(fmt: str, label: str) -> str:
    """A number format whose ZERO section shows the status beside the 0 (P2a A4: the 0 stays a 0, rule 3)."""
    lab = label.replace('"', "'")
    zero = {FMT_PCT: f'"0.0% — {lab}"', FMT_X: f'"0.0x — {lab}"', FMT_USD: f'"$0 — {lab}"'}.get(fmt, f'"0 — {lab}"')
    pos_neg = {FMT_PCT: "0.0%;-0.0%", FMT_X: '0.0"x";-0.0"x"', FMT_USD: "$#,##0;-$#,##0"}.get(fmt, "#,##0;-#,##0")
    return f"{pos_neg};{zero}"


# ============================================================================================ status-driven zeros
def status_note(p: dict, family: str) -> str:
    """The status that makes a 0 in this family of columns (P2a A4) — '' when none. Families: buyback (retired,
    holder value, payout, the A3 key), staking (staking yield, lock rate)."""
    if family == "buyback":
        bp = p.get("buyback_programme") or {}
        st = str(bp.get("status") or "").upper()
        if st in ("HALTED", "ENDED", "PAUSED"):
            when = bp.get("halted_on") or bp.get("ended_on") or bp.get("paused_on")
            return f"{st.lower()} since {when}" if when else st.lower()
    if family == "staking":
        ss = p.get("staking_sunset") or {}
        if ss.get("ended"):
            return f"ended since {ss['ended']}" + (f" ({ss['source']})" if ss.get("source") else "")
    return ""


def is_premint(p: dict, src: Src | None = None) -> bool:
    """A pre-minted supply (P2a A5): no new tokens are minted; releases from the pre-minted pools are dilution."""
    em = p.get("emissions_model") or {}
    if isinstance(em, dict) and em.get("model") == "distributed_from_premint":
        return True
    t = f"{src.text if src else ''} {src.comment if src else ''}".lower()
    return "fixed supply" in t or "pre-mint" in t or "premint" in t or "minted at genesis" in t


# ============================================================================================ the hidden data sheet
class Data:
    """`_data`: one row per project, one column per field; every number the reader shows is here."""

    def __init__(self, ws):
        self.ws = ws
        self.rows: dict = {}
        self.cols: dict = {}
        ws.cell(1, 1, "project")

    def _row(self, project: str) -> int:
        if project not in self.rows:
            self.rows[project] = len(self.rows) + 2
            self.ws.cell(self.rows[project], 1, project)
        return self.rows[project]

    def _col(self, field_id: str, label: str | None = None) -> int:
        if field_id not in self.cols:
            self.cols[field_id] = len(self.cols) + 2
            self.ws.cell(1, self.cols[field_id], label or field_id)
        return self.cols[field_id]

    def put(self, project: str, field_id: str, value) -> str:
        r, c = self._row(project), self._col(field_id)
        self.ws.cell(r, c, value)
        return f"{DATA_SHEET}!${get_column_letter(c)}${r}"

    def ref(self, project: str, field_id: str) -> str | None:
        if project not in self.rows or field_id not in self.cols:
            return None
        return f"{DATA_SHEET}!${get_column_letter(self.cols[field_id])}${self.rows[project]}"

    def value(self, project: str, field_id: str):
        if project not in self.rows or field_id not in self.cols:
            return None
        return self.ws.cell(self.rows[project], self.cols[field_id]).value


# ============================================================================================ the hidden series sheet
def excel_serial(d) -> int:
    return int((pd.Timestamp(d).normalize() - pd.Timestamp("1899-12-30")).days)


class SeriesSheet:
    """`_series`: column A the dates (Excel serials, one row per day to the as-of date), one column per daily series.
    Row 1 holds the selected period in days (B1, from the Summary selector); row 2 each series' key; row 3 its last
    complete day; row 4 its first day. Every period formula reads B1, so one cell drives every growth column."""
    FIRST = 6

    def __init__(self, ws, asof: pd.Timestamp):
        self.ws, self.asof = ws, pd.Timestamp(asof).normalize()
        self.start = self.asof - pd.Timedelta(days=SERIES_YEARS_BACK)
        self.last = self.FIRST + SERIES_YEARS_BACK
        self.cols: dict = {}
        self.meta: dict = {}
        ws.cell(1, 1, "days in the selected period")
        ws.cell(1, 2, _days_formula())
        ws.cell(2, 1, "series")
        ws.cell(3, 1, "last day held")
        ws.cell(4, 1, "first day held")
        ws.cell(self.FIRST - 1, 1, "date")
        for i in range(SERIES_YEARS_BACK + 1):
            ws.cell(self.FIRST + i, 1, excel_serial(self.start + pd.Timedelta(days=i)))

    @property
    def dates(self) -> str:
        return f"{SERIES_SHEET}!$A${self.FIRST}:$A${self.last}"

    def add(self, key: str, s: pd.Series, complete_only: bool = False) -> str | None:
        """The series' column letter (written once per key), or None when nothing is held in the window."""
        if key in self.cols:
            return self.cols[key]
        s = s[(s.index >= self.start) & (s.index <= self.asof)]
        if complete_only:
            s = s[s.index < self.asof]
        if s.empty:
            return None
        c = len(self.cols) + 2
        L = get_column_letter(c)
        self.cols[key] = L
        ws = self.ws
        ws.cell(2, c, key)
        ws.cell(3, c, excel_serial(s.index.max()))
        ws.cell(4, c, excel_serial(s.index.min()))
        for d, v in s.items():
            ws.cell(self.FIRST + (pd.Timestamp(d).normalize() - self.start).days, c, float(v))
        full = pd.date_range(s.index.min(), s.index.max())
        self.meta[key] = {"first": s.index.min(), "last": s.index.max(), "n": len(s),
                          "holes": len(full) - len(s.index.normalize().unique())}
        return L

    def vals(self, L: str) -> str:
        return f"{SERIES_SHEET}!${L}${self.FIRST}:${L}${self.last}"

    def end(self, L: str) -> str:
        return f"{SERIES_SHEET}!${L}$3"

    def flow_growth(self, L: str) -> str:
        """(sum over the last N days to the last day held) ÷ (sum over the N days before) − 1, every day of both windows
        held, else n/a (P1-5) — N from the selector."""
        D, V, E, N = self.dates, self.vals(L), self.end(L), DAYS_REF
        held = f'COUNTIFS({D},">"&({E}-2*{N}),{D},"<="&{E},{V},"<>")'
        cur = f'SUMIFS({V},{D},">"&({E}-{N}),{D},"<="&{E})'
        prev = f'SUMIFS({V},{D},">"&({E}-2*{N}),{D},"<="&({E}-{N}))'
        return f'IF({held}<2*{N},"n/a",IF({prev}=0,"n/a",{cur}/{prev}-1))'

    def stock_change(self, L: str) -> str:
        """value on the last day held ÷ value N days earlier − 1 (both days held), else n/a."""
        D, V, E, N = self.dates, self.vals(L), self.end(L), DAYS_REF
        now = f"INDEX({V},MATCH({E},{D},0))"
        then = f"INDEX({V},MATCH({E}-{N},{D},0))"
        return f'IFERROR(IF(AND(ISNUMBER({now}),ISNUMBER({then})),IF({then}=0,"n/a",{now}/{then}-1),"n/a"),"n/a")'

    def describe(self, key: str) -> str:
        m = self.meta.get(key)
        if not m:
            return "no daily series held"
        holes = f", {m['holes']:,} day(s) with no row" if m["holes"] else ""
        return f"held {m['first'].date()}..{m['last'].date()} ({m['n']:,} days{holes})"


def _days_formula() -> str:
    """The selected period in days: the list's days, or the custom number of days (whole, at least 1)."""
    expr = f"IF(AND(ISNUMBER({CUSTOM_CELL}),{CUSTOM_CELL}>=1),ROUND({CUSTOM_CELL},0),{DEFAULT_CUSTOM_DAYS})"
    for lab, d in reversed(PERIODS):
        if d is None:
            continue
        expr = f'IF({SELECT_CELL}="{lab}",{d},{expr})'
    return "=" + expr


# ============================================================================================ period figures (Python)
def _series(groups: dict, project: str, metric: str) -> pd.Series:
    g = groups.get((project, metric)) if groups else None
    if g is None or getattr(g, "empty", True):
        return pd.Series(dtype=float, index=pd.DatetimeIndex([]))
    s = g.assign(date=pd.to_datetime(g["date"]).dt.normalize()).groupby("date")["value"].last().sort_index()
    return pd.to_numeric(s, errors="coerce").dropna()


def revenue_growth(s: pd.Series, days: int, asof: pd.Timestamp) -> tuple[float | None, str]:
    """(revenue over the P days to the series' last complete day ÷ revenue over the P days before − 1, why not) —
    every day of both windows held (P1-5). The reader computes the same in the workbook (SeriesSheet.flow_growth)."""
    s = s[s.index < asof.normalize()]
    if s.empty:
        return None, "no revenue series"
    end = s.index.max()
    w1 = pd.date_range(end - pd.Timedelta(days=days - 1), end)
    w0 = pd.date_range(w1[0] - pd.Timedelta(days=days), w1[0] - pd.Timedelta(days=1))
    missing = [d for d in w0.append(w1) if d not in s.index]
    if missing:
        return None, (f"needs every day {w0[0].date()}..{end.date()} ({2 * days:,} days); the series is held from "
                      f"{s.index.min().date()}" + (f", {len(missing):,} day(s) missing" if s.index.min() <= w0[0]
                                                    else ""))
    a, b = float(s.loc[w1].sum()), float(s.loc[w0].sum())
    if b == 0:
        return None, f"the earlier {days} days' revenue is 0 — no growth rate"
    return a / b - 1, f"{w1[0].date()}..{end.date()} against {w0[0].date()}..{w0[-1].date()}"


def price_change(s: pd.Series, days: int, asof: pd.Timestamp) -> tuple[float | None, str]:
    s = s[s.index <= asof.normalize()]
    if s.empty:
        return None, "no price series"
    end = s.index.max()
    then = end - pd.Timedelta(days=days)
    if then not in s.index:
        return None, f"no price on {then.date()}; the series is held from {s.index.min().date()}"
    p0 = float(s.loc[then])
    if p0 == 0:
        return None, f"price 0 on {then.date()}"
    return float(s.loc[end]) / p0 - 1, f"{then.date()} to {end.date()}"


def settlement_key(groups: dict, project: str) -> tuple[float | None, str]:
    """A1's key (P1-6): (S365 ÷ circulating, now) ÷ (the same a year earlier) − 1."""
    s = _series(groups, project, "settlement_volume_365d_usd")
    c = _series(groups, project, "circulating_supply_used")
    if s.empty:
        return None, "no trailing-365-day settlement volume series"
    if c.empty:
        return None, "no circulating supply series"
    end = s.index.max()
    then = end - pd.Timedelta(days=365)
    if then not in s.index:
        return None, (f"needs the trailing-365-day settlement volume on {then.date()} (two years of history); held "
                      f"from {s.index.min().date()}")
    cn, ct = c.asof(end), c.asof(then)
    if pd.isna(cn) or pd.isna(ct) or ct <= 0 or cn <= 0:
        return None, f"no circulating supply on {then.date()} or {end.date()}"
    b = float(s.loc[then]) / ct
    if b == 0:
        return None, f"settlement volume 0 on {then.date()}"
    return (float(s.loc[end]) / cn) / b - 1, f"{then.date()} to {end.date()}"


# ============================================================================================ writing cells
@dataclass
class Col:
    """A reader column: ≤35-character header, full definition, format, block, where its numbers come from."""
    header: str
    definition: str
    fmt: str = FMT_TOK
    block: str = ""
    source: str = ""                        # "<audit sheet>, <header>" or how it is computed
    group: int = 0                          # outline level (collapsed group)
    width: float = 14.0
    letter: str = ""
    audit: tuple | None = None              # (sheet, header) it reproduces, for the column map
    family: str = ""                        # status family for a status-driven 0 (P2a A4)
    is_text: bool = False


@dataclass
class Sheet:
    ws: object
    cols: list = field(default_factory=list)
    numbers: int = 0
    texts: int = 0
    codes: int = 0
    value_rows: list = field(default_factory=list)     # project rows (for the blank check)


F_HEAD = Font(name=FONT, size=9, bold=True, color="FFFFFF")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
F_BODY = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True)
F_SUB = Font(name=FONT, size=9, italic=True, color="595959")
F_PARKED = Font(name=FONT, size=10, color="7F7F7F")
FILL_GROUP = PatternFill("solid", fgColor="D9E1F2")
FILL_PARKED = PatternFill("solid", fgColor="EDEDED")
FILL_SELECT = PatternFill("solid", fgColor="FFF2CC")


def _comment(text: str, w: int = 380, h: int = 200) -> Comment:
    c = Comment(text, "token_metrics")
    c.width, c.height = w, h
    return c


def write_headers(sh: Sheet, row: int = 4, block_row: int = 3) -> None:
    ws = sh.ws
    last_block = None
    for i, col in enumerate(sh.cols, start=1):
        assert len(col.header) <= MAX_HEADER, col.header
        col.letter = get_column_letter(i)
        c = ws.cell(row, i, col.header)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
        defin = plain(col.definition) + (f"\n\nSource: {col.source}." if col.source else "")
        c.comment = _comment(defin[:2400])
        if col.block and col.block != last_block:
            b = ws.cell(block_row, i, col.block)
            b.font = Font(name=FONT, size=9, bold=True, color="1F3864")
        last_block = col.block
        ws.column_dimensions[col.letter].width = col.width
        if col.group:
            ws.column_dimensions[col.letter].outlineLevel = col.group
            ws.column_dimensions[col.letter].hidden = True
    ws.row_dimensions[row].height = 42
    ws.freeze_panes = ws.cell(row + 1, 3)


def _idx(sh: Sheet, col: Col) -> int:
    return sh.cols.index(col) + 1


_LEAD: dict = {}                            # this build's as-of day for lead sentences (build_reader sets it)
_PERIOD_WORDS = re.compile(r"\b(days?|quarters?|year|annuali[sz]ed|selected period|period chosen|spot|latest|"
                           r"trailing|month|week|as of)\b", re.I)


def lead_sentence(col: Col, project: str) -> str:
    """P2c-7: one plain sentence before a cell's working reference — what the number is (the column's definition, its
    first sentence, which carries the basis), for which project, and its period (said by the definition, else the
    as-of day)."""
    d = plain(col.definition).strip() or col.header
    first = re.split(r"(?<=[.;])\s", d, maxsplit=1)[0].rstrip(".; ")
    if len(first) > 240:
        first = first[:237].rstrip() + "…"
    asof = _LEAD.get("asof")
    period = "" if _PERIOD_WORDS.search(first) or asof is None else f", as of {asof:%Y-%m-%d}"
    return f"{col.header} for {project}: {first}{period}."


def _with_lead(comment: str, col: Col, project: str) -> str:
    """A comment that is only the working reference gets the lead sentence first (P2c-7)."""
    return f"{lead_sentence(col, project)}\n\n{comment.strip()}" if comment.strip().startswith("Working:") \
        else comment


def put_number(sh: Sheet, data: Data, row: int, col: Col, project: str, field_id: str, value: float,
               fill: PatternFill | None = None, comment: str = "", zero_label: str = "") -> str:
    comment = _with_lead(comment, col, project)
    ref = data.put(project, field_id, float(value))
    c = sh.ws.cell(row, _idx(sh, col), f"={ref}")
    c.number_format, c.font = (_zero_format(col.fmt, zero_label) if zero_label and float(value) == 0 else col.fmt), F_BODY
    if fill is not None:
        c.fill = PatternFill("solid", fgColor=fill.fgColor.rgb)
    if comment:
        c.comment = _comment(comment[:2400])
    sh.numbers += 1
    return c.coordinate


def put_formula(sh: Sheet, row: int, col: Col, formula: str, fill: PatternFill | None = None,
                comment: str = "") -> str:
    c = sh.ws.cell(row, _idx(sh, col), formula)
    c.number_format, c.font = col.fmt, F_BODY
    if fill is not None:
        c.fill = PatternFill("solid", fgColor=fill.fgColor.rgb)
    if comment:
        c.comment = _comment(comment[:2400])
    sh.numbers += 1
    return c.coordinate


def _row_project(sh: Sheet, row: int) -> str:
    return str(sh.ws.cell(row, 1).value or "this project")


def put_code(sh: Sheet, row: int, col: Col, code: str, reason: str, fill: PatternFill | None = None) -> str:
    reason = _with_lead(reason or "", col, _row_project(sh, row)) or reason
    c = sh.ws.cell(row, _idx(sh, col), code)
    c.font = Font(name=FONT, size=9, italic=True, color=CODES[code])
    c.alignment = Alignment(horizontal="right")
    if fill is not None:
        c.fill = PatternFill("solid", fgColor=fill.fgColor.rgb)
    c.comment = _comment(reason or code)
    sh.codes += 1
    return c.coordinate


def put_text(sh: Sheet, row: int, col: Col, text: str, comment: str = "", bold: bool = False) -> str:
    comment = _with_lead(comment, col, _row_project(sh, row)) if comment and col.header != "Project" else comment
    c = sh.ws.cell(row, _idx(sh, col), text)
    c.font = F_BOLD if bold else F_BODY
    c.number_format = TEXT_FMT                     # a descriptor, never a value cell (P1-2)
    if comment:
        c.comment = _comment(comment[:2400])
    sh.texts += 1
    return c.coordinate


def blank_reason(p: dict, src: Src | None, col: Col) -> tuple[str, str]:
    """(code, reason) for an audit cell that is EMPTY by construction (P2a A2: no blank cell on the reader)."""
    import config                                          # noqa: PLC0415
    name = p["name"]
    h = (src.header if src else "") or col.header
    if parked_reason(name):
        return "n/a", parked_reason(name)
    if h.startswith("Payout ratio") or col.header.startswith("Payout ratio"):
        if "holder_value_revenue_usd" not in config.metrics_for_project(p):
            return "by design", (f"{name} returns no revenue-funded value to holders — no buyback, revenue-funded burn "
                                 "or distribution is credited — so there is no payout to divide by revenue.")
    if "TOTAL DILUTION" in h or col.header.startswith("Total dilution"):
        return "n/a", (f"No outflow (security emissions, incentives, grants or unlocks) has a series for {name}, so "
                       "total dilution is not formed.")
    return "by design", (f"'{col.header}' does not apply to {name}: the audit cell is empty by construction (the column "
                         "is formed for other projects on its tab).")


def put_src(sh: Sheet, data: Data, row: int, col: Col, p: dict, src: Src | None, why: str = "",
            extra: str = "") -> str:
    """An audit cell's number as a formula on `_data` (its fill, its comment), or its code. A BLANK audit cell gets a
    code with the reason (P2a A2). A status-driven 0 shows its status beside the 0 (P2a A4)."""
    name = p["name"]
    if src is not None and src.num is None and not str(src.text or "").strip():
        code, reason = blank_reason(p, src, col)
        return put_code(sh, row, col, code, f"{reason}\n\nWorking: {src.where} (empty).")
    if src is not None and src.num is not None:
        note = plain(src.comment) if src.comment else ""
        zl = status_note(p, col.family) if col.family else ""
        lines = [f"{zl[:1].upper()}{zl[1:]}: the figure is 0 because the programme {zl}."] if zl and src.num == 0 else []
        lines += [x for x in (extra, note) if x]
        return put_number(sh, data, row, col, name, f"{src.sheet}|{src.coord}", src.num, src.fill,
                          "\n\n".join(lines + [f"Working: {src.where}."]), zero_label=zl if src.num == 0 else "")
    if src is None:
        code, reason = blank_reason(p, None, col)
        return put_code(sh, row, col, code, why or reason)
    return put_code(sh, row, col, code_of(src.text), _code_reason(src, why), src.fill)


def icon_arrows(ws, rng: str) -> None:
    """▲ / ► / ▼ for a change: down below −1%, flat within ±1%, up from +1% (P1-2)."""
    iset = IconSet(iconSet="3Arrows", showValue=None, reverse=None,
                   cfvo=[FormatObject(type="num", val=-1e12), FormatObject(type="num", val=-0.01),
                         FormatObject(type="num", val=0.01)])
    ws.conditional_formatting.add(rng, Rule(type="iconSet", iconSet=iset))


def code_style(ws, rng: str) -> None:
    """A period formula that returns n/a is shown as a code (grey, italic) — its reason is the cell's comment."""
    first = rng.split(":")[0].replace("$", "")
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"ISTEXT({first})"],
                                                   font=Font(italic=True, color=CODES["n/a"])))


def grey_row(sh: Sheet, row: int) -> None:
    for c in range(1, len(sh.cols) + 1):
        cell = sh.ws.cell(row, c)
        cell.fill = FILL_PARKED
        if cell.font is None or not cell.font.color or cell.font.color.rgb in (None, "FF000000"):
            cell.font = Font(name=FONT, size=10, color="7F7F7F", italic=bool(cell.font and cell.font.i))


# ============================================================================================ the build
def _green(fill) -> bool:
    import build_workbook as bw                            # noqa: PLC0415
    return fill is not None and str(fill.fgColor.rgb)[-6:] == str(bw.FILL_GREEN.fgColor.rgb)[-6:]


def _primary(p: dict) -> int:
    a = p.get("archetypes") or []
    return int(a[0]) if a else 0


def _rank_rows(keys: dict) -> list:
    """[(project, rank or None, key or None)] — measured keys high to low, then the unmeasured by name, UNRANKED
    (P2a A1)."""
    have = sorted((k for k in keys.items() if k[1] is not None), key=lambda kv: -kv[1])
    rest = sorted(k for k, v in keys.items() if v is None)
    out = [(p, i + 1, v) for i, (p, v) in enumerate(have)]
    return out + [(p, None, None) for p in rest]


def build_reader(wb, path, *, asof, groups: dict, headline_cells: list, projects: list,
                 previous: dict | None = None) -> dict:
    """Write the reader workbook at `path` from the built (audit) workbook `wb`. Returns its STATS."""
    import build_workbook as bw                            # noqa: PLC0415
    import config                                          # noqa: PLC0415
    HEADLINE_LOCATIONS.clear()
    STATS.clear()
    COLUMN_MAP.clear()
    asof = pd.Timestamp(asof).normalize()
    _AERO_PARK.clear()
    _AERO_PARK["day"] = asof
    _LEAD["asof"] = asof
    A = Audit(wb)
    out = Workbook()
    ws_about = out.active
    ws_about.title = ABOUT
    tabs = [SUMMARY, A1_TAB, A2_TAB, A3_TAB, A4_TAB, COMPARISON, BUSINESS]
    wss = {t: out.create_sheet(t) for t in tabs}
    ws_charts = out.create_sheet(CHARTS)
    ws_def = out.create_sheet(DEFS)
    ws_data = out.create_sheet(DATA_SHEET)
    ws_data.sheet_state = "hidden"
    ws_ser = out.create_sheet(SERIES_SHEET)
    ws_ser.sheet_state = "hidden"
    data = Data(ws_data)
    ser = SeriesSheet(ws_ser, asof)
    sheets = {t: Sheet(wss[t]) for t in tabs}
    ctx = Ctx(A, data, ser, groups, projects, asof, config, bw, {(h["sheet"], h["project"], h["header"]): h
                                                                  for h in headline_cells}, sheets)
    keys = _archetype_keys(A, groups, projects, config, bw)
    _write_summary(sheets[SUMMARY], ctx, keys)
    for spec in tab_specs(bw):
        _write_arch_tab(sheets[spec.tab], ctx, spec)
    _write_comparison(sheets[COMPARISON], ctx, keys)
    _write_business(sheets[BUSINESS], ctx)
    chart_notes = _write_charts(ws_charts, ctx, keys)
    _write_definitions(ws_def, sheets)
    _write_about(ws_about, asof, sheets, projects)
    blanks = _blank_unexplained(sheets)
    out.save(path)
    import workbook_values                                 # noqa: PLC0415
    vals = workbook_values.store_values(path, out)
    STATS.update({"path": str(path), "tabs": [ws.title for ws in out.worksheets],
                  "columns": {n: len(s.cols) for n, s in sheets.items()},
                  "numbers": {n: s.numbers for n, s in sheets.items()},
                  "texts": {n: s.texts for n, s in sheets.items()},
                  "codes": {n: s.codes for n, s in sheets.items()},
                  "blank_unexplained": len(blanks), "blank_cells": blanks[:20], "charts": chart_notes,
                  "formulas": vals.get("formulas"), "failed": vals.get("failed")})
    _write_column_map(Path(path).with_name("reader_column_map.csv"), A, sheets)
    return dict(STATS)


# P2c-4: the A3 cells a parked project shows as the code "parked" where they hold no number, headline or not
PARKED_CODE_HEADERS = ("Net absorption Q0", "Net absorption incl. unlocks", "RETIREMENT ÷ DILUTION")


@dataclass
class Ctx:
    A: Audit
    data: Data
    ser: SeriesSheet
    groups: dict
    projects: list
    asof: pd.Timestamp
    config: object
    bw: object
    headlines: dict
    sheets: dict = None                     # reader tab -> Sheet (P2c-4: a parked headline's code is rewritten)

    def p(self, name: str) -> dict:
        return next(p for p in self.projects if p["name"] == name)

    def record(self, src: Src | None, sheet: str, cell: str, project: str) -> None:
        """A headline shown on the reader is read there by the guard (P1-9). P2c-4: a PARKED project's headline that
        holds no number reads the code "parked" (not n/a or blocked), the one-line park reason first in its comment."""
        if src is not None and cell and ((src.sheet, project, src.header) in self.headlines
                                         or src.header.startswith(PARKED_CODE_HEADERS)):
            if (src.sheet, project, src.header) in self.headlines:
                HEADLINE_LOCATIONS[(src.sheet, project, src.header)] = (sheet, cell)
            sh = (self.sheets or {}).get(sheet)
            why = parked_reason(project)
            if sh is not None and why:
                c = sh.ws[cell]
                if isinstance(c.value, str) and not c.value.startswith("="):
                    under = c.comment.text if c.comment else ""
                    c.value = "parked"
                    c.font = Font(name=FONT, size=9, italic=True, color=CODES["parked"])
                    c.comment = _comment(f"{why}\n\nBehind it: {under}" if under else why)


def _blank_unexplained(sheets: dict) -> list:
    """Every empty value cell on a project row (P2a E: BLANK-UNEXPLAINED is 0 on the reader)."""
    out = []
    for name, sh in sheets.items():
        for r in sh.value_rows:
            for i, col in enumerate(sh.cols, start=1):
                v = sh.ws.cell(r, i).value
                if v is None or (isinstance(v, str) and not v.strip()):
                    out.append(f"{name}!{get_column_letter(i)}{r} ({col.header})")
    return out


def _archetype_keys(A: Audit, groups, projects, config, bw) -> dict:
    """{project: (primary archetype, key number or None, Src or None, why)} (P1-6)."""
    out = {}
    for p in projects:
        n, arch = p["name"], _primary(p)
        if arch == 1:
            v, why = settlement_key(groups, n)
            out[n] = (arch, v, None, why)
            continue
        sheet, head = {2: (A2_TAB, "Customer revenue ÷ emissions value (x)"),
                       3: (A3_TAB, bw.RETIRE_VS_EMISSIONS_HEAD),
                       4: (A4_TAB, "BURN ÷ ISSUANCE (x)")}.get(arch, (None, None))
        src = A.cell(sheet, head, n) if sheet else None
        out[n] = (arch, src.num if src is not None else None, src,
                  "" if src is not None else f"{n} has no row on {sheet}" if sheet else "no archetype declared")
    return out


def _period_formula_cell(sh: Sheet, ctx: Ctx, row: int, col: Col, name: str, metric: str, kind: str, label: str,
                         fill=None, extra: str = "") -> str:
    """A selected-period growth (flow) or change (stock) of the stored daily series, as a formula on _series."""
    s = _series(ctx.groups, name, metric)
    L = ctx.ser.add(f"{name}|{metric}", s, complete_only=(kind == "flow"))
    drow = ctx.A.data_row(f"{name}|{metric}")
    where = f"audit file, Data!A{drow} ({name}|{metric})" if drow else f"the stored series {name}|{metric}"
    if L is None:
        return put_code(sh, row, col, "n/a", f"No daily {label} series is held for {name}, so there is no change over "
                        f"the selected period.\n\nWorking: {where}.")
    how = ("the sum over the last N days to the last day held ÷ the sum over the N days before − 1; every day of "
           "both windows must hold a row" if kind == "flow" else
           "the value on the last day held ÷ the value N days earlier − 1; both days must hold a row")
    body = ctx.ser.flow_growth(L) if kind == "flow" else ctx.ser.stock_change(L)
    return put_formula(sh, row, col, "=" + body, fill,
                       f"{label[:1].upper()}{label[1:]}, change over the selected period (cell B2 on Summary; N days): "
                       f"{how}. n/a when the history is shorter or a day is missing. The series is "
                       f"{ctx.ser.describe(f'{name}|{metric}')}." + (f"\n\n{extra}" if extra else "")
                       + f"\n\nWorking: {where}.")


def _revenue_metric(config, name: str) -> str:
    return config.revenue_headline_metric(name)


# P2e-3 (Jake 2026-10-28): Summary's revenue column reads revenue_headline_metric — PROTOCOL revenue on every row
# (revenue_usd, or the sum of the protocol revenue streams), never customer revenue; the header says so
REV_HEADER = "Protocol revenue, annualised"
_DATA_Q0 = ("Q0 (flow: trailing 90d sum, or 3 complete months for a monthly series — see Q0 basis · stock: 90d "
            "avg)")
# the source field a stored source string reads (fetch/llama.py: dailyRevenue -> revenue_usd)
_SOURCE_FIELD = {("defillama", "revenue_usd"): "DefiLlama {slug}: dailyRevenue (its 'Revenue' line)",
                 ("defillama:near-intents", "intents_revenue_usd"): "DefiLlama near-intents: dailyRevenue",
                 ("maple_page:island.ALL.revenueUsd", "revenue_usd"): "Maple's own page (island.ALL.revenueUsd)"}


def _series_name(ctx, p: dict, metric: str) -> str:
    """'<label>' from <source> — the stored series a revenue cell reads, as the audit file's Data row names it."""
    A, name = ctx.A, p["name"]
    key = f"{name}|{metric}"
    label = A.data_field(key, "Label") or metric
    src = str(A.data_field(key, "Source (of latest point)") or "").strip()
    where = _SOURCE_FIELD.get((src, metric), src or "no stored row").format(
        slug=p.get("defillama_fees_slug") or name.lower())
    return f"{name}|{metric} '{str(label).split(' — ')[0]}' from {where}"


def revenue_series_note(ctx, p: dict, metric: str, value) -> str:
    """P2e-3: the exact series the protocol-revenue cell reads (each leg where it sums several) and, where it reads 0,
    what the source reports — the stored rows' own count — and the reason ONLY where the project's record (its fee
    split, with its source and date) gives one; otherwise that no reason is published. Never inferred."""
    config, A, name = ctx.config, ctx.A, p["name"]
    legs = list((getattr(config, "REVENUE_TOTAL", {}) or {}).get(name) or ()) if metric == "revenue_total_usd" else []
    parts = [(f"Series: {name}|{metric} = the sum of " + "; ".join(_series_name(ctx, p, m) for m in legs) + ".")
             if legs else f"Series: {_series_name(ctx, p, metric)}."]
    if value == 0:
        held = []
        for m in legs or [metric]:
            k = f"{name}|{m}"
            d = A.data_field(k, "Q0: days of the 90 the series existed for")
            ev = A.data_field(k, "Q0: days with a non-zero value (events)")
            held.append(f"{m}: a stored row on {int(d) if isinstance(d, (int, float)) else '?'} of the 90 days, "
                        f"{int(ev) if isinstance(ev, (int, float)) else '?'} of them non-zero")
        parts.append(f"It reads 0 because the source reports 0 — a measured 0, not a missing figure ({'; '.join(held)}).")
        fs = p.get("fee_split") or {}
        if str(fs.get("status") or "").lower() == "off" and fs.get("note"):
            parts.append(f"The reason, from {name}'s own record (fee split, source {fs.get('source_url', '?')}, read "
                         f"{fs.get('source_date', '?')}): \"{str(fs['note']).split(' See ')[0].strip()}\"")
        else:
            parts.append("The source publishes no reason for the 0.")
        if "customer_revenue_usd" in config.metrics_for_project(p):
            cr = A.data_field(f"{name}|customer_revenue_usd", _DATA_Q0)
            parts.append("Customer revenue (what users pay — the A2 tab's figure) is a different series and is not "
                         "counted here" + (f": {cr:,.0f} over the last 90 days." if isinstance(cr, (int, float))
                                           else "."))
    return " ".join(parts)


def _tvl_metric(p: dict) -> str:
    return "tvl_usd" if _primary(p) == 1 else "protocol_tvl_usd"


# ---------------------------------------------------------------------------------------------- Summary
SUMMARY_SOURCES = {
    "mcap": ("Master", "Market cap ($)"),
    "fdv": ("Master", "FDV ($)"),
    "rev_q0": ("Master", "Revenue Q0 ($)"),
    "hvy": (CP_TAB, "HOLDER VALUE YIELD"),
    "dil": (CP_TAB, "TOTAL DILUTION = outflows a+b+c+d+e (annualised) ÷ tradeable float"),
    "float": (CP_TAB, "TRADEABLE FLOAT"),
    "nhy": (CP_TAB, "NET HOLDER YIELD"),
    "nsc_pct": ("Master", "Net supply change, annualised % of circulating"),
    "nsc_tok": ("Master", "NET SUPPLY CHANGE Q0 (tokens)"),
    "payout": (CP_TAB, "Payout ratio — the project's own net basis"),
    "payout_gross": (CP_TAB, "Payout ratio — gross"),
}
# the figures "What changed" follows over the selected period: (label, metric or a function of the project, kind)
WHAT_CHANGED = (("Price", "price_usd", "stock"), ("Market cap", "market_cap_usd", "stock"),
                ("Revenue", "@revenue", "flow"), ("Fees", "@fees", "flow"), ("TVL", "@tvl", "stock"),
                ("Circulating supply", "circulating_supply_used", "stock"))


def _write_summary(sh: Sheet, ctx: Ctx, keys: dict) -> None:
    ws = sh.ws
    C = {k: Col(*v) for k, v in {
        "project": ("Project", "The project.", FMT_TOK, "", "", 0, 16),
        "token": ("Token", "The token's symbol.", FMT_TOK, "", "", 0, 8),
        "arch": ("Archetypes", "The archetypes the project is assessed under, primary first: A1 infrastructure, A2 "
                 "coordination, A3 revenue buyback, A4 permanent burn.", FMT_TOK, "", "", 0, 11),
        "mcap": ("Market cap", "Price × the circulating supply the token ratios use.", FMT_USD, "Size",
                 "Master, Market cap ($)", 0, 12),
        "fdv": ("FDV", "Fully diluted value: price × the total supply FDV uses.", FMT_USD, "Size", "Master, FDV ($)",
                0, 12),
        # P2e-3: the column is PROTOCOL revenue on every row (revenue_usd — DefiLlama dailyRevenue, Maple's own
        # page — or the sum of the protocol revenue streams for NEAR and Hyperliquid), never customer revenue
        "rev": (REV_HEADER, "Protocol revenue (what the protocol keeps — DefiLlama's 'Revenue', or the project's own "
                "revenue figure where it publishes one; all revenue streams where a project has more than one) over "
                "the last 90 days × 365 ÷ the days the series covers. Not customer revenue (what users pay, e.g. a "
                "lending protocol's borrower interest — the A2 tab's figure). Each cell's comment names its series.",
                FMT_USD, "Business", "Master, Revenue Q0 ($) and Data coverage", 0, 13),
        "growth": ("Revenue growth, selected period", "Revenue over the selected period (cell B2) to the series' last "
                   "complete day ÷ revenue over the same number of days before it − 1. n/a where the history is "
                   "shorter or a day is missing.", FMT_PCT, "Business", "computed from the stored daily revenue series",
                   0, 13),
        "key": ("Archetype key", "The primary archetype's ranking key, higher is better — see Definitions: "
                + " · ".join(f"A{a}: {k[0]}" for a, k in ARCHETYPE_KEYS.items()), FMT_X, "Rank", "A-tab key columns",
                0, 12),
        "rank": ("Rank in archetype", "Position within the primary archetype by its key. A project with no measured "
                 "key reads 'unranked' and sorts last; a parked project reads 'parked' and is outside the ranks.",
                 FMT_RANK, "Rank", "computed", 0, 9),
        "hvy": ("Holder value yield", "Revenue-funded value to holders a year (retired + distributed + liquid voter "
                "bribes) ÷ market cap.", FMT_PCT, "Holders", "CP Comparability, HOLDER VALUE YIELD", 0, 11),
        "dil": ("Total dilution, % of float", "Every outflow a year (security emissions, supply- and demand-side "
                "incentives, grants, unlocks) ÷ the tradeable float.", FMT_PCT, "Holders",
                "CP Comparability, TOTAL DILUTION ÷ tradeable float", 0, 11),
        "dil_tok": ("Total dilution, tokens a year", "Total dilution % × the tradeable float, in tokens a year.",
                    FMT_TOK, "Holders", "computed: CP total dilution × CP tradeable float", 0, 13),
        "nhy": ("Net holder yield", "Holder value yield − every outflow a year at each day's price ÷ market cap.",
                FMT_PCT, "Holders", "CP Comparability, NET HOLDER YIELD", 0, 11),
        "nsc_pct": ("Net supply change, % a year", "Issuance − burn, annualised, % of circulating supply (+ is "
                    "net inflation). For a pre-minted supply nothing is issued: releases from the pre-minted pools "
                    "appear under dilution, not here.", FMT_PCT, "Supply",
                    "Master, Net supply change, annualised % of circulating", 0, 11),
        "nsc_tok": ("Net supply change, 90 days", "Issuance − burn over the last 90 days, in tokens.", FMT_TOK,
                    "Supply", "Master, NET SUPPLY CHANGE Q0 (tokens)", 0, 14),
        "payout": ("Payout ratio", "Credited holder value ÷ revenue on the project's own revenue basis. Above 1.0x the "
                   "comment says what funds the excess.", FMT_PCT, "Holders",
                   "CP Comparability, Payout ratio (net basis)", 0, 10),
        "price": ("Price change, selected period", "Price on the last day held ÷ price the selected number of days "
                  "before − 1.", FMT_PCT, "Market", "computed from the stored daily price series", 0, 12),
        "conf": ("Confidence (share green)", "The share of this row's numbers whose source cell is GREEN "
                 "(verified source, confirmed mechanism, no partial window).", FMT_PCT, "Quality", "computed", 0, 11),
        "changed": ("What changed", "The figures on this row that moved more than 10% over the selected period "
                    "(cell B2): price, market cap, revenue, fees, TVL and circulating supply.", TEXT_FMT, "Quality",
                    "computed from the stored daily series", 0, 40),
    }.items()}
    for k, v in SUMMARY_SOURCES.items():
        if k in C:
            C[k].audit = v
    for k in ("hvy", "payout", "key"):
        C[k].family = "buyback"
    C["changed"].is_text = True
    sh.cols = list(C.values())
    ws.cell(1, 1, "Summary — every project, grouped by primary archetype and ranked by its key").font = F_TITLE
    _selector(ws)
    ws.cell(3, 1, f"As of {ctx.asof.date()} · hover any header for its definition, any cell for its source").font = F_SUB
    write_headers(sh)
    row = 5
    by_arch: dict = {}
    for p in ctx.projects:
        by_arch.setdefault(keys[p["name"]][0], []).append(p)
    for arch in sorted(by_arch):
        g = ws.cell(row, 1, f"{ARCHETYPE_NAMES.get(arch, 'No archetype')} — ranked by "
                    f"{ARCHETYPE_KEYS.get(arch, ('no key',))[0].lower()}")
        g.font, g.fill = F_BOLD, FILL_GROUP
        for c in range(2, len(sh.cols) + 1):
            ws.cell(row, c).fill = FILL_GROUP
        row += 1
        members = [p for p in by_arch[arch] if not parked_reason(p["name"])]
        parked = [p for p in by_arch[arch] if parked_reason(p["name"])]
        for name, rank, _kv in _rank_rows({p["name"]: keys[p["name"]][1] for p in members}):
            _summary_row(sh, ctx, C, row, ctx.p(name), rank, keys[name])
            row += 1
        for p in sorted(parked, key=lambda q: q["name"]):
            _summary_row(sh, ctx, C, row, p, None, keys[p["name"]], parked=True)
            grey_row(sh, row)
            row += 1
    last = row - 1
    for k in ("growth", "price"):
        code_style(ws, f"{C[k].letter}5:{C[k].letter}{last}")
    ws.auto_filter.ref = None


def _selector(ws) -> None:
    """The period selector (P1-5, P2a): B2 picks the period; D2 holds the days when B2 is 'custom'."""
    ws.cell(2, 1, "Period").font = F_BOLD
    sel = ws.cell(2, 2, DEFAULT_PERIOD)
    sel.font, sel.fill = Font(name=FONT, size=11, bold=True, color="1F3864"), FILL_SELECT
    sel.comment = _comment("Choose the period every growth and price-change column on every tab uses: "
                           + ", ".join(p for p, _d in PERIODS) + ". 6M = 182 days, 1Y = 365, 3Y = 1,095; 'custom' "
                           "uses the number of days in D2.")
    dv = DataValidation(type="list", formula1='"' + ",".join(p for p, _d in PERIODS) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B2")
    ws.cell(2, 3, "Custom days").font = F_SUB
    cd = ws.cell(2, 4, DEFAULT_CUSTOM_DAYS)
    cd.fill, cd.font = FILL_SELECT, Font(name=FONT, size=10, bold=True, color="1F3864")
    cd.comment = _comment("Days for the 'custom' period (a whole number, at least 1). Used only when B2 is 'custom'.")
    dv2 = DataValidation(type="whole", operator="between", formula1="1", formula2="3650", allow_blank=False)
    ws.add_data_validation(dv2)
    dv2.add("D2")


def _period_ref(sheet: str) -> str:
    return "the period chosen in Summary!B2" if sheet != SUMMARY else "the period in B2"


def _summary_row(sh, ctx: Ctx, C, row, p, rank, key, parked: bool = False) -> None:
    A, data, config = ctx.A, ctx.data, ctx.config
    name = p["name"]
    sh.value_rows.append(row)
    put_text(sh, row, C["project"], name, parked_reason(name) if parked else "", bold=True)
    put_text(sh, row, C["token"], p.get("symbol") or "—")
    put_text(sh, row, C["arch"], ", ".join(f"A{a}" for a in (p.get("archetypes") or [])) or "—")
    greens, nums = 0, 0
    srcs = {k: A.cell(sheet, head, name) for k, (sheet, head) in SUMMARY_SOURCES.items()}

    def number(k, src, extra=""):
        nonlocal greens, nums
        cell = put_src(sh, data, row, C[k], p, src, extra=extra)
        ctx.record(src, SUMMARY, cell, name)
        if src is not None and src.num is not None:
            nums += 1
            greens += _green(src.fill)

    number("mcap", srcs["mcap"], _mcap_basis_note(config, name))
    number("fdv", srcs["fdv"])
    rq = srcs["rev_q0"]
    metric = _revenue_metric(config, name)
    days = A.data_field(f"{name}|{metric}", "Q0: days of the 90 the series existed for")
    if rq is not None and rq.num is not None and isinstance(days, (int, float)) and days > 0:
        ref_q = data.put(name, "rev_q0", rq.num)
        ref_d = data.put(name, "rev_q0_days", float(days))
        data.put(name, "rev_ann", rq.num * 365 / float(days))
        put_formula(sh, row, C["rev"], f"={ref_q}*365/{ref_d}", rq.fill,
                    f"Protocol revenue over the last 90 days {rq.num:,.0f} × 365 ÷ {int(days)} covered days.\n\n"
                    f"{revenue_series_note(ctx, p, metric, rq.num)}\n\nWorking: {rq.where}; Data {name}|{metric}.")
        nums += 1
        greens += _green(rq.fill)
    elif rq is not None and rq.num is None and not str(rq.text or "").strip():
        code, why = blank_reason(p, rq, C["rev"])
        put_code(sh, row, C["rev"], code, f"{why}\n\n{revenue_series_note(ctx, p, metric, None)}\n\nWorking: "
                 f"{rq.where} (empty).")
    else:
        # a figure with no covered day is unknown, never 0 (P2e-3: 0 is not n/a, unknown is never 0)
        no_days = rq is not None and rq.num is not None
        put_code(sh, row, C["rev"], code_of(rq.text if rq else ""),
                 _code_reason(rq, ("no day of the last 90 is covered by the series (no covered-days count on its "
                                   "Data row)" if no_days else "no revenue figure") + ". "
                              + revenue_series_note(ctx, p, metric, None)),
                 rq.fill if rq else None)
    _period_formula_cell(sh, ctx, row, C["growth"], name, metric, "flow", "revenue")
    _period_formula_cell(sh, ctx, row, C["price"], name, "price_usd", "stock", "price")
    # the archetype key and the rank (P2a A1: no measured key reads "unranked"; a parked project "parked")
    arch, kv, ksrc, kwhy = key
    if kv is not None:
        if ksrc is not None:
            cell = put_src(sh, data, row, C["key"], p, ksrc)
            ctx.record(ksrc, SUMMARY, cell, name)
            nums += 1
            greens += _green(ksrc.fill)
        else:
            put_number(sh, data, row, C["key"], name, "key", kv, None, f"{ARCHETYPE_KEYS[arch][1]}\n\n{kwhy}.\n\n"
                       f"Working: audit file, Data rows {name}|settlement_volume_365d_usd and "
                       f"{name}|circulating_supply_used.")
            nums += 1
    elif ksrc is not None:
        put_src(sh, data, row, C["key"], p, ksrc, f"No measured key: {kwhy}" if kwhy else "No measured key.")
    else:
        put_code(sh, row, C["key"], "n/a", f"No measured {ARCHETYPE_KEYS.get(arch, ('key',))[0].lower()}"
                 + (f": {kwhy}" if kwhy else "") + ".")
    if parked:
        put_code(sh, row, C["rank"], "parked", parked_reason(name))
    elif rank is not None:
        put_number(sh, data, row, C["rank"], name, "rank", rank, None,
                   f"{rank} in {ARCHETYPE_NAMES.get(arch, '')} by {ARCHETYPE_KEYS.get(arch, ('',))[0].lower()}.\n\n"
                   "Working: computed here from the archetype key column.")
    else:
        put_code(sh, row, C["rank"], "unranked",
                 f"Unranked, sorted last in its group: no measured {ARCHETYPE_KEYS.get(arch, ('key',))[0].lower()}"
                 + (f" — {kwhy}" if kwhy else "") + (f" — {plain(ksrc.text)}" if ksrc is not None and ksrc.text
                                                     else "") + ".")
    number("hvy", srcs["hvy"])
    number("dil", srcs["dil"])
    d, f = srcs["dil"], srcs["float"]
    if d is not None and d.num is not None and f is not None and f.num is not None:
        rd, rf = data.put(name, "dil_pct", d.num), data.put(name, "float", f.num)
        put_formula(sh, row, C["dil_tok"], f"={rd}*{rf}", d.fill,
                    f"{d.num:.2%} × tradeable float {f.num:,.0f}.\n\nWorking: {d.where}; {f.where}.")
        nums += 1
        greens += _green(d.fill)
    else:
        bad = d if d is None or d.num is None else f
        if bad is not None and bad.num is None and not str(bad.text or "").strip():
            code, why = blank_reason(p, bad, C["dil"])
            put_code(sh, row, C["dil_tok"], code, f"{why}\n\nWorking: {bad.where} (empty).")
        else:
            put_code(sh, row, C["dil_tok"], code_of(bad.text if bad else ""),
                     _code_reason(bad, "no total dilution or no tradeable float"), bad.fill if bad else None)
    number("nhy", srcs["nhy"])
    nsc_note = _premint_note(p, srcs["nsc_pct"])
    number("nsc_pct", srcs["nsc_pct"], nsc_note)
    number("nsc_tok", srcs["nsc_tok"], _premint_note(p, srcs["nsc_tok"]))
    number("payout", srcs["payout"], _payout_note(p, srcs["payout"], srcs["payout_gross"]))
    if nums:
        put_number(sh, data, row, C["conf"], name, "confidence", greens / nums, None,
                   f"{greens} of {nums} numbers on this row come from GREEN cells.\n\nWorking: computed here from the "
                   "fills of this row's source cells (each named in its own comment).")
    else:
        put_code(sh, row, C["conf"], "n/a", "No number on this row.")
    _what_changed(sh, ctx, row, C["changed"], p)


def _mcap_basis_note(config, name: str) -> str:
    """The circulating basis market cap uses (P2a D6: the reader's market cap comment states it)."""
    note = (getattr(config, "MCAP_BASIS_NOTE", {}) or {}).get(name)
    return note or ""


def _premint_note(p: dict, src: Src | None) -> str:
    if src is None or src.num is None or src.num != 0 or not is_premint(p, src):
        return ""
    return ("A pre-minted supply: no new tokens are minted, so the net supply change is 0. Releases from the "
            "pre-minted pools appear under dilution (total dilution), not under issuance.")


def _payout_note(p: dict, net: Src | None, gross: Src | None) -> str:
    """Above 1.0x: what funds the excess (P2a A3) — said from the gross payout on fees."""
    if net is None or net.num is None or net.num <= 1.0:
        return ""
    g = gross.num if gross is not None else None
    if g is not None and g <= 1.0:
        return (f"Above 1.0x: credited holder value is {net.num:.2f}x revenue on the project's own revenue basis but "
                f"{g:.2f}x gross fees — the excess over revenue is paid out of fees the revenue basis does not count "
                f"(the fees are the funding; the revenue basis is narrower than what the programme spends). Working: "
                f"{gross.where}.")
    if g is not None:
        return (f"Above 1.0x on both bases (revenue {net.num:.2f}x, gross fees {g:.2f}x): the programme spent more than "
                f"the window's fees — the excess is funded from the treasury or from revenue earned before the window. "
                f"Working: {gross.where}.")
    return "Above 1.0x: credited holder value exceeds revenue on the project's own basis; no gross payout to say more."


def _what_changed(sh, ctx: Ctx, row: int, col: Col, p: dict) -> None:
    """The figures that moved more than 10% over the selected period (P2a A7), as one formula on `_data` helpers."""
    name, config, data = p["name"], ctx.config, ctx.data
    parts, used = [], []
    for label, metric, kind in WHAT_CHANGED:
        m = {"@revenue": _revenue_metric(config, name), "@fees": config.fees_headline_metric(name),
             "@tvl": _tvl_metric(p)}.get(metric, metric)
        s = _series(ctx.groups, name, m)
        L = ctx.ser.add(f"{name}|{m}", s, complete_only=(kind == "flow"))
        if L is None:
            continue
        body = ctx.ser.flow_growth(L) if kind == "flow" else ctx.ser.stock_change(L)
        ref = data.put(name, f"chg|{label}", "=" + body)
        parts.append(f'IF(ISNUMBER({ref}),IF(ABS({ref})>{MOVE_THRESHOLD},"{label} "&TEXT({ref},"+0%;-0%")&" · ",""),"")')
        used.append(f"{label} ({name}|{m})")
    if not parts:
        put_text(sh, row, col, "no daily series held", "No daily series is held for this project.")
        return
    joined = "&".join(parts)
    f = (f'=IF(({joined})="","no figure moved more than {MOVE_THRESHOLD:.0%} over "&{SELECT_CELL},'
         f'"over "&{SELECT_CELL}&": "&{joined})')
    c = sh.ws.cell(row, _idx(sh, col), f)
    c.font, c.number_format = F_BODY, "General"
    c.comment = _comment("Moves above 10% over the selected period (Summary!B2), each computed on its stored daily "
                         "series as the period columns are: " + "; ".join(used) + ".\n\nWorking: audit file, Data rows "
                         "of these series.")
    sh.texts += 1


# ---------------------------------------------------------------------------------------------- archetype tabs
@dataclass
class TabSpec:
    tab: str                                # reader tab (= the audit tab of the same name)
    arch: int
    title: str
    key: tuple                              # ("audit", header prefix) | ("a1",)
    key_header: str
    rank_header: str
    a_extra: tuple = ()                     # (kind, ...) columns beside the key
    valuation: tuple = ()                   # (sheet or None, header prefix, short, family)
    traj: tuple = ()                        # (trajectory id, label stem, fmt, definition)
    drivers: tuple = ()                     # (sheet or None, header prefix, short, family) | ("@growth", metric, short, kind)
    blocks_note: str = ""


def tab_specs(bw) -> list:
    """The four archetype tabs, built alike (P2a B): (a) key and rank, (b) valuation, (c) the last four quarters with
    arrows, (d) drivers, (e) everything else collapsed."""
    M = "Master"
    return [
        TabSpec(A1_TAB, 1, "A1 Infrastructure — ranked by settlement volume per token, 1-year change", ("a1",),
                "Settlement per token, 1-yr change", "Rank in A1",
                a_extra=(("interim", "FEES ÷ ISSUANCE (x)"),),
                valuation=((None, "Price — spot ($)", "Price", ""), (None, "Market cap ($)", "Market cap", ""),
                           (M, "FDV ($)", "FDV", ""),
                           (None, "NRR ON SETTLEMENT VOLUME — market cap", "NRR on settlement", ""),
                           (None, "NRR HEADROOM (settlement)", "NRR headroom (settlement)", ""),
                           (None, "NRR ON TRADING THROUGHPUT", "NRR on trading throughput", ""),
                           (M, "Fees ÷ FDV (annualised)", "Fees ÷ FDV, a year", "")),
                traj=("a1_fees_issuance", "Fees ÷ issuance", FMT_X, "FEES ÷ ISSUANCE (x)"),
                drivers=((None, "Settlement volume, trailing 365d", "Settlement volume, 365 days", ""),
                         (None, "Fees Q0 ($)", "Fees, last 90 days", ""),
                         ("@growth", "@fees", "Fees growth, selected period", "flow"),
                         (None, "Fee per transaction ($)", "Fee per transaction", ""),
                         (None, "TVL — chain ($)", "Chain TVL", ""),
                         ("@growth", "tvl_usd", "TVL change, selected period", "stock"),
                         (None, "Stablecoin supply on chain ($)", "Stablecoins on chain", ""),
                         ("@growth", "stablecoin_supply_usd", "Stablecoins, selected period", "stock"),
                         (None, "Gross issuance Q0 (tokens)", "Gross issuance, 90 days (tokens)", ""),
                         (None, "Gross burn Q0 (tokens)", "Gross burn, 90 days (tokens)", ""),
                         (None, "NET ISSUANCE AFTER BURN, annualised %", "Net issuance after burn, % a year", ""),
                         (None, "Issuer releases Q0 (tokens)", "Issuer releases, 90 days (tokens)", ""),
                         (None, "TOTAL YIELD — stakers", "Total yield, stakers", "staking"),
                         (None, "Total yield — non-stakers", "Total yield, non-stakers", ""),
                         (None, "VALIDATOR STAKING YIELD (annual)", "Validator yield", "staking"),
                         (None, "Spread over the 3-month T-bill: Total yield — stakers", "Stakers' yield minus T-bill",
                          ""))),
        TabSpec(A2_TAB, 2, "A2 Coordination — ranked by customer revenue ÷ emissions", ("audit",
                "Customer revenue ÷ emissions value (x)"), "Customer revenue ÷ emissions (x)", "Rank in A2",
                valuation=((None, "Market cap ($)", "Market cap", ""), (M, "FDV ($)", "FDV", ""),
                           (None, "Free float market cap ($, spot)", "Free float market cap", ""),
                           (None, "FREE FLOAT ÷ ARR (x)", "Free float ÷ ARR (x)", "")),
                traj=("a2_rev_per_token", "Revenue ÷ emission", FMT_USD,
                      "CUSTOMER REVENUE PER TOKEN EMITTED ($/token)"),
                drivers=((None, "Customer revenue Q0 ($)", "Customer revenue, 90 days", ""),
                         ("@growth", "customer_revenue_usd", "Customer revenue growth, period", "flow"),
                         (None, "Emissions to suppliers Q0 (tokens)", "Emissions, 90 days (tokens)", ""),
                         (None, "Emissions Q0 ($ at avg price)", "Emissions, 90 days ($)", ""),
                         ("@growth", "emissions_tokens", "Emissions growth, selected period", "flow"),
                         (None, "Supply units (nodes / hotspots / GPUs), latest", "Supply units", ""),
                         ("@growth", "supply_units", "Supply units, selected period", "stock"),
                         (None, "Capacity utilisation (latest)", "Utilisation", ""),
                         (None, "Revenue per supply unit Q0 ($)", "Revenue per supply unit, 90 days", ""),
                         (None, "SUPPLY TRAJECTORY", "Supply trajectory, % a year", ""),
                         (None, "Unlock dilution, annualised % of free float", "Unlocks, % of free float a year", ""),
                         (None, "TOTAL DILUTION = emissions + unlocks", "Total dilution, % of float a year", ""),
                         (None, "PUBLISHED UNLOCK SCHEDULE", "Published unlocks, next 12 months", ""))),
        TabSpec(A3_TAB, 3, "A3 Revenue Buyback — ranked by retired buyback ÷ emissions", ("audit",
                bw.RETIRE_VS_EMISSIONS_HEAD), "Retired ÷ emissions (x)", "Rank in A3",
                a_extra=(("dilution", bw.RETIRE_VS_DILUTION_HEAD, bw.RETIRE_VS_DILUTION_NOW_HEAD),),
                valuation=((None, "Price — spot ($)", "Price", ""), (None, "Market cap ($)", "Market cap", ""),
                           (None, "FDV ($)", "FDV", ""), (None, "Circulating supply", "Circulating supply", ""),
                           (CP_TAB, "HOLDER VALUE YIELD", "Holder value yield", "buyback")),
                traj=("a3_retire_vs_dilution", "Retired ÷ dilution", FMT_X, "RETIREMENT ÷ DILUTION (x)"),
                drivers=((None, "Base for the implied buyback ($)", "Revenue base, last 90 days", ""),
                         ("@growth", "@revenue", "Revenue growth, selected period", "flow"),
                         (None, "Actual buyback Q0 ($)", "Actual buyback, last 90 days", "buyback"),
                         (None, "Actual buyback Q0 RETIRED (tokens)", "Retired: burned + held (tokens)", "buyback"),
                         (None, "Emissions Q0 (tokens)", "Emissions, last 90 days (tokens)", ""),
                         (None, "UNLOCK DILUTION Q0 (tokens)", "Unlocks, last 90 days (tokens)", ""),
                         (CP_TAB, "Payout ratio — the project's own net basis", "Payout ratio", "buyback"),
                         (None, "CIRCULATING RETIREMENT RATE (tokens)", "Retirement rate, % of circulating",
                          "buyback"),
                         (None, "PROTOCOL STAKING YIELD (tokens)", "Staking yield (tokens)", "staking"))),
        TabSpec(A4_TAB, 4, "A4 Permanent Burn — ranked by burn ÷ issuance (burn ÷ release where pre-minted)",
                ("audit", "BURN ÷ ISSUANCE (x)"), "Burn ÷ issuance (x)", "Rank in A4",
                valuation=((None, "Market cap ($)", "Market cap", ""), (M, "FDV ($)", "FDV", ""),
                           (None, "PERMANENT BURN YIELD", "Permanent burn yield", ""),
                           (M, "Fees ÷ FDV (annualised)", "Fees ÷ FDV, a year", "")),
                traj=("a4_net_change", "Net supply change", FMT_TOK, "NET SUPPLY CHANGE (tokens)"),
                drivers=((None, "GROSS BURN Q0 (tokens)", "Gross burn, 90 days (tokens)", ""),
                         ("@growth", "@burn", "Burn growth, selected period", "flow"),
                         ("@issuance", "", "Issuance or release, 90 days", ""),
                         ("@growth", "@issuance", "Issuance growth, selected period", "flow"),
                         (None, "Fees Q0 ($)", "Fees, last 90 days", ""),
                         (None, "Burn as share of fees (measured)", "Burn share of fees", ""),
                         (None, "Revenue required for crossover", "Revenue needed for crossover", ""),
                         (None, "Governance lever", "Governance lever", ""),
                         (None, "Treasury governance could redeploy (tokens)", "Treasury governance can redeploy",
                          ""))),
    ]


A3_TRAJECTORY_ID = "a3_retire_vs_dilution"
# the short header of every (e) column (≤35 characters); a column not listed gets an automatic short header
E_SHORT = {
    "FDV RETIREMENT RATE (tokens)": "Retirement rate, % of FDV supply",
    "Circulating retirement rate ($)": "Retirement rate ($), % of mkt cap",
    "FDV retirement rate ($)": "Retirement rate ($), % of FDV",
    "Materiality": "Materiality",
    "Buyback status": "Buyback status",
    "Programmed?": "Programmed or discretionary",
    "Buyback destination": "Buyback destination",
    "Destination split (share burned)": "Share of buyback burned",
    "Share applied to this window": "Revenue share applied, 90 days",
    "Status of that window's split": "Status of that share",
    "Current documented share": "Current documented share",
    "Per-product split": "Per-product split",
    "Fees Q0 ($)": "Fees, last 90 days",
    "Implied buyback Q0 ($)": "Implied buyback, last 90 days",
    "Price — 90d average ($)": "Price, 90-day average",
    "Implied buyback Q0 (tokens)": "Implied buyback, 90 days (tokens)",
    "Supply figure complete?": "Supply figure complete?",
    "BUYBACK AS % OF SUPPLY (annualised, implied)": "Implied buyback, % of supply a year",
    "Actual buyback Q0 (tokens)": "Actual buyback, 90 days (tokens)",
    "Actual buyback Q0 BURNED (tokens)": "Bought back and burned (tokens)",
    "Actual buyback Q0 HELD (tokens)": "Bought back and held (tokens)",
    "Actual buyback Q0 DISTRIBUTED": "Bought back, paid to stakers",
    "Actual buyback as % of supply (annualised)": "Retired, % of supply a year",
    "Implied − actual ($)": "Implied minus actual buyback ($)",
    "Net absorption Q0 (tokens)": "Net absorption, 90 days (tokens)",
    "Net absorption, implied basis (tokens)": "Net absorption, implied (tokens)",
    "Unlock dilution, annualised % of free float": "Unlocks, % of free float a year",
    "TOTAL DILUTION = emissions + unlocks": "Total dilution, % of float a year",
    "Net absorption incl. unlocks Q0 (tokens)": "Net absorption incl. unlocks",
    "RETIREMENT ÷ DILUTION (x)": "Retired ÷ (emissions + unlocks)",
    "Emissions — the rate IN FORCE": "Emissions in force, a year (tokens)",
    "RETIREMENT ÷ DILUTION at the emission rate IN FORCE": "Retired ÷ dilution, rate in force",
    "Implied SKY purchased, both legs ($)": "Implied purchases, both legs ($)",
    "Implied BURN ($)": "Implied burn ($)",
    "Implied burn (tokens)": "Implied burn (tokens)",
    "Implied burn − ACTUAL burn (tokens)": "Implied minus actual burn (tokens)",
    "Coverage ratio = actual buyback ÷ revenue": "Buyback ÷ revenue (x)",
    "Fees ÷ FDV (annualised)": "Fees ÷ FDV, a year",
    "Tokens locked (ve)": "Tokens locked or staked",
    "Lock rate = locked inside circulating": "Lock rate, % of circulating",
    "Tokens locked — cross-check (protocol dashboard)": "Tokens locked, protocol dashboard",
    "Lock: contract vs dashboard divergence": "Locked: chain vs dashboard gap",
    "Average lock duration (days)": "Average lock duration (days)",
    "3-MONTH T-BILL": "3-month T-bill",
    "Spread over the 3-month T-bill: Total yield — non-stakers": "Non-stakers' yield minus T-bill",
    "Spread over the 3-month T-bill": "Staking yield minus T-bill",
    "Sky Savings Rate": "Sky Savings Rate",
    "  of which BOUGHT on-chain (DEX fills; Q0": "Of which bought on-chain, 90 days",
    "  of which BOUGHT on-chain — FULL HISTORY": "Of which bought on-chain, all time",
    "PROTOCOL STAKING YIELD ($)": "Staking yield ($)",
    "VOTER APR ON VOTING POWER": "Voter APR on voting power",
    "Staking APR as the protocol publishes it": "Staking APR, as published",
    "Share of staker distributions to VIRTUAL": "Share paid to virtual balances",
    "PENDLE still in the OLD vePENDLE": "Still in old vePENDLE (tokens)",
    "sPENDLE + old vePENDLE (tokens)": "sPENDLE + old vePENDLE (tokens)",
    "Effective float": "Effective float (tokens)",
    "Destination effect on float": "Effect on float",
    "Buyback to LOCKED supply": "Buyback to locked supply",
    "Buyback to YIELD PAYOUT": "Buyback paid out as yield",
    "Reserve balance — cross-check": "Reserve balance, dashboard",
    "Contract vs dashboard divergence": "Reserve: chain vs dashboard gap",
    "Umbrella staked $ (Aave only)": "Umbrella staked (Aave, $)",
    "Buyback % supply at Q1": "Implied buyback %, 3 months ago",
    "Buyback % supply at Q2": "Implied buyback %, 6 months ago",
    "Buyback % supply at Q3": "Implied buyback %, 9 months ago",
    "Circulating supply (used, tokens)": "Circulating supply",
    "VALIDATOR STAKING YIELD": "Validator yield",
    "Validator yield — caveat": "Validator yield caveat",
    "Validator yield GROSS of commission": "Validator yield, gross",
    "Consensus part": "Consensus part of yield",
    "Execution part": "Execution part of yield",
    "Issuance ÷ protocol maximum": "Issuance ÷ protocol maximum",
    "Settlement volume: last date of the NRR inputs": "Settlement inputs: last date",
    "NRR ON SETTLEMENT VOLUME — FULL YEAR": "NRR on settlement, full year",
    "Trading throughput, trailing 365d": "Trading throughput, 365 days",
    "NRR HEADROOM (trading throughput)": "NRR headroom (throughput)",
    "Which NRR is meaningful here": "Which NRR is meaningful",
    "Transactions Q0": "Transactions, last 90 days",
    "Fee per transaction — what it divides": "Fee per transaction: basis",
    "Active addresses (latest, low weight)": "Active addresses",
    "RWA on chain": "RWA on chain",
    "Issuance as % of supply (annualised)": "Issuance, % of supply a year",
    "Staked tokens (locked_tokens)": "Staked tokens",
    "Staking rate = staked inside circulating": "Staking rate",
    "Resulting float": "Float after staking",
    "NET ISSUANCE AFTER BURN Q0 (tokens)": "Net issuance after burn (tokens)",
    "Gross issuance Q0 ($ at avg price)": "Gross issuance, 90 days ($)",
    "FEES ÷ ISSUANCE (x) — key ratio": "Fees ÷ issuance (x)",
    "FREE FLOAT = circulating − locked": "Free float (tokens)",
    "Materiality ": "Materiality",
    "Customer revenue Q0 ($)": "Customer revenue, 90 days",
    "  of which payments collected at the fee aggregator": "Of which aggregator payments",
    "  of which Direct Requests": "Of which Direct Requests",
    "  of which CCIP 1.2/1.5 fees": "Of which CCIP fees",
    "  of which VRF v2.5 payments": "Of which VRF payments",
    "  of which Automation v2.3 premium": "Of which Automation premium",
    "  Customer revenue — not measured": "Customer revenue not measured",
    "  Reserve inflow ÷ payments collected": "Reserve inflow ÷ aggregator",
    "  Emissions basis — what feeds": "Emissions basis",
    "  of which checker-node rewards": "Of which checker-node rewards",
    "  of which cloud host": "Of which cloud host rewards",
    "  of which edge rewards": "Of which edge rewards",
    "  Emissions RELEASED, % of circulating": "Emissions released, % a year",
    "  Emissions EARNED (incl. vesting)": "Emissions earned, % a year",
    "CUSTOMER REVENUE PER TOKEN EMITTED": "Revenue per token emitted ($)",
    "Supplier earnings from emissions": "Supplier earnings: emissions ($)",
    "Supplier earnings from customers": "Supplier earnings: customers ($)",
    "Customer share of supplier earnings": "Customer share of earnings",
    "Publisher Conviction": "Publisher Conviction ($)",
    "Staked tokens (float metric, not demand)": "Staked tokens",
    "Issuance basis for the crossover": "Issuance basis",
    "Release ÷ burn (x)": "Burn ÷ release (x), inverse shown",
    "Q0 WINDOW CAVEAT": "90-day window caveat",
    "NET SUPPLY CHANGE Q0 (tokens)": "Net supply change, 90 days",
    "NET SUPPLY CHANGE, annualised % of circulating": "Net supply change, % a year",
    "Burn execution": "Burn execution",
    "Burn status": "Burn status",
    "Documented share of fees burned": "Documented share of fees burned",
    "Gross burn Q0 ($ at avg price)": "Gross burn, 90 days ($)",
    "GROSS ISSUANCE Q0 (tokens)": "Gross issuance, 90 days (tokens)",
    "Pool release Q0 (tokens)": "Pool release, 90 days (tokens)",
    "Net mint Q0 (SELF-REPORTED by the protocol)": "Net mint, self-reported",
    "Self-reported figure preferred?": "Self-reported figure preferred?",
    "Derived net supply change": "Derived net supply change",
    "Self-reported − derived": "Self-reported minus derived",
    "Net supply change ($ at avg price)": "Net supply change ($)",
    "Cumulative burn ÷ max supply": "Cumulative burn ÷ max supply",
    "Burn split Q0 (tokens)": "Burn split (tokens)",
    "Burn split — by source": "Burn split by source",
    "One-off burn Q0..Y1": "One-off burns, 12 months (tokens)",
    "One-off burn events": "One-off burn events",
    "Treasury governance could redeploy ($, spot)": "Treasury can redeploy ($)",
    "Implied burn Q0 (tokens)": "Implied burn, 90 days (tokens)",
    "Actual − implied burn (tokens)": "Actual minus implied burn (tokens)",
    "Cross-check: Δ implied circulating supply": "Implied circulating change check",
}
# audit columns the reader shows elsewhere or by another route (P2a B): fixed-window changes (the selector drives
# every growth column) and the quarter columns (block (c) shows the four quarters)
_FIXED_WINDOW = re.compile(r"Δ(30d|3m|6m|9m)$")
_QUARTER_COL = re.compile(r"^(Fees ÷ issuance at|Rev ÷ emission at|Net supply change) Q[123]\b")


def short_header(audit_header: str) -> str:
    for k, v in E_SHORT.items():
        if audit_header.startswith(k):
            return v
    h = plain(re.split(r" = | — |\(", audit_header)[0]).strip().rstrip(",")
    h = h[:1].upper() + h[1:].lower() if h.isupper() else h
    return h[:MAX_HEADER].rstrip()


def _trajectory(A: Audit) -> dict:
    """{key "tab|project|id": {"vals": [Q0, Q-1, Q-2, Q-3], "row", "break", "modelled", "trend", "t12m", "t12m_prev"}}"""
    import sanity                                          # noqa: PLC0415
    if "Trajectory" not in A.wb.sheetnames:
        return {}
    ws = A.wb["Trajectory"]
    hdr = {ws.cell(4, c).value: c for c in range(1, ws.max_column + 1)}
    qc = [hdr.get(q) for q in ("Q0", "Q-1", "Q-2", "Q-3")]

    def num(r, c):
        if c is None:
            return None
        try:
            v = sanity.stated_number(A.X.value("Trajectory", r, c))
        except Exception:  # noqa: BLE001
            v = None
        return None if v is None else float(v)

    out = {}
    for r in range(5, ws.max_row + 1):
        k = ws.cell(r, 2).value
        if isinstance(k, str) and k.count("|") == 2 and None not in qc:
            out[k] = {"vals": [num(r, c) for c in qc], "row": r,
                      "break": str(ws.cell(r, hdr.get("Series break (basis change)", 0) or 1).value or "")
                      if hdr.get("Series break (basis change)") else "",
                      "modelled": str(ws.cell(r, hdr.get("MODELLED quarters (excluded from any scoring)", 0)).value or "")
                      if hdr.get("MODELLED quarters (excluded from any scoring)") else "",
                      "trend": str(ws.cell(r, hdr.get("Trend", 0)).value or "") if hdr.get("Trend") else "",
                      "t12m": num(r, hdr.get("T12M")), "t12m_prev": num(r, hdr.get("T12M a quarter earlier (Q-1..Q-4)"))}
    return out


def _quarter_change(t: dict | None, asof: pd.Timestamp) -> tuple[str, object, str]:
    """(kind, value, reason) for the change of the latest quarter on the one before — inheriting the state of what it
    compares (P2a A6): a series break inside the two quarters → blocked; a rolling-T12M series → the T12M change; a
    MODELLED quarter → the change, said MODELLED."""
    if not t:
        return "n/a", None, "No row on the Trajectory tab."
    vals, row = t["vals"], t["row"]
    where = f"Working: audit file, Trajectory row {row}."
    dates = [pd.Timestamp(d) for d in re.findall(r"\b(20\d\d-\d\d-\d\d)\b", t["break"])]
    inside = [d for d in dates if asof - pd.Timedelta(days=180) < d <= asof]
    if inside:
        return "blocked", None, (f"Blocked: the two quarters compared straddle a series break (a change of measuring "
                                 f"point) on {', '.join(str(d.date()) for d in inside)} — a change across it is not a "
                                 f"change of the business. {plain(t['break'])}\n\n{where}")
    if "rolling T12M" in t["trend"]:
        a, b = t["t12m"], t["t12m_prev"]
        if a is None or b is None or b == 0:
            return "n/a", None, ("A lumpy series, read on a trailing-12-month basis: no T12M change "
                                 "(a T12M figure is missing or the earlier one is 0).\n\n" + where)
        return "t12m", (a - b) / abs(b), (f"A lumpy series, so the change is read on a TRAILING-12-MONTH basis: T12M "
                                          f"{a:,.6g} against {b:,.6g} a quarter earlier.\n\n{where}")
    q0, q1 = vals[0], vals[1]
    if q0 is None or q1 is None or q1 == 0:
        return "n/a", None, ("No change: " + ("the quarter before is 0" if q1 == 0 else
                                             "a quarter has no value on the Trajectory tab") + f".\n\n{where}")
    mod = [q for q in ("Q0", "Q-1") if re.search(rf"(?<![\w-]){re.escape(q)}(?![\d])", t["modelled"])]
    note = (f" MODELLED: {', '.join(mod)} is modelled, not measured — the change is shown, not scored." if mod else "")
    return ("modelled" if mod else "ok"), (q0 - q1) / abs(q1), (f"Latest quarter {q0:,.6g} against {q1:,.6g}.{note}"
                                                                 f"\n\n{where}")


FILL_MODELLED = PatternFill("solid", fgColor="FFE699")


def _write_arch_tab(sh: Sheet, ctx: Ctx, spec: TabSpec) -> None:
    A, data, config = ctx.A, ctx.data, ctx.config
    ws, tab = sh.ws, spec.tab
    pmap = {p["name"]: p for p in ctx.projects}
    names = [n for n in A.rows(tab) if n in pmap] if tab in A.wb.sheetnames else []
    if not names:
        ws.cell(1, 1, f"{spec.title} — no project in scope").font = F_TITLE
        return
    audit_cols = A.cols(tab)
    used: set = set()
    cols: list = [Col("Project", "The project.", FMT_TOK, "", "", 0, 16), Col("Token", "The token's symbol.", FMT_TOK,
                                                                                  "", "", 0, 8)]
    how: list = [None, None]

    def add(col: Col, h):
        cols.append(col)
        how.append(h)

    def audit_col(prefix: str):
        c = A.col(tab, prefix)
        if c is not None:
            used.add(c)
        return c

    # (a) key and rank
    if spec.key[0] == "a1":
        add(Col(spec.key_header, ARCHETYPE_KEYS[1][1], FMT_PCT, "(a) Rank key", "computed from the stored "
                "trailing-365-day settlement volume and circulating supply", 0, 13), ("a1key",))
    else:
        kc = audit_col(spec.key[1])
        kh = A.header(tab, kc) if kc else spec.key[1]
        add(Col(spec.key_header, ARCHETYPE_KEYS[spec.arch][1], FMT_X, "(a) Rank key", f"{tab}, {kh}", 0, 12,
                audit=(tab, kh), family="buyback" if spec.arch == 3 else ""), ("audit", kc))
    add(Col(spec.rank_header, "Position on this tab by the rank key, higher first. No measured key reads 'unranked' and "
            "sorts last; a parked project reads 'parked' and is outside the ranks.", FMT_RANK, "(a) Rank key",
            "computed", 0, 9), ("rank",))
    for extra in spec.a_extra:
        if extra[0] == "dilution":
            dc, dn = audit_col(extra[1]), audit_col(extra[2])
            add(Col("Retired ÷ (emissions + unlocks)", "The same retirement against emissions AND unlocks (at the rate "
                    "in force where one is declared) — the unlock-inclusive figure beside the key.", FMT_X,
                    "(a) Rank key", f"{tab}, RETIREMENT ÷ DILUTION (at the rate in force where declared)", 0, 13,
                    audit=(tab, extra[1]), family="buyback"), ("dilution", dc, dn))
        elif extra[0] == "interim":
            ic = audit_col(extra[1])
            ih = A.header(tab, ic) if ic else extra[1]
            add(Col("Fees ÷ issuance, Q0 (interim)", "INTERIM, not the rank key: fees ÷ gross issuance over the last 90 "
                    "days, shown beside the rank key while settlement per token, 1-year change forms for one project at "
                    "most. The rank key changes only on Jake's go-ahead.", FMT_X, "(a) Rank key", f"{tab}, {ih}", 0, 13,
                    audit=(tab, ih)), ("interim", ic))
    if spec.arch == 4:
        audit_col("Release ÷ burn (x)")              # its reciprocal IS the key for a pre-minted supply (header comment)
    # (b) valuation
    for sheet, prefix, short, fam in spec.valuation:
        h = _src_header(A, sheet or tab, prefix)
        if sheet is None:
            c = audit_col(prefix)
            if c is None:
                continue
            fmt = FMT_PRICE if prefix.startswith("Price") else reader_format(A.wb[tab].cell(5, c).number_format)
            add(Col(short, h, fmt, "(b) Valuation", f"{tab}, {h}", 0, 12, audit=(tab, h), family=fam), ("audit", c))
        else:
            c = A.col(sheet, prefix)
            if c is None:
                continue
            fmt = reader_format(A.wb[sheet].cell(5, c).number_format)
            add(Col(short, h, fmt, "(b) Valuation", f"{sheet}, {h}", 0, 12, audit=(sheet, h), family=fam),
                ("other", (sheet, prefix)))
    # (c) the last four quarters and the change
    tid, stem, tfmt, tdef = spec.traj
    labels = ("latest quarter", "1 quarter back", "2 quarters back", "3 quarters back")
    for i, lab in enumerate(labels):
        add(Col(f"{stem}, {lab}"[:MAX_HEADER], f"{tdef} for the 90 days ending {i * 90} days before the as-of date "
                "(Trajectory tab, the same formula on each shifted quarter).", tfmt, "(c) Trajectory",
                f"Trajectory tab, {tid}", 0, 11), ("traj", tid, i))
    add(Col("Change vs previous quarter", "(latest quarter − the quarter before) ÷ |the quarter before|; ▲ up from "
            "+1%, ► within ±1%, ▼ down. A change across a series break reads 'blocked'; a lumpy series is compared "
            "on a trailing-12-month basis; a MODELLED quarter is said so.", FMT_PCT, "(c) Trajectory",
            "computed from the Trajectory tab", 0, 11), ("trend", tid))
    # (d) drivers
    for d in spec.drivers:
        if d[0] == "@growth":
            add(Col(d[2][:MAX_HEADER], f"{d[2]}: the change of the stored daily series over the period chosen in "
                    "Summary!B2 (custom: the days in Summary!D2). Flows: the period's sum against the same number of "
                    "days before; stocks: the last value against the value that many days earlier.", FMT_PCT,
                    "(d) Drivers", "computed from the stored daily series", 0, 12), ("growth", d[1], d[3]))
            continue
        if d[0] == "@issuance":
            ci, cr = audit_col("GROSS ISSUANCE Q0 (tokens)"), audit_col("Pool release Q0 (tokens)")
            add(Col(d[2], "GROSS ISSUANCE over the last 90 days, or — for a pre-minted supply — the POOL RELEASE "
                    "(released from pre-minted pools, not minted), on the basis the key uses.", FMT_TOK, "(d) Drivers",
                    f"{tab}, GROSS ISSUANCE Q0 or Pool release Q0", 0, 13,
                    audit=(tab, A.header(tab, ci)) if ci else None), ("issuance", ci, cr))
            continue
        sheet, prefix, short, fam = d
        if sheet is None:
            c = audit_col(prefix)
            if c is None:
                continue
            h = A.header(tab, c)
            add(Col(short, h, reader_format(A.wb[tab].cell(5, c).number_format), "(d) Drivers", f"{tab}, {h}", 0, 13,
                    audit=(tab, h), family=fam), ("audit", c))
        else:
            c = A.col(sheet, prefix)
            if c is None:
                continue
            h = A.header(sheet, c)
            add(Col(short, h, reader_format(A.wb[sheet].cell(5, c).number_format), "(d) Drivers", f"{sheet}, {h}", 0,
                    13, audit=(sheet, h), family=fam), ("other", (sheet, prefix)))
    # (e) everything else: audit order, collapsed, in four groups by axis
    note_c, flag_c = A.col(tab, "Notes"), A.col(tab, "Data flags")
    used |= {1, 2, note_c, flag_c}
    traj_of = {}
    for h, c in audit_cols.items():
        if h.startswith("Trend — "):
            m = re.match(r"Trend — (\S+)", h)
            traj_of[c] = m.group(1) if m else ""
    axis_label = {"CAPTURE": "(e) More: capture", "ALIGNMENT": "(e) More: supply and dilution",
                  "BUSINESS": "(e) More: business", "CONTEXT": "(e) More: context"}
    buckets: dict = {k: [] for k in axis_label}
    for h, c in sorted(audit_cols.items(), key=lambda kv: kv[1]):
        if (c in used or h.startswith("Detail — ") or c in traj_of or _FIXED_WINDOW.search(h)
                or _QUARTER_COL.search(h)):
            continue
        axis = str(A.wb[tab].cell(3, c).value or "CONTEXT")
        axis = axis if axis in axis_label else "CONTEXT"
        lab = axis_label[axis]
        fmt = reader_format(A.wb[tab].cell(5, c).number_format)
        is_text = A.wb[tab].cell(5, c).number_format == "@" or "status" in h.lower() or h.endswith("?") \
            or h.startswith(("Which ", "Issuance basis", "Fee per transaction — what", "Settlement volume: last date",
                             "  Emissions basis", "  Customer revenue — not measured", "Burn split — by source",
                             "One-off burn events", "Burn execution", "Q0 WINDOW CAVEAT", "Validator yield — caveat"))
        buckets[axis].append((Col(short_header(h), h, TEXT_FMT if is_text else fmt, lab, f"{tab}, {h}", 1, 13,
                                  audit=(tab, h), is_text=is_text), ("text" if is_text else "audit", c)))
        nxt = c + 1
        if nxt in traj_of:
            tid2 = traj_of[nxt]
            buckets[axis].append((Col(TREND_SHORT.get(tid2, "Quarter change")[:MAX_HEADER], f"The quarter-on-quarter "
                                      f"change of the column to the left: (latest quarter − the quarter before) ÷ "
                                      f"|the quarter before| (Trajectory, {tid2}); ▲ up from +1%, ► within ±1%, ▼ "
                                      "down. A change across a series break reads 'blocked'.", FMT_PCT, lab,
                                      f"Trajectory tab, {tid2}", 1, 11, audit=(tab, A.header(tab, nxt))),
                                  ("trend", tid2)))
    for axis in axis_label:
        for col, h in buckets[axis]:
            add(col, h)
    sh.cols = cols
    ws.cell(1, 1, spec.title).font = F_TITLE
    ws.cell(2, 1, "Blocks: (a) rank key · (b) valuation · (c) the last four quarters · (d) drivers · (e) more, "
            "collapsed (click + above the columns). Growth columns follow the period chosen on Summary (B2). "
            "Hover a header for its definition, a cell for its source.").font = F_SUB
    write_headers(sh)
    traj = _trajectory(A)
    keyvals = {}
    for n in names:
        if parked_reason(n):
            continue
        if spec.key[0] == "a1":
            keyvals[n] = settlement_key(ctx.groups, n)[0]
        else:
            kc = how[2][1]
            r = A.row(tab, n)
            keyvals[n] = A.at(tab, r, kc).num if (r and kc) else None
    order = _rank_rows(keyvals) + [(n, None, None) for n in sorted(n for n in names if parked_reason(n))]
    trend_cols, period_cols = [], []
    ranked = sum(1 for _n, rk, _v in order if rk is not None)
    for i, (name, rank, kv) in enumerate(order):
        row = 5 + i
        p = pmap[name]
        parked = bool(parked_reason(name))
        sh.value_rows.append(row)
        r_a = A.row(tab, name)
        notes = [audit_notes(A.at(tab, r_a, c).text) for c in (note_c, flag_c) if c]
        if parked:
            notes.insert(0, parked_reason(name))
        put_text(sh, row, cols[0], name, "\n\n".join(n for n in notes if n), bold=True)
        put_text(sh, row, cols[1], p.get("symbol") or "—")
        for j, (col, h) in enumerate(zip(cols, how)):
            if h is None:
                continue
            kind = h[0]
            if kind == "audit":
                src = A.at(tab, r_a, h[1]) if h[1] else None
                extra = _premint_note(p, src) if src is not None and "NET SUPPLY CHANGE" in src.header.upper() else ""
                cell = put_src(sh, data, row, col, p, src, extra=extra)
                ctx.record(src, tab, cell, name)
            elif kind == "interim":
                src = A.at(tab, r_a, h[1]) if h[1] else None
                cell = put_src(sh, data, row, col, p, src, extra=(
                    f"INTERIM — not the rank key and not ranked. {ranked} of {len(names)} projects on this tab hold the "
                    f"rank key ({spec.key_header.lower()}) on this build; this column is shown until more do. The rank "
                    "key changes only on Jake's go-ahead."))
                ctx.record(src, tab, cell, name)
            elif kind == "a1key":
                v, why = settlement_key(ctx.groups, name)
                if v is not None:
                    put_number(sh, data, row, col, name, "a1key", v, None, f"{ARCHETYPE_KEYS[1][1]}\n\n{why}.\n\n"
                               f"Working: audit file, Data rows {name}|settlement_volume_365d_usd and "
                               f"{name}|circulating_supply_used.")
                else:
                    put_code(sh, row, col, "n/a", f"No measured key: {why}.")
            elif kind == "inverse":
                src = A.at(tab, r_a, h[1])
                if src.num is not None and src.num != 0:
                    put_number(sh, data, row, col, name, f"inv|{tab}|{src.coord}", 1.0 / src.num, src.fill,
                               f"1 ÷ release ÷ burn {src.num:,.6g} — the reciprocal, so higher is better.\n\nWorking: "
                               f"{src.where}.")
                else:
                    put_src(sh, data, row, col, p, src, "No burn ÷ release: " + ("the release ÷ burn is 0"
                                                                                 if src.num == 0 else "not formed"))
            elif kind == "text":
                src = A.at(tab, r_a, h[1])
                put_text(sh, row, col, plain(src.text)[:120] if src.text else "—",
                         _code_reason(src) if src.text and len(plain(src.text)) > 120
                         else _with_lead(f"Working: {src.where}.", col, name))
            elif kind == "rank":
                if parked:
                    put_code(sh, row, col, "parked", parked_reason(name))
                elif rank is not None:
                    put_number(sh, data, row, col, name, f"rank|{tab}", rank, None,
                               f"{rank} of {ranked} ranked on this tab by {spec.key_header.lower()}.\n\nWorking: "
                               "computed here from the rank key column.")
                else:
                    put_code(sh, row, col, "unranked", f"Unranked, sorted last: no measured {spec.key_header.lower()} "
                             "— the key cell's comment says why.")
            elif kind == "dilution":
                use = h[2] if (name in config.CURRENT_EMISSION_RATE and h[2]) else h[1]
                src = A.at(tab, r_a, use) if use else None
                ctx.record(src, tab, put_src(sh, data, row, col, p, src), name)
            elif kind == "other":
                src = A.cell(h[1][0], h[1][1], name)
                cell = put_src(sh, data, row, col, p, src,
                               extra=_payout_note(p, src, A.cell(CP_TAB, "Payout ratio — gross", name))
                               if col.header == "Payout ratio" else "")
                ctx.record(src, tab, cell, name)
            elif kind == "issuance":
                basis = config.issuance_basis(name)
                c_use = h[2] if basis == "pool_release_tokens" and h[2] else h[1]
                src = A.at(tab, r_a, c_use) if c_use else None
                put_src(sh, data, row, col, p, src, extra=("A pre-minted supply: the pool release is shown (released "
                                                           "from pre-minted pools, not minted)."
                                                           if c_use == h[2] else ""))
            elif kind == "growth":
                metric = {"@fees": config.fees_headline_metric(name), "@revenue": _revenue_metric(config, name),
                          "@burn": config.a4_burn_metric(name), "@issuance": config.issuance_basis(name)
                          }.get(h[1], h[1])
                _period_formula_cell(sh, ctx, row, col, name, metric, h[2], col.header.split(",")[0].lower())
                period_cols.append(j + 1)
            elif kind == "traj":
                t = traj.get(f"{tab[:2]}|{name}|{h[1]}")
                v = t["vals"][h[2]] if t else None
                if v is not None:
                    q = ("Q0", "Q-1", "Q-2", "Q-3")[h[2]]
                    mod = bool(t and re.search(rf"(?<![\w-]){re.escape(q)}(?![\d])", t["modelled"]))
                    put_number(sh, data, row, col, name, f"traj|{tab}|{h[1]}|{h[2]}", v,
                               FILL_MODELLED if mod else None,
                               ("MODELLED, not measured.\n\n" if mod else "")
                               + f"Working: audit file, Trajectory row {t['row']}.")
                else:
                    put_code(sh, row, col, "n/a", "No value for this quarter on the Trajectory tab"
                             + (f" (row {t['row']}): " + plain(t["break"] or "a quarter not held") if t else
                                " (no row).") + ".")
            elif kind == "trend":
                t = traj.get(f"{tab[:2]}|{name}|{h[1]}")
                k2, v, why = _quarter_change(t, ctx.asof)
                if k2 in ("ok", "modelled", "t12m"):
                    put_number(sh, data, row, col, name, f"qchg|{tab}|{h[1]}", v,
                               FILL_MODELLED if k2 == "modelled" else None, why)
                    trend_cols.append(j + 1)
                else:
                    put_code(sh, row, col, k2, why)
        if parked:
            grey_row(sh, row)
    last = 4 + len(order)
    for c in sorted(set(trend_cols)):
        L = get_column_letter(c)
        icon_arrows(ws, f"{L}5:{L}{last}")
    for c in sorted(set(period_cols)):
        L = get_column_letter(c)
        code_style(ws, f"{L}5:{L}{last}")
    ws.sheet_properties.outlinePr.summaryRight = False


TREND_SHORT = {
    "a3_circ_retirement": "Quarter change: retirement rate",
    "a3_fdv_retirement": "Quarter change: FDV retirement",
    "a3_implied_buyback_pct": "Quarter change: implied buyback",
    "a3_retired_tokens": "Quarter change: retired tokens",
    "a3_actual_buyback_pct": "Quarter change: retired %",
    "a3_net_absorption": "Quarter change: net absorption",
    "a3_retire_vs_dilution": "Quarter change: retired ÷ dilution",
    "a3_retire_vs_dilution_now": "Quarter change: at rate in force",
    "a3_protocol_yield": "Quarter change: staking yield",
    "a3_buyback_locked": "Quarter change: buyback locked",
}


def _src_header(A: Audit, sheet: str, prefix: str) -> str:
    c = A.col(sheet, prefix)
    return A.header(sheet, c) if c else prefix


def _a3_key_value(A: Audit, tab, key_c, name):
    if key_c is None:
        return None
    r = A.row(tab, name)
    return A.at(tab, r, key_c).num if r else None


# ---------------------------------------------------------------------------------------------- Comparison
COMPARISON_COLS = (
    # (key, header, source (sheet, prefix) | "@...", format, definition, family)
    ("hvy", "Holder value yield", (CP_TAB, "HOLDER VALUE YIELD"), FMT_PCT,
     "Revenue-funded value to holders a year (retired + distributed + liquid voter bribes) ÷ market cap.", "buyback"),
    ("dil", "Total dilution, % of float", (CP_TAB, "TOTAL DILUTION = outflows a+b+c+d+e (annualised) ÷ tradeable float"),
     FMT_PCT, "Every outflow a year (security emissions, incentives, grants, unlocks) ÷ the tradeable float.", ""),
    ("nhy", "Net holder yield", (CP_TAB, "NET HOLDER YIELD"), FMT_PCT,
     "Holder value yield − every outflow a year at each day's price ÷ market cap. The sort key of this tab.", ""),
    ("payout", "Payout ratio", (CP_TAB, "Payout ratio — the project's own net basis"), FMT_PCT,
     "Credited holder value ÷ revenue on the project's own revenue basis.", "buyback"),
    ("subsidy", "Subsidy cover (x)", (CP_TAB, "SUBSIDY COVER"), FMT_X,
     "Revenue ÷ the $ of supply-side incentives, per covered day: above 1.0x the business pays for its incentives.", ""),
    ("leak", "Leakage (ii)", (CP_TAB, "LEAKAGE (ii)"), FMT_PCT,
     "Incentive spend not returned as revenue = max(0, 1 − subsidy cover). A proposed definition.", ""),
    ("fees_fdv", "Fees ÷ FDV, a year", ("Master", "Fees ÷ FDV (annualised)"), FMT_PCT,
     "Fees over the last 90 days, annualised, ÷ fully diluted value.", ""),
    ("rev_mcap", "Revenue ÷ market cap, a year", "@rev_mcap", FMT_PCT,
     "Protocol revenue, annualised (Summary's figure) ÷ market cap.", ""),
    ("lock", "Lock rate, % of circulating", (CP_TAB, "Holding demand — lock rate"), FMT_PCT,
     "Tokens staked or locked inside the circulating supply ÷ circulating supply.", "staking"),
    ("stake_tbill", "Staking yield minus T-bill", "@stake_tbill", FMT_PCT,
     "The staker's yield (the validator yield for an A1 chain, the protocol's token yield otherwise) minus the "
     "3-month T-bill.", "staking"),
    ("insider", "Insider share of stake", (CP_TAB, "Insider share of staked / locked supply"), FMT_PCT,
     "Insider-held tokens' share of the staked or locked supply.", ""),
)


def _write_comparison(sh: Sheet, ctx: Ctx, keys: dict) -> None:
    A, data = ctx.A, ctx.data
    ws = sh.ws
    cols = [Col("Project", "The project.", FMT_TOK, "", "", 0, 16), Col("Token", "The token's symbol.", FMT_TOK, "",
                                                                         "", 0, 8),
            Col("Primary archetype", "The archetype the project is ranked under on Summary.", TEXT_FMT, "", "", 0, 18,
                is_text=True)]
    spec = {}
    for k, h, src, fmt, d, fam in COMPARISON_COLS:
        col = Col(h, d, fmt, "Holders" if k in ("hvy", "dil", "nhy", "payout") else "Subsidy and value"
                  if k in ("subsidy", "leak", "fees_fdv", "rev_mcap") else "Holding", (
                      f"{src[0]}, {_src_header(A, src[0], src[1])}" if isinstance(src, tuple) else "computed"), 0, 12,
                  audit=(src[0], _src_header(A, src[0], src[1])) if isinstance(src, tuple) else None, family=fam)
        cols.append(col)
        spec[k] = (col, src)
    sh.cols = cols
    ws.cell(1, 1, "Comparison — every project on one row, on figures that mean the same for each, sorted by net holder "
            "yield").font = F_TITLE
    ws.cell(2, 1, "Only like-for-like figures. Parked projects are shown greyed at the bottom. Hover a header for its "
            "definition, a cell for its source.").font = F_SUB
    write_headers(sh)
    nhy = {p["name"]: (A.cell(CP_TAB, "NET HOLDER YIELD", p["name"]) or Src("", "")).num for p in ctx.projects}
    live = [p for p in ctx.projects if not parked_reason(p["name"])]
    order = sorted([p for p in live if nhy[p["name"]] is not None], key=lambda p: -nhy[p["name"]]) + \
        sorted([p for p in live if nhy[p["name"]] is None], key=lambda p: p["name"]) + \
        sorted([p for p in ctx.projects if parked_reason(p["name"])], key=lambda p: p["name"])
    for i, p in enumerate(order):
        row, name = 5 + i, p["name"]
        sh.value_rows.append(row)
        put_text(sh, row, cols[0], name, parked_reason(name), bold=True)
        put_text(sh, row, cols[1], p.get("symbol") or "—")
        arch = _primary(p)
        put_text(sh, row, cols[2], ARCHETYPE_NAMES.get(arch, "none"))
        for k, (col, src) in spec.items():
            if isinstance(src, tuple):
                s = A.cell(src[0], src[1], name)
                extra = _payout_note(p, s, A.cell(CP_TAB, "Payout ratio — gross", name)) if k == "payout" else ""
                cell = put_src(sh, data, row, col, p, s, extra=extra)
                ctx.record(s, COMPARISON, cell, name)
            elif src == "@rev_mcap":
                rv, mc = data.value(name, "rev_ann"), A.cell("Master", "Market cap ($)", name)
                if isinstance(rv, (int, float)) and mc is not None and mc.num:
                    rr, rm = data.ref(name, "rev_ann"), data.put(name, "mcap", mc.num)
                    put_formula(sh, row, col, f"={rr}/{rm}", None, f"Revenue, annualised {rv:,.0f} ÷ market cap "
                                f"{mc.num:,.0f}.\n\nWorking: Summary's revenue (audit file, Master, Revenue Q0 and its "
                                f"Data coverage); {mc.where}.")
                else:
                    put_code(sh, row, col, "n/a", "No annualised revenue or no market cap for this project (see "
                             "Summary).")
            elif src == "@stake_tbill":
                head = ("Spread over the 3-month T-bill: Staker yield — validator" if arch == 1 else
                        "Spread over the 3-month T-bill: Staker yield — protocol, tokens")
                s = A.cell(CP_TAB, head, name)
                put_src(sh, data, row, col, p, s, extra="The validator yield's spread (an A1 chain)." if arch == 1
                        else "The protocol's token staking yield's spread.")
        if parked_reason(name):
            grey_row(sh, row)


# ---------------------------------------------------------------------------------------------- Business
def _write_business(sh: Sheet, ctx: Ctx) -> None:
    A, data, config = ctx.A, ctx.data, ctx.config
    ws = sh.ws
    C = {k: Col(*v) for k, v in {
        "project": ("Project", "The project.", FMT_TOK, "", "", 0, 16),
        "arch": ("Primary archetype", "The archetype the project is ranked under on Summary.", TEXT_FMT, "", "", 0, 16),
        "fees": ("Fees, last 90 days", "Fees over the last 90 days (Master's fees column).", FMT_USD, "Income",
                 "Master, Fees Q0 ($)", 0, 12),
        "rev": ("Protocol revenue, last 90 days", "Protocol revenue over the last 90 days (the series Summary's "
                "protocol revenue reads; its comment names it).", FMT_USD,
                "Income", "Master, Revenue Q0 ($)", 0, 12),
        "rev_ann": (REV_HEADER, "Protocol revenue over the last 90 days × 365 ÷ the days the series covers (as on "
                    "Summary).", FMT_USD, "Income", "Master, Revenue Q0 and Data coverage", 0, 12),
        "growth": ("Revenue growth, selected period", "Revenue over the selected period ÷ the same number of days before "
                   "− 1, every day held.", FMT_PCT, "Income", "computed from the stored daily revenue series", 0, 12),
        "tvl": ("TVL", "Chain TVL for an A1 chain, protocol TVL otherwise (Master's TVL column).", FMT_USD, "Assets",
                "Master, TVL ($)", 0, 12),
        "tvl_g": ("TVL change, selected period", "TVL on the last day held ÷ TVL the selected number of days before − "
                  "1, nominal (US dollars).", FMT_PCT, "Assets", "computed from the stored daily TVL series", 0, 12),
        "tvl_cp": ("TVL change at constant prices", "For a protocol: the change of its TVL at CONSTANT PRICES (each "
                   "token's quantity at fixed prices, DefiLlama's token quantities), so a rising token price cannot "
                   "pass for growth. For a chain: 'by design' — chains stay in US dollars and the price change is the "
                   "price-effect line.", FMT_PCT, "Assets", "computed from the stored constant-price TVL series", 0,
                   13),
        "price": ("Price change, selected period", "The token's price change over the selected period — for a chain, the "
                  "price-effect line beside its US-dollar figures.", FMT_PCT, "Assets",
                  "computed from the stored daily price series", 0, 12),
        "unit": ("Unit of demand", "The project's own unit of demand (transactions, settlement, supply units, loans) — "
                 "the series its CP 'BUSINESS 1' slot reads.", TEXT_FMT, "Demand", "CP Comparability, Business 1 — "
                 "series, unit, source", 0, 28),
        "unit_v": ("Unit of demand, value", "The unit of demand: a stock now, or a flow over the last 90 days "
                   "annualised over its covered days.", FMT_TOK, "Demand", "CP Comparability, BUSINESS 1", 0, 13),
        "rev_unit": ("Revenue per unit of demand", "Revenue, annualised ÷ the unit of demand's value.", FMT_PRICE,
                     "Demand", "computed: Business revenue annualised ÷ CP BUSINESS 1", 0, 12),
    }.items()}
    C["arch"].is_text = C["unit"].is_text = True
    C["fees"].audit, C["rev"].audit = ("Master", "Fees Q0 ($)"), ("Master", "Revenue Q0 ($)")
    sh.cols = list(C.values())
    ws.cell(1, 1, "Business — what each project earns and the demand behind it").font = F_TITLE
    ws.cell(2, 1, "Growth columns follow the period chosen on Summary (B2). Protocols' TVL is also shown at constant "
            "prices; chains stay in US dollars beside their price change.").font = F_SUB
    write_headers(sh)
    order = sorted(ctx.projects, key=lambda p: (bool(parked_reason(p["name"])), _primary(p), p["name"]))
    period_cols = [C["growth"], C["tvl_g"], C["tvl_cp"], C["price"]]
    for i, p in enumerate(order):
        row, name = 5 + i, p["name"]
        sh.value_rows.append(row)
        put_text(sh, row, C["project"], name, parked_reason(name), bold=True)
        put_text(sh, row, C["arch"], ARCHETYPE_NAMES.get(_primary(p), "none"))
        for k, head in (("fees", "Fees Q0 ($)"), ("rev", "Revenue Q0 ($)"), ("tvl", "TVL ($)")):
            s = A.cell("Master", head, name)
            cell = put_src(sh, data, row, C[k], p, s)
            ctx.record(s, BUSINESS, cell, name)
        rv = data.value(name, "rev_ann")
        if isinstance(rv, (int, float)):
            put_formula(sh, row, C["rev_ann"], f"={data.ref(name, 'rev_ann')}", None,
                        f"Revenue over the last 90 days × 365 ÷ covered days = {rv:,.0f} (as on Summary).\n\nWorking: "
                        "audit file, Master, Revenue Q0 ($) and its Data coverage.")
        else:
            s = A.cell("Master", "Revenue Q0 ($)", name)
            put_src(sh, data, row, C["rev_ann"], p, s, "No annualised revenue: the 90-day revenue is not a number")
        _period_formula_cell(sh, ctx, row, C["growth"], name, _revenue_metric(config, name), "flow", "revenue")
        _period_formula_cell(sh, ctx, row, C["tvl_g"], name, _tvl_metric(p), "stock", "TVL")
        if _primary(p) == 1:
            put_code(sh, row, C["tvl_cp"], "by design", "A chain: its figures stay in US dollars (no constant-price "
                     "basis is built for a chain's TVL); the price change beside it is the price-effect line.")
        elif "protocol_tvl_constant_price_usd" in config.metrics_for_project(p):
            _period_formula_cell(sh, ctx, row, C["tvl_cp"], name, "protocol_tvl_constant_price_usd", "stock",
                                 "TVL at constant prices")
        else:
            put_code(sh, row, C["tvl_cp"], "not wired", f"No constant-price TVL series is wired for {name}.")
        _period_formula_cell(sh, ctx, row, C["price"], name, "price_usd", "stock", "price")
        ut = A.cell(CP_TAB, "Business 1 — series, unit, source", name)
        put_text(sh, row, C["unit"], _unit_text(config, ut.text) if ut is not None and ut.text else "—",
                 ((plain(ut.text) or f"No unit of demand is declared for {name} on the audit file's CP tab "
                   "(Business 1).") + f"\n\nWorking: {ut.where}.") if ut is not None else "No CP row.")
        uv = A.cell(CP_TAB, "BUSINESS 1 — core usage", name)
        put_src(sh, data, row, C["unit_v"], p, uv)
        if isinstance(rv, (int, float)) and uv is not None and uv.num:
            ru = data.put(name, "unit_v", uv.num)
            put_formula(sh, row, C["rev_unit"], f"={data.ref(name, 'rev_ann')}/{ru}", uv.fill,
                        f"Revenue, annualised {rv:,.0f} ÷ {uv.num:,.6g}.\n\nWorking: {uv.where}; Master, Revenue Q0.")
        else:
            put_code(sh, row, C["rev_unit"], "n/a", "No revenue per unit: " + (
                "no annualised revenue" if not isinstance(rv, (int, float)) else "no unit-of-demand value") + ".")
        if parked_reason(name):
            grey_row(sh, row)
    last = 4 + len(order)
    for col in period_cols:
        code_style(ws, f"{col.letter}5:{col.letter}{last}")


def _unit_text(config, text: str) -> str:
    """The unit of demand in plain words: the series' own label and its unit, never the internal key."""
    m = re.match(r"\s*([a-z0-9_]+)\s*\(([^)]*)\)", str(text or ""))
    if not m:
        return plain(text)[:60]
    label = str((config.METRICS.get(m.group(1)) or {}).get("label") or m.group(1).replace("_", " "))
    label = re.sub(r" per day$", "", re.split(r" — | \(|: | = ", label)[0].strip())
    return f"{label[:1].upper()}{label[1:]} ({m.group(2).split(';')[0].strip()})"[:60]


# ---------------------------------------------------------------------------------------------- Charts
def _write_charts(ws, ctx: Ctx, keys: dict) -> dict:
    """Four native charts on `_data` (P2a B): the key against 1.0x; holder value yield against total dilution; revenue
    against market cap (log–log); net supply change sorted. Parked projects are left out and listed under each."""
    from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
    from openpyxl.chart.label import DataLabelList
    A, data = ctx.A, ctx.data
    dws = data.ws
    start = max(len(data.rows) + 4, 60)
    notes = {}
    ws.cell(1, 1, "Charts — every chart reads the hidden _data sheet and recalculates with the file").font = F_TITLE
    ws.cell(2, 1, "Parked projects are left out of every chart and listed under it. Values are as of "
            f"{ctx.asof.date()}.").font = F_SUB
    parked = sorted(p["name"] for p in ctx.projects if parked_reason(p["name"]))
    live = [p for p in ctx.projects if not parked_reason(p["name"])]

    def table(r0: int, title: str, heads: list, rows: list) -> int:
        dws.cell(r0, 1, title)
        for j, h in enumerate(heads):
            dws.cell(r0 + 1, 1 + j, h)
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                dws.cell(r0 + 2 + i, 1 + j, v)
        return r0 + 2

    # 1. the key per project against 1.0x, grouped by archetype (A2-A4: ratios; A1's key is a change, its own panel)
    rows1, out1 = [], []
    for arch in (2, 3, 4):
        for p in sorted((q for q in live if keys[q["name"]][0] == arch), key=lambda q: q["name"]):
            kv = keys[p["name"]][1]
            if kv is None:
                out1.append(f"{p['name']} (no measured key)")
                continue
            rows1.append((f"A{arch} · {p['name']}", kv, 1.0))
    r1 = table(start, "chart 1: archetype key", ["project", "key (x)", "crossover"], rows1)
    a1_rows = [(p["name"], keys[p["name"]][1]) for p in live if keys[p["name"]][0] == 1 and keys[p["name"]][1] is not None]
    out1 += [f"{p['name']} (A1: its key is a 1-year change, not a ratio — "
             + ("no measured key" if keys[p['name']][1] is None else "shown in panel 1b") + ")"
             for p in live if keys[p["name"]][0] == 1]
    bar = BarChart()
    bar.type, bar.title = "col", "1. Archetype key against the 1.0x crossover (A2–A4)"
    bar.y_axis.title, bar.x_axis.title = "key (x)", "project, grouped by archetype"
    if rows1:
        bar.add_data(Reference(dws, min_col=2, min_row=r1 - 1, max_row=r1 + len(rows1) - 1), titles_from_data=True)
        bar.set_categories(Reference(dws, min_col=1, min_row=r1, max_row=r1 + len(rows1) - 1))
        line = LineChart()
        line.add_data(Reference(dws, min_col=3, min_row=r1 - 1, max_row=r1 + len(rows1) - 1), titles_from_data=True)
        bar += line
    bar.width, bar.height = 26, 12
    ws.add_chart(bar, "A4")
    notes["1"] = out1 + [f"{n} (parked)" for n in parked]
    ws.cell(29, 1, "Not shown: " + ("; ".join(notes["1"]) or "none")).font = F_SUB
    # 1b. A1's key
    r1b = table(r1 + len(rows1) + 3, "chart 1b: A1 key", ["project", "settlement per token, 1-year change"], a1_rows)
    if a1_rows:
        b1 = BarChart()
        b1.type, b1.title = "col", "1b. A1 key: settlement per token, 1-year change"
        b1.add_data(Reference(dws, min_col=2, min_row=r1b - 1, max_row=r1b + len(a1_rows) - 1), titles_from_data=True)
        b1.set_categories(Reference(dws, min_col=1, min_row=r1b, max_row=r1b + len(a1_rows) - 1))
        b1.width, b1.height = 14, 8
        ws.add_chart(b1, "R4")
    else:
        ws.cell(4, 18, "1b. A1 key: no A1 project has two years of settlement history yet, so there is nothing to "
                "plot.").font = F_SUB
    # 2. holder value yield against total dilution (÷ circulating, the basis net holder yield uses), with y = x
    rows2, out2 = [], []
    for p in live:
        h = A.cell(CP_TAB, "HOLDER VALUE YIELD", p["name"])
        d = A.cell(CP_TAB, "TOTAL DILUTION = outflows a+b+c+d+e (annualised) ÷ circulating", p["name"])
        if h is None or d is None or h.num is None or d.num is None:
            out2.append(p["name"])
            continue
        rows2.append((p["name"], d.num, h.num))
    r2 = table(r1b + len(a1_rows) + 3, "chart 2: holder value yield vs total dilution",
               ["project", "total dilution", "holder value yield"], rows2)
    sc = ScatterChart()
    sc.title, sc.style = "2. Holder value yield against total dilution (break-even: net holder yield = 0)", 13
    sc.x_axis.title, sc.y_axis.title = "total dilution, % of circulating a year", "holder value yield"
    for i in range(len(rows2)):
        s = Series(Reference(dws, min_col=3, min_row=r2 + i), Reference(dws, min_col=2, min_row=r2 + i),
                   title=rows2[i][0])
        s.marker.symbol, s.graphicalProperties.line.noFill = "circle", True
        s.dLbls = _name_labels(DataLabelList)
        sc.series.append(s)
    mx = max([0.01] + [max(abs(x[1]), abs(x[2])) for x in rows2])
    rd = r2 + len(rows2) + 1
    dws.cell(rd - 1, 1, "break-even")
    dws.cell(rd, 2, 0.0)
    dws.cell(rd, 3, 0.0)
    dws.cell(rd + 1, 2, mx)
    dws.cell(rd + 1, 3, mx)
    diag = Series(Reference(dws, min_col=3, min_row=rd, max_row=rd + 1), Reference(dws, min_col=2, min_row=rd,
                                                                                    max_row=rd + 1),
                  title="break-even (holder value = dilution)")
    sc.series.append(diag)
    sc.width, sc.height = 26, 14
    ws.add_chart(sc, "A32")
    notes["2"] = [f"{n} (no holder value yield or no total dilution)" for n in out2] + [f"{n} (parked)" for n in parked]
    ws.cell(61, 1, "Not shown: " + ("; ".join(notes["2"]) or "none")).font = F_SUB
    # 3. annualised revenue against market cap, log–log, labelled
    rows3, out3 = [], []
    for p in live:
        rv, mc = data.value(p["name"], "rev_ann"), A.cell("Master", "Market cap ($)", p["name"])
        if not isinstance(rv, (int, float)) or rv <= 0 or mc is None or not mc.num or mc.num <= 0:
            out3.append(p["name"])
            continue
        rows3.append((p["name"], mc.num, rv))
    r3 = table(rd + 4, "chart 3: revenue vs market cap", ["project", "market cap", "revenue, annualised"], rows3)
    sc3 = ScatterChart()
    sc3.title, sc3.style = "3. Annualised revenue against market cap (log–log)", 13
    sc3.x_axis.title, sc3.y_axis.title = "market cap ($, log)", "revenue a year ($, log)"
    sc3.x_axis.scaling.logBase = 10
    sc3.y_axis.scaling.logBase = 10
    for i in range(len(rows3)):
        s = Series(Reference(dws, min_col=3, min_row=r3 + i), Reference(dws, min_col=2, min_row=r3 + i),
                   title=rows3[i][0])
        s.marker.symbol, s.graphicalProperties.line.noFill = "circle", True
        s.dLbls = _name_labels(DataLabelList)
        sc3.series.append(s)
    sc3.width, sc3.height = 26, 14
    ws.add_chart(sc3, "A64")
    notes["3"] = [f"{n} (no positive annualised revenue or market cap)" for n in out3] + [f"{n} (parked)"
                                                                                          for n in parked]
    ws.cell(93, 1, "Not shown: " + ("; ".join(notes["3"]) or "none")).font = F_SUB
    # 4. net supply change, % a year, sorted
    rows4, out4 = [], []
    for p in live:
        s = A.cell("Master", "Net supply change, annualised % of circulating", p["name"])
        if s is None or s.num is None:
            out4.append(p["name"])
            continue
        rows4.append((p["name"], s.num))
    rows4.sort(key=lambda x: x[1])
    r4 = table(r3 + len(rows3) + 3, "chart 4: net supply change", ["project", "net supply change, % a year"], rows4)
    b4 = BarChart()
    b4.type, b4.title = "bar", "4. Net supply change, % of circulating a year (sorted; + is net inflation)"
    if rows4:
        b4.add_data(Reference(dws, min_col=2, min_row=r4 - 1, max_row=r4 + len(rows4) - 1), titles_from_data=True)
        b4.set_categories(Reference(dws, min_col=1, min_row=r4, max_row=r4 + len(rows4) - 1))
    b4.width, b4.height = 26, 14
    ws.add_chart(b4, "A96")
    notes["4"] = [f"{n} (no net supply change)" for n in out4] + [f"{n} (parked)" for n in parked]
    ws.cell(125, 1, "Not shown: " + ("; ".join(notes["4"]) or "none")).font = F_SUB
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.column_dimensions["A"].width = 12
    for ch in (bar, sc, sc3, b4):
        for ax in (ch.x_axis, ch.y_axis):
            ax.delete = False                      # openpyxl hides axes unless told (Excel then shows none)
    for ch in (sc, sc3):
        ch.legend = None                           # each point carries its project's name
    sc.x_axis.number_format = sc.y_axis.number_format = "0%"
    b4.y_axis.number_format = "0%"
    sc3.x_axis.number_format = sc3.y_axis.number_format = '[>=1000000000]$#,##0,,,"bn";$#,##0,,"m"'
    bar.y_axis.number_format = '0.0"x"'
    for r in (29, 61, 93, 125):
        ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=16)
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 30
    return {k: len(v) for k, v in notes.items()}


def _name_labels(DataLabelList):
    """A point's label: its project's name only (no value, no legend key)."""
    d = DataLabelList()
    d.showSerName, d.showVal, d.showCatName, d.showLegendKey, d.showPercent = True, False, False, False, False
    return d


# ---------------------------------------------------------------------------------------------- Definitions, About
def _write_definitions(ws, sheets: dict) -> None:
    ws.cell(1, 1, "Definitions — every reader column").font = F_TITLE
    heads = ("Tab", "Column", "Header", "Definition", "Format", "Source")
    for i, h in enumerate(heads, start=1):
        c = ws.cell(3, i, h)
        c.font, c.fill = F_HEAD, FILL_HEAD
    r = 4
    for name, sh in sheets.items():
        for col in sh.cols:
            fmt = {FMT_USD: "$ ($m / $bn)", FMT_PCT: "%", FMT_X: "x", FMT_TOK: "number", FMT_RANK: "rank",
                   FMT_PRICE: "$", TEXT_FMT: "text"}.get(col.fmt, col.fmt)
            for i, v in enumerate((name, col.letter, col.header, plain(col.definition), fmt, col.source), start=1):
                c = ws.cell(r, i, v)
                c.font = F_BODY
                c.alignment = Alignment(wrap_text=i == 4, vertical="top")
            r += 1
    for L, w in zip("ABCDEF", (18, 8, 36, 90, 12, 50)):
        ws.column_dimensions[L].width = w
    ws.freeze_panes = "A4"
    r += 1
    ws.cell(r, 1, "Codes in place of a number").font = F_BOLD
    for code, why in (("n/a", "no value — the comment says why (a missing reading, a short history, an unformed "
                       "window)"), ("blocked", "a rule refused the figure (a stale input, a basis change, a lockup "
                                    "read too old) — the comment names the rule"),
                      ("not wired", "no route to the figure yet"),
                      ("by design", "chased and closed, or suppressed on purpose, or the column does not apply to "
                       "the project — not missing data"),
                      ("parked", "the project is parked: shown greyed with its numbers, outside the ranks and the "
                       "charts"),
                      ("unranked", "in a rank column: the project has no measured key, so it sorts last")):
        r += 1
        c = ws.cell(r, 1, code)
        c.font = Font(name=FONT, size=10, italic=True, color=CODES[code])
        ws.cell(r, 3, why).font = F_BODY


def _write_about(ws, asof, sheets: dict, projects: list) -> None:
    parked = [f"{p['name']} — {parked_reason(p['name'])}" for p in projects if parked_reason(p["name"])]
    tab_lines = {
        SUMMARY: "every project on one row, grouped by its primary archetype and ranked by that archetype's key; the "
                 "period selector is in B2.",
        A1_TAB: "the infrastructure chains: settlement and fees against issuance, NRR, yields.",
        A2_TAB: "the coordination networks: customer revenue against the emissions that pay suppliers, supply units, "
                "dilution.",
        A3_TAB: "the revenue-buyback tokens: what is retired against what is emitted, valuation, payout.",
        A4_TAB: "the permanent-burn tokens: burn against issuance (or against release for a pre-minted supply).",
        COMPARISON: "every project on the figures that mean the same for each, sorted by net holder yield.",
        BUSINESS: "fees, revenue, TVL and each project's own unit of demand, with growth over the selected period.",
        CHARTS: "four charts on the same numbers: the keys against 1.0x, holder value against dilution, revenue "
                "against market cap, net supply change.",
        DEFS: "every column's full definition and where its numbers come from.",
    }
    lines = [
        ("Token metrics", F_TITLE),
        (f"Data as of {pd.Timestamp(asof).date()}. Every figure is checked in {AUDIT_NAME}, built in the same run: it "
         "holds every working column, the stored data, the checks and the run log.", F_BODY),
        ("", F_BODY),
        ("What this file measures", F_BOLD),
        ("It measures whether a token's holders are paid by the business behind it, or diluted to keep it going. For "
         "each project it sets what the business earns and returns to holders against the new tokens it hands out. A "
         "token whose business returns more than it dilutes is building value for its holders; one that dilutes more "
         "is spending them.", F_BODY),
        ("", F_BODY),
        ("The four archetypes, and the key each is ranked by (higher is better)", F_BOLD),
        ("A1 Infrastructure — a chain whose token secures it: ranked by settlement volume per circulating token, "
         "1-year change.", F_BODY),
        ("A2 Coordination — a network paying suppliers in tokens: ranked by customer revenue ÷ emissions (1.0x is the "
         "crossover).", F_BODY),
        ("A3 Revenue buyback — revenue buys tokens back: ranked by tokens retired ÷ tokens emitted (1.0x is the "
         "crossover).", F_BODY),
        ("A4 Permanent burn — usage burns tokens: ranked by burn ÷ issuance, or burn ÷ release for a pre-minted supply "
         "(1.0x is the crossover).", F_BODY),
        ("", F_BODY),
        ("How to read a cell", F_BOLD),
        ("Fill is confidence: GREEN verified and usable · AMBER a real number with a caveat (hover it) · RED not usable "
         "· no fill: computed here from the cells it names.", F_BODY),
        ("A word in place of a number: n/a (no value), blocked (a rule refused it), not wired (no route yet), by design "
         "(closed on purpose, or the column does not apply), parked (see below); 'unranked' in a rank column means no "
         "measured key. A missing figure is never shown as 0; a real 0 says why beside it.", F_BODY),
        ("Comments: hover a header for its definition; hover a number for its source cell in the audit file.", F_BODY),
        ("Formats: $m and $bn for dollars; % for shares and rates; 0.0x for ratios (1.0x is the crossover). 'Last 90 "
         "days' is the trailing window to the as-of date.", F_BODY),
        ("", F_BODY),
        ("The period selector", F_BOLD),
        ("Summary!B2 sets the period for every growth and price-change column on every tab: 1D, 7D, 30D, 90D, 6M, 1Y, "
         "3Y or custom (the days in Summary!D2). A change needs every day of both windows; a shorter history reads "
         "n/a with the reason in the comment.", F_BODY),
        ("", F_BODY),
        ("The tabs", F_BOLD),
    ]
    for name, sh in sheets.items():
        lines.append((f"{name} — {tab_lines.get(name, '')}", F_BODY))
    lines.append((f"{CHARTS} — {tab_lines[CHARTS]}", F_BODY))
    lines.append((f"{DEFS} — {tab_lines[DEFS]}", F_BODY))
    lines += [
        ("", F_BODY),
        ("Parked projects", F_BOLD),
        ("A parked project is shown greyed with its numbers, outside the ranks and the charts, until work on it "
         "resumes.", F_BODY),
    ]
    lines += [(f"Parked: {x}", F_BODY) for x in parked] or [("No project is parked on this build.", F_BODY)]
    for i, (t, f) in enumerate(lines, start=1):
        c = ws.cell(i, 1, t)
        c.font = f
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120
    if SUMMARY in sheets:
        _write_coverage(ws, len(lines) + 2, sheets[SUMMARY].ws)


COVERAGE_CODES = (("blocked", "a rule refused the key: unequal coverage of its two legs, a stale input or a basis "
                               "change; the key cell's comment names the rule"),
                  ("n/a", "the key has no value: a leg has no reading, the history is too short, or a ratio over 0"),
                  ("not wired", "no route to one of the key's inputs yet"),
                  ("by design", "the key does not apply to the project (for example a fixed supply that emits "
                                "nothing)"))


def _summary_sections(ws) -> list:
    """[(archetype label, first row, last row)] of the Summary tab's project rows, by its section header rows."""
    out, cur = [], None
    for r in range(5, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if a and b is None and re.match(r"A[1-4] ", str(a)):
            cur = [str(a).split(" — ")[0], None, None]
            out.append(cur)
        elif a and b is not None and cur is not None:
            cur[1] = cur[1] or r
            cur[2] = r
    return [tuple(x) for x in out if x[1]]


def _write_coverage(ws, top: int, sws) -> None:
    """P2c-8: the honest-coverage block — per archetype, how many projects are ranked, unranked (by the key's code) and
    parked, as live formulas on the Summary tab. Its text is fixed; only its figures move."""
    hdr = {sws.cell(4, c).value: c for c in range(1, sws.max_column + 1)}
    kL, rL = get_column_letter(hdr["Archetype key"]), get_column_letter(hdr["Rank in archetype"])
    ws.cell(top, 1, "Honest coverage — counted from the Summary tab (the figures follow the data)").font = F_BOLD
    heads = ("Archetype", "Projects", "Ranked", "Unranked", "Parked") + tuple(f"Unranked: {c}" for c, _ in
                                                                           COVERAGE_CODES)
    for j, h in enumerate(heads, start=1):
        c = ws.cell(top + 1, j, h)
        c.font, c.fill = F_HEAD, FILL_HEAD
    r = top + 2
    first = r
    for label, a, b in _summary_sections(sws):
        rng = lambda L: f"Summary!${L}${a}:${L}${b}"                     # noqa: E731
        ws.cell(r, 1, label).font = F_BODY
        forms = [f"=COUNTA({rng('A')})", f'=COUNTIF({rng(rL)},">0")', f'=COUNTIF({rng(rL)},"unranked")',
                 f'=COUNTIF({rng(rL)},"parked")']
        forms += [f'=COUNTIFS({rng(rL)},"unranked",{rng(kL)},"{code}")' for code, _ in COVERAGE_CODES]
        for j, f in enumerate(forms, start=2):
            c = ws.cell(r, j, f)
            c.number_format, c.font = "0", F_BODY
        r += 1
    ws.cell(r, 1, "All archetypes").font = F_BOLD
    for j in range(2, len(heads) + 1):
        L = get_column_letter(j)
        c = ws.cell(r, j, "=" + "+".join(f"{L}{i}" for i in range(first, r)) if r > first else "=0")
        c.number_format, c.font = "0", F_BOLD
    r += 2
    ws.cell(r, 1, "Why a project is unranked, by the code in its key cell").font = F_BOLD
    for code, why in COVERAGE_CODES:
        r += 1
        ws.cell(r, 1, f"{code} — {why}.").font = F_BODY
    r += 1
    ws.cell(r, 1, "parked — the project is parked: its row is greyed, its numbers shown, outside the ranks; the reason "
                  "is listed under Parked projects above.").font = F_BODY
    for j in range(2, len(heads) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 12


# ---------------------------------------------------------------------------------------------- the column map
MAPPED_TABS = ("Master", A1_TAB, A2_TAB, A3_TAB, A4_TAB, CP_TAB)


def _write_column_map(path: Path, A: Audit, sheets: dict) -> None:
    """Every audit column → its reader place, or why it stays audit-only (P1-9)."""
    placed: dict = {}
    for name, sh in sheets.items():
        for col in sh.cols:
            if col.audit:
                sheet, head = col.audit
                c = A.col(sheet, head) if sheet in A.wb.sheetnames else None
                if c is not None:
                    placed.setdefault((sheet, c), []).append((name, col.letter, col.header))
    rows = []
    for sheet in A.wb.sheetnames:
        if sheet in MAPPED_TABS:
            for h, c in sorted(A.cols(sheet).items(), key=lambda kv: kv[1]):
                L = get_column_letter(c)
                hits = placed.get((sheet, c))
                if hits:
                    for rt, rl, rh in hits:
                        rows.append((AUDIT_NAME, sheet, L, h, READER_NAME, rt, rl, rh, "moved (also kept "
                                     "on the audit file)"))
                    continue
                if h.startswith("Detail — "):
                    why = "audit file; a detail line (to Project View)"
                elif h.startswith("Trend — "):
                    why = "audit file (a trend string); the reader shows it as a quarter change where its figure is shown"
                elif _FIXED_WINDOW.search(h):
                    why = ("audit file; a fixed-window change — the reader shows the selected-period growth of the same "
                           "series (the period selector drives every growth column)")
                elif _QUARTER_COL.search(h):
                    why = "audit file; a past quarter — the reader's block (c) shows the last four quarters"
                elif h in ("Project", "Symbol"):
                    why = "the row label"
                elif h in ("Notes", "Data flags (stale / manual / missing)"):
                    why = "the reader's project cell comment"
                else:
                    why = "audit file"
                rows.append((AUDIT_NAME, sheet, L, h, "", "", "", "", why))
        else:
            for h, c in sorted(A.cols(sheet).items(), key=lambda kv: kv[1]):
                hits = placed.get((sheet, c))
                L = get_column_letter(c)
                if hits:
                    for rt, rl, rh in hits:
                        rows.append((AUDIT_NAME, sheet, L, h, READER_NAME, rt, rl, rh,
                                     "shown on the reader (kept on the audit file)"))
            rows.append((AUDIT_NAME, sheet, "(all)", "(every column)", "", "", "", "",
                         "audit file, unchanged (same tab, same columns)"))
    COLUMN_MAP[:] = rows
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(("audit_file", "audit_tab", "audit_column", "audit_header", "reader_file", "reader_tab",
                    "reader_column", "reader_header", "note"))
        w.writerows(rows)
