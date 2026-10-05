"""
xlcalc.py — the workbook's formulas, evaluated in Python (Jake, 2026-10-05).

credibility_report.py needs the verdicts, and the verdicts are formulas over the headline cells. On a
machine without LibreOffice (Jake's Windows box) nothing recalculates an openpyxl-written file, so this
evaluates them directly. It covers exactly the subset build_workbook.py writes — checked by
`unsupported()`, which names any function outside it rather than guessing:

    INDEX  MATCH(…,0)  IF  IFERROR  ISNUMBER  AND  OR  NOT  LEFT  ABS  ROUND  SQRT  NA
    COUNTIF  COUNTIFS  COUNTA  TEXT(…,"#,##0")
    + - * / ^ & = <> < > <= >=, unary minus, %, sheet-qualified cells and ranges

Excel's rules where they matter: errors propagate through arithmetic and comparisons; IF and IFERROR
evaluate only the branch they take; an empty cell is 0 in arithmetic and "" in text; text compares
case-insensitively; a number sorts below text, text below a boolean; MATCH(…,0) is an exact,
case-insensitive match. Dates are Excel serial numbers. Read-only: it never writes the workbook.
"""
from __future__ import annotations

import datetime as _dt
import fnmatch
import math
import re
import sys

from openpyxl.utils import column_index_from_string, get_column_letter

SUPPORTED = {"INDEX", "MATCH", "IF", "IFERROR", "ISNUMBER", "AND", "OR", "NOT", "LEFT", "ABS", "ROUND",
             "SQRT", "NA", "COUNTIF", "COUNTIFS", "COUNTA", "TEXT"}


class XLError(Exception):
    """An Excel error value (#N/A, #VALUE!, #DIV/0!, #REF!, #NUM!, #NAME?). Returned, not raised,
    except inside the evaluator where raising is the propagation."""

    def __init__(self, code: str):
        super().__init__(code)
        self.code = code

    def __repr__(self):
        return self.code

    def __eq__(self, other):
        return isinstance(other, XLError) and other.code == self.code

    def __hash__(self):
        return hash(self.code)


class Range:
    __slots__ = ("sheet", "r1", "c1", "r2", "c2")

    def __init__(self, sheet, r1, c1, r2, c2):
        self.sheet, self.r1, self.c1, self.r2, self.c2 = sheet, min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2)

    def cells(self):
        for r in range(self.r1, self.r2 + 1):
            for c in range(self.c1, self.c2 + 1):
                yield r, c


_EPOCH = _dt.datetime(1899, 12, 30)


def _serial(v):
    if isinstance(v, _dt.datetime):
        return (v - _EPOCH).total_seconds() / 86400.0
    if isinstance(v, _dt.date):
        return float((v - _EPOCH.date()).days)
    if isinstance(v, _dt.time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400.0
    return v


# ---------------------------------------------------------------------------- tokens
_TOKEN = re.compile(r"""
    (?P<ws>\s+)
  | (?P<str>"(?:[^"]|"")*")
  | (?P<err>\#N/A|\#VALUE!|\#DIV/0!|\#REF!|\#NUM!|\#NAME\?|\#NULL!)
  | (?P<ref>(?:(?:'(?:[^']|'')+'|[A-Za-z_][\w\.]*)!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)(?![\w(])
  | (?P<num>\d+(?:\.\d*)?(?:[eE][+-]?\d+)?|\.\d+(?:[eE][+-]?\d+)?)
  | (?P<func>[A-Z][A-Z0-9\.]*)(?=\()
  | (?P<bool>TRUE|FALSE)(?![\w(])
  | (?P<op><>|<=|>=|[-+*/^&=<>%(),])
""", re.X)


def tokenize(src: str) -> list[tuple[str, str]]:
    out, i = [], 0
    while i < len(src):
        m = _TOKEN.match(src, i)
        if not m:
            raise SyntaxError(f"cannot read the formula at {src[i:i + 30]!r}")
        i = m.end()
        kind = m.lastgroup
        if kind != "ws":
            out.append((kind, m.group(kind)))
    return out


# ---------------------------------------------------------------------------- parser (to tuples)
class _Parser:
    # Excel precedence, loosest first: comparison, &, + -, * /, ^, unary -, %
    def __init__(self, toks):
        self.t, self.i = toks, 0

    def peek(self):
        return self.t[self.i] if self.i < len(self.t) else (None, None)

    def take(self, val=None):
        tok = self.peek()
        if val is not None and tok[1] != val:
            raise SyntaxError(f"expected {val!r}, found {tok[1]!r}")
        self.i += 1
        return tok

    def parse(self):
        node = self.comparison()
        if self.i != len(self.t):
            raise SyntaxError(f"unexpected {self.peek()[1]!r}")
        return node

    def comparison(self):
        node = self.concat()
        while self.peek()[0] == "op" and self.peek()[1] in ("=", "<>", "<", ">", "<=", ">="):
            op = self.take()[1]
            node = ("cmp", op, node, self.concat())
        return node

    def concat(self):
        node = self.additive()
        while self.peek() == ("op", "&"):
            self.take()
            node = ("cat", node, self.additive())
        return node

    def additive(self):
        node = self.term()
        while self.peek()[0] == "op" and self.peek()[1] in "+-" and self.peek()[1]:
            op = self.take()[1]
            node = ("bin", op, node, self.term())
        return node

    def term(self):
        node = self.power()
        while self.peek()[0] == "op" and self.peek()[1] in ("*", "/"):
            op = self.take()[1]
            node = ("bin", op, node, self.power())
        return node

    def power(self):
        node = self.unary()
        while self.peek() == ("op", "^"):
            self.take()
            node = ("bin", "^", node, self.unary())
        return node

    def unary(self):
        if self.peek()[0] == "op" and self.peek()[1] in ("-", "+"):
            op = self.take()[1]
            inner = self.unary()
            return ("neg", inner) if op == "-" else inner
        return self.percent()

    def percent(self):
        node = self.primary()
        while self.peek() == ("op", "%"):
            self.take()
            node = ("bin", "/", node, ("num", 100.0))
        return node

    def primary(self):
        kind, val = self.take()
        if kind == "num":
            return ("num", float(val))
        if kind == "str":
            return ("str", val[1:-1].replace('""', '"'))
        if kind == "bool":
            return ("bool", val == "TRUE")
        if kind == "err":
            return ("err", val)
        if kind == "ref":
            return ("ref", val)
        if kind == "func":
            self.take("(")
            args = []
            if self.peek() != ("op", ")"):
                while True:
                    if self.peek()[1] in (",", ")"):
                        args.append(("missing",))
                    else:
                        args.append(self.comparison())
                    if self.peek() == ("op", ","):
                        self.take()
                        continue
                    break
            self.take(")")
            return ("fn", val, args)
        if (kind, val) == ("op", "("):
            node = self.comparison()
            self.take(")")
            return node
        raise SyntaxError(f"unexpected {val!r}")


def parse(formula: str):
    return _Parser(tokenize(formula[1:] if formula.startswith("=") else formula)).parse()


def functions_in(node, acc=None) -> set:
    acc = set() if acc is None else acc
    if isinstance(node, tuple):
        if node and node[0] == "fn":
            acc.add(node[1])
        for x in node[1:]:
            if isinstance(x, (tuple, list)):
                for y in (x if isinstance(x, list) else [x]):
                    functions_in(y, acc)
    return acc


# ---------------------------------------------------------------------------- coercions
def _num(v):
    v = _serial(v)
    if isinstance(v, XLError):
        raise v
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip())
    except ValueError:
        raise XLError("#VALUE!") from None


def _text(v):
    v = _serial(v)
    if isinstance(v, XLError):
        raise v
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e15:
            return str(int(v))
        return repr(v)
    return str(v)


def _bool(v):
    v = _serial(v)
    if isinstance(v, XLError):
        raise v
    if isinstance(v, bool):
        return v
    if v is None:
        return False
    if isinstance(v, (int, float)):
        return v != 0
    s = str(v).upper()
    if s in ("TRUE", "FALSE"):
        return s == "TRUE"
    raise XLError("#VALUE!")


def _rank(v):
    """Excel's cross-type order: numbers < text < booleans; an empty cell compares as its partner's blank."""
    if isinstance(v, bool):
        return 2, v
    if isinstance(v, (int, float)):
        return 0, float(v)
    return 1, str(v).lower()


def _compare(op, a, b):
    a, b = _serial(a), _serial(b)
    for v in (a, b):
        if isinstance(v, XLError):
            raise v
    if a is None:
        a = "" if isinstance(b, str) else (False if isinstance(b, bool) else 0.0)
    if b is None:
        b = "" if isinstance(a, str) else (False if isinstance(a, bool) else 0.0)
    ra, rb = _rank(a), _rank(b)
    return {"=": ra == rb, "<>": ra != rb, "<": ra < rb, ">": ra > rb, "<=": ra <= rb, ">=": ra >= rb}[op]


def _format_thousands(v):
    return f"{round(_num(v)):,.0f}"


# ---------------------------------------------------------------------------- the evaluator
class Workbook:
    def __init__(self, path_or_wb):
        if isinstance(path_or_wb, (str, bytes)) or hasattr(path_or_wb, "__fspath__"):
            import openpyxl
            path_or_wb = openpyxl.load_workbook(path_or_wb)
        self.wb = path_or_wb
        self.sheets = {ws.title: ws for ws in self.wb.worksheets}
        self._raw: dict[str, dict] = {}
        self._cache: dict[tuple, object] = {}
        self._ast: dict[str, object] = {}
        self._busy: set = set()
        self._match_idx: dict[tuple, dict] = {}

    # raw cell contents, read once per sheet
    def _sheet(self, name):
        raw = self._raw.get(name)
        if raw is None:
            ws = self.sheets.get(name)
            if ws is None:
                raise XLError("#REF!")
            raw = {}
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None:
                        raw[(c.row, c.column)] = c.value
            self._raw[name] = raw
        return raw

    def value(self, sheet: str, row: int, col: int):
        key = (sheet, row, col)
        if key in self._cache:
            return self._cache[key]
        raw = self._sheet(sheet).get((row, col))
        if isinstance(raw, str) and raw.startswith("=") and len(raw) > 1:
            if key in self._busy:
                return XLError("#REF!")          # a circular reference; LibreOffice shows Err:522
            self._busy.add(key)
            try:
                out = self.formula(raw, sheet)
            finally:
                self._busy.discard(key)
        else:
            out = _serial(raw) if not isinstance(raw, str) else raw
            if isinstance(raw, int) and not isinstance(raw, bool):
                out = float(raw)
        self._cache[key] = out
        return out

    def cell(self, sheet: str, coord: str):
        m = re.fullmatch(r"\$?([A-Z]{1,3})\$?(\d+)", coord)
        return self.value(sheet, int(m.group(2)), column_index_from_string(m.group(1)))

    def formula(self, text: str, sheet: str):
        ast = self._ast.get(text)
        if ast is None:
            try:
                ast = parse(text)
            except SyntaxError as e:
                ast = ("syntax", str(e))
            self._ast[text] = ast
        if ast[0] == "syntax":
            return XLError("#NAME?")
        try:
            v = self._eval(ast, sheet)
            if isinstance(v, Range):
                v = self._scalar(v)
            return 0.0 if v is None else v           # a formula that lands on an empty cell shows 0
        except XLError as e:
            return e
        except RecursionError:
            return XLError("#REF!")

    # references
    def _range(self, ref: str, sheet: str) -> Range:
        if "!" in ref:
            sh, ref = ref.rsplit("!", 1)
            sheet = sh[1:-1].replace("''", "'") if sh.startswith("'") else sh
        parts = ref.replace("$", "").split(":")
        m1 = re.fullmatch(r"([A-Z]+)(\d+)", parts[0])
        m2 = re.fullmatch(r"([A-Z]+)(\d+)", parts[-1])
        return Range(sheet, int(m1.group(2)), column_index_from_string(m1.group(1)),
                     int(m2.group(2)), column_index_from_string(m2.group(1)))

    def _scalar(self, rng: Range):
        if rng.r1 == rng.r2 and rng.c1 == rng.c2:
            return self.value(rng.sheet, rng.r1, rng.c1)
        raise XLError("#VALUE!")          # implicit intersection is never used by this workbook

    def _val(self, node, sheet):
        v = self._eval(node, sheet)
        if isinstance(v, Range):
            v = self._scalar(v)
        if isinstance(v, XLError):
            raise v
        return v

    def _eval(self, node, sheet):
        kind = node[0]
        if kind == "num" or kind == "str" or kind == "bool":
            return node[1]
        if kind == "err":
            raise XLError(node[1])
        if kind == "ref":
            return self._range(node[1], sheet)
        if kind == "missing":
            return None
        if kind == "neg":
            return -_num(self._val(node[1], sheet))
        if kind == "cat":
            return _text(self._val(node[1], sheet)) + _text(self._val(node[2], sheet))
        if kind == "cmp":
            return _compare(node[1], self._val(node[2], sheet), self._val(node[3], sheet))
        if kind == "bin":
            a, b = _num(self._val(node[2], sheet)), _num(self._val(node[3], sheet))
            op = node[1]
            if op == "+":
                return a + b
            if op == "-":
                return a - b
            if op == "*":
                return a * b
            if op == "/":
                if b == 0:
                    raise XLError("#DIV/0!")
                return a / b
            try:
                return float(a ** b)
            except (OverflowError, ZeroDivisionError, ValueError):
                raise XLError("#NUM!") from None
        if kind == "fn":
            return self._call(node[1], node[2], sheet)
        raise XLError("#NAME?")

    # functions
    def _call(self, name, args, sheet):
        f = getattr(self, "_f_" + name, None)
        if f is None:
            raise XLError("#NAME?")
        return f(args, sheet)

    def _f_IF(self, args, sheet):
        cond = _bool(self._val(args[0], sheet))
        if cond:
            return self._eval(args[1], sheet) if len(args) > 1 else True
        return self._eval(args[2], sheet) if len(args) > 2 and args[2] != ("missing",) else False

    def _f_IFERROR(self, args, sheet):
        try:
            return self._val(args[0], sheet)
        except XLError:
            return self._val(args[1], sheet)

    def _f_ISNUMBER(self, args, sheet):
        try:
            v = self._val(args[0], sheet)
        except XLError:
            return False
        return isinstance(_serial(v), (int, float)) and not isinstance(v, bool)

    def _f_AND(self, args, sheet):
        return all([_bool(self._val(a, sheet)) for a in args])

    def _f_OR(self, args, sheet):
        return any([_bool(self._val(a, sheet)) for a in args])

    def _f_NOT(self, args, sheet):
        return not _bool(self._val(args[0], sheet))

    def _f_LEFT(self, args, sheet):
        s = _text(self._val(args[0], sheet))
        n = int(_num(self._val(args[1], sheet))) if len(args) > 1 else 1
        if n < 0:
            raise XLError("#VALUE!")
        return s[:n]

    def _f_ABS(self, args, sheet):
        return abs(_num(self._val(args[0], sheet)))

    def _f_SQRT(self, args, sheet):
        x = _num(self._val(args[0], sheet))
        if x < 0:
            raise XLError("#NUM!")
        return math.sqrt(x)

    def _f_ROUND(self, args, sheet):
        x, d = _num(self._val(args[0], sheet)), int(_num(self._val(args[1], sheet)))
        q = 10.0 ** d
        return math.floor(abs(x) * q + 0.5) / q * (1 if x >= 0 else -1)    # half away from zero

    def _f_NA(self, args, sheet):
        raise XLError("#N/A")

    def _f_TEXT(self, args, sheet):
        fmt = _text(self._val(args[1], sheet))
        v = self._val(args[0], sheet)
        if isinstance(v, str):
            try:
                v = float(v)
            except ValueError:
                return v                         # TEXT of text is the text itself, as in Excel
        if fmt == "#,##0":
            return _format_thousands(v)
        raise XLError("#VALUE!")                 # a format this evaluator does not know: said, not guessed

    def _f_INDEX(self, args, sheet):
        rng = self._eval(args[0], sheet)
        if not isinstance(rng, Range):
            raise XLError("#VALUE!")
        r = int(_num(self._val(args[1], sheet))) if len(args) > 1 and args[1] != ("missing",) else 0
        c = int(_num(self._val(args[2], sheet))) if len(args) > 2 and args[2] != ("missing",) else 0
        rows, cols = rng.r2 - rng.r1 + 1, rng.c2 - rng.c1 + 1
        if c == 0 and cols == 1:
            c = 1
        if r == 0 and rows == 1:
            r = 1
        if not (1 <= r <= rows and 1 <= c <= cols):
            raise XLError("#REF!")
        v = self.value(rng.sheet, rng.r1 + r - 1, rng.c1 + c - 1)
        if isinstance(v, XLError):
            raise v
        return v

    def _f_MATCH(self, args, sheet):
        key = self._val(args[0], sheet)
        rng = self._eval(args[1], sheet)
        kind = _num(self._val(args[2], sheet)) if len(args) > 2 else 1.0
        if not isinstance(rng, Range) or (rng.r1 != rng.r2 and rng.c1 != rng.c2):
            raise XLError("#N/A")
        if kind != 0:
            raise XLError("#N/A")                # only exact matches are written; anything else is refused
        ik = (rng.sheet, rng.r1, rng.c1, rng.r2, rng.c2)
        idx = self._match_idx.get(ik)
        if idx is None:
            idx = {}
            for i, (r, c) in enumerate(rng.cells(), start=1):
                v = self.value(rng.sheet, r, c)
                if v is None or isinstance(v, XLError):
                    continue
                idx.setdefault(_rank(v), i)
            self._match_idx[ik] = idx
        if key is None:
            raise XLError("#N/A")
        pos = idx.get(_rank(_serial(key)))
        if pos is None:
            raise XLError("#N/A")
        return float(pos)

    def _criterion(self, crit):
        crit = _serial(crit)
        if isinstance(crit, (int, float)) and not isinstance(crit, bool):
            return lambda v: isinstance(_serial(v), (int, float)) and not isinstance(v, bool) and float(_serial(v)) == crit
        s = _text(crit)
        m = re.match(r"(<>|<=|>=|=|<|>)?(.*)$", s, re.S)
        op, rest = m.group(1), m.group(2)
        try:
            n = float(rest)
        except ValueError:
            n = None
        if op in ("<", ">", "<=", ">=") and n is not None:
            return lambda v: isinstance(_serial(v), (int, float)) and not isinstance(v, bool) and \
                _compare(op, float(_serial(v)), n)
        pat = rest.lower()
        wild = any(ch in pat for ch in "*?")

        def eq(v):
            if v is None:
                return pat == ""
            if isinstance(v, XLError):
                return False
            if n is not None and isinstance(_serial(v), (int, float)) and not isinstance(v, bool):
                return float(_serial(v)) == n
            t = _text(v).lower()
            return fnmatch.fnmatchcase(t, pat.replace("[", "[[]")) if wild else t == pat
        return (lambda v: not eq(v)) if op == "<>" else eq

    def _countifs(self, pairs, sheet):
        ranges, tests = [], []
        for rng_node, crit_node in pairs:
            rng = self._eval(rng_node, sheet)
            if not isinstance(rng, Range):
                raise XLError("#VALUE!")
            ranges.append(list(rng.cells()))
            tests.append((rng.sheet, self._criterion(self._val(crit_node, sheet))))
        if len({len(r) for r in ranges}) != 1:
            raise XLError("#VALUE!")
        n = 0
        for i in range(len(ranges[0])):
            if all(t(self.value(sh, *ranges[k][i])) for k, (sh, t) in enumerate(tests)):
                n += 1
        return float(n)

    def _f_COUNTIFS(self, args, sheet):
        return self._countifs(list(zip(args[0::2], args[1::2])), sheet)

    def _f_COUNTIF(self, args, sheet):
        return self._countifs([(args[0], args[1])], sheet)

    def _f_COUNTA(self, args, sheet):
        n = 0
        for a in args:
            v = self._eval(a, sheet)
            if isinstance(v, Range):
                n += sum(1 for r, c in v.cells() if self.value(v.sheet, r, c) is not None)
            elif v is not None:
                n += 1
        return float(n)


def unsupported(wb) -> dict[str, list[str]]:
    """{function: [first few cells]} for every function a formula uses that this evaluator lacks."""
    out: dict[str, list[str]] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("=") and len(v) > 1:
                    try:
                        names = functions_in(parse(v))
                    except SyntaxError:
                        names = {"<unparsable>"}
                    for f in names - SUPPORTED:
                        out.setdefault(f, [])
                        if len(out[f]) < 3:
                            out[f].append(f"{ws.title}!{c.coordinate}")
    return out


def display(v):
    """A computed value as the report prints it: an error as its code."""
    return v.code if isinstance(v, XLError) else v


sys.setrecursionlimit(max(sys.getrecursionlimit(), 20000))
__all__ = ["Workbook", "XLError", "parse", "unsupported", "display", "get_column_letter"]
