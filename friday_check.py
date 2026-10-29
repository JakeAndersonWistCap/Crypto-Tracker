"""friday_check.py — the Friday checks on the two workbooks, in one command (P2f-2). READ-ONLY.

    python friday_check.py                      # token_metrics.xlsx + token_metrics_audit.xlsx in this folder
    python friday_check.py --reader R --audit A --store metrics.db --asof 2026-10-30
    python friday_check.py --compare            # the guard on the committed fixtures instead of the store

It opens the two workbooks, runs nine checks and prints one line per check:

    PASS  <name>
    FAIL  <name> — <first reasons>              (at most 10 reasons, then "and N more")

then `FRIDAY CHECK: <n> PASS, <m> FAIL`. Exit code 1 on any FAIL, else 0. It writes nothing, fetches nothing and never
writes the store: the guard (check 7) is the repo's own `measuring_guard.py`, run as a subprocess (it builds on a copy).

    1  ranks identical on Summary and on each project's primary archetype tab
    2  About's coverage counts equal the Summary counts
    3  no header over reader.MAX_HEADER (35) characters
    4  no blank cell without a code
    5  every parked row reads "parked" in its rank and key cells
    6  every value-cell comment opens with a plain sentence
    7  measuring_guard: 0 violation(s); the ACCEPTED lines listed
    8  coverage scoreboard: BROKEN and NOT WIRED, with their projects
    9  the workbook's as-of date equals today

Each check is a function (workbook(s) in, list of reasons out; empty = PASS), so it can be tested on a fake workbook.
"""
from __future__ import annotations

import argparse
import os
import re
import subprocess
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

import reader  # noqa: E402

MAX_REASONS = 10
VALUE_TABS = (reader.SUMMARY, reader.A1_TAB, reader.A2_TAB, reader.A3_TAB, reader.A4_TAB, reader.COMPARISON,
              reader.BUSINESS)
ARCH_TABS = {1: reader.A1_TAB, 2: reader.A2_TAB, 3: reader.A3_TAB, 4: reader.A4_TAB}
GUARD_BASELINE = ROOT / "tests" / "fixtures" / "guard_baseline.json"
GUARD_TIMEOUT = 3600                         # seconds; the guard builds the 16-project workbook on a copy


# ======================================================================================================== helpers
def _project_rows(ws) -> dict:
    """{project: row} of a reader tab's project rows (an archetype group row has nothing in column 2)."""
    return {ws.cell(r, 1).value: r for r in range(5, ws.max_row + 1)
            if ws.cell(r, 1).value and ws.cell(r, 2).value is not None}


def _col(ws, header: str) -> int | None:
    return next((c for c in range(1, ws.max_column + 1) if ws.cell(4, c).value == header), None)


def _rank_col(ws) -> int | None:
    return next((c for c in range(1, ws.max_column + 1) if str(ws.cell(4, c).value or "").startswith("Rank in")),
                None)


def _block_of(ws, c: int) -> str:
    """The block label over column c (row 3 holds it on the first column of each block)."""
    for j in range(c, 0, -1):
        v = ws.cell(3, j).value
        if v:
            return str(v)
    return ""


def _primary_archetypes(sws) -> dict:
    """{project: primary archetype number} from Summary's section rows."""
    out = {}
    for label, lo, hi in reader._summary_sections(sws):
        for r in range(lo, hi + 1):
            out[sws.cell(r, 1).value] = int(label[1])
    return out


def _is_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _is_formula(v) -> bool:
    return isinstance(v, str) and v.startswith("=") and len(v) > 1


def _asof_of_reader(wb):
    """The as-of day the reader states on About ('Data as of YYYY-MM-DD.'), or None."""
    if "About" not in wb.sheetnames:
        return None
    ws = wb["About"]
    for r in range(1, min(ws.max_row, 30) + 1):
        m = re.search(r"Data as of (\d{4}-\d{2}-\d{2})", str(ws.cell(r, 1).value or ""))
        if m:
            return date.fromisoformat(m.group(1))
    return None


def _asof_of_audit(wb):
    """The as-of day on the audit file's Run Log title line ('… data as of YYYY-MM-DD (UTC)'), or None."""
    if "Run Log" not in wb.sheetnames:
        return None
    for row in wb["Run Log"].iter_rows(min_row=1, max_row=6, max_col=1, values_only=True):
        m = re.search(r"data as of (\d{4}-\d{2}-\d{2})", str(row[0] or ""), re.I)
        if m:
            return date.fromisoformat(m.group(1))
    return None


# ======================================================================================================== 1 ranks
def check_ranks(wb, X=None) -> list:
    """The rank in Summary's 'Rank in archetype' equals the project's rank on its primary tab; ranks on a tab run
    1..n in row order and only a primary holds one; 'secondary' reads only on a tab that is not the project's primary."""
    import xlcalc                                                        # noqa: PLC0415
    X = X or xlcalc.Workbook(wb)
    s = wb[reader.SUMMARY]
    sr = _col(s, "Rank in archetype")
    if sr is None:
        return ["Summary has no 'Rank in archetype' column"]
    srows, prim, bad = _project_rows(s), _primary_archetypes(s), []
    if not prim:
        return ["Summary has no archetype sections"]
    for name, arch in prim.items():
        ws = wb[ARCH_TABS[arch]]
        rows, rc = _project_rows(ws), _rank_col(ws)
        if name not in rows or rc is None:
            continue
        want, got = X.value(reader.SUMMARY, srows[name], sr), X.value(ws.title, rows[name], rc)
        if got != want:
            bad.append(f"{name}: rank {got!r} on {ws.title}, {want!r} on Summary")
    for arch, tab in ARCH_TABS.items():
        ws = wb[tab]
        rc = _rank_col(ws)
        if rc is None:
            bad.append(f"{tab} has no 'Rank in' column")
            continue
        ranked = []
        for n, r in sorted(_project_rows(ws).items(), key=lambda kv: kv[1]):
            v = X.value(tab, r, rc)
            if _is_number(v):
                ranked.append(v)
                if prim.get(n) != arch:
                    bad.append(f"{n}: holds rank {v:g} on {tab} but its primary archetype is "
                               f"A{prim.get(n, '?')}")
            elif v == "secondary" and prim.get(n) == arch:
                bad.append(f"{n}: reads 'secondary' on {tab}, its own primary tab")
        if ranked != [float(i) for i in range(1, len(ranked) + 1)]:
            bad.append(f"{tab}: ranks run {[int(x) for x in ranked]}, not 1..{len(ranked)} in row order")
    return bad


# ======================================================================================================== 2 About
def check_about_counts(wb, X=None) -> list:
    """About's honest-coverage block (COUNTA / COUNTIF / COUNTIFS over Summary) equals the counts recomputed here from
    the Summary cells."""
    import xlcalc                                                        # noqa: PLC0415
    X = X or xlcalc.Workbook(wb)
    s, a = wb[reader.SUMMARY], wb["About"]
    top = next((r for r in range(1, a.max_row + 1) if str(a.cell(r, 1).value or "").startswith("Honest coverage")),
               None)
    if top is None:
        return ["About has no 'Honest coverage' block"]
    rk, kk = _col(s, "Rank in archetype"), _col(s, "Archetype key")
    if rk is None or kk is None:
        return ["Summary lacks 'Rank in archetype' or 'Archetype key'"]
    sections = reader._summary_sections(s)
    if not sections:
        return ["Summary has no archetype sections"]
    heads = ["Projects", "Ranked", "Unranked", "Parked"] + [f"Unranked: {c}" for c, _ in reader.COVERAGE_CODES]
    bad, totals = [], [0] * len(heads)

    def compare(row: int, label: str, want: list) -> None:
        for j, (h, w) in enumerate(zip(heads, want), start=2):
            got = X.value("About", row, j)
            if got != float(w):
                bad.append(f"{label} {h}: About says {got!r}, Summary has {w}")

    for i, (label, lo, hi) in enumerate(sections):
        row = top + 2 + i
        if str(a.cell(row, 1).value or "") != label:
            bad.append(f"About row {row} is {a.cell(row, 1).value!r}, expected {label!r}")
        ranks = [X.value(reader.SUMMARY, r, rk) for r in range(lo, hi + 1)]
        keys = [X.value(reader.SUMMARY, r, kk) for r in range(lo, hi + 1)]
        want = [hi - lo + 1, sum(_is_number(v) and v > 0 for v in ranks), ranks.count("unranked"),
                ranks.count("parked")]
        want += [sum(1 for k, rv in zip(keys, ranks) if rv == "unranked" and k == code)
                 for code, _w in reader.COVERAGE_CODES]
        compare(row, label, want)
        totals = [t + w for t, w in zip(totals, want)]
    compare(top + 2 + len(sections), "All archetypes", totals)
    return bad


# ======================================================================================================== 3 headers
def check_headers(wb) -> list:
    """No header over reader.MAX_HEADER characters on any reader tab's header row (row 4)."""
    bad = []
    for tab in VALUE_TABS:
        if tab not in wb.sheetnames:
            bad.append(f"tab {tab} is missing")
            continue
        ws = wb[tab]
        for c in range(1, ws.max_column + 1):
            h = ws.cell(4, c).value
            if h is not None and len(str(h)) > reader.MAX_HEADER:
                bad.append(f"{tab}!{ws.cell(4, c).coordinate}: {len(str(h))} characters ({str(h)!r})")
    return bad


# ======================================================================================================== 4 blanks
def check_no_blank(wb) -> list:
    """A value cell on a project row is never empty: it holds a number, a formula or a code from reader.CODES, and a code
    carries its reason as a comment."""
    bad = []
    for tab in VALUE_TABS:
        if tab not in wb.sheetnames:
            continue
        ws = wb[tab]
        for name, r in _project_rows(ws).items():
            for c in range(3, ws.max_column + 1):
                cell = ws.cell(r, c)
                v = cell.value
                head = ws.cell(4, c).value
                if v is None or (isinstance(v, str) and not v.strip()):
                    bad.append(f"{tab} / {name} / {head}: blank, no code")
                elif isinstance(v, str) and v in reader.CODES and not (cell.comment and cell.comment.text.strip()):
                    bad.append(f"{tab} / {name} / {head}: code {v!r} without a comment")
    return bad


# ======================================================================================================== 5 parked
def parked_projects(names, asof=None) -> list:
    """The projects reader.parked_reason parks on the workbook's as-of day (the day the build used)."""
    if asof is not None:
        reader._AERO_PARK.clear()
        reader._AERO_PARK["day"] = asof
    return [n for n in names if reader.parked_reason(n)]


def _is_key_trajectory(ws, c: int) -> bool:
    """A key-quarter column: in (c) Trajectory, before the key's change column."""
    if _block_of(ws, c) != "(c) Trajectory":
        return False
    hdrs = [str(ws.cell(4, k).value or "") for k in range(c, ws.max_column + 1)]
    return reader.KEY_CHANGE_HEADER in hdrs and str(ws.cell(4, c).value) != reader.KEY_CHANGE_HEADER


def check_parked(wb, asof=None) -> list:
    """Every non-number cell of a parked row's rank and key cells reads 'parked' (Summary: key and rank; a tab: its (a)
    Rank key block, the key's quarters and the key's change). A number stays shown."""
    s = wb[reader.SUMMARY]
    parked = parked_projects(list(_project_rows(s)), asof)
    if not parked:
        return []
    bad = []
    for tab in (reader.SUMMARY,) + tuple(ARCH_TABS.values()):
        ws = wb[tab]
        rows, rc = _project_rows(ws), _rank_col(ws)
        if rc is None:
            bad.append(f"{tab} has no 'Rank in' column")
            continue
        block = ([rc - 1, rc] if tab == reader.SUMMARY else
                 [c for c in range(3, ws.max_column + 1)
                  if _block_of(ws, c) == "(a) Rank key" or str(ws.cell(4, c).value or "") == reader.KEY_CHANGE_HEADER
                  or _is_key_trajectory(ws, c)])
        for n in parked:
            if n not in rows:
                continue
            for c in block:
                v = ws.cell(rows[n], c).value
                if _is_formula(v) or _is_number(v):
                    continue                                   # a number: parked figures are shown, not ranked
                if v != "parked":
                    bad.append(f"{tab} / {n} / {ws.cell(4, c).value}: reads {v!r}, not 'parked'")
    return bad


# ======================================================================================================== 6 comments
_CELL_REF = re.compile(r"^(?:(?:'[^']+'|[A-Za-z_][\w.]*)!\$?[A-Z]{1,3}\$?\d+"                 # Sheet!A1
                       r"|\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+"                            # A1:B2
                       r"|\$?[A-Z]{1,3}\$?\d+\s*(?:$|[=+\-*/(]))")                            # A1 / A1 = ...


def comment_problem(text: str) -> str:
    """'' when a comment opens with a plain sentence, else what it opens with instead (P2c-7)."""
    t = (text or "").strip()
    if t.startswith("Working:"):
        return "starts with 'Working:'"
    if t.startswith("="):
        return "starts with a formula"
    if _CELL_REF.match(t):
        return "starts with a cell reference"
    return ""


def check_comments(wb) -> list:
    """Every value-cell comment opens with a plain sentence: not 'Working:', a cell reference or a formula."""
    bad = []
    for tab in VALUE_TABS:
        if tab not in wb.sheetnames:
            continue
        ws = wb[tab]
        for name, r in _project_rows(ws).items():
            for c in range(3, ws.max_column + 1):
                cell = ws.cell(r, c)
                if cell.comment is None:
                    continue
                why = comment_problem(cell.comment.text)
                if why:
                    bad.append(f"{tab}!{cell.coordinate} ({name} / {ws.cell(4, c).value}): comment {why}")
    return bad


# ======================================================================================================== 7 guard
_VIOLATIONS = re.compile(r"violation\(s\)")
_VCOUNT = re.compile(r"[—-]\s*(\d+) violation\(s\)")


def guard_command(mode: str, store: str) -> tuple:
    """(argv, the path it runs on) of the repo's own guard entry point."""
    if mode == "compare":
        return [sys.executable, str(ROOT / "measuring_guard.py"), "--compare", str(GUARD_BASELINE)], str(GUARD_BASELINE)
    return [sys.executable, str(ROOT / "measuring_guard.py"), "--newer", str(store)], str(store)


def judge_guard(returncode: int, out: str) -> tuple:
    """(reasons, ACCEPTED lines) from the guard's output: PASS only when it prints '0 violation(s)' and exits 0."""
    accepted = []
    for ln in out.splitlines():
        t = ln.strip()
        if t.startswith("ACCEPTED:") and t not in accepted:
            accepted.append(t)
    counts = [int(m.group(1)) for ln in out.splitlines() for m in [_VCOUNT.search(ln)] if m and "GUARD" in ln]
    reasons = []
    if not counts:
        reasons.append(f"the guard printed no 'N violation(s)' line (exit code {returncode}); last output: "
                       + (out.strip().splitlines()[-1][:200] if out.strip() else "(none)"))
    elif counts[0] != 0:
        viol = []
        grab = False
        for ln in out.splitlines():
            if _VIOLATIONS.search(ln):
                grab = True
                continue
            if grab:
                if not ln.startswith("  ") or ln.strip().startswith("ACCEPTED"):
                    break
                viol.append(ln.strip())
        reasons.append(f"{counts[0]} violation(s)")
        reasons += viol
    elif returncode != 0:
        reasons.append(f"0 violation(s) printed but the guard exited with code {returncode}")
    return reasons, accepted


def check_guard(mode: str, store: str, run=subprocess.run) -> tuple:
    """Run `measuring_guard.py --newer <store>` (or `--compare <baseline>`) as a subprocess. Returns
    (reasons, notes): notes = 'ran on: <path>' and every ACCEPTED line."""
    argv, ran_on = guard_command(mode, store)
    notes = [f"ran on: {ran_on}"]
    if mode != "compare" and not Path(store).exists():
        return [f"store not found: {store}"], notes
    try:
        p = run(argv, cwd=str(ROOT), capture_output=True, timeout=GUARD_TIMEOUT,
                env={**os.environ, "PYTHONIOENCODING": "utf-8"})
    except Exception as e:                                                # a crash or a timeout is a FAIL
        return [f"the guard did not run: {type(e).__name__}: {e}"], notes
    out = (p.stdout or b"").decode("utf-8", "replace")
    err = (p.stderr or b"").decode("utf-8", "replace")
    reasons, accepted = judge_guard(p.returncode, out)
    if reasons and err.strip() and not _VCOUNT.search(out):
        reasons.append("stderr: " + err.strip().splitlines()[-1][:200])
    head = [ln.strip() for ln in out.splitlines() if ln.startswith("ACCEPTED LOSSES")]     # the guard's own count
    return reasons, notes + head[:1] + accepted


# ======================================================================================================== 8 scoreboard
def read_scoreboard(wb) -> dict:
    """{project: {column: n}} of the audit file's Coverage tab scoreboard (ALL first), or {} when it is not found."""
    if "Coverage" not in wb.sheetnames:
        return {}
    head, board = None, {}
    for row in wb["Coverage"].iter_rows(min_row=1, max_row=200, values_only=True):
        if head is None:
            if row and row[0] == "Project" and "BROKEN" in row and "NOT WIRED" in row:
                head = list(row)
            continue
        if not row or not row[0]:
            break
        board[str(row[0])] = {h: v for h, v in zip(head[1:], row[1:]) if h}
    return board


def scoreboard_verdict(board: dict) -> tuple:
    """(reasons, notes) from a scoreboard. BROKEN > 0 FAILS the line (a reference or check that disagrees). NOT WIRED
    is the standing, declared count (a cell labelled 'not wired'): it is PRINTED with its projects, never a FAIL —
    a check that is red every day tells Jake nothing (P2f-2)."""
    if "ALL" not in board:
        return ["the Coverage tab holds no scoreboard (no 'Project' header row with BROKEN and NOT WIRED)"], []
    n = {k: int(board["ALL"].get(k) or 0) for k in ("BROKEN", "NOT WIRED")}
    notes = [f"BROKEN {n['BROKEN']}, NOT WIRED {n['NOT WIRED']}"]
    named_by = {}
    for k in ("BROKEN", "NOT WIRED"):
        named = {p: int(b.get(k) or 0) for p, b in board.items() if p != "ALL" and int(b.get(k) or 0)}
        rest = n[k] - sum(named.values())
        if rest:
            named["(no project: SPEC list)"] = rest
        named_by[k] = named
    if named_by["NOT WIRED"]:
        notes.append("NOT WIRED (declared, informational): " + ", ".join(f"{p}: {v}" for p, v in named_by["NOT WIRED"].items()))
    if n["BROKEN"]:
        return [f"{n['BROKEN']} BROKEN: " + ", ".join(f"{p}: {v}" for p, v in named_by["BROKEN"].items())], notes
    return [], notes


def check_scoreboard(audit_wb) -> tuple:
    return scoreboard_verdict(read_scoreboard(audit_wb))


# ======================================================================================================== 9 as-of
def check_asof(reader_asof, audit_asof, today: date) -> list:
    """The workbook's as-of day (reader About; audit Run Log) equals today (or --asof)."""
    bad = []
    for label, d in (("reader (About)", reader_asof), ("audit (Run Log)", audit_asof)):
        if d is None:
            bad.append(f"{label}: no as-of date found (today is {today})")
        elif d != today:
            bad.append(f"{label}: as-of {d}, today {today}")
    return bad


# ======================================================================================================== runner
def format_line(status: str, name: str, reasons: list) -> str:
    if status == "PASS":
        return f"PASS  {name}"
    shown = "; ".join(reasons[:MAX_REASONS]) + (f"; and {len(reasons) - MAX_REASONS} more"
                                                if len(reasons) > MAX_REASONS else "")
    return f"FAIL  {name} — {shown}"


def run_checks(reader_path, audit_path, store, asof, guard_mode) -> list:
    """[(name, reasons, notes)] for the nine checks. A check that cannot run reports why and FAILs."""
    import openpyxl                                                       # noqa: PLC0415
    import xlcalc                                                         # noqa: PLC0415
    today = date.fromisoformat(asof) if asof else date.today()
    results = []

    def add(name, fn):
        try:
            out = fn()
        except Exception as e:                                            # a check that crashes is a FAIL, said
            out = [f"check crashed: {type(e).__name__}: {e}"]
        reasons, notes = out if isinstance(out, tuple) else (out, [])
        results.append((name, reasons, notes))

    rwb = awb = X = None
    rerr = aerr = ""
    try:
        rwb = openpyxl.load_workbook(reader_path)
        X = xlcalc.Workbook(rwb)
    except Exception as e:
        rerr = f"cannot open {reader_path}: {type(e).__name__}: {e}"
    try:
        awb = openpyxl.load_workbook(audit_path, read_only=True)
    except Exception as e:
        aerr = f"cannot open {audit_path}: {type(e).__name__}: {e}"

    def on_reader(fn):
        return (lambda: [rerr]) if rwb is None else fn

    r_asof = _asof_of_reader(rwb) if rwb is not None else None
    a_asof = _asof_of_audit(awb) if awb is not None else None
    add("ranks identical on Summary and the primary tab", on_reader(lambda: check_ranks(rwb, X)))
    add("About coverage counts equal Summary", on_reader(lambda: check_about_counts(rwb, X)))
    add(f"no header over {reader.MAX_HEADER} characters", on_reader(lambda: check_headers(rwb)))
    add("no blank cell without a code", on_reader(lambda: check_no_blank(rwb)))
    add("parked rows read 'parked'", on_reader(lambda: check_parked(rwb, r_asof)))
    add("comments open with a plain sentence", on_reader(lambda: check_comments(rwb)))
    add("guard: 0 violations", lambda: check_guard(guard_mode, store))
    add("scoreboard: no BROKEN (NOT WIRED listed)", (lambda: [aerr]) if awb is None else (lambda: check_scoreboard(awb)))
    add("as-of date is today", lambda: ([rerr] if rerr else []) + ([aerr] if aerr else [])
        or check_asof(r_asof, a_asof, today))
    return results


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--reader", default=reader.READER_NAME, help=f"the reader workbook (default {reader.READER_NAME})")
    ap.add_argument("--audit", default=reader.AUDIT_NAME, help=f"the audit workbook (default {reader.AUDIT_NAME})")
    ap.add_argument("--store", default="metrics.db", help="the store the guard reads (default metrics.db)")
    ap.add_argument("--asof", help="YYYY-MM-DD: compare the workbook's as-of with this day instead of today")
    ap.add_argument("--guard-mode", choices=("newer", "compare"), default="newer",
                    help="newer: measuring_guard.py --newer <store> (default); compare: --compare on the fixtures")
    ap.add_argument("--compare", action="store_true", help="same as --guard-mode compare (the fixture build)")
    a = ap.parse_args(argv)
    if a.asof:
        try:
            date.fromisoformat(a.asof)
        except ValueError:
            ap.error(f"--asof must be YYYY-MM-DD, got {a.asof!r}")
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")                          # a console that cannot show a character
    mode = "compare" if a.compare else a.guard_mode
    results = run_checks(a.reader, a.audit, a.store, a.asof, mode)
    npass = nfail = 0
    for name, reasons, notes in results:
        if reasons:
            nfail += 1
            print(format_line("FAIL", name, reasons))
        else:
            npass += 1
            print(format_line("PASS", name, reasons))
        for n in notes:
            print("      " + n)
    print(f"FRIDAY CHECK: {npass} PASS, {nfail} FAIL")
    return 1 if nfail else 0


if __name__ == "__main__":
    import progress                                    # stdout line-buffered when piped (2026-10-11)
    progress.line_buffered()
    sys.exit(main())
