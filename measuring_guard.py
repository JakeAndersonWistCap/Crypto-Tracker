"""
measuring_guard.py — THE MEASURING-POINT GUARD (Round G8 item 0, Jake 2026-10-20).

Two rounds in a row a change of basis shipped without a declared handover — NEAR's issuance in G6, Hyperliquid's burn
total in G7 — and each blanked headline cells. The fixtures could not show it: the series that moved (NEAR's measured
epoch mint, Hyperliquid's HyperEVM priority burn) were stored after the fixtures were dumped (2026-10-13).

The guard builds the workbook on the fixture store WITH PROBE ROWS — one week of synthetic rows, source "probe:<metric>",
for every series the project reads that the fixtures do not hold — so a view that changes when a newer series arrives is
exercised too. It records, per (project, series), the set of distinct source BASES (trajectory.basis_of, the basis the
series-break and measuring-point rules read) and every headline cell that holds a value. `compare` fails on:
  - a series whose set of bases gained a basis that is neither a declared handover point (config series_handover
    ordered_points) nor a declared component (config.DECLARED_COMPONENTS);
  - a headline that held a value before and holds none after (config.REGRESSION_ALLOWLIST answers it, as the build's
    own continuity guard does);
  - a headline whose text names measuring_point_changed, even one blank in the baseline too (G9 C3) — no allowance.
"Before" is the baseline committed beside the tests (tests/fixtures/guard_baseline.json), written by the last accepted
code; "after" is the code under test. A legitimate change of basis is declared in config, then the baseline rewritten:

    python measuring_guard.py --write tests/fixtures/guard_baseline.json
    python measuring_guard.py --compare tests/fixtures/guard_baseline.json
    python measuring_guard.py --newer metrics.db        # the checks on the live store (a copy): violations or OK
"""
from __future__ import annotations

import argparse
import json
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent
PROBE_DAYS = 7
MPC = "measuring_point_changed"


def _payloads() -> list[dict]:
    return [json.loads(f.read_text()) for f in sorted((ROOT / "fixtures").glob("real_*.json"))]


def probe_rows(payload: dict) -> list[dict]:
    """One week of probe rows for every stored (not view-only) series of the project the dump does not hold."""
    import pandas as pd
    import config
    name, asof = payload["meta"]["project"], pd.Timestamp(payload["meta"]["asof"])
    p = config.PROJECT_BY_NAME.get(name)
    if p is None:
        return []
    have = {r["metric"] for r in payload["metrics"]}
    out = []
    for m in sorted(config.metrics_for_project(p)):
        d = config.METRICS.get(m) or {}
        if m in have or d.get("view_only"):
            continue
        for k in range(PROBE_DAYS, 0, -1):
            out.append({"date": str((asof - pd.Timedelta(days=k)).date()), "metric": m, "value": 1.0,
                        "source": f"probe:{m}", "tier": 3, "fetched_at": f"{asof.date()}T00:00:00Z"})
    return out


def build(out_dir: Path, probes: bool = True, previous: dict | None = None) -> dict:
    """Build on the fixture store (+ probes); the snapshot: bases per series, headline values, probed series — and, given
    a previous snapshot's headlines, sanity's own lost-headline guard (sanity.lost_headlines) on this build."""
    import pandas as pd
    sys.path.insert(0, str(ROOT / "tests"))
    sys.modules.setdefault("recalc", None)                 # no LibreOffice recalculation in a guard build
    import openpyxl
    import build_workbook as bw
    import store as sm
    import trajectory
    import xlcalc
    from real_dump import to_store
    import sanity
    out_dir = Path(out_dir)
    db = out_dir / "metrics.db"
    if db.exists():
        db.unlink()
    asofs, probed = [], {}
    for pl in _payloads():
        rows = probe_rows(pl) if probes else []
        if rows:
            probed[pl["meta"]["project"]] = sorted({r["metric"] for r in rows})
        to_store(dict(pl, metrics=pl["metrics"] + rows), db)
        asofs.append(pl["meta"]["asof"])
    sm.DB_PATH = str(db)
    st = sm.Store(str(db))
    path = out_dir / "wb.xlsx"
    try:
        bw.build_workbook(st, path, asof=pd.Timestamp(max(asofs)))
    finally:
        st.close()
    bases = {}
    for (proj, metric), g in bw._LAST_GROUPS.items():
        if g is None or getattr(g, "empty", True) or "source" not in g:
            continue
        bases[f"{proj}|{metric}"] = sorted({trajectory.basis_of(metric, s) for s in g["source"].astype(str).unique()})
    wb = openpyxl.load_workbook(path)
    X = xlcalc.Workbook(wb)
    heads: dict = {}
    mpc: list = []
    for hc in bw._HEADLINE_CELLS:
        cell = wb[hc["sheet"]][hc["cell"]]
        try:
            v = X.value(hc["sheet"], cell.row, cell.column)
        except Exception:  # noqa: BLE001
            v = None
        n = sanity.stated_number(v)
        if n is not None:
            heads.setdefault(hc["project"], {})[f"{hc['sheet'][:2]} | {hc['header']}"] = float(n)
        elif MPC in str(v or ""):
            mpc.append([hc["project"], f"{hc['sheet'][:2]} | {hc['header']}", str(v)[:200]])
    snap = {"asof": str(max(asofs)), "bases": bases, "headlines": heads, "probed": probed, "mpc_headlines": mpc}
    if previous is not None:
        snap["sanity_lost"] = [list(h) for h in sanity.lost_headlines(wb, previous, bw._HEADLINE_CELLS, "baseline")]
    return snap


def newer_series(db_path: str) -> list[tuple[str, str, str, str]]:
    """(project, metric, newest stored, newest in the fixtures or "none") for every series a store holds past the
    fixtures' cut — the series a fixture build can only exercise through probe rows (Round G8 item 0)."""
    import sqlite3
    fx = {}
    for pl in _payloads():
        name = pl["meta"]["project"]
        for r in pl["metrics"]:
            k = (name, r["metric"])
            fx[k] = max(fx.get(k, ""), str(r["date"])[:10])
    cut = max(str(pl["meta"]["asof"])[:10] for pl in _payloads())
    con = sqlite3.connect(db_path)
    try:
        got = con.execute("SELECT project, metric, MAX(date) FROM metrics GROUP BY project, metric").fetchall()
    finally:
        con.close()
    names = {pl["meta"]["project"] for pl in _payloads()}
    return sorted((p, m, str(d)[:10], fx.get((p, m), "none")) for p, m, d in got
                  if p in names and str(d)[:10] > max(cut, fx.get((p, m), "")))


def live_check(db_path: str, baseline: dict) -> tuple[list[str], dict]:
    """ROUND G10 5d (Jake: `--newer metrics.db` printed 570 series and no verdict): the basis and lost-headline checks ON
    THE LIVE STORE. The workbook is built on a COPY of the store (the store itself is never written) at its newest day;
    then, per series, a basis outside the fixture baseline's real bases and the declared handover points / components,
    on a series that now holds more than one basis, is a violation (a series the fixtures held only as probe rows has
    no fixture basis: its bases must all be declared once there are two); a headline the PREVIOUS RUN'S SNAPSHOT held
    that this build does not (config.regression_allowance answers it), and a headline blanked by
    measuring_point_changed, are violations. Returns (violations, snapshot)."""
    import shutil
    import sqlite3
    import pandas as pd
    sys.path.insert(0, str(ROOT / "tests"))
    sys.modules.setdefault("recalc", None)
    import openpyxl
    import build_workbook as bw
    import config
    import sanity
    import store as sm
    import trajectory
    import xlcalc
    with tempfile.TemporaryDirectory() as d:
        db = Path(d) / "metrics.db"
        shutil.copyfile(db_path, db)
        con = sqlite3.connect(str(db))
        asof = pd.Timestamp(con.execute("SELECT MAX(date) FROM metrics").fetchone()[0]).normalize()
        con.close()
        sm.DB_PATH = str(db)
        st = sm.Store(str(db))
        path = Path(d) / "wb.xlsx"
        try:
            bw.build_workbook(st, path, asof=asof)
        finally:
            st.close()
        bases = {}
        for (proj, metric), g in bw._LAST_GROUPS.items():
            if g is None or getattr(g, "empty", True) or "source" not in g:
                continue
            bases[f"{proj}|{metric}"] = sorted({trajectory.basis_of(metric, s) for s in g["source"].astype(str).unique()})
        wb = openpyxl.load_workbook(path)
        X = xlcalc.Workbook(wb)
        heads, mpc = {}, []
        for hc in bw._HEADLINE_CELLS:
            cell = wb[hc["sheet"]][hc["cell"]]
            try:
                v = X.value(hc["sheet"], cell.row, cell.column)
            except Exception:  # noqa: BLE001
                v = None
            n = sanity.stated_number(v)
            if n is not None:
                heads.setdefault(hc["project"], {})[f"{hc['sheet'][:2]} | {hc['header']}"] = float(n)
            elif MPC in str(v or ""):
                mpc.append([hc["project"], f"{hc['sheet'][:2]} | {hc['header']}", str(v)[:200]])
    bad = []
    for key, now in sorted(bases.items()):
        if len(now) < 2:
            continue
        proj, metric = key.split("|", 1)
        was = {b for b in (baseline.get("bases") or {}).get(key) or () if not str(b).startswith("probe:")}
        new = set(now) - was - _declared(proj, metric)
        if new:
            bad.append(f"BASIS {proj} {metric}: {sorted(new)} beside {sorted(set(now) - new)} with no declared handover "
                       f"or component")
    # LOST HEADLINES against the PREVIOUS RUN'S SNAPSHOT (export_snapshot, written at the end of every run) — the fixture
    # baseline's headlines are built with probe rows (a probed BUSINESS slot reads 1), so they are no "before" for a
    # live store
    import export_snapshot as _es
    prev_date, prev = _es.load_previous(asof - pd.Timedelta(days=1))
    for proj, cells in (prev or {}).items():
        for head, v in cells.items():
            if not isinstance(v, float) or head in (heads.get(proj) or {}) or config.regression_allowance(proj, head):
                continue
            bad.append(f"HEADLINE LOST {proj} {head}: {v:,.6g} in the {prev_date} snapshot, no value on the live store")
    for proj, head, text in mpc:
        bad.append(f"HEADLINE BLANKED BY {MPC} {proj} {head}: {text}")
    return bad, {"asof": str(asof.date()), "bases": bases, "headlines": heads, "mpc_headlines": mpc,
                 "previous": prev_date}


def _declared(project: str, metric: str) -> set:
    """Bases a series may gain: its declared handover points and declared components."""
    import config
    import trajectory
    out = set()
    for pt in ((config.declared_handover(project, metric) or {}).get("ordered_points") or ()):
        out.add(trajectory.basis_of(metric, pt))
    for c in (getattr(config, "DECLARED_COMPONENTS", {}) or {}).get((project, metric), ()):
        out.add(trajectory.basis_of(metric, c.get("source", "")))
    # a DECLARED ZERO (config.DECLARED_ZERO, G10 3b: Sky's Stage 2 burn before its first row) is the series' own
    # declaration — the build never counts it as a measuring point (aggregate drops declared:zero from the points)
    if (project, metric) in (getattr(config, "DECLARED_ZERO", {}) or {}):
        out.add(trajectory.basis_of(metric, "declared:zero[x]"))
    return out


def compare(before: dict, after: dict) -> list[str]:
    """Every undeclared new basis of an existing series, and every headline lost."""
    import config
    bad = []
    for key, now in after["bases"].items():
        was = before["bases"].get(key)
        if was is None:
            continue                                      # a new series is not a change of basis
        proj, metric = key.split("|", 1)
        new = set(now) - set(was) - _declared(proj, metric)
        if new:
            bad.append(f"BASIS {proj} {metric}: {sorted(new)} added to {was} with no declared handover or component")
    for proj, cells in before["headlines"].items():
        for head, v in cells.items():
            if head in (after["headlines"].get(proj) or {}):
                continue
            if config.regression_allowance(proj, head):
                continue
            bad.append(f"HEADLINE LOST {proj} {head}: {v:,.6g} before, no value after")
    # G9 C3: a headline blanked by a change of measuring point fails whatever the baseline held — blank before too
    for proj, head, text in after.get("mpc_headlines") or ():
        bad.append(f"HEADLINE BLANKED BY {MPC} {proj} {head}: {text}")
    return bad


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--write", help="build and write the baseline snapshot to this path")
    g.add_argument("--compare", help="build and compare with this baseline snapshot")
    g.add_argument("--newer", metavar="DB", help="the basis and lost-headline checks on this (live) store, built on a "
                                                 "copy; with --list, the series it holds past the fixtures' cut")
    ap.add_argument("--list", action="store_true", help="with --newer: list every series newer than the fixtures")
    a = ap.parse_args(argv)
    if a.newer:
        rows = newer_series(a.newer)
        print(f"SERIES NEWER THAN THE FIXTURES — {len(rows)} (the count only; the checks below are the verdict)")
        if a.list:
            for p, m, d, f in rows:
                print(f"  {p:<14} {m:<48} stored to {d}  fixtures to {f}")
        # G10 5d: the verdict — the basis and lost-headline checks run on (a copy of) the live store
        base = json.loads((ROOT / "tests" / "fixtures" / "guard_baseline.json").read_text(encoding="utf-8"))
        bad, snap = live_check(a.newer, base)
        print(f"LIVE STORE (built on a copy, as of {snap['asof']}): {len(snap['bases'])} series, "
              f"{sum(len(v) for v in snap['headlines'].values())} headlines with a value; lost headlines judged "
              + (f"against the {snap['previous']} snapshot" if snap["previous"] else
                 "against NO snapshot (none on file before the store's last day — export_snapshot writes one at the end "
                 "of every run): only the basis and measuring_point_changed checks ran"))
        print(f"MEASURING-POINT GUARD ON THE LIVE STORE — {len(bad)} violation(s)")
        for b in bad:
            print("  " + b)
        print("VERDICT: " + ("OK — no undeclared basis, no lost headline" if not bad else
                             f"{len(bad)} VIOLATION(S) — each line above"))
        return 1 if bad else 0
    base = json.loads(Path(a.compare).read_text(encoding="utf-8")) if a.compare else None
    with tempfile.TemporaryDirectory() as d:
        snap = build(Path(d), previous=(base or {}).get("headlines"))
    if a.write:
        Path(a.write).write_text(json.dumps(snap, indent=1, sort_keys=True), encoding="utf-8")
        print(f"baseline written: {len(snap['bases'])} series, "
              f"{sum(len(v) for v in snap['headlines'].values())} headlines with a value, "
              f"{sum(len(v) for v in snap['probed'].values())} probed series")
        return 0
    bad = compare(base, snap)
    print(f"MEASURING-POINT GUARD — {len(bad)} violation(s)")
    for b in bad:
        print("  " + b)
    lost = snap.get("sanity_lost") or []
    print(f"SANITY lost_headlines on the same build — {len(lost)} hit(s)")
    for p, _rule, msg in lost:
        print(f"  {p}: {msg[:200]}")
    return 1 if bad else 0


if __name__ == "__main__":
    import progress
    progress.line_buffered()
    sys.exit(main())
