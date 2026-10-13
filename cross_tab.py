#!/usr/bin/env python3
"""
cross_tab.py — Master must agree with the A-tabs (Jake's review of the 97b29ae run, 2026-10-11, item 16).

Master is the headline and the daily snippet reads it, so every Master column that summarises an A-tab figure must
show the SAME number as that A-tab cell for the project: same series, same window, same supply basis, same coverage
rule. And the same figure must not differ between two A-tabs (item 17: "the same metric for the same project differs
between any two tabs").

    python cross_tab.py                       build today's workbook from metrics.db (a temp file), evaluate it with
                                              xlcalc and list every mismatch: project, column, Master value, A-tab
                                              value, and the formula (the series) each read
    python cross_tab.py --workbook X.xlsx     check an existing workbook

MASTER_SOURCES maps a Master column to the A-tab column(s) it summarises, in the order tried: the project's own
archetype tabs first (its primary archetype, then the others). build_workbook writes those Master cells as a
reference to the A-tab cell (so they agree by construction) and refuses a build whose Master cell for a mapped column
is anything else (`structural_mismatches`, raised as CrossTabMismatch). This module's value check is the second
guard: it evaluates both cells and compares the numbers.
"""
from __future__ import annotations

import argparse
import sys

# Master header -> [(A-tab sheet prefix, A-tab header), ...] in priority order. A prefix "A*" means: the project's
# own archetype tabs, primary first.
MASTER_SOURCES = {
    "Buyback status": [("A3", "Buyback status")],
    "Burn status": [("A4", "Burn status")],
    "Market cap ($)": [("A*", "Market cap ($)")],
    "Circulating supply (used)": [("A3", "Circulating supply")],
    "Fees Q0 ($)": [("A*", "Fees Q0 ($)")],
    "Implied buyback Q0 ($) = revenue × the split that applied to this window": [("A3", "Implied buyback Q0 ($) = revenue × share")],
    "Actual buyback Q0 ($) — observed": [("A3", "Actual buyback Q0 ($) — observed")],
    # EVERY BURN (config.a1_burn_metric, Phase 2 amendment 1): A1's total first; A4 where the project has no A1 row
    "Gross burn Q0 (tokens)": [("A1", "Gross burn Q0 (tokens) — every burn, fee-market included"),
                               ("A4", "GROSS BURN Q0 (tokens)")],
    "Gross issuance Q0 (tokens)": [("A4", "GROSS ISSUANCE Q0 (tokens)"), ("A1", "Gross issuance Q0 (tokens)")],
    # the fee-burn chains' net change is on A1 since Phase 2 amendment 1 (Jake 2026-10-12)
    "NET SUPPLY CHANGE Q0 (tokens) = issuance − burn": [
        ("A4", "NET SUPPLY CHANGE Q0 (tokens) — self-reported where published, else issuance − burn, on the basis to the left"),
        ("A1", "NET ISSUANCE AFTER BURN Q0 (tokens) — issuance − every burn, on the basis A4 uses")],
    "Net supply change, annualised % of circulating": [
        ("A4", "NET SUPPLY CHANGE, annualised % of circulating (signed; + is net inflation)"),
        ("A1", "NET ISSUANCE AFTER BURN, annualised % of circulating (+ is net inflation; fee-market burns included)")],
    "Fees ÷ issuance ($, Q0; issuance at 90d avg price)": [("A1", "FEES ÷ ISSUANCE (x) — key ratio")],
    "Fees ÷ FDV (annualised)": [("A3", "Fees ÷ FDV (annualised)")],
    "Buyback % of supply (annualised, implied, split as at this window)": [
        ("A3", "BUYBACK AS % OF SUPPLY (annualised, implied)")],
}
A_TABS = ("A1", "A2", "A3", "A4", "CP")       # CP: the Phase 2 comparability tab (2026-10-12)
EPS = 1e-9


POOL_RELEASE = "Pool release Q0 (tokens) — measured, premint distribution (where issuance is n/a)"


def sources_for(head: str, project: str) -> list:
    """MASTER_SOURCES for one project. FIX ROUND 3 (R9, Jake 2026-10-13: "Master Gross issuance shows releases for
    GEODNET (11.79M) and Hyperliquid (3.84M) — pre-minted supply = 0, releases on their own line"): a premint project's
    Master gross issuance reads A4's "0 — fixed supply" like every other tab; its pool release is Master's own "Issuer
    releases Q0" column. (Was: A4's pool-release column, Jake's review of 97b29ae item 3.)"""
    return MASTER_SOURCES.get(head) or []


def tab_order(archetypes) -> list[str]:
    """The project's own A-tabs, primary archetype first."""
    return [f"A{a}" for a in archetypes if f"A{a}" in A_TABS]


def resolve(sources, archetypes, cells: dict, project: str):
    """The first (sheet, header, coordinate) among `sources` that the project has a cell for; `cells` maps
    (sheet prefix, project, header) -> (sheet title, coordinate)."""
    for prefix, header in sources:
        for tab in (tab_order(archetypes) if prefix == "A*" else [prefix]):
            hit = cells.get((tab, project, header))
            if hit:
                return hit + (header,)
    return None


def _same(a, b) -> bool:
    na = lambda v: v is None or v == "" or (isinstance(v, str) and v.strip().lower() in ("n/a", "#n/a"))  # noqa: E731
    if na(a) and na(b):
        return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= EPS * max(1.0, abs(a), abs(b))
    return str(a).strip() == str(b).strip()


def _trend_na(v) -> bool:
    """A trend cell that says nothing: empty, an error, 'n/a — …', or an arrowless text with no figure in it."""
    if v is None or v == "":
        return True
    if not isinstance(v, str):
        return not isinstance(v, (int, float))            # an evaluator error value (#N/A)
    t = v.strip()
    return t.lower().startswith(("n/a", "#n/a")) or (t.startswith("–") and "%" not in t)


def _grid(wb):
    """{(sheet prefix, project, header): (sheet title, row, col)} for Master and the A-tabs (header row = the row whose
    first cell is 'Project')."""
    out = {}
    for ws in wb.worksheets:
        pre = "Master" if ws.title == "Master" else ws.title[:2]
        if pre not in ("Master",) + A_TABS:
            continue
        hr = next((r for r in range(1, 12) if ws.cell(r, 1).value == "Project"), None)
        if hr is None:
            continue
        heads = {c: ws.cell(hr, c).value for c in range(1, ws.max_column + 1) if ws.cell(hr, c).value}
        for r in range(hr + 1, ws.max_row + 1):
            proj = ws.cell(r, 1).value
            if not proj:
                break                                     # the first table only (A1 has a monthly table below it)
            for c, h in heads.items():
                out[(pre, str(proj), str(h))] = (ws.title, r, c)
    return out


def _trend_beside(wb, loc) -> bool:
    """Is the column right of this cell a trajectory trend column (build_workbook._trend_spec)?"""
    ws = wb[loc[0]]
    hr = next((r for r in range(1, 12) if ws.cell(r, 1).value == "Project"), None)
    return bool(hr) and str(ws.cell(hr, loc[2] + 1).value or "").startswith("Trend — ")


def mismatches(wb, value) -> list[dict]:
    """Every Master-vs-A-tab and A-tab-vs-A-tab disagreement. `wb`: an openpyxl workbook (formulas); `value(sheet, row,
    col)`: the evaluated value."""
    import config
    grid = _grid(wb)
    cells = {k: (v[0], v[1], v[2]) for k, v in grid.items()}
    found = []
    projects = sorted({p for (pre, p, _h) in grid if pre == "Master"})
    for p in projects:
        arch = (config.PROJECT_BY_NAME.get(p) or {}).get("archetypes") or []
        for mh in MASTER_SOURCES:
            sources = sources_for(mh, p)
            m = grid.get(("Master", p, mh))
            if not m:
                continue
            hit = None
            for prefix, header in sources:
                for tab in (tab_order(arch) if prefix == "A*" else [prefix]):
                    if (tab, p, header) in cells:
                        hit = cells[(tab, p, header)] + (header,)
                        break
                if hit:
                    break
            if not hit:
                continue
            mv, av = value(m[0], m[1], m[2]), value(hit[0], hit[1], hit[2])
            if not _same(mv, av):
                found.append({"project": p, "column": mh, "master": mv, "atab": av, "atab_cell": f"{hit[0]}!{hit[3]}",
                              "master_formula": str(wb[m[0]].cell(m[1], m[2]).value)[:300],
                              "atab_formula": str(wb[hit[0]].cell(hit[1], hit[2]).value)[:300]})
            # THE TREND BESIDE IT (theory layer, Phase 1): a Master trend cell beside a column that references an A-tab
            # headline reads that headline's trajectory, so it must equal the A-tab's own trend cell
            mt, at = _trend_beside(wb, m), _trend_beside(wb, hit)
            if mt and at:
                tv, ta = value(m[0], m[1], m[2] + 1), value(hit[0], hit[1], hit[2] + 1)
                # TWO EMPTY TRENDS AGREE however each says it (Jake's Phase 1 review 2026-10-12, item 9: PancakeSwap,
                # parked, read as a mismatch with n/a on both sides)
                if not _same(tv, ta) and not (_trend_na(tv) and _trend_na(ta)):
                    found.append({"project": p, "column": f"{mh} (trend)", "master": tv, "atab": ta,
                                  "atab_cell": f"{hit[0]}!{hit[3]} (trend)",
                                  "master_formula": str(wb[m[0]].cell(m[1], m[2] + 1).value)[:300],
                                  "atab_formula": str(wb[hit[0]].cell(hit[1], hit[2] + 1).value)[:300]})
        # THE SAME HEADER ON TWO A-TABS must read the same for the project
        by_head: dict = {}
        for (pre, proj, h), loc in grid.items():
            if proj == p and pre in A_TABS:
                by_head.setdefault(h, []).append(loc)
        for h, locs in by_head.items():
            if len(locs) < 2 or h in ("Project", "Symbol", "Notes", "Data flags (stale / manual / missing)"):
                continue
            vals = [(loc, value(*loc)) for loc in locs]
            for (l0, v0), (l1, v1) in zip(vals, vals[1:]):
                if not _same(v0, v1):
                    found.append({"project": p, "column": h, "master": None, "atab": None,
                                  "pair": f"{l0[0]} {v0!r} vs {l1[0]} {v1!r}",
                                  "master_formula": str(wb[l0[0]].cell(l0[1], l0[2]).value)[:300],
                                  "atab_formula": str(wb[l1[0]].cell(l1[1], l1[2]).value)[:300]})
    return found


def check_workbook(path) -> list[dict]:
    import openpyxl
    import xlcalc
    wb = openpyxl.load_workbook(path)
    X = xlcalc.Workbook(wb)
    return mismatches(wb, lambda s, r, c: xlcalc.display(X.value(s, r, c)))


def describe(m: dict) -> str:
    if m.get("pair"):
        return (f"{m['project']} | {m['column'][:70]} | A-TABS DISAGREE: {m['pair']}\n"
                f"      reads: {m['master_formula'][:160]}\n      vs:    {m['atab_formula'][:160]}")
    return (f"{m['project']} | {m['column'][:70]} | Master {m['master']!r} vs {m['atab_cell'][:60]} {m['atab']!r}\n"
            f"      Master reads: {m['master_formula'][:160]}\n      A-tab reads:  {m['atab_formula'][:160]}")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workbook")
    a = ap.parse_args(argv)
    path = a.workbook
    if not path:
        import tempfile
        from pathlib import Path
        import store as store_mod
        from build_workbook import build_workbook
        path = str(Path(tempfile.mkdtemp()) / "cross_tab.xlsx")
        build_workbook(store_mod.Store(store_mod.DB_PATH), path)
    found = check_workbook(path)
    print(f"CROSS-TAB CONSISTENCY — {len(found)} mismatch(es)")
    for m in found:
        print("  " + describe(m).encode("ascii", "replace").decode("ascii"))
    return 1 if found else 0


if __name__ == "__main__":
    import progress                                    # stdout line-buffered when piped (2026-10-11)
    progress.line_buffered()
    sys.exit(main())
